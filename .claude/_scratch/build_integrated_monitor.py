#!/usr/bin/env python3
"""EE Integrated Monitor — CDR delivery + GWP amortization in one view"""

import pandas as pd
import numpy as np
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE   = '/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/.playwright-mcp'
OUTPUT = '/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/EE_Monitor_v1.xlsx'

SNAPSHOT = pd.Timestamp('2026-04-30')
LAST_INJ_MONTH = pd.Timestamp('2026-04-01')   # last month with closed INJ RP

# ── 1. Projection (INJ, deduplicated) ────────────────────────────────────────
proj_raw = pd.read_csv(f'{BASE}/Master-EE-UoP-Amortization-CO-Injection---Projection.csv',
    header=None, skiprows=5, usecols=[0,1], names=['month','mt'])
proj_raw = proj_raw[proj_raw['month'].astype(str).str.match(r'^\d{4}-\d{2}', na=False)].copy()
proj_raw['month'] = pd.to_datetime(proj_raw['month'])
proj_raw['mt'] = pd.to_numeric(proj_raw['mt'].astype(str).str.replace(',',''), errors='coerce').fillna(0)
proj_raw = proj_raw.drop_duplicates(subset=['month'], keep='first')
proj_dict = dict(zip(proj_raw['month'], proj_raw['mt']))

# ── 2. Assets — portfolio effective rate ─────────────────────────────────────
def sum_cdr(start, end, praw=proj_raw):
    if pd.isna(start) or pd.isna(end): return 0.0
    m1 = pd.Timestamp(start.year, start.month, 1)
    m2 = pd.Timestamp(end.year, end.month, 1)
    if m2 < m1: return 0.0
    return float(praw.loc[(praw['month'] >= m1) & (praw['month'] <= m2), 'mt'].sum())

raw_a = pd.read_csv(f'{BASE}/Master-EE-UoP-Amortization-CO-Injection---Combined-Amortization-History.csv',
    header=None, skiprows=3, low_memory=False)
raw_a.columns = raw_a.iloc[0]
df_a = raw_a.iloc[1:].reset_index(drop=True).copy()
df_a = df_a[df_a['Asset ID'].astype(str).str.match(r'^EE-', na=False)].copy()
for c in ['Total GWP','Drawn to Date']:
    df_a[c] = pd.to_numeric(df_a[c].astype(str).str.replace(r'[,\s$]','',regex=True), errors='coerce').fillna(0)
df_a['Commission']  = pd.to_datetime(df_a['Commission'], errors='coerce')
df_a['Disposal']    = pd.to_datetime(df_a['Disposal'],   errors='coerce')
df_a['is_disposed'] = df_a['Disposal'].notna()
df_a['window_close']= df_a['Commission'].apply(lambda d: d+pd.DateOffset(months=59) if pd.notna(d) else pd.NaT)

active = df_a[~df_a['is_disposed']].copy()
active['proj_full']  = active.apply(lambda r: sum_cdr(r['Commission'], r['window_close']), axis=1)
active['rate_i']     = np.where(active['proj_full']>0, active['Total GWP']/active['proj_full'], 0)
portfolio_rate = active['rate_i'].sum()   # kg GWP charged per 1 MT CDR across entire portfolio

active['remaining_gwp']  = active['Total GWP'] - active['Drawn to Date']
active['proj_remaining'] = active.apply(
    lambda r: sum_cdr(SNAPSHOT+pd.DateOffset(months=1), r['window_close']), axis=1)

total_gwp      = active['Total GWP'].sum()
total_drawn    = active['Drawn to Date'].sum()
total_remaining= active['remaining_gwp'].sum()
total_proj_rem = active['proj_remaining'].sum()

print(f"Portfolio effective rate: {portfolio_rate:.5f} kg GWP / MT CDR")
print(f"Total GWP (active):       {total_gwp:,.0f} kg")
print(f"Drawn to date:            {total_drawn:,.0f} kg")
print(f"Remaining:                {total_remaining:,.0f} kg")

# ── 3. Monthly GWP actual from Amortization History ──────────────────────────
month_map = {'Jan':1,'Feb':2,'Mar':3,'Apr':4,'May':5,'Jun':6,
             'Jul':7,'Aug':8,'Sep':9,'Oct':10,'Nov':11,'Dec':12}

def parse_mon_col(col):
    m = re.match(r'^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)-(\d{2})$', str(col))
    if not m: return None
    mo, yr = m.groups()
    return pd.Timestamp(2000+int(yr), month_map[mo], 1)

monthly_gwp_actual = {}
for col in raw_a.columns:
    ts = parse_mon_col(col)
    if ts is None or ts > LAST_INJ_MONTH: continue   # skip projected-forward columns
    vals = pd.to_numeric(
        df_a[col].astype(str).str.replace(r'[,\s$]','',regex=True), errors='coerce').fillna(0)
    s = vals.sum()
    if s > 0:
        monthly_gwp_actual[ts] = round(s, 2)

# ── 4. INJ CDR actuals → calendar months ─────────────────────────────────────
rp_data = [
    ('Mini RP',    '2025-08-11','2025-08-15', 162.188),
    ('Aug-Sep RP', '2025-08-16','2025-09-08', 280.091),
    ('Sep RP',     '2025-09-09','2025-09-30', 351.137),
    ('Oct-Nov RP', '2025-10-01','2025-11-06', 574.730),
    ('Nov-Jan RP', '2025-11-07','2026-01-31', 157.040),
    ('Feb-Mar RP', '2026-02-01','2026-03-02', 864.000),
    ('Mar RP',     '2026-03-03','2026-03-31', 711.760),
    ('Apr RP',     '2026-04-01','2026-04-30', 578.837),
]
rps = pd.DataFrame(rp_data, columns=['rp','start','end','mt'])
rps['start'] = pd.to_datetime(rps['start'])
rps['end']   = pd.to_datetime(rps['end'])
rps['days']  = (rps['end']-rps['start']).dt.days + 1

def month_cdr(m):
    m_end = m + pd.offsets.MonthEnd(0)
    t = 0.0
    for _, rp in rps.iterrows():
        ov_s = max(rp['start'], m)
        ov_e = min(rp['end'],   m_end)
        if ov_e >= ov_s:
            t += rp['mt'] * ((ov_e-ov_s).days+1) / rp['days']
    return round(t, 3) if t > 0 else None

# ── 5. Build rows ─────────────────────────────────────────────────────────────
# Start Aug 2024 (first CO production) so pre-injection CO draws are visible
all_months = pd.date_range('2024-08-01', '2030-12-01', freq='MS')
rows = []
for m in all_months:
    sched_cdr  = proj_dict.get(m, 0.0)
    actual_cdr = month_cdr(m) if m <= LAST_INJ_MONTH else None
    delta_cdr  = round(actual_cdr - sched_cdr, 3) if actual_cdr is not None else None

    sched_gwp  = round(sched_cdr * portfolio_rate, 2)
    actual_gwp = monthly_gwp_actual.get(m, None)
    delta_gwp  = round(actual_gwp - sched_gwp, 2) if actual_gwp is not None else None

    rows.append({'month':m, 'year':m.year,
                 'sched_cdr':sched_cdr,  'actual_cdr':actual_cdr,  'delta_cdr':delta_cdr,
                 'sched_gwp':sched_gwp,  'actual_gwp':actual_gwp,  'delta_gwp':delta_gwp})

df = pd.DataFrame(rows)

# Running totals
cs_c, ca_c, cs_g, ca_g = 0., 0., 0., 0.
cumul = {'cs_c':[], 'ca_c':[], 'cd_c':[], 'cs_g':[], 'ca_g':[], 'cd_g':[]}
for _, r in df.iterrows():
    cs_c += r['sched_cdr'];  cs_g += r['sched_gwp']
    cumul['cs_c'].append(round(cs_c,1));  cumul['cs_g'].append(round(cs_g,1))
    if pd.notna(r['actual_cdr']):
        ca_c += r['actual_cdr']
        cumul['ca_c'].append(round(ca_c,1));  cumul['cd_c'].append(round(ca_c-cs_c,1))
    else:
        cumul['ca_c'].append(None);  cumul['cd_c'].append(None)
    if r['actual_gwp'] is not None and pd.notna(r['actual_gwp']):
        ca_g += r['actual_gwp']
        cumul['ca_g'].append(round(ca_g,1));  cumul['cd_g'].append(round(ca_g-cs_g,1))
    else:
        cumul['ca_g'].append(None);  cumul['cd_g'].append(None)

for k, v in cumul.items():
    df[k] = v

# ── 6. Build Excel ────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = 'Monitor'

HDR_FILL  = PatternFill('solid', fgColor='1F3864')
GRP_FILL  = PatternFill('solid', fgColor='2E75B6')
YR_FILL   = PatternFill('solid', fgColor='D6DCE4')
HIST_FILL = PatternFill('solid', fgColor='FFFFFF')
FUT_FILL  = PatternFill('solid', fgColor='F7F7F7')
HDR_FONT  = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
GRP_FONT  = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
YR_FONT   = Font(name='Calibri', bold=True, size=10)
NORM      = Font(name='Calibri', size=10)
FUTURE    = Font(name='Calibri', size=10, color='AAAAAA')
L = Alignment(horizontal='left',   vertical='center')
C = Alignment(horizontal='center', vertical='center')
R = Alignment(horizontal='right',  vertical='center')
N1 = '#,##0.0'

# Row 1: summary banner
summary = (
    f'Portfolio rate: {portfolio_rate:.4f} kg GWP per MT CDR  |  '
    f'Total GWP (active): {total_gwp:,.0f} kg  |  '
    f'Drawn to date: {total_drawn:,.0f} kg  |  '
    f'Remaining: {total_remaining:,.0f} kg  |  '
    f'Snapshot: {SNAPSHOT.strftime("%b %d, %Y")}'
)
ws.merge_cells('A1:M1')
ws.cell(1,1, summary).font = Font(name='Calibri', italic=True, size=9, color='444444')
ws.cell(1,1).alignment = L
ws.row_dimensions[1].height = 14

# Row 2: group headers
for col, (label, span) in enumerate([
    ('', 1), ('', 1),
    ('GROSS CDR (MT bio-oil)', 5),
    ('GWP AMORTIZATION (kg)', 5),
], 1):
    if label:
        end_col = col + span - 1
        # Find start col by counting
        pass

# Manual group headers
ws.merge_cells('C2:G2')
ws.cell(2,3, 'GROSS CDR  (MT bio-oil)').font = GRP_FONT
ws.cell(2,3).fill = GRP_FILL
ws.cell(2,3).alignment = C

ws.merge_cells('H2:M2')
ws.cell(2,8, 'GWP AMORTIZATION  (kg)').font = GRP_FONT
ws.cell(2,8).fill = HDR_FILL
ws.cell(2,8).alignment = C

ws.cell(2,1).fill = HDR_FILL; ws.cell(2,2).fill = HDR_FILL
ws.row_dimensions[2].height = 18

# Row 3: column headers
COLS = [
    'Month', 'Year',
    'Scheduled', 'Actual', 'Delta', 'Cumul Scheduled', 'Cumul Delta',
    'Expected Draw', 'Actual Draw', 'GWP Delta', 'Cumul Expected', 'Cumul Actual', 'Cumul GWP Delta',
]
for i, h in enumerate(COLS, 1):
    c = ws.cell(3, i, h)
    c.font = HDR_FONT; c.fill = HDR_FILL
    c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
ws.row_dimensions[3].height = 30
ws.freeze_panes = 'A4'

NFMT = ['@','@', N1,N1,N1,N1,N1, N1,N1,N1,N1,N1,N1]

prev_year = None
row_i = 4
for _, r in df.iterrows():
    is_future = r['actual_cdr'] is None and r['month'] > LAST_INJ_MONTH
    is_pre_inj = r['month'] < pd.Timestamp('2025-08-01')  # CO tDBM only period

    font = FUTURE if is_future else NORM
    fill = FUT_FILL if is_future else HIST_FILL

    # Year separator
    if r['year'] != prev_year:
        for ci in range(1, 14):
            ws.cell(row_i, ci).fill = YR_FILL
            ws.cell(row_i, ci).font = YR_FONT
        ws.cell(row_i, 1, str(r['year'])).alignment = L
        ws.row_dimensions[row_i].height = 13
        row_i += 1
        prev_year = r['year']

    vals = [
        r['month'].strftime('%b %Y'), None,
        r['sched_cdr'], r['actual_cdr'], r['delta_cdr'], r['cs_c'], r['cd_c'],
        r['sched_gwp'], r['actual_gwp'], r['delta_gwp'], r['cs_g'], r['ca_g'], r['cd_g'],
    ]
    for ci, (val, fmt) in enumerate(zip(vals, NFMT), 1):
        c = ws.cell(row_i, ci, val)
        c.font = font; c.fill = fill
        c.alignment = L if ci == 1 else R
        if fmt != '@' and val is not None and not isinstance(val, str):
            c.number_format = fmt

    row_i += 1

# Column widths
for i, w in enumerate([13,6, 11,10,10,16,13, 14,13,12,16,14,16], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ── Annual rollup tab ─────────────────────────────────────────────────────────
ws2 = wb.create_sheet('Annual Rollup')
AH = ['Year','CDR Sched (MT)','CDR Actual (MT)','CDR %','GWP Expected (kg)','GWP Actual (kg)','GWP %']
for i, h in enumerate(AH, 1):
    c = ws2.cell(1,i,h); c.font=HDR_FONT; c.fill=HDR_FILL; c.alignment=C
ws2.row_dimensions[1].height = 24

for ri, (yr, grp) in enumerate(df.groupby('year'), 2):
    sc = grp['sched_cdr'].sum()
    ac = grp[grp['actual_cdr'].notna()]['actual_cdr'].sum()
    ac = ac if ac > 0 else None
    sg = grp['sched_gwp'].sum()
    ag = grp[grp['actual_gwp'].notna()]['actual_gwp'].sum()
    ag = ag if ag > 0 else None
    partial = grp['actual_cdr'].notna().sum() < len(grp)
    note = ' *' if partial and ac is not None else ''
    vals = [str(yr)+note, round(sc,1), ac, round(ac/sc*100,1) if ac else None,
            round(sg,1), ag, round(ag/sg*100,1) if ag else None]
    is_fut = ac is None and ag is None
    for ci, val in enumerate(vals, 1):
        c = ws2.cell(ri, ci, val)
        c.font = FUTURE if is_fut else NORM
        c.alignment = L if ci==1 else R
        if ci in (2,3,5,6) and val is not None: c.number_format = N1
        if ci in (4,7) and val is not None: c.number_format = '0.0"%"'
ws2.cell(len(df.groupby('year'))+3, 1, '* partial year').font = Font(name='Calibri', italic=True, size=9, color='888888')
for i, w in enumerate([12,18,17,10,20,18,10],1):
    ws2.column_dimensions[get_column_letter(i)].width = w

wb.save(OUTPUT)
print(f"\nSaved → {OUTPUT}")
print("\n=== ANNUAL ROLLUP ===")
for yr, grp in df.groupby('year'):
    sc = grp['sched_cdr'].sum()
    ac_vals = grp[grp['actual_cdr'].notna()]['actual_cdr']
    ac = ac_vals.sum() if len(ac_vals) > 0 else 0
    sg = grp['sched_gwp'].sum()
    ag_vals = grp[grp['actual_gwp'].notna()]['actual_gwp']
    ag = ag_vals.sum() if len(ag_vals) > 0 else 0
    print(f"  {yr}: CDR {ac:>8,.1f} / {sc:>10,.1f} sched ({ac/sc*100:.0f}%)   GWP {ag:>9,.1f} / {sg:>10,.1f} sched ({ag/sg*100:.0f}% )" if sc>0 else f"  {yr}: CDR {ac:,.1f}")
