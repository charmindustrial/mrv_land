#!/usr/bin/env python3
"""
MRV BI Dashboard — build + data inject.
Structural updates vs prior version:
  - Injection-side transport pulled OUT of source sections → new Downstream Transport section
    between Kerry and Well Facility, following actual system flow:
    Charm prod → AECN prod → Kerry prod → Downstream Transport → Well Facility
  - Apr-26 production left blank (PP not yet closed per Max)
  - Charts added to each tab (time-series line charts)
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel

OUT = "/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/mrv-operational-documents/MRV_BI_Dashboard_Shell.xlsx"

NAVY="1F3864"; BD="2E5FA3"; BM="4472C4"; BL="DEEAF1"
GR="F5F5F5"; WH="FFFFFF"; AM="FFF2CC"; SITE_BG="EDF2F7"
CERT_BG="E2EFDA"; PP_BG="DEEAF1"; PEND_BG="FFF2CC"

def F(c): return PatternFill("solid", fgColor=c)
def K(b=False, c="000000", s=10, i=False): return Font(bold=b,color=c,size=s,italic=i,name="Calibri")
def A(h="left", v="center", w=False): return Alignment(horizontal=h,vertical=v,wrap_text=w)

MO=["Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25",
    "Jan-26","Feb-26","Mar-26","Apr-26","May-26","Jun-26"]
SRC=["AECN","Charm","Kerry"]
CA=1; CB=2; NM=12; CY=CB+NM; CF=CY+1; CL=CF
YS,YE=6,11

def mc(i): return CB+i
def yr(row): return f"{get_column_letter(mc(YS))}{row}:{get_column_letter(mc(YE))}{row}"

def col_hdr(ws, r, lbl="Metric"):
    def hc(col,v,bg):
        c=ws.cell(r,col,v); c.fill=F(bg); c.font=K(b=True,c=WH)
        c.alignment=A(h="center" if col>CA else "left")
    hc(CA,lbl,NAVY)
    for i,m in enumerate(MO): hc(mc(i),m,BD)
    hc(CY,"YTD 2026",NAVY); hc(CF,"FY 2025*",NAVY)
    ws.row_dimensions[r].height=18

def sec(ws, r, title, bg=BD):
    ws.merge_cells(start_row=r,start_column=CA,end_row=r,end_column=CL)
    c=ws.cell(r,CA,title); c.fill=F(bg); c.font=K(b=True,c=WH,s=10); c.alignment=A()
    ws.row_dimensions[r].height=16

def dr(ws, r, lbl, ind=0, tot=False, pct=False, days=False,
        ytd=True, nytd=False, bg_ov=None, data_fill=None):
    fmt="0.0%" if pct else ("#,##0" if days else "#,##0.00")
    bg=bg_ov if bg_ov else (BL if tot else (GR if r%2==0 else WH))
    c=ws.cell(r,CA,"  "*ind+lbl); c.fill=F(bg); c.font=K(b=tot); c.alignment=A()
    for i in range(NM):
        x=ws.cell(r,mc(i)); x.fill=F(data_fill if data_fill else bg)
        x.alignment=A(h="center"); x.number_format=fmt
    if nytd:
        ws.cell(r,CY).fill=F(bg)
    else:
        y=ws.cell(r,CY)
        if ytd: y.value=f"=SUM({yr(r)})"
        y.fill=F(bg); y.font=K(b=tot); y.alignment=A(h="center"); y.number_format=fmt
    f=ws.cell(r,CF); f.fill=F(bg); f.alignment=A(h="center"); f.number_format=fmt
    ws.row_dimensions[r].height=15

def sp(ws, r):
    for col in range(CA,CL+1): ws.cell(r,col).fill=F(WH)
    ws.row_dimensions[r].height=5

def nt(ws, r, txt):
    ws.merge_cells(start_row=r,start_column=CA,end_row=r,end_column=CL)
    c=ws.cell(r,CA,txt); c.font=K(c="888888",s=9,i=True); c.alignment=A(w=True)
    ws.row_dimensions[r].height=14

def title_block(ws, r, title, subtitle):
    ws.merge_cells(start_row=r,start_column=CA,end_row=r,end_column=CL)
    ws.merge_cells(start_row=r+1,start_column=CA,end_row=r+1,end_column=CL)
    t=ws.cell(r,CA,title); t.font=K(b=True,c=WH,s=13); t.fill=F(NAVY); t.alignment=A(h="center")
    ws.row_dimensions[r].height=22
    s=ws.cell(r+1,CA,subtitle); s.font=K(c="555555",s=9,i=True); s.fill=F(BL); s.alignment=A(h="center")
    ws.row_dimensions[r+1].height=13

def set_w(ws, lw=34, fp="B4"):
    ws.column_dimensions["A"].width=lw
    for i in range(NM): ws.column_dimensions[get_column_letter(mc(i))].width=8.5
    ws.column_dimensions[get_column_letter(CY)].width=11
    ws.column_dimensions[get_column_letter(CF)].width=12
    ws.freeze_panes=fp

def chart_hdr(ws, r, title):
    ws.merge_cells(start_row=r,start_column=CA,end_row=r,end_column=CL)
    c=ws.cell(r,CA,f"▸ Chart: {title}")
    c.fill=F(BM); c.font=K(b=True,c=WH,s=10); c.alignment=A()
    ws.row_dimensions[r].height=16

def add_line_chart(ws, anchor_row, title, y_ws, data_rows, label_row,
                   label_ws=None, ytitle="", bold_idx=None, width=22, height=11):
    """Add a LineChart below anchor_row. data_rows = [(label, row_number), ...]"""
    chart=LineChart()
    chart.title=title; chart.style=10
    chart.y_axis.title=ytitle; chart.x_axis.title="Month"
    chart.width=width; chart.height=height
    if label_ws is None: label_ws=y_ws
    for i,(lbl,row) in enumerate(data_rows):
        ref=Reference(y_ws,min_col=mc(0),max_col=mc(NM-1),min_row=row,max_row=row)
        chart.add_data(ref,from_rows=True,titles_from_data=False)
        chart.series[-1].title=SeriesLabel(v=lbl)
        if bold_idx is not None and i==bold_idx:
            chart.series[-1].graphicalProperties.line.width=25000
    cats=Reference(label_ws,min_col=mc(0),max_col=mc(NM-1),min_row=label_row,max_row=label_row)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{anchor_row}")
    return anchor_row + 24   # approximate rows consumed

def put(ws, row, col, value, fill=None, fmt=None):
    c=ws.cell(row=row,column=col,value=value)
    if fill: c.fill=F(fill)
    if fmt: c.number_format=fmt

wb=openpyxl.Workbook()
refs={}

# ── 1. DASHBOARD ──────────────────────────────────────────────────────────────
ws=wb.active; ws.title="Dashboard"; ws.sheet_properties.tabColor=NAVY
title_block(ws,1,"MRV Business Intelligence  |  Charm Industrial",
            "Rolling 12-month view (Jul-25 → Jun-26)  |  updated each reporting close")
KPI=[
    ("YTD Net CDR (tCO₂e)","blended total — all sources",BD),
    ("Net Efficiency — YTD","net / gross CDR",BD),
    ("Last Injection → Credits (days)","from last ton of reporting period",BD),
    ("Pipeline: injected, unverified (tCO₂e)","batches in Certify draft status",BM),
    ("Penalties / Gross CDR — YTD","all sources + downstream + well facility",BM),
    ("Source Concentration","largest single-source % of net CDR",BM),
]
tc=[(1,5),(6,10),(11,15)]
for idx,(tt,st,bg) in enumerate(KPI):
    tr=3 if idx<3 else 6; cs,ce=tc[idx%3]
    ws.merge_cells(start_row=tr,start_column=cs,end_row=tr,end_column=ce)
    ws.merge_cells(start_row=tr+1,start_column=cs,end_row=tr+1,end_column=ce)
    lc=ws.cell(tr,cs,tt); lc.fill=F(bg); lc.font=K(b=True,c=WH,s=9); lc.alignment=A(h="center")
    vc=ws.cell(tr+1,cs,st); vc.fill=F(BL); vc.font=K(c="555555",s=9,i=True); vc.alignment=A(h="center")
    ws.row_dimensions[tr].height=14; ws.row_dimensions[tr+1].height=14
sp(ws,8)
sec(ws,9,"Monthly Summary — Key Metrics  (see detail tabs for drilldowns)")
col_hdr(ws,10,"Metric")
r=11
for lbl,tot,pct,nytd in [
    ("Net CDR — total operation (tCO₂e)",False,False,False),
    ("  AECN",False,False,False),("  Charm",False,False,False),("  Kerry",False,False,False),
    ("Gross CDR — total operation (tCO₂e)",False,False,False),
    ("Net efficiency (net / gross CDR)",False,True,True),
    ("Injected in Certify draft (pipeline, tCO₂e)",False,False,False),
    ("Last injection → credits (days)",False,False,True),
    ("Penalties as % of gross CDR — total",False,True,True),
    ("AECN share of net CDR (%)",False,True,True),("Charm share of net CDR (%)",False,True,True),
    ("Kerry share of net CDR (%)",False,True,True),
    ("Charm bio-oil CI (tCO₂e / tCO₂ stored)",False,False,True),
    ("Charm biochar CI (tCO₂e / tCO₂ stored)",False,False,True),
]:
    ind=1 if lbl.startswith("  ") else 0
    dr(ws,r,lbl.strip(),ind=ind,tot=tot,pct=pct,nytd=nytd); r+=1
sp(ws,r); r+=1
nt(ws,r,"ℹ  Last injection → credits = OKR. FY 2025* = credits issued in calendar year 2025.")
set_w(ws,lw=36,fp="B11")

# ── 2. RUN RATE ───────────────────────────────────────────────────────────────
ws=wb.create_sheet("Run Rate"); ws.sheet_properties.tabColor=BD
title_block(ws,1,"Run Rate — Net CDR Credited & Pipeline",
            "Credits attributed to month of issuance (revenue timing). Rolling 12-month. tCO₂e.")
col_hdr(ws,3,"Metric"); r=4

sec(ws,r,"Credited — Net CDR (tCO₂e)"); r+=1
dr(ws,r,"Total operation",tot=True,data_fill=CERT_BG); refs["rr_net_total"]=r; r+=1
for s in SRC: dr(ws,r,s,ind=1); r+=1
sp(ws,r); r+=1

sec(ws,r,"Gross CDR — for reference (tCO₂e, not directly credited)"); r+=1
dr(ws,r,"Total operation",tot=True); refs["rr_gross_total"]=r; r+=1
for s in SRC: dr(ws,r,s,ind=1); r+=1
sp(ws,r); r+=1

sec(ws,r,"Pipeline — Batches in Certify Draft Status (tCO₂e)"); r+=1
dr(ws,r,"Total operation",tot=True); r+=1
for s in SRC: dr(ws,r,s,ind=1); r+=1
sp(ws,r); r+=1

sec(ws,r,"Verification Timing"); r+=1
dr(ws,r,"Last injection → credits — total operation (days) [OKR metric]",tot=True,days=True,nytd=True); r+=1
dr(ws,r,"  Last injection of period → period close (days)",ind=1,days=True,nytd=True); r+=1
dr(ws,r,"  Period close → credit issuance (days)",ind=1,days=True,nytd=True); r+=1
sp(ws,r); r+=1
dr(ws,r,"Blended avg injection-to-credit lag — total (days, ton-weighted)",days=True,nytd=True); r+=1
for s in SRC: dr(ws,r,f"{s} blended avg lag (days)",ind=1,days=True,nytd=True); r+=1
nt(ws,r+1,"ℹ  FY 2025* = credits issued in calendar year 2025 = 4,461.21 tCO₂e. "
          "Oct-25 cell = =166.40+291.85 (two periods both issued Oct 2025). "
          "Source breakdown (AECN/Charm/Kerry) pending — requires drilling individual GHG statements.")
chart_row = r + 3

sp(ws,chart_row-1)
chart_hdr(ws,chart_row,"Net CDR Credited — Monthly (tCO₂e)")
refs["rr_chart_anchor"]=chart_row+1
add_line_chart(ws,chart_row+1,"Net CDR Credited by Month (tCO₂e)",ws,
               [("Total Net CDR",refs["rr_net_total"])],3,ytitle="tCO₂e")
set_w(ws)

# ── 3. EFFICIENCY ─────────────────────────────────────────────────────────────
ws=wb.create_sheet("Efficiency"); ws.sheet_properties.tabColor=BD
title_block(ws,1,"Efficiency — Net CDR Yield","Net/gross CDR and net CDR per tonne injected")
col_hdr(ws,3,"Metric"); r=4
sec(ws,r,"Net CDR / Gross CDR (carbon efficiency ratio)"); r+=1
dr(ws,r,"Total operation",tot=True,pct=True,nytd=True); refs["eff_ratio_total"]=r; r+=1
for s in SRC:
    dr(ws,r,s,ind=1,pct=True,nytd=True); refs[f"eff_ratio_{s.lower()}"]=r; r+=1
sp(ws,r); r+=1
sec(ws,r,"Net CDR per Tonne Bio-Oil Injected (tCO₂e / tonne)"); r+=1
dr(ws,r,"Total operation",tot=True,nytd=True); refs["eff_yield_total"]=r; r+=1
for s in SRC: dr(ws,r,s,ind=1,nytd=True); r+=1
sp(ws,r); r+=1
sec(ws,r,"Memo: Volumes"); r+=1
dr(ws,r,"Net CDR — total (tCO₂e)",tot=True); r+=1
dr(ws,r,"Gross CDR — total (tCO₂e)",tot=True); r+=1
dr(ws,r,"Bio-oil injected — total (tonnes)",tot=True); r+=1
for s in SRC: dr(ws,r,f"{s} bio-oil injected (tonnes)",ind=1); r+=1
nt(ws,r,"ℹ  Efficiency ratios: derive YTD from YTD memo row totals.")
chart_row=r+3
chart_hdr(ws,chart_row,"Net/Gross CDR Efficiency — Monthly")
add_line_chart(ws,chart_row+1,"Net / Gross CDR (carbon efficiency ratio)",ws,
    [("Total",refs["eff_ratio_total"]),("AECN",refs["eff_ratio_aecn"]),
     ("Charm",refs["eff_ratio_charm"]),("Kerry",refs["eff_ratio_kerry"])],3,
    ytitle="Net/Gross ratio",bold_idx=0)
set_w(ws)

# ── 4. PENALTIES ──────────────────────────────────────────────────────────────
ws=wb.create_sheet("Penalties"); ws.sheet_properties.tabColor=BD
ws_pen=ws
title_block(ws,1,"Penalties — Emission Deductions from Gross CDR",
            "System flow: Charm prod → AECN prod → Kerry prod → Downstream Transport → Well Facility")
col_hdr(ws,3,"Metric"); r=4

# Total section
sec(ws,r,"Total Penalties — All Sources (tCO₂e)"); r+=1
dr(ws,r,"Total operation",tot=True); refs["pen_total"]=r; r+=1
dr(ws,r,"  Charm (production-side)",ind=1); refs["pen_charm"]=r; r+=1
dr(ws,r,"  AECN (production-side)",ind=1); refs["pen_aecn"]=r; r+=1
dr(ws,r,"  Kerry (production-side)",ind=1); refs["pen_kerry"]=r; r+=1
dr(ws,r,"  Downstream transport",ind=1); refs["pen_downstream"]=r; r+=1
dr(ws,r,"  Well facility (site emissions)",ind=1,bg_ov=SITE_BG); refs["pen_site"]=r; r+=1
sp(ws,r); r+=1

# Charm production penalties — ordered by materiality
CT=["Energy — propane (pyrolysis / drying)",
    "Capital Goods EE (equipment amortization)",
    "Transportation — truck (feedstock to facility)",
    "SP+C EE (spare parts & consumables)",
    "Energy — electricity",
    "Other / residual (gasoline, tailgas, water, waste — each <5% avg)"]
sec(ws,r,"Charm — Production Penalties (tCO₂e)"); r+=1
refs["charm_comp_rows"]=[]; refs["charm_comp_labels"]=CT
for pt in CT: dr(ws,r,pt,ind=1); refs["charm_comp_rows"].append(r); r+=1
sp(ws,r); r+=1

# AECN production penalties — no injection-side transport here
AT=["Transportation — truck (feedstock to facility)",
    "Energy — electricity","N₂O / process emissions","Other / residual"]
sec(ws,r,"AECN — Production Penalties (tCO₂e)"); r+=1
refs["aecn_rows"]=[]
for pt in AT: dr(ws,r,pt,ind=1); refs["aecn_rows"].append(r); r+=1
sp(ws,r); r+=1

# Kerry production penalties
KT=["Transportation — truck (feedstock to facility)",
    "Energy — electricity","N₂O / process emissions","Other / residual"]
sec(ws,r,"Kerry — Production Penalties (tCO₂e)"); r+=1
refs["kerry_rows"]=[]
for pt in KT: dr(ws,r,pt,ind=1); refs["kerry_rows"].append(r); r+=1
sp(ws,r); r+=1

# Downstream Transport — injection-side, all sources together
DT=["Total all sources",
    "  AECN — rail to injection well","  AECN — final truck to injection well",
    "  Charm — rail to injection well","  Charm — final truck to injection well",
    "  Kerry — rail to injection well","  Kerry — final truck to injection well"]
sec(ws,r,"Downstream Transport — Oil to Injection Well (tCO₂e)"); r+=1
refs["downstream_total_row"]=r
for i,pt in enumerate(DT):
    is_tot=(i==0); dr(ws,r,pt,tot=is_tot); r+=1
sp(ws,r); r+=1

# Well facility
ST=["Energy — electricity (injection site operations)",
    "Energy — methanol (well operations)",
    "Fugitive / venting emissions","Other site emissions"]
sec(ws,r,"Well Facility — Site Emissions (tCO₂e)",bg=BM); r+=1
refs["site_rows"]=[]
for pt in ST: dr(ws,r,pt,ind=1,bg_ov=SITE_BG); refs["site_rows"].append(r); r+=1
sp(ws,r); r+=1

# % of total section
sec(ws,r,"Penalties as % of All-Source Total (formula cells)"); r+=1
for lbl,src_row,bg in [
    ("Total operation (100%)",refs["pen_total"],BL),
    ("Charm",refs["pen_charm"],GR),
    ("AECN",refs["pen_aecn"],WH),
    ("Kerry",refs["pen_kerry"],GR),
    ("Downstream transport",refs["pen_downstream"],WH),
    ("Well facility",refs["pen_site"],SITE_BG),
]:
    tot_r=refs["pen_total"]
    for i in range(NM):
        col=mc(i); tc=get_column_letter(col)
        fml = 1 if src_row==tot_r else f"=IF({tc}{tot_r}>0,{tc}{src_row}/{tc}{tot_r},\"\")"
        c=ws.cell(r,col,fml); c.fill=F(bg); c.number_format="0.0%"; c.alignment=A(h="center")
    s_ytd=f"SUM({get_column_letter(mc(YS))}{src_row}:{get_column_letter(mc(YE))}{src_row})"
    t_ytd=f"SUM({get_column_letter(mc(YS))}{tot_r}:{get_column_letter(mc(YE))}{tot_r})"
    yc=ws.cell(r,CY); yc.fill=F(bg); yc.number_format="0.0%"; yc.alignment=A(h="center")
    yc.value = 1 if src_row==tot_r else f"=IF({t_ytd}>0,{s_ytd}/{t_ytd},\"\")"
    ws.cell(r,CF).fill=F(bg)
    ws.cell(r,CA,lbl).fill=F(bg); ws.cell(r,CA).font=K(); ws.cell(r,CA).alignment=A()
    ws.row_dimensions[r].height=15; r+=1

nt(ws,r,"ℹ  Charm = production-side only (propane, Cap EE, SP+C EE, biomass transport, electricity, other). "
        "Injection-side transport in Downstream section. Blue = PP summary data. Pending: AECN/Kerry/Downstream/Site.")
r+=2

# Charts
chart_hdr(ws,r,"Total Penalties by Source — Monthly (tCO₂e)"); r+=1
add_line_chart(ws,r,"Penalties by Source (tCO₂e)",ws,
    [("Total",refs["pen_total"]),("Charm",refs["pen_charm"]),
     ("AECN",refs["pen_aecn"]),("Kerry",refs["pen_kerry"]),
     ("Downstream",refs["pen_downstream"]),("Well Facility",refs["pen_site"])],
    3,ytitle="tCO₂e",bold_idx=0); r+=25
chart_hdr(ws,r,"Charm Penalty Components — Monthly (tCO₂e)"); r+=1
add_line_chart(ws,r,"Charm Penalty Components (tCO₂e)",ws,
    list(zip(CT,refs["charm_comp_rows"])),3,ytitle="tCO₂e"); r+=25
set_w(ws)

# ── 5. SCALING ANALYSIS ───────────────────────────────────────────────────────
ws=wb.create_sheet("Scaling Analysis"); ws.sheet_properties.tabColor=BM
ws_sc=ws
title_block(ws,1,"Scaling Analysis — Penalty Trendlines vs. Volume",
            "Slope & R² per penalty type vs. bio-oil volume. Diagnosis: R²≥0.85=Linear | 0.60-0.84=Investigate | <0.60=Non-linear.")
r=3
SH=[("Penalty Type",38),("Source",10),("N (months)",12),("Slope vs Volume",16),("Intercept",12),("R²",10),("Diagnosis",16)]
for j,(h,w) in enumerate(SH):
    ws.column_dimensions[get_column_letter(j+1)].width=w
    c=ws.cell(r,j+1,h); c.fill=F(NAVY); c.font=K(b=True,c=WH); c.alignment=A(h="center" if j>0 else "left")
ws.row_dimensions[r].height=18; r+=1
PSN="Penalties"
for types,slbl,rk in [(CT,"Charm","charm_comp_rows"),(AT,"AECN","aecn_rows"),(KT,"Kerry","kerry_rows"),(ST,"Site","site_rows")]:
    row_list=refs[rk]
    for pt,pen_r in zip(types,row_list):
        bg=GR if r%2==0 else WH
        ws.cell(r,1,pt).fill=F(bg); ws.cell(r,1).alignment=A(w=True); ws.cell(r,1).font=K(s=10)
        ws.cell(r,2,slbl).fill=F(bg); ws.cell(r,2).alignment=A(h="center"); ws.cell(r,2).font=K(s=10)
        mr=f"'{PSN}'!{get_column_letter(mc(0))}{pen_r}:{get_column_letter(mc(NM-1))}{pen_r}"
        ws.cell(r,3,f"=COUNTA({mr})").fill=F(bg); ws.cell(r,3).alignment=A(h="center")
        for col in [4,5,6,7]:
            cx=ws.cell(r,col,"link vol row"); cx.fill=F(bg); cx.font=K(c="999999",i=True,s=9); cx.alignment=A(h="center")
        ws.row_dimensions[r].height=15; r+=1
sp(ws,r); r+=1
nt(ws,r,"ℹ  Wire D=SLOPE, E=INTERCEPT, F=RSQ against Production (Charm) bio-oil volume row after data is connected."); r+=2
chart=LineChart()
chart.title="Penalty Trendlines vs. Time (rolling 12 months)"; chart.style=10
chart.y_axis.title="tCO₂e"; chart.x_axis.title="Month"; chart.width=24; chart.height=14
for key,lbl,bold in [("pen_total","Total",True),("pen_charm","Charm",False),
                      ("pen_aecn","AECN",False),("pen_kerry","Kerry",False),
                      ("pen_downstream","Downstream",False),("pen_site","Well Facility",False)]:
    pr=refs.get(key,5)
    ref=Reference(ws_pen,min_col=mc(0),max_col=mc(NM-1),min_row=pr,max_row=pr)
    chart.add_data(ref)
    chart.series[-1].title=SeriesLabel(v=lbl)
    if bold: chart.series[-1].graphicalProperties.line.width=25000
cats=Reference(ws_pen,min_col=mc(0),max_col=mc(NM-1),min_row=3)
chart.set_categories(cats)
ws_sc.add_chart(chart,f"A{r}")
set_w(ws,lw=38,fp="A4")

# ── 6. PRODUCTION (Charm) ─────────────────────────────────────────────────────
ws=wb.create_sheet("Production (Charm)"); ws.sheet_properties.tabColor=BD
title_block(ws,1,"Production — Charm Bio-Oil & Biochar (Charm-produced oil only)",
            "Source: Charm PP close summary sheets. Jul-25 not separate (combined Jun-Jul period). Apr-26 pending PP close.")
col_hdr(ws,3,"Metric"); r=4

sec(ws,r,"Inputs — Biomass Processed by Charm"); r+=1
dr(ws,r,"Total dry biomass processed (dry MT)",tot=True); refs["prod_biomass"]=r; r+=1
sp(ws,r); r+=1

sec(ws,r,"Bio-Oil (WODO) — Output & Yield"); r+=1
dr(ws,r,"Bio-oil produced (MT)",tot=True); refs["prod_bio_oil_vol"]=r; r+=1
dr(ws,r,"Bio-oil carbon mass (tC)"); refs["prod_bio_oil_c"]=r; r+=1
dr(ws,r,"Carbon yield: bio-oil C / biomass C input (%)",pct=True,nytd=True); r+=1
sp(ws,r); r+=1

sec(ws,r,"Bio-Oil Carbon Intensity (tCO₂e / tCO₂ net CDR stored)"); r+=1
dr(ws,r,"Bio-oil CI",tot=True,nytd=True); refs["prod_bio_oil_ci"]=r; r+=1
sp(ws,r); r+=1

sec(ws,r,"Biochar — Output & Yield"); r+=1
dr(ws,r,"Biochar produced (MT)",tot=True); refs["prod_biochar_vol"]=r; r+=1
dr(ws,r,"Biochar carbon mass (tC)"); refs["prod_biochar_c"]=r; r+=1
dr(ws,r,"Carbon yield: biochar C / biomass C input (%)",pct=True,nytd=True); r+=1
sp(ws,r); r+=1

sec(ws,r,"Biochar Carbon Intensity (tCO₂e / tCO₂ net CDR stored)"); r+=1
dr(ws,r,"Biochar CI",tot=True,nytd=True); refs["prod_biochar_ci"]=r; r+=1

nt(ws,r+1,"ℹ  CI pending (needs net CDR from GHG statements). For final values use Charm MRV folder in RP packages.")
r+=3

chart_hdr(ws,r,"Production Volumes — Monthly (MT)"); r+=1
add_line_chart(ws,r,"Charm Production Volumes (MT)",ws,
    [("Bio-oil (WODO)",refs["prod_bio_oil_vol"]),("Biochar",refs["prod_biochar_vol"]),
     ("Biomass input",refs["prod_biomass"])],3,ytitle="MT"); r+=25

chart_hdr(ws,r,"Carbon Mass — Monthly (tC)"); r+=1
add_line_chart(ws,r,"Carbon Mass in Products (tC)",ws,
    [("Bio-oil C mass",refs["prod_bio_oil_c"]),("Biochar C mass",refs["prod_biochar_c"])],
    3,ytitle="tC")
set_w(ws)

# ── 7. SOURCE MIX ─────────────────────────────────────────────────────────────
ws=wb.create_sheet("Source Mix"); ws.sheet_properties.tabColor=BM
title_block(ws,1,"Source Mix — Concentration & Efficiency by Oil Source",
            "Share of net CDR from each source + per-source efficiency roll-up")
col_hdr(ws,3,"Metric"); r=4

sec(ws,r,"Net CDR (tCO₂e) by Source — Absolute"); r+=1
dr(ws,r,"Total operation",tot=True); refs["sm_net_total"]=r; r+=1
for s in SRC: dr(ws,r,s,ind=1); refs[f"sm_net_{s.lower()}"]=r; r+=1
sp(ws,r); r+=1

sec(ws,r,"Net CDR Share by Source (% of total net CDR)"); r+=1
for s in SRC:
    dr(ws,r,s,pct=True,nytd=True); refs[f"sm_share_{s.lower()}"]=r; r+=1
sp(ws,r); r+=1

sec(ws,r,"Gross CDR (tCO₂e) by Source — Absolute"); r+=1
dr(ws,r,"Total operation",tot=True); r+=1
for s in SRC: dr(ws,r,s,ind=1); r+=1
sp(ws,r); r+=1

sec(ws,r,"Efficiency Roll-Up by Source (Net / Gross CDR)"); r+=1
dr(ws,r,"Total operation",tot=True,pct=True,nytd=True); refs["sm_eff_total"]=r; r+=1
for s in SRC:
    dr(ws,r,s,ind=1,pct=True,nytd=True); refs[f"sm_eff_{s.lower()}"]=r; r+=1
sp(ws,r); r+=1

sec(ws,r,"Concentration Risk"); r+=1
dr(ws,r,"Largest single-source share of net CDR (%)",pct=True,nytd=True); r+=1
dr(ws,r,"Number of active sources this month",days=True,nytd=True); r+=1
nt(ws,r+1,"ℹ  Share and efficiency: no YTD auto-sum — derive from YTD absolute rows.")
r+=3

chart_hdr(ws,r,"Net CDR by Source — Monthly (tCO₂e)"); r+=1
add_line_chart(ws,r,"Net CDR by Source (tCO₂e)",ws,
    [("Total",refs["sm_net_total"]),("AECN",refs["sm_net_aecn"]),
     ("Charm",refs["sm_net_charm"]),("Kerry",refs["sm_net_kerry"])],
    3,ytitle="tCO₂e",bold_idx=0); r+=25

chart_hdr(ws,r,"Efficiency by Source — Monthly"); r+=1
add_line_chart(ws,r,"Net/Gross CDR by Source",ws,
    [("Total",refs["sm_eff_total"]),("AECN",refs["sm_eff_aecn"]),
     ("Charm",refs["sm_eff_charm"]),("Kerry",refs["sm_eff_kerry"])],
    3,ytitle="ratio",bold_idx=0)
set_w(ws)

# ── 8. PROTOCOL HEADROOM (WiP) ────────────────────────────────────────────────
ws=wb.create_sheet("Protocol Headroom"); ws.sheet_properties.tabColor="FFC000"
title_block(ws,1,"⚠  Protocol Headroom  [WORK IN PROGRESS — thresholds not confirmed]",
            "Verify all thresholds against BSSE v1.2 and PDD before use.")
PC=[("Parameter",36),("Protocol / Source",22),("Threshold (⚠ unconfirmed)",22),
    ("Current Value",14),("Prior Period",14),("Trend",10),("Headroom",12),("Status",12)]
for j,(h,w) in enumerate(PC):
    ws.column_dimensions[get_column_letter(j+1)].width=w
    c=ws.cell(3,j+1,h); c.fill=F(NAVY); c.font=K(b=True,c=WH)
    c.alignment=A(h="center" if j>0 else "left")
ws.row_dimensions[3].height=18; ws.freeze_panes=ws.cell(4,1)
WIP=["Bio-oil H:C atomic ratio","Bio-oil carbon content (%)","Bio-oil oxygen content (%)",
     "Biomass moisture at sampling (%)","Mass balance closure (%)","BOL vs. scale mass discrepancy (%)",
     "Net CDR / Gross CDR ratio","Largest single penalty type (% gross CDR)",
     "Biochar H:C atomic ratio","Biochar carbon content (%)","Biochar yield (% dry biomass mass)"]
for i,p in enumerate(WIP):
    row=4+i; bg=GR if i%2==0 else WH
    for j,v in enumerate([p,"⚠ confirm source","⚠ confirm threshold","—","—","—","—","—"]):
        is_wip="⚠" in str(v)
        c=ws.cell(row,j+1,v); c.fill=F(AM if is_wip else bg)
        c.font=K(s=10,c="7F6000" if is_wip else "000000")
        c.alignment=A(h="center" if j>0 else "left",w=(j==0))
    ws.row_dimensions[row].height=15
nr=4+len(WIP)+1
ws.merge_cells(start_row=nr,start_column=1,end_row=nr,end_column=len(PC))
n=ws.cell(nr,1,"⚠  WiP — confirm thresholds against BSSE v1.2 and PDD before operational use.")
n.font=K(c="7F6000",s=9,i=True); n.alignment=A(w=True); n.fill=F(AM); ws.row_dimensions[nr].height=20

# ── 9. DATA SOURCES ───────────────────────────────────────────────────────────
ws=wb.create_sheet("Data Sources"); ws.sheet_properties.tabColor="888888"
title_block(ws,1,"Data Sources — Wiring Guide","Maps each metric to its primary source for live-feed connection")
DC=[("Metric / Tab",34),("Source System",20),("Specific Location",38),("Update Frequency",18),("Notes",26)]
for j,(h,w) in enumerate(DC):
    ws.column_dimensions[get_column_letter(j+1)].width=w
    c=ws.cell(3,j+1,h); c.fill=F(NAVY); c.font=K(b=True,c=WH); c.alignment=A()
ws.row_dimensions[3].height=18; ws.freeze_panes=ws.cell(4,1)
DSR=[
    ("Run Rate — net CDR total","Certify GHG statements","Statement list → Net CO₂e removed","Each reporting close","Credit issuance month attribution"),
    ("Run Rate — source breakdown","Certify GHG statements","Individual statement → removal detail by source","Each reporting close","AECN/Charm/Kerry split"),
    ("Run Rate — pipeline","Certify removal data","Drafts page — removal entries not in a closed statement","Weekly","Draft status = not yet submitted"),
    ("Efficiency ratios","Derived","From Run Rate volume memo rows","Derived",""),
    ("Penalties — Charm production","Charm PP summary → RP Charm MRV folder","Consolidated emissions inventory","Each reporting close","Use RP Charm MRV folder for final"),
    ("Penalties — Downstream transport","Certify individual removals","COBB calc sheet transport rows per batch","Each reporting close","Rail + final truck per source"),
    ("Penalties — AECN/Kerry production","Certify GHG statements","Individual statement removal detail","Each reporting close",""),
    ("Penalties — Well facility","Certify GHG statements","Site emissions batch","Each reporting close","Methanol well-specific"),
    ("Production — Charm","PP summary → RP Charm MRV folder","Output type table (WODO/Biochar/Aqueous)","Each reporting close","Use RP folder for final verified values"),
    ("Source Mix — tCDR by source","Certify GHG statements","Filter per-statement removals by source","Each reporting close",""),
    ("Protocol Headroom — thresholds","BSSE v1.2 + PDD","Verify manually from protocol text","Static / on update","WiP"),
]
for i,rd in enumerate(DSR):
    row=4+i; bg=GR if i%2==0 else WH
    for j,v in enumerate(rd):
        c=ws.cell(row,j+1,v); c.fill=F(bg); c.font=K(s=10); c.alignment=A(w=(j in (0,2)))
    ws.row_dimensions[row].height=15

# ═══════════════════════════════════════════════════════════════════════════════
#  DATA INJECTION
# ═══════════════════════════════════════════════════════════════════════════════

# Run Rate — Total Net CDR (Certify authoritative)
NET_CDR={
    0: None,
    1: 555.95,           # Aug-25
    2: 135.39,           # Sep-25
    3: "=166.40+291.85", # Oct-25 — two periods (formula for auditability)
    4: 454.71,           # Nov-25
    5: None,
    6: None,
    7: 43.47,            # Feb-26
    8: 764.14,           # Mar-26
    9: 612.99,           # Apr-26
    10: 490.82,          # May-26
    11: None,
}
ws_rr=wb["Run Rate"]
for idx,val in NET_CDR.items():
    if val is not None:
        put(ws_rr,refs["rr_net_total"],mc(idx),val,fill=CERT_BG,fmt="#,##0.00")
put(ws_rr,refs["rr_net_total"],CF,4461.21,fill=CERT_BG,fmt="#,##0.00")

# Production (Charm) — 8 months (Apr-26 left blank, PP not yet closed)
PROD={
    1: (84.84, 20.928, 10.353,  9.186, 7.931),  # Aug-25
    2: (64.58, 12.630,  6.573,  3.868, 3.304),  # Sep-25
    3: (47.45, 14.227,  8.247,  8.085, 7.399),  # Oct-25
    4: (25.00,  9.984,  5.684,  4.580, 3.891),  # Nov-25
    5: (28.95,  6.481,  3.157,  4.074, 3.619),  # Dec-25
    6: (44.93,  8.888,  5.051,  6.560, 5.047),  # Jan-26
    7: (47.98, 10.966,  5.744,  6.766, 5.228),  # Feb-26
    8: (65.40, 16.307,  8.326,  6.259, 4.920),  # Mar-26
}
ws_prod=wb["Production (Charm)"]
for idx,(bm,bov,boc,bcv,bcc) in PROD.items():
    col=mc(idx)
    for row,val in [(refs["prod_biomass"],bm),(refs["prod_bio_oil_vol"],bov),
                    (refs["prod_bio_oil_c"],boc),(refs["prod_biochar_vol"],bcv),(refs["prod_biochar_c"],bcc)]:
        put(ws_prod,row,col,val,fill=PP_BG,fmt="#,##0.000")

# Penalties — Charm totals and component breakdown (8 months, Apr-26 omitted)
CHARM_TOTAL={
    1:25.021, 2:16.508, 3:20.105, 4:16.396, 5:13.523, 6:16.552, 7:17.724, 8:20.497
}
CHARM_COMP={
    #   propane  cap_ee  bm_trans  spce    elec   other
    1: (16.023, 4.231,  1.299,   1.768,  1.340, 0.355),
    2: ( 8.638, 3.091,  2.218,   1.015,  0.985, 0.559),
    3: (12.905, 2.474,  1.174,   1.286,  1.157, 1.119),
    4: (10.574, 1.202,  1.368,   0.906,  1.002, 1.339),
    5: ( 8.205, 1.527,  1.760,   0.658,  0.747, 0.627),
    6: (10.232, 2.280,  1.734,   0.879,  1.004, 0.416),
    7: (11.194, 2.458,  1.332,   0.994,  1.033, 0.709),
    8: (12.269, 3.379,  1.754,   1.354,  1.110, 0.645),
}
ws_pen=wb["Penalties"]
for idx,val in CHARM_TOTAL.items():
    put(ws_pen,refs["pen_charm"],mc(idx),val,fill=PP_BG,fmt="#,##0.00")
for idx,comps in CHARM_COMP.items():
    col=mc(idx)
    for row,val in zip(refs["charm_comp_rows"],comps):
        put(ws_pen,row,col,val,fill=PP_BG,fmt="#,##0.00")

wb.save(OUT)
print(f"Saved: {OUT}")
