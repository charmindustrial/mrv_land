#!/usr/bin/env python3
"""
Inject bio-oil injected volumes (t) and derived Net CDR per tonne metrics.
Source: Certify removal records, all 9 credited statements + 4 drafts, 2026-06-05.
"""

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

FILE = "/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/mrv-operational-documents/MRV_BI_Dashboard_Shell.xlsx"

CERT="E2EFDA"; CALC="EBF1DE"; PP="DEEAF1"
def F(c): return PatternFill("solid", fgColor=c)
def put(ws, row, col, value, fill=None, fmt="#,##0.000"):
    c = ws.cell(row=row, column=col, value=value)
    if fill: c.fill = F(fill)
    c.number_format = fmt

CB=2; NM=12; CY=CB+NM; CF=CY+1
def mc(i): return CB+i
def cl(i): return get_column_letter(mc(i))
MO=["Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25",
    "Jan-26","Feb-26","Mar-26","Apr-26","May-26","Jun-26"]

wb = openpyxl.load_workbook(FILE)

# ─────────────────────────────────────────────────────────────────────────────
# Bio-oil injected volumes by credit-issuance month
# ─────────────────────────────────────────────────────────────────────────────
# Mapped to credit month (when credits were issued, not injection month)
# Oct-25 combines two periods: Aug 16-Sep 8 (264.139 t) + Sep 9-30 (284.436 t)

BIO_OIL_TOTAL = {
    1:  553.599,                    # Aug-25: Jun-Jul 2025, all AECN
    2:  118.562,                    # Sep-25: Aug 1-15 2025, all Kerry
    3:  "=264.139+284.436",         # Oct-25: two periods (formula)
    4:  432.060,                    # Nov-25: Oct-Nov 2025, all AECN
    7:  98.867,                     # Feb-26: Nov 2025-Jan 2026, all AECN
    8:  568.558,                    # Mar-26: Feb-Mar 2026, all AECN
    9:  462.251,                    # Apr-26: Mar 2026, all AECN
   10:  407.596,                    # May-26: Apr 2026, mixed
   11:  198.977,                    # Jun-26: pipeline (4 drafts)
}

# AECN bio-oil by month
# Pure AECN months: total = AECN
# Mixed Oct-25: Sep9-30 (284.436, 100% AECN) + Aug16-Sep8 AECN share (59.55/166.40 × 264.139)
# Mixed May-26: proportional by net CDR share (approximate)
BIO_OIL_AECN = {
    1:  553.599,                    # Aug-25: all AECN
    2:  0,                          # Sep-25: all Kerry
    3:  "=284.436+264.139*59.55/166.40",  # Oct-25: formula
    4:  432.060,                    # Nov-25: all AECN
    7:  98.867,                     # Feb-26: all AECN
    8:  568.558,                    # Mar-26: all AECN
    9:  462.251,                    # Apr-26: all AECN
   10:  "=407.596*475.09/490.82",   # May-26: proportional (~394.5 t)
   11:  "=121.287+22.835+50.297",   # Jun-26 pipeline: 3 AECN drafts
}

# Kerry bio-oil by month
BIO_OIL_KERRY = {
    1:  0,
    2:  118.562,                    # Sep-25: all Kerry
    3:  "=264.139*106.85/166.40",   # Oct-25: Kerry share of Aug16-Sep8
    4:  0,
    7:  0,
    8:  0,
    9:  0,
   10:  "=407.596*8.80/490.82",     # May-26: proportional (~7.3 t)
   11:  4.558,                      # Jun-26 pipeline: 1 Kerry draft
}

# Charm bio-oil: from Production tab (already populated); add May-26 proportional estimate
# Charm May-26 from Production tab row 8 col mc(9) = Apr-26 production...
# But the Apr 2026 GHG statement only credited 6.93 tCO₂e Charm net.
# Charm injected in Apr period ≈ proportional: 407.596 × 6.93/490.82 ≈ 5.75 t
BIO_OIL_CHARM_MAY26 = 407.596 * 6.93 / 490.82  # ≈ 5.75 t

# ─────────────────────────────────────────────────────────────────────────────
# EFFICIENCY TAB
# Rows: 19=Bio-oil total, 20=AECN, 21=Charm, 22=Kerry
#       11=Net CDR per tonne total, 12=AECN, 13=Charm, 14=Kerry
# ─────────────────────────────────────────────────────────────────────────────
eff = wb["Efficiency"]

for i, v in BIO_OIL_TOTAL.items():
    put(eff, 19, mc(i), v, fill=CERT)

for i, v in BIO_OIL_AECN.items():
    put(eff, 20, mc(i), v, fill=CALC)

# Charm bio-oil row 21 — already partially populated; add May-26 estimate
put(eff, 21, mc(10), BIO_OIL_CHARM_MAY26, fill=CALC)

for i, v in BIO_OIL_KERRY.items():
    put(eff, 22, mc(i), v, fill=CALC)

# Row 11: Total Net CDR per tonne = formula net/bio-oil
for i in range(NM):
    c_col = cl(i)
    c = eff.cell(11, mc(i))
    c.value = f"=IF({c_col}19>0,{c_col}17/{c_col}19,\"\")"
    c.fill = F(CALC); c.number_format = "#,##0.000"
# YTD
s = get_column_letter(mc(6)); e = get_column_letter(mc(11))
yc = eff.cell(11, CY)
yc.value = f"=IF(SUM({s}19:{e}19)>0,SUM({s}17:{e}17)/SUM({s}19:{e}19),\"\")"
yc.fill = F(CALC); yc.number_format = "#,##0.000"

# Row 12: AECN per tonne = net AECN / AECN bio-oil
for i in range(NM):
    c_col = cl(i)
    c = eff.cell(12, mc(i))
    c.value = f"=IF({c_col}20>0,'Run Rate'!{c_col}6/{c_col}20,\"\")"
    c.fill = F(CALC); c.number_format = "#,##0.000"

# Row 13: Charm per tonne = Charm net CDR / Charm bio-oil
for i in range(NM):
    c_col = cl(i)
    c = eff.cell(13, mc(i))
    c.value = f"=IF({c_col}21>0,'Run Rate'!{c_col}7/{c_col}21,\"\")"
    c.fill = F(CALC); c.number_format = "#,##0.000"

# Row 14: Kerry per tonne
for i in range(NM):
    c_col = cl(i)
    c = eff.cell(14, mc(i))
    c.value = f"=IF({c_col}22>0,'Run Rate'!{c_col}8/{c_col}22,\"\")"
    c.fill = F(CALC); c.number_format = "#,##0.000"

# FY25 bio-oil totals (in-window 2025: Aug+Sep+Oct+Nov)
# Aug 553.599 + Sep 118.562 + Oct 548.575 + Nov 432.060 = 1,652.796
put(eff, 19, CF, 1652.796, fill=CERT)
put(eff, 20, CF, 553.599 + 0 + (284.436 + 264.139*59.55/166.40) + 432.060, fill=CALC)  # AECN FY25
put(eff, 22, CF, 118.562 + 264.139*106.85/166.40, fill=CALC)  # Kerry FY25

# ─────────────────────────────────────────────────────────────────────────────
# RUN RATE — add bio-oil memo rows (rows 19-22 in Efficiency, but also
# need to fill the Efficiency "memo: volumes" section)
# Also fix the Efficiency row 19 (bio-oil total) for total
# which already done above.
# ─────────────────────────────────────────────────────────────────────────────

# ─────────────────────────────────────────────────────────────────────────────
# SOURCE MIX — add bio-oil injected rows if applicable
# ─────────────────────────────────────────────────────────────────────────────
# Source Mix gross CDR rows already populated. Nothing to add here specifically.

# ─────────────────────────────────────────────────────────────────────────────
# DASHBOARD — update bio-oil / pipeline rows now that Efficiency has data
# ─────────────────────────────────────────────────────────────────────────────
db = wb["Dashboard"]

# Row 17 (pipeline) already wired to Run Rate row 17.
# The bio-oil total in pipeline Jun-26 (198.977 t) is in Efficiency row 19 mc(11).
# Dashboard already pulls from Run Rate row 17 for tCO₂e pipeline.

wb.save(FILE)
print("Saved.")

# Print net CDR per tonne for verification
print("\nNet CDR per tonne bio-oil injected:")
wb2 = openpyxl.load_workbook(FILE)
eff2 = wb2["Efficiency"]
for i in range(NM):
    bio = eff2.cell(19, mc(i)).value
    net_raw = eff2.cell(17, mc(i)).value
    if bio and net_raw:
        try:
            net = float(str(net_raw).replace("=",""))
        except:
            net = None
        if net:
            print(f"  {MO[i]}: {net:.2f} net tCO₂e / {float(str(bio).split('+')[0].split('*')[0].replace('=','')) if isinstance(bio,str) else bio:.1f} t bio-oil")
        else:
            print(f"  {MO[i]}: formula cell, bio-oil={bio}")
    elif bio:
        print(f"  {MO[i]}: bio-oil={bio}, net=blank")
