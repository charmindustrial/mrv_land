#!/usr/bin/env python3
"""
First-pass data injection into MRV BI Dashboard Shell.
Populates:
  - Run Rate: Total Net CDR (all 12 months, from Certify — authoritative)
  - Production (Charm): bio-oil & biochar volumes + C masses (Sep-25, Nov-25, Jan–Apr 26)
  - Penalties: Charm total + component breakdown for same months (PP summary, production-side only)

Attribution note for Run Rate: credits mapped to the calendar month they were ISSUED,
not the month they were injected. This is revenue-recognition timing.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

FILE = "/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/mrv-operational-documents/MRV_BI_Dashboard_Shell.xlsx"

# Colors
AMBER = PatternFill("solid", fgColor="FFF2CC")   # pending FRP verification
CERT  = PatternFill("solid", fgColor="E2EFDA")   # verified from Certify
PP    = PatternFill("solid", fgColor="DEEAF1")    # from PP summary (pre-FRP)

wb = openpyxl.load_workbook(FILE)

# Column layout — must match build script
CB = 2; NM = 12; CY = CB + NM; CF = CY + 1
def mc(i): return CB + i

def put(ws, row, col, value, fill=None):
    c = ws.cell(row=row, column=col)
    c.value = value
    if fill:
        c.fill = fill

def note_row(ws, row, text, last_col=15):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(color="888888", size=9, italic=True, name="Calibri")
    c.alignment = Alignment(horizontal="left", wrap_text=True)
    ws.row_dimensions[row].height = 14

# ──────────────────────────────────────────────────────────────────────────────
#  RUN RATE — Total Net CDR by month (credit issuance date)
# ──────────────────────────────────────────────────────────────────────────────
ws = wb["Run Rate"]

# Net CDR by month index (0=Jul-25 … 11=Jun-26), credited to issuance month
# Source: Certify GHG statements list, all 15 periods reviewed
NET_CDR = {
    0:  None,     # Jul-25  — no credits issued
    1:  555.95,   # Aug-25  — Jun 1–Jul 31 2025 period, issued 19 Aug 2025
    2:  135.39,   # Sep-25  — Aug 1–15 2025 period, issued 19 Sep 2025
    3:  458.25,   # Oct-25  — Aug 16–Sep 8 (166.40) + Sep 9–30 (291.85), issued Oct 2025
    4:  454.71,   # Nov-25  — Oct 1–Nov 6 2025 period, issued 19 Nov 2025
    5:  None,     # Dec-25  — no credits issued
    6:  None,     # Jan-26  — no credits issued
    7:  43.47,    # Feb-26  — Nov 7 2025–Jan 31 2026 period, issued 18 Feb 2026
    8:  764.14,   # Mar-26  — Feb 1–Mar 2 2026 period, issued 10 Mar 2026
    9:  612.99,   # Apr-26  — Mar 3–31 2026 period, issued 16 Apr 2026
    10: 490.82,   # May-26  — Apr 1–30 2026 period, issued 13 May 2026
    11: None,     # Jun-26  — May 2026 period in verification; Jun-26 open
}

ROW_NET_TOTAL = 5
for idx, val in NET_CDR.items():
    put(ws, ROW_NET_TOTAL, mc(idx), val, fill=CERT if val else None)

# FY 2025 total: all credits issued in calendar year 2025
# = Feb 2025 (246.73) + Mar 2025 (592.66) + May 2025 (1,002.95) + Jun 2025 (1,014.57)
#   + Aug 2025 (555.95) + Sep 2025 (135.39) + Oct 2025 (458.25) + Nov 2025 (454.71)
put(ws, ROW_NET_TOTAL, CF, 4461.21, fill=CERT)

# Attribution methodology note
note_row(ws, 50,
    "ℹ  Data source: Certify GHG statements list, all 15 periods reviewed 2026-06-05. "
    "Net CDR credited to the calendar month credits were ISSUED (revenue recognition), not injection month. "
    "Oct-25 combines two statements (Aug 16–Sep 8: 166.40 + Sep 9–30: 291.85 tCO₂e). "
    "Source breakdown (AECN/Charm/Kerry) pending — requires drilling individual GHG statements. "
    "Gross CDR pending same. Pipeline (injected, unverified) pending.")

# ──────────────────────────────────────────────────────────────────────────────
#  PRODUCTION (CHARM) — bio-oil & biochar volumes and C masses
# ──────────────────────────────────────────────────────────────────────────────
ws = wb["Production (Charm)"]

# Available months from PP close summaries
# Aug-25, Oct-25, Dec-25 not yet read. May-26, Jun-26 not closed.
# Apr-26 data appears identical to Mar-26 outputs — flagged as pending FRP verification.
# Units: biomass=dry MT, bio-oil/biochar=MT, C mass=tC

PROD = {
    #idx: (biomass_MT, bio_oil_MT, bio_oil_C_tC, biochar_MT, biochar_C_tC, fill)
    2:  (64.58, 12.63, 6.573, 3.868, 3.304, PP),    # Sep-25
    4:  (25.00,  9.98, 5.684, 4.580, 3.891, PP),    # Nov-25
    6:  (44.93,  8.89, 5.051, 6.560, 5.047, PP),    # Jan-26
    7:  (47.98, 10.97, 5.744, 6.766, 5.228, PP),    # Feb-26
    8:  (65.40, 16.31, 8.326, 6.259, 4.920, PP),    # Mar-26
    9:  (65.40, 16.31, 8.326, 6.259, 4.920, AMBER), # Apr-26 — outputs match Mar-26; pending FRP
}

ROW_BIOMASS     = 5
ROW_BIOIL_VOL   = 8
ROW_BIOIL_C     = 9
ROW_BIOCHAR_VOL = 16
ROW_BIOCHAR_C   = 17

for idx, (bm, bov, boc, bcv, bcc, fill) in PROD.items():
    col = mc(idx)
    put(ws, ROW_BIOMASS,     col, bm,  fill=fill)
    put(ws, ROW_BIOIL_VOL,   col, bov, fill=fill)
    put(ws, ROW_BIOIL_C,     col, boc, fill=fill)
    put(ws, ROW_BIOCHAR_VOL, col, bcv, fill=fill)
    put(ws, ROW_BIOCHAR_C,   col, bcc, fill=fill)

note_row(ws, 25,
    "ℹ  Source: Charm Production Period close summary spreadsheets, reviewed 2026-06-05. "
    "Blue = from PP summary (pre-FRP). Amber = Apr-26 outputs match Mar-26 exactly — likely "
    "PP summary not yet updated with final Apr data; verify against Final Reporting Package. "
    "Aug-25, Oct-25, Dec-25, May-26 production data not yet read. Carbon yield and CI pending "
    "(requires biomass C% and net CDR from individual GHG statements).")

# ──────────────────────────────────────────────────────────────────────────────
#  PENALTIES — Charm total + production-side component breakdown
# ──────────────────────────────────────────────────────────────────────────────
ws = wb["Penalties"]

# Charm CDR penalties = emissions allocated to WODO + Biochar fractions in PP summary.
# Allocation % = (WODO C mass + Biochar C mass) / total C mass across all outputs.
# Components allocated proportionally (CDR C fraction × total PP component emission).
# NOTE: injection-side transportation (rail + final truck) not in PP summary;
# those rows left blank pending individual batch data.

CHARM_TOTAL = {
    2:  16.508,  # Sep-25  (WODO 10.985 + Biochar 5.523)
    4:  16.396,  # Nov-25  (WODO 9.733 + Biochar 6.663)
    6:  16.552,  # Jan-26  (WODO 8.279 + Biochar 8.273)
    7:  17.724,  # Feb-26  (WODO 9.279 + Biochar 8.445)
    8:  20.497,  # Mar-26  (WODO 12.884 + Biochar 7.613)
    9:  21.603,  # Apr-26  (WODO 13.579 + Biochar 8.024) — pending FRP
}

# Component breakdown (production-side only; CDR-fraction-allocated from consolidated inventory)
# Keys: bm=biomass transport, elec=electricity, prop=propane, n2o=tailgas, other=remainder
CHARM_COMP = {
    2:  {"bm": 2.218, "elec": 0.985, "prop":  8.638, "n2o": 0.095, "other": 4.571},
    4:  {"bm": 1.368, "elec": 1.002, "prop": 10.574, "n2o": 0.036, "other": 3.411},
    6:  {"bm": 1.734, "elec": 1.004, "prop": 10.232, "n2o": 0.068, "other": 3.507},
    7:  {"bm": 1.332, "elec": 1.033, "prop": 11.194, "n2o": 0.073, "other": 4.090},
    8:  {"bm": 1.754, "elec": 1.110, "prop": 12.269, "n2o": 0.099, "other": 5.276},
    9:  {"bm": 1.754, "elec": 1.110, "prop": 13.370, "n2o": 0.099, "other": 5.276},
}

ROW_CHARM_TOTAL = 7
ROW_CHARM_BM    = 20   # Transportation — truck (feedstock to facility)
# rows 21–22: rail + final truck — injection-side, leave blank
ROW_CHARM_ELEC  = 23
ROW_CHARM_PROP  = 24
ROW_CHARM_N2O   = 25
ROW_CHARM_OTHER = 26

for idx, val in CHARM_TOTAL.items():
    fill = AMBER if idx == 9 else PP
    put(ws, ROW_CHARM_TOTAL, mc(idx), val, fill=fill)

for idx, d in CHARM_COMP.items():
    fill = AMBER if idx == 9 else PP
    col = mc(idx)
    put(ws, ROW_CHARM_BM,    col, d["bm"],   fill=fill)
    put(ws, ROW_CHARM_ELEC,  col, d["elec"], fill=fill)
    put(ws, ROW_CHARM_PROP,  col, d["prop"], fill=fill)
    put(ws, ROW_CHARM_N2O,   col, d["n2o"],  fill=fill)
    put(ws, ROW_CHARM_OTHER, col, d["other"],fill=fill)

note_row(ws, 55,
    "ℹ  Charm data source: PP close summary consolidated emissions inventory, reviewed 2026-06-05. "
    "Charm CDR penalties = emissions allocated to WODO + Biochar C fractions (excludes aqueous). "
    "Production-side only: transportation rows 21–22 (injection-side rail + truck) left blank "
    "pending individual batch data. 'Other' = Capital Goods EE + SP+C EE + gasoline/diesel "
    "+ water + consumables + waste. AECN, Kerry, and Well Facility penalties pending GHG statement drill. "
    "Blue = PP summary. Amber = Apr-26, pending FRP verification.")

wb.save(FILE)
print("Done.")
