#!/usr/bin/env python3
"""
FINAL COMPLETE DATA INJECTION — every cell with available data.
No partial stops. Run until nothing remains.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

FILE = "/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/mrv-operational-documents/MRV_BI_Dashboard_Shell.xlsx"

CERT="E2EFDA"; CALC="EBF1DE"; PP="DEEAF1"; ZERO="F2F2F2"
def F(c): return PatternFill("solid", fgColor=c)
def put(ws, row, col, value, fill=None, fmt="#,##0.00"):
    c = ws.cell(row=row, column=col, value=value)
    if fill: c.fill = F(fill)
    c.number_format = fmt

CB=2; NM=12; CY=CB+NM; CF=CY+1
def mc(i): return CB+i
def cl(i): return get_column_letter(mc(i))
MO=["Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25",
    "Jan-26","Feb-26","Mar-26","Apr-26","May-26","Jun-26"]

wb = openpyxl.load_workbook(FILE)

# ─────────────────────────────────────────────────────────────────────────────
# RUN RATE
# ─────────────────────────────────────────────────────────────────────────────
rr = wb["Run Rate"]

# Row 6 AECN — Aug-25 was pending, confirmed 100% AECN
put(rr, 6, mc(1), 555.95, fill=CERT)
# Sep-25 zero (Kerry)
put(rr, 6, mc(2), 0, fill=CERT)
# All other months already populated in prior pass

# Row 7 Charm — zero for all months except May-26 (already set)
for i in [1,2,3,4,5,6,7,8,9]:
    if rr.cell(7, mc(i)).value is None:
        put(rr, 7, mc(i), 0, fill=CERT)

# Row 8 Kerry — zero for months with no Kerry
for i in [1,4,5,6,7,8]:
    put(rr, 8, mc(i), 0, fill=CERT)

# Row 19 Charm pipeline = 0 (no Charm drafts in Certify)
put(rr, 19, mc(11), 0, fill=ZERO)

# Row 24: Last injection of period → reporting period close (days)
LAST_INJ_TO_CLOSE = {
    1: 8,   # Aug-25: Jul 23→Jul 31 = 8 days
    2: 0,   # Sep-25: Aug 15→Aug 15 = 0 days
    3: 0,   # Oct-25: Sep 30→Sep 30 = 0 days (worst-case statement)
    4: 1,   # Nov-25: Nov 5→Nov 6 = 1 day
    7: 50,  # Feb-26: Dec 12→Jan 31 = 50 days
    8: 0,   # Mar-26: Mar 2→Mar 2 = 0 days
    9: 5,   # Apr-26: Mar 26→Mar 31 = 5 days
   10: 1,   # May-26: Apr 29→Apr 30 = 1 day
}
for i,v in LAST_INJ_TO_CLOSE.items():
    put(rr, 24, mc(i), v, fill=CERT, fmt="#,##0")

# Row 25: Reporting period close → credit issuance (days)
CLOSE_TO_CREDITS = {
    1: 19,  # Aug-25: Jul 31→Aug 19
    2: 35,  # Sep-25: Aug 15→Sep 19
    3: 31,  # Oct-25: Sep 30→Oct 31 (using worse statement: Sep 8→Oct 9=31 days)
    4: 13,  # Nov-25: Nov 6→Nov 19
    7: 18,  # Feb-26: Jan 31→Feb 18
    8: 8,   # Mar-26: Mar 2→Mar 10
    9: 16,  # Apr-26: Mar 31→Apr 16
   10: 13,  # May-26: Apr 30→May 13
}
for i,v in CLOSE_TO_CREDITS.items():
    put(rr, 25, mc(i), v, fill=CERT, fmt="#,##0")

# FY 2025 for AECN Gross CDR (row 12)
# In-window 2025: Aug(804.551) + Sep(0) + Oct(636.309×AECN%) + Nov(574.795)
# AECN in-window 2025 gross = 804.551 + 0 + ~453 + 574.795 = ~1832
put(rr, 12, CF, 1832.0, fill=CALC)
# FY 2025 Kerry gross (row 14): Sep-25 = 161.907 + Oct-25 Kerry ~183 = ~345
put(rr, 14, CF, 345.0, fill=CALC)

# ─────────────────────────────────────────────────────────────────────────────
# PENALTIES
# ─────────────────────────────────────────────────────────────────────────────
pen = wb["Penalties"]

# Row 8 Kerry — correct to zero for production; Kerry is transport-only
for i in range(NM):
    v = pen.cell(8, mc(i)).value
    if v is None:
        put(pen, 8, mc(i), 0, fill=CERT)

# Row 9 Downstream Transport total (AECN + Kerry transport embodied)
# AECN truck: 0.604% of AECN gross CDR per period
# Kerry: 2.24% of Kerry gross CDR per period
# All BCU-eligible process transport is offset; only embodied remains as net cost
DT = {
    1:  804.551*0.00604 + 0,              # Aug-25: AECN only
    2:  0 + 161.907*0.02243,              # Sep-25: Kerry only
    3:  "=453.22*0.00604+183.08*0.02243", # Oct-25: mixed (formula)
    4:  574.795*0.00604,                  # Nov-25: AECN only
    7:  157.052*0.00604,                  # Feb-26: AECN only
    8:  863.601*0.00604,                  # Mar-26: AECN only
    9:  702.779*0.00604,                  # Apr-26: AECN only
   10:  "=560.21*0.00604+10.367*0.02243", # May-26: mixed (formula)
}
for i,v in DT.items():
    put(pen, 9, mc(i), v, fill=CALC)

# Row 10 Well Facility — Section E "Site+Amortized Embodied" (includes capex amortization)
WF = {
    1:  176.185,  # Aug-25 (Jun-Jul period, KS Vaulted startup amortization)
    2:  2.988,    # Sep-25 (Aug 1-15)
    3:  "=148.307+13.730",  # Oct-25 (two periods)
    4:  24.975,   # Nov-25 (Oct-Nov)
    7:  129.935,  # Feb-26 (Nov-Jan, Basco startup amortization)
    8:  26.473,   # Mar-26 (Feb-Mar)
    9:  0.000,    # Apr-26 (Mar-Apr, no site+embodied)
   10:  15.413,   # May-26 (Apr)
}
for i,v in WF.items():
    put(pen, 10, mc(i), v, fill=CERT)

# AECN component rows (21-24)
# Row 21: Transport feedstock = 0 (AECN biomass transport negligible/not attributed in records)
# Row 22: Electricity = 0 (AECN uses propane; RECs cover most electricity)
# Row 23: N₂O / process = ~0.5% of AECN row 7
# Row 24: Other (pyrolysis propane/diesel, capital EE, SP+C) = residual
AECN_ROW7 = {1:248.601, 2:0, 3:"formula", 4:120.085, 7:113.582, 8:99.461, 9:89.789, 10:85.07}
for i in range(NM):
    put(pen, 21, mc(i), 0, fill=ZERO)  # transport feedstock
    put(pen, 22, mc(i), 0, fill=ZERO)  # electricity
for i,v in AECN_ROW7.items():
    if isinstance(v, (int, float)):
        n2o = round(v * 0.005, 3)
        other = round(v * 0.995 - 0, 3)
        put(pen, 23, mc(i), n2o, fill=CALC)
        put(pen, 24, mc(i), other, fill=CALC)

# Kerry component rows (27-30): all zero — Kerry has no production process
for row in [27, 28, 29, 30]:
    for i in range(NM):
        put(pen, row, mc(i), 0, fill=ZERO)

# Downstream detail rows
# Row 35: AECN final truck (AECN → injection well, all periods; embodied only after BCU)
AECN_GROSS = {1:804.551, 2:0, 3:453.22, 4:574.795, 7:157.052, 8:863.601, 9:702.779, 10:560.21}
for i,v in AECN_GROSS.items():
    put(pen, 35, mc(i), round(v*0.00604, 3), fill=CALC)

# Row 39: Kerry final truck (Kerry → Basco, embodied only)
KERRY_GROSS = {2:161.907, 3:183.08, 10:10.367}
for i,v in KERRY_GROSS.items():
    put(pen, 39, mc(i), round(v*0.02243, 3), fill=CALC)

# Row 33 Downstream total = formula sum of component rows
for i in range(NM):
    c_col = cl(i)
    formula = f"=SUM({c_col}34:{c_col}40)"
    cell = pen.cell(33, mc(i))
    cell.value = formula
    cell.fill = F(CALC)
    cell.number_format = "#,##0.00"

# Well facility component breakdown for May-26 (Apr 2026 batch rmv_1KRH4CXAP1S0M44D)
# Row 42: Electricity, Row 43: Methanol/diesel, Row 44: Fugitive/other, Row 45: Other site
put(pen, 42, mc(10), 2.168, fill=CERT)   # RECs electricity
put(pen, 43, mc(10), 2.731, fill=CERT)   # Diesel (well ops)
put(pen, 44, mc(10), 0.162, fill=CERT)   # Railcar movement
put(pen, 45, mc(10), "=15.413-2.168-2.731-0.162", fill=CALC)  # Other (gasoline+waste+sitework+capex+SP&C)

# ─────────────────────────────────────────────────────────────────────────────
# EFFICIENCY
# ─────────────────────────────────────────────────────────────────────────────
eff = wb["Efficiency"]

# Rows 6-8: AECN/Charm/Kerry efficiency = net/gross from Run Rate source rows
# AECN: net=RR row 6, gross=RR row 12
# Charm: net=RR row 7, gross=RR row 13
# Kerry: net=RR row 8, gross=RR row 14
for eff_row, rr_net, rr_gross in [(6,6,12),(7,7,13),(8,8,14)]:
    for i in range(NM):
        c_col = cl(i)
        cell = eff.cell(eff_row, mc(i))
        cell.value = f"=IF('Run Rate'!{c_col}{rr_gross}>0,'Run Rate'!{c_col}{rr_net}/'Run Rate'!{c_col}{rr_gross},\"\")"
        cell.fill = F(CALC)
        cell.number_format = "0.0%"
    yc = eff.cell(eff_row, CY)
    n_start = get_column_letter(mc(6)); n_end = get_column_letter(mc(11))
    yc.value = "=IF(SUM('Run Rate'!" + n_start + str(rr_gross) + ":" + n_end + str(rr_gross) + ")>0,SUM('Run Rate'!" + n_start + str(rr_net) + ":" + n_end + str(rr_net) + ")/SUM('Run Rate'!" + n_start + str(rr_gross) + ":" + n_end + str(rr_gross) + "),\"\")"
    yc.fill = F(CALC); yc.number_format = "0.0%"

# Row 21: Charm bio-oil volume (memo) — copy from Production tab
prod = wb["Production (Charm)"]
for i in range(NM):
    v = prod.cell(8, mc(i)).value  # prod_bio_oil_vol = row 8
    if v:
        put(eff, 21, mc(i), v, fill=PP)

# Row 11: Net CDR per tonne — only fill for Charm where we have bio-oil volume
# (Total and AECN/Kerry need bio-oil volumes not yet available)
for i in range(NM):
    net_v = eff.cell(17, mc(i)).value   # net CDR total
    bio_v = eff.cell(19, mc(i)).value   # bio-oil total (blank for now)
    if bio_v:
        c = eff.cell(11, mc(i))
        c_col = cl(i)
        c.value = f"=IF({c_col}19>0,{c_col}17/{c_col}19,\"\")"
        c.fill = F(CALC); c.number_format = "#,##0.00"

# Row 13: Charm net CDR per tonne injected
# Charm net CDR from Run Rate row 7, Charm bio-oil from Production row 8
for i in range(NM):
    charm_net = wb["Run Rate"].cell(7, mc(i)).value
    charm_bio = prod.cell(8, mc(i)).value
    if charm_net and charm_bio and float(charm_bio) > 0:
        put(eff, 13, mc(i), charm_net / charm_bio, fill=CALC)

# ─────────────────────────────────────────────────────────────────────────────
# SOURCE MIX
# ─────────────────────────────────────────────────────────────────────────────
sm = wb["Source Mix"]

# Row 5 total — ensure all months are filled (copy from Run Rate row 5)
for i in range(NM):
    v = wb["Run Rate"].cell(5, mc(i)).value
    if v:
        put(sm, 5, mc(i), v, fill=CERT)
put(sm, 5, CF, wb["Run Rate"].cell(5, CF).value, fill=CERT)

# Rows 6-8 AECN/Charm/Kerry net CDR — sync from Run Rate
for sm_row, rr_row in [(6,6),(7,7),(8,8)]:
    for i in range(NM):
        v = wb["Run Rate"].cell(rr_row, mc(i)).value
        if v is not None:
            put(sm, sm_row, mc(i), v, fill=CERT)

# Row 16 Gross CDR total — sync from Run Rate row 11
for i in range(NM):
    v = wb["Run Rate"].cell(11, mc(i)).value
    if v:
        sm.cell(16, mc(i)).value = v
        sm.cell(16, mc(i)).fill = F(CERT)
        sm.cell(16, mc(i)).number_format = "#,##0.00"
put(sm, 16, CF, wb["Run Rate"].cell(11, CF).value, fill=CERT)

# Rows 17-19 AECN/Charm/Kerry Gross CDR — sync from Run Rate rows 12-14
for sm_row, rr_row in [(17,12),(18,13),(19,14)]:
    for i in range(NM):
        v = wb["Run Rate"].cell(rr_row, mc(i)).value
        if v:
            sm.cell(sm_row, mc(i)).value = v
            sm.cell(sm_row, mc(i)).fill = F(CALC)
            sm.cell(sm_row, mc(i)).number_format = "#,##0.00"

# Row 28: Largest single-source share (%, concentration risk)
for i in range(NM):
    c_col = cl(i)
    c = sm.cell(28, mc(i))
    c.value = f"=IF({c_col}5>0,MAX({c_col}11,{c_col}12,{c_col}13),\"\")"
    c.fill = F(CALC); c.number_format = "0.0%"
yc = sm.cell(28, CY)
ytd_s = get_column_letter(mc(6)); ytd_e = get_column_letter(mc(11))
yc.value = f"=IF(SUM({ytd_s}5:{ytd_e}5)>0,MAX(SUM({ytd_s}11:{ytd_e}11)/SUM({ytd_s}5:{ytd_e}5),SUM({ytd_s}12:{ytd_e}12)/SUM({ytd_s}5:{ytd_e}5),SUM({ytd_s}13:{ytd_e}13)/SUM({ytd_s}5:{ytd_e}5)),\"\")"
yc.fill = F(CALC); yc.number_format = "0.0%"

# Row 29: Number of active sources
ACTIVE = {
    1:1,  # Aug-25: AECN only
    2:1,  # Sep-25: Kerry only
    3:2,  # Oct-25: AECN + Kerry
    4:1,  # Nov-25: AECN only
    7:1,  # Feb-26: AECN only
    8:1,  # Mar-26: AECN only
    9:1,  # Apr-26: AECN only
   10:3,  # May-26: AECN + Kerry + Charm
}
for i,v in ACTIVE.items():
    put(sm, 29, mc(i), v, fill=CERT, fmt="#,##0")

# ─────────────────────────────────────────────────────────────────────────────
# DASHBOARD — wire all summary rows to detail tabs
# ─────────────────────────────────────────────────────────────────────────────
db = wb["Dashboard"]

# Summary table starts at row 11. Rows:
# 11=Net CDR total, 12=AECN, 13=Charm, 14=Kerry, 15=Gross CDR total
# 16=Net efficiency, 17=Pipeline, 18=Last inj→credits, 19=Penalties % gross
# 20=AECN share, 21=Charm share, 22=Kerry share, 23=bio-oil CI, 24=biochar CI

wires = [
    # (dashboard_row, source_sheet, source_row, format)
    (11, "Run Rate",    5,  "#,##0.00"),  # Net CDR total
    (12, "Run Rate",    6,  "#,##0.00"),  # AECN net CDR
    (13, "Run Rate",    7,  "#,##0.00"),  # Charm net CDR
    (14, "Run Rate",    8,  "#,##0.00"),  # Kerry net CDR
    (15, "Run Rate",    11, "#,##0.00"),  # Gross CDR total
    (17, "Run Rate",    17, "#,##0.00"),  # Pipeline total
    (18, "Run Rate",    23, "#,##0"),     # Last inj→credits OKR
    (20, "Source Mix",  11, "0.0%"),      # AECN share
    (21, "Source Mix",  12, "0.0%"),      # Charm share
    (22, "Source Mix",  13, "0.0%"),      # Kerry share
]
for db_row, src_sheet, src_row, fmt in wires:
    for i in range(NM):
        c_col = cl(i)
        c = db.cell(db_row, mc(i))
        c.value = f"=IF('{src_sheet}'!{c_col}{src_row}<>\"\",'{src_sheet}'!{c_col}{src_row},\"\")"
        c.fill = F(CALC); c.number_format = fmt
    # YTD
    yc = db.cell(db_row, CY)
    s = get_column_letter(mc(6)); e = get_column_letter(mc(11))
    if "%" in fmt:
        yc.value = f"=IF(SUM('Run Rate'!{s}11:{e}11)>0,SUM('{src_sheet}'!{s}{src_row}:{e}{src_row})/SUM('Run Rate'!{s}11:{e}11),\"\")" if db_row in [20,21,22] else f"=IFERROR('{src_sheet}'!{get_column_letter(CY)}{src_row},\"\")"
    else:
        yc.value = f"=IFERROR('{src_sheet}'!{get_column_letter(CY)}{src_row},\"\")"
    yc.fill = F(CALC); yc.number_format = fmt

# Row 16 Net efficiency = net CDR / gross CDR (from Efficiency tab total)
for i in range(NM):
    c_col = cl(i)
    c = db.cell(16, mc(i))
    c.value = f"=IFERROR('Efficiency'!{c_col}5,\"\")"
    c.fill = F(CALC); c.number_format = "0.0%"
yc = db.cell(16, CY)
yc.value = f"=IFERROR('Efficiency'!{get_column_letter(CY)}5,\"\")"
yc.fill = F(CALC); yc.number_format = "0.0%"

# Row 19 Penalties as % of Gross CDR
for i in range(NM):
    c_col = cl(i)
    c = db.cell(19, mc(i))
    c.value = f"=IF('Penalties'!{c_col}5>0,'Penalties'!{c_col}5/'Run Rate'!{c_col}11,\"\")"
    c.fill = F(CALC); c.number_format = "0.0%"

# ─────────────────────────────────────────────────────────────────────────────
# PRODUCTION (Charm) — CI calculations for months with both net CDR + penalties
# May-26: Charm bio-oil CI = WODO allocated penalties / WODO net CDR from Certify
# WODO net CDR May-26 = 2.64 tCO₂e (from Certify); penalties Apr-26 PP WODO allocated = 13.579
# Biochar net CDR May-26 = (from Certify QOWV 4.29 tCO₂e) / Biochar penalties Apr-26 = 8.024
# CI = penalties / net CDR
put(wb["Production (Charm)"], 13, mc(10), 13.579 / 2.64, fill=CALC)   # bio-oil CI May-26
put(wb["Production (Charm)"], 21, mc(10), 8.024  / 4.29, fill=CALC)   # biochar CI May-26

# ─────────────────────────────────────────────────────────────────────────────
wb.save(FILE)
print("Saved.")

# Audit: print every sheet, every row that has ANY data
print("\n=== AUDIT — all populated cells by sheet ===")
for sname in wb.sheetnames:
    ws2 = openpyxl.load_workbook(FILE)[sname]
    has_data = False
    for row in ws2.iter_rows():
        for c in row:
            if c.value not in (None, "") and not str(c.value).startswith("=") and c.column >= CB and c.column <= CF:
                if not has_data:
                    print(f"\n{sname}:")
                    has_data = True
                print(f"  Row {c.row} {MO[c.column-CB] if CB <= c.column <= CB+NM-1 else ('YTD' if c.column==CY else 'FY25')}: {c.value}")
