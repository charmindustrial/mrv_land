#!/usr/bin/env python3
"""
Inject source breakdown + verification timing data from all 9 GHG statements.
Source: Certify removals tab, reviewed 2026-06-05 by subagent.

Run Rate row mapping (from build script trace):
  Row 5 = Net CDR Total (already populated)
  Row 6 = AECN
  Row 7 = Charm
  Row 8 = Kerry

  Row 23 = OKR: last injection → credits (total, days)
  Row 28 = AECN blended lag
  Row 29 = Charm blended lag
  Row 30 = Kerry blended lag

Source Mix row mapping:
  Row 5 = Net CDR total
  Row 6 = AECN
  Row 7 = Charm
  Row 8 = Kerry
"""

import openpyxl
from openpyxl.styles import PatternFill

FILE = "/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/mrv-operational-documents/MRV_BI_Dashboard_Shell.xlsx"

CERT_BG = "E2EFDA"    # verified from Certify
PEND_BG = "FFF2CC"    # pending / not separable

def F(c): return PatternFill("solid", fgColor=c)

def put(ws, row, col, value, fill=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    if fill: c.fill = F(fill)
    if fmt:  c.number_format = fmt

CB = 2
def mc(i): return CB + i

wb = openpyxl.load_workbook(FILE)

# ── Month index mapping ────────────────────────────────────────────────────────
# idx 0=Jul-25, 1=Aug-25, 2=Sep-25, 3=Oct-25, 4=Nov-25, 5=Dec-25,
#     6=Jan-26, 7=Feb-26, 8=Mar-26, 9=Apr-26, 10=May-26, 11=Jun-26

# ── RUN RATE: source breakdown ────────────────────────────────────────────────
ws = wb["Run Rate"]

# Row 6 = AECN net CDR by credit-issuance month
# Aug-25 (1): Jun-Jul 2025 period = mixed/legacy template — leave pending
# Oct-25 (3): two Oct statements combined: Aug 16-Sep 8 AECN 59.55 + Sep 9-30 AECN 291.85 = formula
AECN = {
    # 1: pending (Jun-Jul legacy template)
    2:  0,                        # Sep-25: 100% Kerry
    3:  "=59.55+291.85",          # Oct-25: two statements
    4:  454.71,                   # Nov-25: 100% AECN
    7:  43.47,                    # Feb-26: 100% AECN
    8:  764.14,                   # Mar-26: 100% AECN
    9:  612.99,                   # Apr-26: 100% AECN
    10: 475.09,                   # May-26: mixed
}
for idx, val in AECN.items():
    put(ws, 6, mc(idx), val, fill=CERT_BG, fmt="#,##0.00")

# Row 7 = Charm net CDR
CHARM_NET = {
    2:  0,
    3:  0,
    4:  0,
    7:  0,
    8:  0,
    9:  0,
    10: 6.93,     # May-26: QOWV 4.29 + WODO 2.64
}
for idx, val in CHARM_NET.items():
    put(ws, 7, mc(idx), val, fill=CERT_BG, fmt="#,##0.00")

# Row 8 = Kerry net CDR
KERRY = {
    1:  None,                     # Aug-25: mixed — pending
    2:  135.39,                   # Sep-25: 100% Kerry
    3:  106.85,                   # Oct-25: Kerry from Aug 16-Sep 8 statement
    4:  0,
    7:  0,
    8:  0,
    9:  0,
    10: 8.80,                     # May-26: mixed
}
for idx, val in KERRY.items():
    if val is not None:
        put(ws, 8, mc(idx), val, fill=CERT_BG, fmt="#,##0.00")

# Aug-25 (idx=1): source not separable from legacy template — mark pending
for row in [6, 7, 8]:
    put(ws, row, mc(1), None, fill=PEND_BG)

# ── RUN RATE: OKR timing (row 23) ─────────────────────────────────────────────
# Last injection → credit issuance (days)
# For Oct-25 (two statements): conservative = max of the two (31 days for Aug 16-Sep 8 statement)
OKR_DAYS = {
    1:  27,   # Aug-25: Jul 23 → Aug 19 = 27 days
    2:  35,   # Sep-25: Aug 15 → Sep 19 = 35 days
    3:  31,   # Oct-25: worst case — Sep 8 → Oct 9 (31 vs 17 for the other)
    4:  14,   # Nov-25: Nov 5 → Nov 19 = 14 days
    7:  68,   # Feb-26: Dec 12 → Feb 18 = 68 days (Nov-Jan 3-removal period)
    8:  8,    # Mar-26: Mar 2 → Mar 10 = 8 days
    9:  21,   # Apr-26: Mar 26 → Apr 16 = 21 days
    10: 14,   # May-26: Apr 29 → May 13 = 14 days
}
for idx, val in OKR_DAYS.items():
    put(ws, 23, mc(idx), val, fill=CERT_BG, fmt="#,##0")

# ── SOURCE MIX: same source breakdown ────────────────────────────────────────
# Source Mix rows: 5=total (already from Net CDR), 6=AECN, 7=Charm, 8=Kerry
ws = wb["Source Mix"]

for idx, val in AECN.items():
    put(ws, 6, mc(idx), val, fill=CERT_BG, fmt="#,##0.00")
for idx, val in CHARM_NET.items():
    put(ws, 7, mc(idx), val, fill=CERT_BG, fmt="#,##0.00")
for idx, val in KERRY.items():
    if val is not None:
        put(ws, 8, mc(idx), val, fill=CERT_BG, fmt="#,##0.00")
# Aug-25 pending
for row in [6, 7, 8]:
    put(ws, row, mc(1), None, fill=PEND_BG)

# Total Net CDR in Source Mix = same as Run Rate row 5
ws_rr = wb["Run Rate"]
ws_sm = wb["Source Mix"]
for i in range(12):
    val = ws_rr.cell(row=5, column=mc(i)).value
    if val is not None:
        put(ws_sm, 5, mc(i), val, fill=CERT_BG, fmt="#,##0.00")
fy25 = ws_rr.cell(row=5, column=2+12+1).value
if fy25:
    put(ws_sm, 5, 2+12+1, fy25, fill=CERT_BG, fmt="#,##0.00")

wb.save(FILE)
print("Done.")

# Verify
wb2 = openpyxl.load_workbook(FILE)
ws2 = wb2["Run Rate"]
MO = ["Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25","Jan-26","Feb-26","Mar-26","Apr-26","May-26","Jun-26"]
print("\nRun Rate source rows:")
for label, row in [("AECN",6),("Charm",7),("Kerry",8)]:
    vals = [ws2.cell(row=row, column=mc(i)).value for i in range(12)]
    filled = [(MO[i], vals[i]) for i in range(12) if vals[i] is not None]
    print(f"  {label}: {filled}")
print("\nOKR timing row 23:")
vals = [ws2.cell(row=23, column=mc(i)).value for i in range(12)]
print(f"  {[(MO[i], vals[i]) for i in range(12) if vals[i] is not None]}")
