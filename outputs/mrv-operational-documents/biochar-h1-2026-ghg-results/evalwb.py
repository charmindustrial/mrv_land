"""Resolve the formula chain in an Isometric Certify export workbook.

The export ships formulas with no cached results, so openpyxl reads them as
blank. This walks the references (Sheet!A1 style), arithmetic and the two
functions the export uses (SUM, ROUND) and returns real numbers.
"""
import re, math
import openpyxl

import os
PATH = os.environ.get('CERTIFY_EXPORT') or os.path.join(os.path.dirname(os.path.abspath(__file__)), 'certify_export_2026-08-07.xlsx')

wb = openpyxl.load_workbook(PATH, data_only=False)
cells = {}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                cells[(ws.title, c.coordinate)] = c.value

_cache, _stack = {}, set()

REF = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_ ]*))!(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?"
                 r"|(?<![A-Za-z0-9_!$])(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?(?![A-Za-z0-9_(])")


def _num(v):
    if v is None or v == '' or isinstance(v, str) and not v.strip():
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    try:                      # the export stores some results as text
        return float(str(v).strip().replace(',', ''))
    except ValueError:
        return 0.0


def cell(sheet, coord):
    coord = coord.replace('$', '')
    key = (sheet, coord)
    if key in _cache:
        return _cache[key]
    if key in _stack:
        raise RuntimeError(f'circular reference at {key}')
    raw = cells.get(key)
    if not (isinstance(raw, str) and raw.startswith('=')):
        val = _num(raw)
        _cache[key] = val
        return val
    _stack.add(key)
    try:
        val = evaluate(raw[1:], sheet)
    finally:
        _stack.discard(key)
    _cache[key] = val
    return val


def rng(sheet, a, b):
    a, b = a.replace('$', ''), b.replace('$', '')
    ca, ra = re.match(r'([A-Z]+)(\d+)', a).groups()
    cb, rb = re.match(r'([A-Z]+)(\d+)', b).groups()
    ia, ib = openpyxl.utils.column_index_from_string(ca), openpyxl.utils.column_index_from_string(cb)
    out = []
    for ci in range(min(ia, ib), max(ia, ib) + 1):
        letter = openpyxl.utils.get_column_letter(ci)
        for ri in range(min(int(ra), int(rb)), max(int(ra), int(rb)) + 1):
            out.append(cell(sheet, f'{letter}{ri}'))
    return out


def evaluate(body, sheet):
    """Rewrite an Excel expression into Python and evaluate it."""
    parts, pos = [], 0
    for m in REF.finditer(body):
        parts.append(body[pos:m.start()])
        qs, bs, c1, c2, lc1, lc2 = m.groups()
        tgt = qs or bs or sheet
        start, end = (c1, c2) if c1 else (lc1, lc2)
        if end:
            parts.append(f'_R({tgt!r},{start!r},{end!r})')
        else:
            parts.append(f'_C({tgt!r},{start!r})')
        pos = m.end()
    parts.append(body[pos:])
    expr = ''.join(parts)
    expr = re.sub(r'\bSUM\(', '_SUM(', expr)
    expr = re.sub(r'\bROUND\(', '_ROUND(', expr)
    env = {
        '_C': lambda s, c: cell(s, c),
        '_R': lambda s, a, b: rng(s, a, b),
        '_SUM': lambda x: sum(x) if isinstance(x, list) else float(x),
        '_ROUND': lambda x, n=0: _xlround(x, int(n)),
    }
    return eval(expr, {'__builtins__': {}}, env)


def _xlround(x, n):
    """Excel rounds half away from zero; Python rounds half to even."""
    f = 10 ** n
    y = x * f
    return (math.floor(y + 0.5) if y >= 0 else math.ceil(y - 0.5)) / f


def sheet_rows(name):
    ws = wb[name]
    hdr = [c.value for c in ws[1]]
    out = []
    for r in range(2, ws.max_row + 1):
        rec = {}
        for i, h in enumerate(hdr, start=1):
            if h is None:
                continue
            coord = f'{openpyxl.utils.get_column_letter(i)}{r}'
            raw = cells.get((name, coord))
            if isinstance(raw, str) and raw.startswith('='):
                try:
                    rec[h] = cell(name, coord)
                except Exception as e:
                    rec[h] = f'ERR:{e}'
            else:
                rec[h] = raw
        if any(v not in (None, '', ' ') for v in rec.values()):
            out.append(rec)
    return out
