#!/usr/bin/env python3
"""Build Batch 2-187 QA Checklist xlsx with FAIL/FLAG/N/A formatting and [log:KEY] citations."""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

FAIL_FILL = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
FAIL_FONT = Font(color='9C0006', bold=True)
FLAG_FILL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
FLAG_FONT = Font(color='996600')
NA_FONT = Font(color='808080', italic=True)
HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=12)
SECTION_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
SECTION_FONT = Font(bold=True, size=11)

THIN = Side(border_style='thin', color='808080')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='left')
WRAP_CENTER = Alignment(wrap_text=True, vertical='center', horizontal='center')

ROWS = [
    ('HEADER',),
    ('META', 'Batch ID', '2-187'),
    ('META', 'Reporting Period', 'Apr 2026 (Certify removal date 03 Apr 2026)'),
    ('META', 'Reviewer', 'Claude QA Agent'),
    ('META', 'Date', '2026-05-01'),
    ('META', 'Injection Site', 'Basco 6'),
    ('META', 'Oil Type', 'AECN'),
    ('META', 'Delivery Mode', 'Rail (railcar TIMX201060, consumed_lot offload_2 of 4 — NOT the last batch on the railcar; 2-186=1st pull, 2-187=2nd pull THIS BATCH, 2-188/2-189 remain)'),
    ('META', 'Net CDR (Certify)', '29.48 tCO2e'),
    ('SPACER',),
    ('COLUMNS',),

    ('SECTION', '1. Batch Folder Completeness'),
    ('1.1', 'Batch folder exists with correct naming', 'Pass',
     'Drive folder "2-187" at id 1FZKcMtZYmalLB8J3QAcBtR1ecvfHOxkO; contains "2-187 Ops Notes Rail" sheet, Testing subfolder (SR26-02264.pdf), Scale Tickets subfolder (full + empty JPGs). [log:drive_folder_url, drive_inventory.batch_folder]',
     'VR Appendix 1'),
    ('1.2', 'Scale Tickets subfolder present and populated', 'Pass',
     'Drive > Scale Tickets contains 2-187 Full Scale Ticket.jpg and 2-187 Empty Scale Ticket.jpg. One offload pair (consistent with batch 2-187 = consumed_lot offload_2 only; sister batches 2-186/2-188/2-189 each contribute their own ticket pairs). Both files also uploaded to Certify Sources as "(1)" copies. [log:scale_tickets.full_ticket.filename, scale_tickets.empty_ticket.filename, drive_inventory.Scale_Tickets_subfolder]',
     'VR Appendix 1, Sec 5.1'),
    ('1.3', 'Testing subfolder present with required lab results', 'Pass',
     'Drive > Testing contains SR26-02264.pdf (7-page SGS Certificate of Analysis covering batches 2-186 through 2-189 + 3-55 through 3-57). Page 2 covers batch 2-187 explicitly (Certificate SR26-02264.002, Client ID = "BATCH 2-187"). [log:lab_results.filename_drive, lab_results.page_for_2_187, lab_results.certificate_id_for_2_187, lab_results.client_id_on_page]',
     'VR Appendix 1'),
    ('1.4', 'Ops Notes spreadsheet present and complete', 'Pass',
     '"2-187 Ops Notes Rail" Google Sheet (id 1RnVQ7KR1E-u5VBtdlvM2o3B-2sfBHE25yZJLRnRJD7w) present in Drive folder; owner Garrett Lutz; modified 2026-05-01 09:04. Batch Composition (railcar TIMX201060, sparging 0.0129), BOLs Loaded (3 inbound BOLs totaling 80.61 MT corrected), and offload table shows offload_1 + offload_2 rows populated as a running railcar tally; offload_2 is THIS batch (Cargo 21,736.12864 kg, allocation fraction 0.26964556). Offload_3/4 rows zero by design. [log:ops_notes.filename, ops_notes.url, ops_notes.trucks_unloaded]',
     'VR Appendix 1'),
    ('1.5', 'Injection site confirmed Basco 6, no Vaulted refs', 'Pass',
     'CHN PDF page 2 LOCATION = "BASCO LOUISIANA"; Certify project = "Charm Bio-oil Geologic Storage" (Basco). Ops Notes references Opelousas->Basco truck leg. No Vaulted references in any source. [log:lab_results.location, drive_inventory._note_no_vaulted_refs]',
     'PDD'),

    ('SECTION', '2. Scale Tickets & Mass Determination'),
    ('2.1', 'Empty and full scale ticket images present', 'Pass',
     'Both JPG images Read directly (NOT Drive OCR snippet, per OCR-vs-image rule). Full ticket weigh #0762 dated 3-31-26 at PILOT TRAVEL CENTER I-49 EXIT 40 BUNKIE LA scale #3122, axle sum 10,200 + 33,920 + 33,780 = 77,900 lb. Empty ticket weigh #1778 dated 4-01-26 same location, axle sum 9,500 + 11,820 + 8,660 = 29,980 lb. Both internally consistent (gross = axle sum). [log:scale_tickets.full_ticket, scale_tickets.empty_ticket]',
     'VR Sec 5.1'),
    ('2.2', 'Dates within 2 weeks of injection', 'Pass',
     'Full ticket 2026-03-31; empty ticket 2026-04-01; injection completion 2026-04-03. Within 3 days. [log:scale_tickets.full_ticket.date, scale_tickets.empty_ticket.date, ops_notes.completion_date]',
     'VR Sec 5.1'),
    ('2.3', 'Tractor and trailer numbers match between empty and full', 'Pass',
     'Full ticket: Tractor CH1, Trailer KL-192. Empty ticket: Tractor CH1, Trailer KL-192. Same vehicle weighed at both events. [log:scale_tickets.tractor_match, scale_tickets.trailer_match]',
     'VR Sec 5.1'),
    ('2.4', 'Net mass correctly calculated in Ops Notes', 'Pass',
     'NOT the last offload of the railcar (consumed_lot offload_2 of 4) — the rail last-offload conservatism policy does NOT apply here. Cargo (kg) is raw scale-ticket math using Ops Notes lbs->kg factor 0.453592: (77,900 - 29,980) lb × 0.453592 = 47,920 lb × 0.453592 = 21,736.12864 kg. Ops Notes offload_2 row Cargo Mass (kg) = 21,736.12864 — exact match. [log:scale_tickets.net_load_lb, scale_tickets.net_load_kg, ops_notes.lbs_to_kg_conversion, ops_notes.trucks_unloaded.1.cargo_kg, independent_calculations.scale_ticket_net_kg]',
     'VR Sec 5.1'),
    ('2.5', 'Injectate mass uses conservative MIN(scale net, BOL mass)', 'Pass',
     'Per-offload rail batch (non-last): scale-ticket math is authoritative for this batch; BOL aggregate (80,610 kg corrected) is at railcar level, not per-offload. The 21,736.12864 kg from scale tickets is THIS batch\'s injectate basis. Sparging deduction (next item) applied to that base. [log:scale_tickets.net_load_kg, ops_notes.trucks_unloaded.1.cargo_kg]',
     'VR Sec 5.1, GHG Stmt Sec C'),
    ('2.6', 'Sparging deduction applied correctly', 'Pass',
     'Sparging factor 0.0129 applied: 21,736.12864 × 0.0129 = 280.3960595 kg sparging loss. Mass of Bio-Oil Injected = 21,736.12864 - 280.3960595 = 21,455.73258 kg. Ops Notes shows exactly these values; arithmetic correct. [log:ops_notes.sparging_mass_loss_factor, ops_notes.trucks_unloaded.1.cargo_kg, ops_notes.trucks_unloaded.1.sparging_loss_kg, ops_notes.trucks_unloaded.1.mass_of_bio_oil_injected_kg, independent_calculations.sparging_loss_kg, independent_calculations.mass_of_bio_oil_injected_kg]',
     'VR Sec 5.1'),
    ('2.7', 'Mass values consistent across scale tickets, Ops Notes, Certify', 'Pass',
     'Scale ticket net = 21,736.13 kg = Ops Notes Cargo (kg) = COBB offloaded_quantity_kg = Certify "Allocated Bio-Oil Mass - AECN via Rail" = 21,736.12864 kg. Post-sparging Mass of Injectant (Ops Notes) = 21,455.73258 kg shown in Certify Datapoints as 21.456 t. All four sources align. [log:scale_tickets.net_load_kg, ops_notes.trucks_unloaded.1.cargo_kg, ops_notes.trucks_unloaded.1.mass_of_bio_oil_injected_kg, cobb.offloaded_quantity_kg, certify_calculation_view.datapoints.Allocated Bio-Oil Mass - AECN via Rail, certify_calculation_view.datapoints.Mass of Injectant (per injection)]',
     'VR Sec 5.1'),

    ('SECTION', '3. Bills of Lading & Transportation'),
    ('3.1', 'Truck BOLs present for each shipment', 'Pass',
     'All 3 inbound truck BOLs present in Certify Sources: Documents camion 2317-R.pdf, Billing Truck 2316-R.pdf, Billing Truck 2314-T Corrections (1).pdf. Per Document Location Rule, presence in Certify Sources is sufficient. Filename typo "2314-T" — see Section 9.9 hygiene flag. [log:certify_sources_tab, qa_findings_preliminary.flags.2]',
     'VR Appendix 1'),
    ('3.2', 'Rail batch: 3 truck BOLs from AECN, destination = SOPOR', 'Pass',
     '3 BOLs (2314-R, 2316-R, 2317-R) listed in COBB row 64 BOL_number field for batch 2-187 and in Ops Notes BOLs Loaded section. Destinations = SOPOR Quebec railyard (rail-batch convention). [log:ops_notes.BOLs_loaded, cobb.BOL_number]',
     'VR Appendix 1'),
    ('3.3', 'BOL mass >= reported injectate mass at railcar level', 'Pass',
     'Sum of 3 BOL masses (26.05 + 29.06 + 25.50) = 80.61 MT = 80,610 kg (Ops Notes Total Loaded Mass). This is the AECN-CORRECTED inbound total per Max + Garrett confirmation 2026-04-30 documented on sister batch 2-186 (original BOL was short by 9.6 MT; AECN sent correction documented as "Billing Truck 2314-T Corrections (1).pdf"). This batch\'s offload mass = 21,736.13 kg, well below 80,610 kg railcar inbound. Rail-batch BOL mass check applied at railcar level (per SKILL.md). [log:ops_notes.BOLs_loaded, ops_notes.total_loaded_mass_MT, ops_notes.total_loaded_mass_kg, ops_notes.trucks_unloaded.1.cargo_kg, ops_notes._note_total_loaded]',
     'Transport Emissions Acctg v1.1'),
    ('3.4', 'Origin and destination correct on BOLs', 'Pass',
     'Carrier route: AECN Quebec -> SOPOR (truck, 3 BOLs labeled with "-R" rail-load suffix) -> Opelousas LA (rail leg via railcar TIMX201060) -> Basco 6 (truck, this batch consumed_lot offload_2). Three legs documented across Drive + Certify Sources. [log:cobb.BOL_number, ops_notes.BOLs_loaded, certify_components.components_detail]',
     'VR Sec 3.5'),
    ('3.5', 'Transportation distances consistent with known AECN rail route', 'Pass',
     'Three legs verified: (1) AECN->SOPOR truck full single distance 173 km (= 107.5 mi) per Certify Datapoint "Distance from AECN to Sopor Quebec (truck)", and Allocated Distance AECN to Sopor Quebec 86.96 mi (139.95 km) computed as 322.5 mi total trucking × 0.26964556 batch share; (2) SOPOR->Opelousas rail full 2,199 mi per "Distance from SOPOR to Opelousas (RAIL)", Allocated 592.95062 mi (2,199 × 0.26964556); (3) Opelousas->Basco truck per-batch 42.9 mi (full per-batch, no allocation since each offload uses dedicated truck). All match Ops Notes pre-computed values for offload_2 row. [log:certify_calculation_view.datapoints, ops_notes.distances, ops_notes.trucks_unloaded.1.upstream_allocation_fraction, ops_notes.trucks_unloaded.1.rail_miles_allocated, ops_notes.trucks_unloaded.1.truck_leg_1_miles_allocated]',
     'VR Sec 3.5'),
    ('3.6', 'Transport mode = 6 components (3 process + 3 embodied) for rail', 'Pass',
     '6 transport components in Certify: AECN to SOPOR Tanker Truck Process Emissions (0.267 t), SOPOR to Opelousas Rail Process Emissions (1.238 t), Opelousas to Basco - Tanker Truck Process Emissions (0.107 t), AECN to SOPOR Tanker Truck Embodied Emissions - AECN (0.021 t), SOPOR to Opelousas Rail Embodied Emissions (0.018 t), Opelousas to Basco Tanker Truck Embodied Emissions (0.010 t). [log:certify_components.components_detail]',
     'PDD'),
    ('3.7', 'Billing documents present and consistent with BOLs', 'Pass',
     'Documents camion 2317-R.pdf, Billing Truck 2316-R.pdf, Billing Truck 2314-T Corrections (1).pdf in Certify Sources serve as both BOLs and billing for the 3 inbound truck legs (AECN convention). Charm-A5-2024-10-24-62-1976.pdf (BCU certificate) also present. [log:certify_sources_tab]',
     'VR Appendix 1'),

    ('SECTION', '4. Bio-Oil Composition & Testing'),
    ('4.1', 'CHN PDF present in Testing subfolder', 'Pass',
     'SR26-02264.pdf (7-page SGS Certificate, in Drive Testing) AND CHN-186-189.pdf (Certify Sources alias). Page 2 covers batch 2-187 (Certificate SR26-02264.002, "BATCH 2-187"). [log:lab_results.filename_drive, lab_results.filename_certify, lab_results.page_for_2_187, lab_results.certificate_id_for_2_187]',
     'VR Appendix 1'),
    ('4.2', 'Carbon content within expected AECN range', 'Pass',
     'Carbon content for batch 2-187 = 41.2% (m/m) (ASTM D5291 Method C). Within AECN expected range (mean 39.42%, +/-2 SD ~27.2%-51.6%). [log:lab_results.carbon_pct, lab_results.carbon_in_AECN_2SD_range_27_2_to_51_6_pct]',
     'Biomass Feedstock Acctg v1.2'),
    ('4.3', 'Hydrogen and nitrogen values reported and reasonable', 'Pass',
     'Hydrogen 6.9%, Nitrogen <0.75%, Viscosity 22.04 cSt @ 40C, Acid Number 46 mg KOH/g, Density 1.212 g/cm3 @ 15C. All within typical AECN ranges. [log:lab_results.hydrogen_pct, lab_results.nitrogen_pct_lt_0_75, lab_results.viscosity_40C_cSt, lab_results.acid_number_mgKOHpg, lab_results.density_15C_gcm3]',
     'Biomass Feedstock Acctg v1.2'),
    ('4.4', 'Lab report references correct batch ID', 'Pass',
     'Certificate SR26-02264.002 explicitly labeled "BATCH 2-187" on page 2. Sampled 2026-04-02 04:45 at BASCO LOUISIANA (after offload_2 truck arrived 2026-04-01). Sample 1 day before injection completion 2026-04-03. [log:lab_results.certificate_id_for_2_187, lab_results.client_id_on_page, lab_results.sampled_datetime, lab_results.location]',
     'VR Appendix 1'),
    ('4.5', 'pH/conductivity (aqueous phase)', 'N/A',
     'AECN is not aqueous phase. pH of injectate is recorded in Ops Notes Section 5.5. [log:oil_type]',
     'N/A'),

    ('SECTION', '5. Ops Notes Accuracy & Internal Consistency'),
    ('5.1', 'Correct batch ID and injection date', 'Pass',
     'Ops Notes Rail header shows Batch number = 2-187, Completion Date = 4/3/2026. Matches Certify removal date (03 Apr 2026) and COBB injection_completion_date (2026-04-03). [log:ops_notes.batch_number_in_file, ops_notes.completion_date, cobb.injection_completion_date]',
     'VR Sec 5.1'),
    ('5.2', 'Scale ticket values correctly transcribed to Ops Notes', 'Pass',
     'Full ticket axle sum 77,900 lb = Ops Notes offload_2 Loaded Mass (lbs) 77,900 (exact). Empty ticket axle sum 29,980 lb = Ops Notes offload_2 Unloaded Mass (lbs) 29,980 (exact). Both values verified by direct image Read of JPGs (per Drive-OCR-not-primary-source rule). Cargo (kg) using Ops Notes factor 0.453592: 47,920 × 0.453592 = 21,736.12864 = Ops Notes offload_2 Cargo Mass (kg) 21736.12864 (exact). [log:scale_tickets.full_ticket.axle_sum_lb, scale_tickets.empty_ticket.axle_sum_lb, ops_notes.lbs_to_kg_conversion, ops_notes.trucks_unloaded.1.loaded_lb, ops_notes.trucks_unloaded.1.unloaded_lb, ops_notes.trucks_unloaded.1.cargo_kg]',
     'VR Sec 5.1'),
    ('5.3', 'Oil and aqueous phase mass reconciliation', 'N/A',
     'AECN-only batch; no aqueous phase. [log:oil_type]',
     'N/A'),
    ('5.4', 'COBB values present and match Ops Notes', 'Flag',
     'COBB row 64 for "2-187" verified: cell A = "2-187" ✓, BOL_number = "2314-R, 2316-R, 2317-R" matches Ops Notes BOLs Loaded ✓, offloaded_quantity_kg = 21,736.12864 = Ops Notes offload_2 Cargo (kg) ✓, injection_completion_date = 2026-04-03 = Ops Notes 4/3/2026 ✓. **Two COBB-side data-provenance issues for Garrett (NOT 2-187-impacting):** (a) COBB row 64 origin_quantity_kg = 71,010 (stale pre-correction). Post-correction inbound aggregate = 80,610 kg per Ops Notes Total Loaded Mass — consistent with the 9.6 MT AECN BOL correction first surfaced on 2-186 (2026-04-30). Sister-batch sum across rows 63-66 of COBB = 21,854.06 + 21,736.13 + 5,620.18 + 21,799.63 = 71,010.00 = pre-correction aggregate exactly (the 5,620 kg on 2-188 row reflects the original pre-correction policy back-calc). COBB has not been refreshed for the correction. (b) COBB row 64 pH field shows "33" instead of "3.3" (Ops Notes offload_2 pH = 3.3 correctly). Both items are operational data-provenance for Garrett to clean up — verifier reads Ops Notes / Certify (both correct). [log:cobb.cobb_row_id_verified, cobb.cobb_cell_A_value, cobb.row_number, cobb.consumed_lot, cobb.BOL_number, cobb.offloaded_quantity_kg, cobb.origin_quantity_kg, cobb._note_origin_qty_stale, cobb.pH_in_cobb_raw, cobb._note_pH_typo, cobb.sister_batches_railcar_TIMX201060, cobb.sister_batches_offloaded_sum_kg, cobb._note_sum_matches_stale_origin]',
     'VR Sec 5.1'),
    ('5.5', 'pH measurement noted', 'Pass',
     'Ops Notes offload_2 row pH = 3.3. [log:ops_notes.pH_offload_2, ops_notes.trucks_unloaded.1.pH]',
     'VR Sec 5.1'),
    ('5.6', 'Completion date present and matches COBB', 'Pass',
     'All three sources align: Ops Notes Completion Date 4/3/2026, COBB row 64 injection_completion_date 2026-04-03, Certify removal date 03 Apr 2026. [log:ops_notes.completion_date, cobb.injection_completion_date, certify_components.removal_date]',
     'VR Sec 5.1'),
    ('5.7', 'No erroneous density measurements', 'N/A',
     'Density field expected blank for AECN per protocol. Ops Notes Batch Composition does not include density entry — expected behavior. [log:ops_notes.density_field, ops_notes._note_density]',
     'N/A'),

    ('SECTION', '6. AECN Feedstock Emissions & Allocation'),
    ('6.1', 'AECN process emissions LCA CI spreadsheet present', 'Pass',
     '"Bioenergy AE Cote-Nord Canada - LCA CI calculation - 1Jul2025 to 31Dec2025 S.xlsx" present in Certify Sources, linked to Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail) datapoint. Vintage covers H2 2025 (correct for batch sampled in Apr 2026 since EFs are H1/H2 — H2 2025 carries through Q1 2026 reporting). [log:certify_sources_tab.14]',
     'VR Appendix 1'),
    ('6.2', 'LCA with stack emissions file present', 'Pass',
     '"BH 2025 LCA w Stack Emissions.xlsx" present in Certify Sources, linked to Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail). [log:certify_sources_tab.16]',
     'VR Appendix 1'),
    ('6.3', 'Supporting GHG data file present', 'Pass',
     '"Supporting GHG data 31Dec2025 S.xlsx" present in Certify Sources, linked to Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail). [log:certify_sources_tab.17]',
     'VR Appendix 1'),
    ('6.4', 'AECN pyrolysis process EF matches current value', 'Pass',
     'Modal value 0.104 tCO2e/t (display) for "Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail)". KB current native value 0.10402 tCO2e/t (H2 2025) — within rounding. Independent calc: 21.736 t × 0.10402 = 2.261 tCO2e ✓ matches Certify "AECN pyrolysis (from rail loads) - process emissions" 2.261 t. [log:certify_components.components_detail.AECN pyrolysis (from rail loads) - process emissions, independent_calculations.pyrolysis_process_t]',
     'Standard EFs 2026 tab'),
    ('6.5', 'AECN pyrolysis embodied EF applied', 'Pass',
     'Modal value 0.02116 tCO2e/t for "Carbon Intensity - AECN Pyrolysis Embodied Emissions (via Rail)". KB current value 0.0212 tCO2e/t (precise per Certify modal slightly more granular at 0.02116). Independent calc: 21.736 t × 0.02116 = 0.460 tCO2e ✓ matches Certify "AECN pyrolysis (from rail loads) - embodied emissions" 0.460 t. [log:certify_components.components_detail.AECN pyrolysis (from rail loads) - embodied emissions, independent_calculations.pyrolysis_embodied_t]',
     'Standard EFs 2026 tab'),
    ('6.6', 'Mass basis for pyrolysis EFs is Origin Bio-Oil Mass (pre-sparging)', 'Pass',
     'Pyrolysis components apply 0.104 / 0.02116 to "Allocated Bio-Oil Mass - AECN via Rail" = 21,736.13 kg = pre-sparging Cargo (kg). Sparging is applied at injection (Section 2.6) and only affects Mass of Injectant for sequestration. Methodology correct. [log:certify_calculation_view.datapoints.Allocated Bio-Oil Mass - AECN via Rail, certify_components.components_detail.AECN pyrolysis (from rail loads) - process emissions]',
     'Bio-oil Geol Storage v1.1'),

    ('SECTION', '7. Gross-to-Net Calculation'),
    ('7.1', 'Gross CO2e = injectate mass x carbon wt% (decimal) x 3.667', 'Pass',
     'Independent calc: 21.456 t (Mass of Injectant, post-sparging) × 0.412 (Carbon Content from CHN page 2) × 3.667 = 32.415 tCO2e. Certify Sequestrations 32.415 tCO2e — exact match. [log:independent_calculations.gross_co2e_t, certify_components.sequestrations_total_t, certify_components.components_detail.Injection Sequestration]',
     'GHG Stmt Sec E'),
    ('7.2', 'Process emissions correctly deducted', 'Pass',
     'All 8 activity components deducted: 3 transport process (AECN->SOPOR truck 0.267, SOPOR->Opelousas rail 1.238, Opelousas->Basco truck 0.107) + 3 transport embodied (0.021 + 0.018 + 0.010) + AECN pyrolysis process 2.261 + AECN pyrolysis embodied 0.460. Sum = 4.382 tCO2e ✓ matches Certify Activities total. Each component math reconciles independently. [log:certify_components.components_detail, certify_components.activities_total_t, independent_calculations.activities_sum_t]',
     'GHG Stmt Sec E'),
    ('7.3', 'Uncertainty discount correctly applied', 'Pass',
     'Discount 170 kg present in Calculation View. Per Isometric Standard S3.7.3 Option B (Variance Propagation), Certify computes the aggregate from per-component +/- inputs — Isometric\'s aggregation is authoritative; QA scope is verifying the INPUTS (covered in Section 7a Layer 1+2: ±72.57 kg per discrete-weight datapoint, ±0.4% carbon). 0.170 t is ~0.5% of gross sequestration (32.415 t), consistent with typical AECN rail batch ratios (sister batch 2-186 was 0.5%). [log:certify_components.uncertainty_discount_kg, certify_components.sequestrations_total_t]',
     'Isometric Standard v2.0 S3.7'),
    ('7.4', 'Net removals consistent with internal calc and BCU offset', 'Pass',
     'Net = 32.415 (sequestrations) + 1.612 (BCU adjustments) - 4.382 (activities) - 0.170 (uncertainty) = 29.475 tCO2e. Certify shows Net CO2e Removed = 29.476 t (displayed as 29.48). Reconciles within rounding. [log:certify_components.net_co2e_t, certify_components.displayed_net_co2e_rounded_t, independent_calculations.net_t]',
     'GHG Stmt Sec E'),

    ('SECTION', '7a. Per-Component Uncertainty Verification'),
    ('7a.1', 'Layer 1 - Completeness: every required component has uncertainty input', 'Pass',
     'All mass-input components show ±72.57 kg in INPUTS section (verified by opening each component modal individually): Injection Sequestration (Mass of product 21,455.73 ±72.57), AECN->SOPOR Truck Process (Mass of load 21,736.13 ±72.57), SOPOR->Opelousas Rail Process (Mass of load 21,736.13 ±72.57), Opelousas->Basco Truck Process (Mass of load 21,736.13 ±72.57), AECN Pyrolysis Process (Mass 21,736.13 ±72.57), AECN Pyrolysis Embodied (Mass 21,736.13 ±72.57). Embodied transport components correctly have NO Mass-of-load datapoint (per-mile EFs, not per-T-mi). Carbon Content uncertainty ±0.1648% (= 41.2 × 0.004 = 0.1648, matches 0.4% relative). [log:certify_components.components_detail]',
     'Isometric Standard S3.7.3'),
    ('7a.2', 'Layer 2 - Correctness: reported +/- matches expected value', 'Pass',
     'Per FC9 per-removal framing: a per-batch offload removal is fixed by ONE outbound discrete weight at Basco offload, so ±72.57 kg per discrete weight is the correct value on every leg (truck process + rail process + truck process). Verified by opening each component modal and reading the INPUTS section directly. Carbon content ±0.1648% = 41.2 × 0.004 confirms standard 0.4% relative uncertainty. Conservatism precondition (N_offloads ≥ N_inbound_BOLs) for railcar TIMX201060 will be formally validated at the railcar\'s last batch QA (2-188 or 2-189 depending on physical-pull-vs-suffix mapping per Max\'s 2026-04-30 note on 2-186). [log:certify_components.components_detail, qa_findings_preliminary.section_8_rail_last_batch_check]',
     'KB Section 7a / FC9'),
    ('7a.3', 'Layer 3 - Discount presence in net calc and consistency', 'Pass',
     'Aggregate uncertainty discount 170 kg in Calculation View. Visible in net formula: 32.415 sequestrations + 1.612 BCU - 4.382 activities - 170 kg uncertainty = 29.476 t Net. Magnitude reasonable given mass and carbon uncertainty inputs. [log:certify_components.uncertainty_discount_kg, certify_components.net_co2e_t, certify_components.sequestrations_total_t, certify_components.activities_total_t, certify_components.adjustments_BCU_total_t]',
     'Isometric Standard S3.7.3'),

    ('SECTION', '8. Data Integrity & Anomaly Check'),
    ('8.1', 'No unexplained outliers in mass, carbon content, or calculated values', 'Pass',
     'Carbon (41.2%), distances (107.5/2199/42.9 mi full-leg + 86.96/592.95 mi allocated), EFs (0.1143 truck process / 0.0259 rail process / 0.23756 truck embodied / 0.03113 rail embodied 2026 corrected value — same value sister batch 2-186 carries; 2-182\'s stale 0.0000352240 was previously updated). Mass 21,736.13 kg pre-sparging / 21,455.73 kg post-sparging within expected per-truck offload range. Both scale tickets internally consistent (gross = axle sum). All math reconciles to Certify within rounding. [log:lab_results.carbon_pct, certify_components.components_detail, independent_calculations]',
     'GHG Stmt Sec C'),
    ('8.2', 'Spills/losses documented; unrecovered mass deducted', 'Pass',
     'No spills/losses for this offload. Ops Notes "Oil Left in Railcar (MT/kg)" cells correctly blank for 2-187 (this is the 2nd pull; railcar TIMX201060 still has remaining oil for sister batches 2-188/2-189). Last-offload mass reconciliation belongs to whichever batch is the railcar\'s physical last pull (per Max\'s 2026-04-30 note on 2-186, that\'s 2-189; per consumed_lot suffix it would be 2-188). [log:ops_notes.trucks_unloaded, ops_notes._note_total_unloaded]',
     'GHG Stmt Sec C'),
    ('8.3', 'Sister-batch sanity check (rail batch precondition for FC9 framing)', 'Pass',
     'Railcar TIMX201060 is mid-drawdown at this batch. Sister-batch state per COBB (showing pre-correction values): 2-186 row 63 = consumed_lot offload_1, Cargo 21,854.06 kg (1st pull, raw); 2-187 row 64 = consumed_lot offload_2, Cargo 21,736.13 kg (THIS BATCH, 2nd pull, raw); 2-188 row 65 = consumed_lot offload_4, Cargo 5,620.18 kg (still pre-correction policy back-calc on COBB); 2-189 row 66 = consumed_lot offload_3, Cargo 21,799.63 kg. **Section 8.3 last-offload reconciliation DEFERRED to the railcar\'s final batch QA** (per Max 2026-04-30: 2-189 is the actual physical 4th/last pull; consumed_lot suffix doesn\'t match physical sequence). For 2-187, the FC9 per-offload uncertainty framing (±72.57 kg per leg) is conservative iff N_offloads ≥ N_inbound_BOLs at last batch — that precondition check will be formally validated when the railcar is fully drawn down. 2-187 is independent of last-offload mechanics: this is a non-last raw-scale-ticket pull. [log:cobb.sister_batches_railcar_TIMX201060, cobb._note_sum_matches_stale_origin, qa_findings_preliminary.section_8_rail_last_batch_check, ops_notes.trucks_unloaded.1.upstream_allocation_fraction]',
     'KB Section 8 / FC9'),

    ('SECTION', '9. Isometric Certify - Removal Component Mapping'),
    ('9.1', 'SEQUESTRATIONS: Gross injection sequestration matches internal gross', 'Pass',
     'Certify Sequestration component 32.415 tCO2e. Independent calc: 21.456 t (Mass of Injectant) × 0.412 × 3.667 = 32.415 t. Exact match. Carbon content matches CHN PDF page 2 (41.2%). [log:certify_components.components_detail.Injection Sequestration, lab_results.carbon_pct, independent_calculations.gross_co2e_t]',
     'Certify Components'),
    ('9.2', 'ACTIVITIES - Transport: 6 components correct for AECN rail route', 'Pass',
     '6 transport components present with correct EFs and distances. Rail embodied EF = 0.03113 kgCO2e/mi (2026 corrected value 0.0000311307 MTCO2e/mi in native units, NOT stale 0.0000352240 carried by sister batch 2-182). Rail-embodied stale-EF finding from 2-182 remains RESOLVED on 2-187 — same shared datapoint as 2-186. Allocated distances correctly use the post-correction BOL inbound (80,610 kg) as denominator (allocation fraction 0.26964556 = 21,736.13 / 80,610). All 6 transport components math-reconcile independently to Certify values within rounding. [log:certify_components.components_detail, certify_calculation_view.datapoints.Emissions Factor - Vehicle Embodied Emissions -- Railcar, independent_calculations]',
     'Certify Components'),
    ('9.3', 'ACTIVITIES - Injection: LCS and sample transport', 'N/A',
     'AECN does not require LCS pre-treatment. Sample transport not separately accounted (typical for AECN rail batches). [log:oil_type]',
     'N/A'),
    ('9.4', 'ACTIVITIES - Pyrolysis: process + embodied entries present', 'Pass',
     'Both "AECN pyrolysis (from rail loads) - process emissions" (2.261 t) and "AECN pyrolysis (from rail loads) - embodied emissions" (0.460 t) present. Rail-batch naming convention used (deliberate signal to VVB to expect scale tickets as authoritative mass source). [log:certify_components.components_detail]',
     'Certify Components'),
    ('9.5', 'REDUCTIONS: applicable BCU categories populated', 'Pass',
     '"BCUs - Applied to Bio-Oil Transport Process Emissions" (1.612 t) is the only applicable category for typical AECN rail batches without site-emission burdens. Other categories (injection diesel, injection waste, pre-treat diesel, pre-treat waste) appear only for site-emission-burdened removals. [log:certify_components.components_detail.BCUs - Applied to Bio-Oil Transport Process Emissions]',
     'Certify Components'),
    ('9.6', 'REDUCTIONS - BCU quantity >= corresponding activity emission', 'Pass',
     'BCU 1.612 t = exact sum of 3 transport process activities (0.267 + 1.238 + 0.107 = 1.612). BCU >= activity (here exactly equal). Embodied transport emissions (0.021 + 0.018 + 0.010 t) are NOT BCU-eligible per protocol — confirmed not double-claimed. [log:independent_calculations.BCU_offset_t, certify_components.components_detail]',
     'Bio-oil Geol Storage v1.1'),
    ('9.7', 'COUNTERFACTUALS: zero per protocol', 'Pass',
     'Total removal counterfactuals amount = 0 tCO2e (Components tab > Removal counterfactuals section). Correct for bio-oil geological storage. [log:certify_components.counterfactuals_t]',
     'Bio-oil Geol Storage v1.1'),
    ('9.8', 'Net removal matches internal calculation', 'Pass',
     'Certify Net = 29.476 t (displayed 29.48 t) = 32.415 - 4.382 + 1.612 - 0.170 ✓. Each input verified independently. [log:certify_components.net_co2e_t, certify_components.displayed_net_co2e_rounded_t, independent_calculations.net_t]',
     'Certify Calculation View'),
    ('9.9', 'All Datapoints have descriptive source names', 'Flag',
     '4 non-descriptive filenames in Certify Sources (same shared rail-batch datapoints as sister batches 2-185, 2-186): "Screenshot 2026-04-10 at 3.19.33 PM.png" (Truck Transport EF source), "Screenshot 2026-02-23 at 11.29.10 AM.png" (Pyrolysis Process EF support), "Screenshot 2026-01-27 125934.png" (Opelousas-Basco distance), "image (44).png" (SOPOR-Opelousas rail Watco distance email). Hygiene issue — content correct but filenames non-descriptive. **Additionally:** "Billing Truck 2314-T Corrections (1).pdf" has "-T" suffix where COBB lists "2314-R" — likely a filename typo (rail-load BOLs use "-R" suffix; this is the AECN BOL correction document referenced on sister batch 2-186 flag). Same flag as 2-185/2-186 — propagates because these are SHARED rail-batch datapoint sources. [log:certify_sources_tab.12, certify_sources_tab.13, certify_sources_tab.16, certify_sources_tab.19, certify_sources_tab.20]',
     'Certify Sources'),
    ('9.10', 'All Sources match files in batch folder on Drive', 'Pass',
     'Sources tab shows 23 source rows. All required Appendix 1 docs present in Certify Sources: 3 inbound BOLs (2314-R, 2316-R, 2317-R), AECN LCA CI spreadsheet, LCA with stack emissions, Supporting GHG data, AECN Process Emissions Calculation, Standard EFs Railcar Embodied, Charm-A5 BCU certificate, scale tickets, CHN PDF, Ops Notes, route maps, Truck Embodied EF source. Per Document Location Rule, presence in Certify Sources is sufficient. [log:certify_sources_tab_total, certify_sources_tab]',
     'Certify Sources'),
]


def build_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'QA Checklist'

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 100
    ws.column_dimensions['E'].width = 30

    row = 1

    for entry in ROWS:
        kind = entry[0]
        if kind == 'HEADER':
            ws.cell(row, 1, 'Charm Industrial — AECN Injection Batch QA Checklist')
            ws.cell(row, 1).font = Font(bold=True, size=14)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            ws.cell(row, 1).alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[row].height = 22
            row += 1
        elif kind == 'META':
            ws.cell(row, 1, entry[1]).font = Font(bold=True)
            ws.cell(row, 2, entry[2])
            row += 1
        elif kind == 'SPACER':
            row += 1
        elif kind == 'COLUMNS':
            headers = ['Item #', 'Check Item', 'Status', 'Notes / Evidence', 'Verifier Reference']
            for c, h in enumerate(headers, start=1):
                cell = ws.cell(row, c, h)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = WRAP_CENTER
                cell.border = BORDER
            ws.row_dimensions[row].height = 28
            row += 1
        elif kind == 'SECTION':
            ws.cell(row, 1, entry[1]).font = SECTION_FONT
            ws.cell(row, 1).fill = SECTION_FILL
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            ws.cell(row, 1).alignment = Alignment(horizontal='left', vertical='center', indent=1)
            ws.row_dimensions[row].height = 20
            row += 1
        else:
            item, check, status, evidence, verifier = entry
            ws.cell(row, 1, item)
            ws.cell(row, 2, check)
            ws.cell(row, 3, status)
            ws.cell(row, 4, evidence)
            ws.cell(row, 5, verifier)

            for c in range(1, 6):
                cell = ws.cell(row, c)
                cell.border = BORDER
                cell.alignment = WRAP

            if status == 'Fail':
                for c in range(1, 6):
                    ws.cell(row, c).fill = FAIL_FILL
                    ws.cell(row, c).font = FAIL_FONT
            elif status == 'Flag':
                for c in range(1, 6):
                    ws.cell(row, c).fill = FLAG_FILL
                    ws.cell(row, c).font = FLAG_FONT
            elif status == 'N/A':
                for c in range(1, 6):
                    ws.cell(row, c).font = NA_FONT

            ev_len = len(evidence or '')
            est_lines = max(1, ev_len // 90 + evidence.count('\n'))
            ws.row_dimensions[row].height = max(20, min(est_lines * 14, 240))

            row += 1

    row += 1
    pass_count = sum(1 for e in ROWS if len(e) >= 3 and e[0] not in ('HEADER','META','SPACER','COLUMNS','SECTION') and e[2] == 'Pass')
    fail_count = sum(1 for e in ROWS if len(e) >= 3 and e[0] not in ('HEADER','META','SPACER','COLUMNS','SECTION') and e[2] == 'Fail')
    flag_count = sum(1 for e in ROWS if len(e) >= 3 and e[0] not in ('HEADER','META','SPACER','COLUMNS','SECTION') and e[2] == 'Flag')
    na_count = sum(1 for e in ROWS if len(e) >= 3 and e[0] not in ('HEADER','META','SPACER','COLUMNS','SECTION') and e[2] == 'N/A')
    ws.cell(row, 1, 'SUMMARY').font = Font(bold=True, size=12)
    row += 1
    ws.cell(row, 1, 'Pass'); ws.cell(row, 2, pass_count)
    row += 1
    ws.cell(row, 1, 'Fail'); ws.cell(row, 2, fail_count)
    ws.cell(row, 1).fill = FAIL_FILL; ws.cell(row, 2).fill = FAIL_FILL
    row += 1
    ws.cell(row, 1, 'Flag'); ws.cell(row, 2, flag_count)
    ws.cell(row, 1).fill = FLAG_FILL; ws.cell(row, 2).fill = FLAG_FILL
    row += 1
    ws.cell(row, 1, 'N/A'); ws.cell(row, 2, na_count)
    row += 2
    ws.cell(row, 1, 'Net CDR (Certify, displayed)').font = Font(bold=True); ws.cell(row, 2, '29.48 tCO2e')
    row += 1
    ws.cell(row, 1, 'Material findings').font = Font(bold=True); ws.cell(row, 2, 'NONE — Net CDR 29.48 tCO2e is correctly calculated. 2-187 is the 2nd offload of railcar TIMX201060 (consumed_lot offload_2 of 4); raw-scale-ticket-driven and reconciles cleanly across scale tickets, Ops Notes, COBB offloaded_quantity_kg, and Certify Allocated Bio-Oil Mass (all = 21,736.13 kg). Independent gross-to-net calc reproduces Certify exactly: 21.456 t × 0.412 × 3.667 = 32.415 t Sequestrations; 4.382 t Activities; 1.612 t BCU offset; 0.170 t uncertainty discount → Net 29.476 t (displayed 29.48). Section 8.3 last-offload precondition check is DEFERRED to the railcar\'s final batch QA per FC9.')
    row += 1
    ws.cell(row, 1, 'Hygiene flags').font = Font(bold=True); ws.cell(row, 2, '5.4 (COBB-side: row 64 origin_quantity_kg=71,010 stale [pre-correction; corrected = 80,610 per Ops Notes], AND pH field shows "33" instead of "3.3" [typo; Ops Notes correctly shows 3.3] — both Garrett-side data-provenance, verifier-irrelevant); 9.9 (4 non-descriptive Certify filenames [image (44).png, Screenshot 2026-04-10 PM.png, Screenshot 2026-02-23 AM.png, Screenshot 2026-01-27 125934.png] propagating across rail batches, plus "Billing Truck 2314-T Corrections (1).pdf" with "-T" suffix where COBB lists "2314-R")')
    print(f'Pass={pass_count}, Fail={fail_count}, Flag={flag_count}, N/A={na_count}')

    out_path = '/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/qa/batch-2-187/Batch_2-187_QA_Checklist.xlsx'
    wb.save(out_path)
    print(f'Saved {out_path}')


if __name__ == '__main__':
    build_workbook()
