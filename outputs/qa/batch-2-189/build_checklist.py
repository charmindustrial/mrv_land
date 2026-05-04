#!/usr/bin/env python3
"""Build Batch 2-189 QA Checklist xlsx with FAIL/FLAG/N/A formatting and [log:KEY] citations."""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

FAIL_FILL = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
FAIL_FONT = Font(color='9C0006', bold=True)
FLAG_FILL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
FLAG_FONT = Font(color='996600')
NA_FONT = Font(color='808080', italic=True)
HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=12)
SECTION_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
SECTION_FONT = Font(bold=True, size=11)

THIN = Side(border_style='thin', color='808080')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='left')
WRAP_CENTER = Alignment(wrap_text=True, vertical='center', horizontal='center')

ROWS = [
    ('HEADER',),
    ('META', 'Batch ID', '2-189'),
    ('META', 'Reporting Period', 'Apr 2026 (Certify removal date 03 Apr 2026)'),
    ('META', 'Reviewer', 'Claude QA Agent'),
    ('META', 'Date', '2026-05-01'),
    ('META', 'Injection Site', 'Basco 6'),
    ('META', 'Oil Type', 'AECN'),
    ('META', 'Delivery Mode', 'Rail (railcar TIMX201060, consumed_lot offload_3 of 4 — 3rd physical pull, NOT the last offload of the railcar; sister batches 2-186=offload_1, 2-187=offload_2, 2-188=offload_4 [last])'),
    ('META', 'Net CDR (Certify)', '29.64 tCO2e'),
    ('SPACER',),
    ('COLUMNS',),

    ('SECTION', '1. Batch Folder Completeness'),
    ('1.1', 'Batch folder exists with correct naming', 'Pass',
     'Drive folder "2-189" at id 18atbX2cP4OyxvIHO7jZ7GryvHRFG226Y; contains "2-189 Ops Notes Rail" sheet, Testing subfolder (SR26-02264.pdf), Scale Tickets subfolder (full + empty JPGs). [log:drive_folder_url]',
     'VR Appendix 1'),
    ('1.2', 'Scale Tickets subfolder present and populated', 'Pass',
     'Drive > Scale Tickets contains 2-189 Full Scale Ticket.jpg (2.4 MB) and 2-189 Empty Scale Ticket.jpg (3.2 MB). One offload pair (consistent with 2-189 = consumed_lot offload_3 only; sister batches each contribute their own ticket pairs). Both also uploaded to Certify Sources as "(1)" copies. [log:scale_tickets.full_ticket.filename, scale_tickets.empty_ticket.filename]',
     'VR Appendix 1, Sec 5.1'),
    ('1.3', 'Testing subfolder present with required lab results', 'Pass',
     'Drive > Testing contains SR26-02264.pdf (7-page SGS Certificate of Analysis covering batches 2-186 through 2-189 + 3-55 through 3-57). Page 4 covers batch 2-189 (Certificate SR26-02264.004, Client ID = "BATCH 2-189"). [log:lab_results.certificate_number, lab_results.page, lab_results.batch_id_on_certificate]',
     'VR Appendix 1'),
    ('1.4', 'Ops Notes spreadsheet present and complete', 'Pass',
     '"2-189 Ops Notes Rail" Google Sheet (id 18wiHZwNa-fyg4b24xxroRAEgciDXoqVeTnLPSIzr6BM) present in Drive folder; owner Garrett Lutz; modified 2026-05-01 09:04. Batch Composition (railcar TIMX201060, sparging 0.0129), BOLs Loaded (3 inbound BOLs totaling 80.61 MT corrected), and offload table shows offload_1 + offload_2 + offload_3 rows populated as a running railcar tally; offload_3 is THIS batch (Cargo 21,799.63152 kg, allocation fraction 0.27043334). offload_4 row pending (sister batch 2-188 last-offload reconciliation). [log:ops_notes.filename, ops_notes.modified_at, ops_notes.offload_table.2.associated_batch_id]',
     'VR Appendix 1'),
    ('1.5', 'Injection site confirmed Basco 6, no Vaulted refs', 'Pass',
     'CHN PDF page 4 LOCATION = "BASCO LOUISIANA"; Certify project = "Charm Bio-oil Geologic Storage" (Basco). Ops Notes references Opelousas->Basco truck leg. No Vaulted references in any source. [log:lab_results.location_on_cert]',
     'PDD'),

    ('SECTION', '2. Scale Tickets & Mass Determination'),
    ('2.1', 'Empty and full scale ticket images present', 'Pass',
     'Both JPG images Read directly from pixels (NOT Drive OCR snippet, per OCR-vs-image rule). Full ticket #1312226091779 dated 4-01-26 13:16 at PILOT TRAVEL CENTER I-49 EXIT 40 BUNKIE LA scale #3122, axle sum 10,000 + 33,460 + 33,660 = 77,120 lb. Empty ticket #1312226091785 dated 4-01-26 15:24 same location, axle sum 9,380 + 11,500 + 8,180 = 29,060 lb. Both internally consistent (gross = axle sum, no axle-vs-cert discrepancy). [log:scale_tickets.full_ticket, scale_tickets.empty_ticket]',
     'VR Sec 5.1'),
    ('2.2', 'Dates within 2 weeks of injection', 'Pass',
     'Both tickets 2026-04-01; injection completion 2026-04-03. 2 days. [log:scale_tickets.full_ticket.date, scale_tickets.empty_ticket.date, ops_notes.completion_date]',
     'VR Sec 5.1'),
    ('2.3', 'Tractor and trailer numbers match between empty and full', 'Pass',
     'Full ticket: Tractor CH2, Trailer KL-188. Empty ticket: Tractor CH2, Trailer KL-188. Same vehicle weighed at both events. [log:scale_tickets.full_ticket.tractor, scale_tickets.empty_ticket.tractor, scale_tickets.tractor_trailer_match]',
     'VR Sec 5.1'),
    ('2.4', 'Net mass correctly calculated in Ops Notes', 'Pass',
     '2-189 is consumed_lot offload_3 (3rd pull), NOT the last offload of railcar TIMX201060 — the rail last-offload conservatism policy does NOT apply here. Cargo (kg) is raw scale-ticket math: (77,120 - 29,060) lb × 0.453592 = 48,060 lb × 0.453592 = 21,799.633 kg. Ops Notes offload_3 row Cargo Mass (kg) = 21,799.63152 — exact match. [log:scale_tickets.computed_net_lb, scale_tickets.computed_net_kg, ops_notes.batch_composition.lbs_to_kg_factor, ops_notes.offload_table.2.cargo_mass_kg]',
     'VR Sec 5.1'),
    ('2.5', 'Injectate mass uses conservative MIN(scale net, BOL mass)', 'Pass',
     'Per-offload rail batch (non-last): scale-ticket math is authoritative for THIS batch; BOL aggregate (80,610 kg corrected) is at railcar level, not per-offload. The 21,799.63 kg from scale tickets is 2-189\'s injectate basis. Sparging deduction (next item) applied to that base. [log:scale_tickets.computed_net_kg, ops_notes.offload_table.2.cargo_mass_kg]',
     'VR Sec 5.1, GHG Stmt Sec C'),
    ('2.6', 'Sparging deduction applied correctly', 'Pass',
     'Sparging factor 0.0129 applied: 21,799.63152 × 0.0129 = 281.2152466 kg sparging loss. Mass of Bio-Oil Injected = 21,799.63152 - 281.2152466 = 21,518.41627 kg. Ops Notes shows exactly these values; arithmetic correct. [log:ops_notes.batch_composition.sparging_mass_loss_rate, ops_notes.offload_table.2.cargo_mass_kg, ops_notes.offload_table.2.sparging_loss_kg, ops_notes.offload_table.2.mass_of_bio_oil_injected_kg]',
     'VR Sec 5.1'),
    ('2.7', 'Mass values consistent across scale tickets, Ops Notes, Certify', 'Pass',
     'Scale ticket net = 21,799.63 kg = Ops Notes Cargo (kg) = COBB offloaded_quantity_kg = Certify "Allocated Bio-Oil Mass - AECN via Rail" = 21,799.63152 kg. Post-sparging Mass of Injectant (Ops Notes) = 21,518.41627 kg shown in Certify Datapoints as 21.518 t. All four sources align. [log:scale_tickets.computed_net_kg, ops_notes.offload_table.2.cargo_mass_kg, ops_notes.offload_table.2.mass_of_bio_oil_injected_kg, cobb_2-189_row.offloaded_quantity_kg, certify_datapoints.list]',
     'VR Sec 5.1'),

    ('SECTION', '3. Bills of Lading & Transportation'),
    ('3.1', 'Truck BOLs present for each shipment', 'Pass',
     'All 3 inbound truck BOLs present in Certify Sources: Documents camion 2317-R.pdf, Billing Truck 2316-R.pdf, Billing Truck 2314-T Corrections (1).pdf. Per Document Location Rule, presence in Certify Sources is sufficient. Filename "-T" inconsistency on the corrections PDF — see Section 9.9. [log:certify_sources_full_enumeration]',
     'VR Appendix 1'),
    ('3.2', 'Rail batch: 3 truck BOLs from AECN, destination = SOPOR', 'Pass',
     '3 BOLs (2314-R, 2316-R, 2317-R) listed in COBB row 66 BOL_number field for batch 2-189 and in Ops Notes BOLs Loaded section. Destinations = SOPOR Quebec railyard (rail-batch convention). [log:ops_notes.bols_loaded_section, cobb_2-189_row.bol_number]',
     'VR Appendix 1'),
    ('3.3', 'BOL mass >= reported injectate mass at railcar level', 'Pass',
     'Sum of 3 BOL masses (26.05 + 29.06 + 25.50) = 80.61 MT = 80,610 kg (Ops Notes Total Loaded Mass). This is the AECN-CORRECTED inbound total (original BOL aggregate was 71,010 kg; AECN sent corrected BOL adding 9.6 MT, documented as "Billing Truck 2314-T Corrections (1).pdf"). This batch\'s offload mass = 21,799.63 kg, well below 80,610 kg railcar inbound. Rail-batch BOL mass check applied at railcar level. [log:bol_billing_documents, ops_notes.total_loaded_mass_mt, ops_notes.offload_table.2.cargo_mass_kg]',
     'Transport Emissions Acctg v1.1'),
    ('3.4', 'Origin and destination correct on BOLs', 'Pass',
     'Carrier route: AECN Quebec -> SOPOR (truck, 3 BOLs labeled with "-R" rail-load suffix) -> Opelousas LA (rail leg via railcar TIMX201060) -> Basco 6 (truck, this batch consumed_lot offload_3). Three legs documented across Drive + Certify Sources. [log:cobb_2-189_row.bol_number, ops_notes.bols_loaded_section, certify_components]',
     'VR Sec 3.5'),
    ('3.5', 'Transportation distances consistent with known AECN rail route', 'Pass',
     'Three legs verified: (1) AECN->SOPOR truck full single distance 173 km (= 107.5 mi) per Certify Datapoint "Distance from AECN to Sopor Quebec (truck)", and Allocated Distance AECN to Sopor Quebec 140.35 km (87.21 mi) computed as 322.5 mi total trucking × 0.27043334 batch share; (2) SOPOR->Opelousas rail full 2,199 mi per "Distance from SOPOR to Opelousas (RAIL)", Allocated 594.68 mi (2,199 × 0.27043334); (3) Opelousas->Basco truck per-batch 42.9 mi (full per-batch, no allocation since each offload uses dedicated truck). All match Ops Notes pre-computed values for offload_3 row. [log:certify_datapoints.list, ops_notes.batch_composition, ops_notes.offload_table.2.upstream_allocation, ops_notes.offload_table.2.rail_miles_allocated, ops_notes.offload_table.2.truck_leg_1_miles_allocated]',
     'VR Sec 3.5'),
    ('3.6', 'Transport mode = 6 components (3 process + 3 embodied) for rail', 'Pass',
     '6 transport components in Certify: AECN to SOPOR Tanker Truck Process (0.268 t), SOPOR to Opelousas Rail Process (1.242 t), Opelousas to Basco Tanker Truck Process (0.107 t), AECN to SOPOR Tanker Truck Embodied (0.021 t), SOPOR to Opelousas Rail Embodied (0.019 t), Opelousas to Basco Tanker Truck Embodied (0.010 t). [log:certify_components]',
     'PDD'),
    ('3.7', 'Billing documents present and consistent with BOLs', 'Pass',
     'Documents camion 2317-R.pdf, Billing Truck 2316-R.pdf, Billing Truck 2314-T Corrections (1).pdf in Certify Sources serve as both BOLs and billing for the 3 inbound truck legs (AECN convention). Charm-A3-2024-10-24-64-1976.pdf (BCU certificate) also present. [log:certify_sources_full_enumeration]',
     'VR Appendix 1'),

    ('SECTION', '4. Bio-Oil Composition & Testing'),
    ('4.1', 'CHN PDF present in Testing subfolder', 'Pass',
     'SR26-02264.pdf (7-page SGS Certificate, in Drive Testing) AND CHN-186-189.pdf (Certify Sources alias). Page 4 covers batch 2-189 (Certificate SR26-02264.004, "BATCH 2-189"). [log:lab_results.filename, lab_results.certify_filename_alias, lab_results.page]',
     'VR Appendix 1'),
    ('4.2', 'Carbon content within expected AECN range', 'Pass',
     'Carbon content for batch 2-189 = 41.3% (m/m) (ASTM D5291 Method C). Within AECN expected range (mean 39.42%, ±2 SD ~27.2%-51.6%). [log:lab_results.carbon_pct, lab_results.carbon_within_aecn_normal_range]',
     'Biomass Feedstock Acctg v1.2'),
    ('4.3', 'Hydrogen and nitrogen values reported and reasonable', 'Pass',
     'Hydrogen 6.8%, Nitrogen <0.75%, Viscosity 19.46 cSt @ 40°C, Acid Number 45 mg KOH/g, Density 1.211 g/cm³ @ 15°C. All within typical AECN ranges. [log:lab_results.hydrogen_pct, lab_results.nitrogen_pct, lab_results.viscosity_cSt_at_40C, lab_results.acid_number_mg_KOH_per_g, lab_results.density_g_per_cm3_at_15C]',
     'Biomass Feedstock Acctg v1.2'),
    ('4.4', 'Lab report references correct batch ID', 'Pass',
     'Certificate SR26-02264.004 explicitly labeled "BATCH 2-189" on page 4. Sampled 2026-04-02 23:28 at BASCO LOUISIANA (after offload_3 truck arrived 2026-04-01). Sample ~1 day before injection completion 2026-04-03. [log:lab_results.batch_id_on_certificate, lab_results.sampled_datetime, lab_results.location_on_cert]',
     'VR Appendix 1'),
    ('4.5', 'pH/conductivity (aqueous phase)', 'N/A',
     'AECN is not aqueous phase. pH of injectate is recorded in Ops Notes Section 5.5. [log:oil_type]',
     'N/A'),

    ('SECTION', '5. Ops Notes Accuracy & Internal Consistency'),
    ('5.1', 'Correct batch ID and injection date', 'Pass',
     'Ops Notes Rail header shows Batch number = 2-189, Completion Date = 4/3/2026. Matches Certify removal date (03 Apr 2026) and COBB injection_completion_date (2026-04-03). [log:ops_notes.batch_number_on_sheet, ops_notes.completion_date, cobb_2-189_row.injection_completion_date]',
     'VR Sec 5.1'),
    ('5.2', 'Scale ticket values correctly transcribed to Ops Notes', 'Pass',
     'Full ticket gross 77,120 lb = Ops Notes offload_3 Loaded Mass (lbs) 77,120 (exact). Empty ticket gross 29,060 lb = Ops Notes offload_3 Unloaded Mass (lbs) 29,060 (exact). Both values verified by direct image Read of JPGs (per Drive-OCR-not-primary-source rule). Cargo (kg) using Ops Notes factor 0.453592: 48,060 × 0.453592 = 21,799.63152 = Ops Notes offload_3 Cargo Mass (kg) 21799.63152 (exact). [log:scale_tickets.full_ticket.certified_gross_lb, scale_tickets.empty_ticket.certified_gross_lb, ops_notes.batch_composition.lbs_to_kg_factor, ops_notes.offload_table.2.loaded_mass_lb, ops_notes.offload_table.2.unloaded_mass_lb, ops_notes.offload_table.2.cargo_mass_kg]',
     'VR Sec 5.1'),
    ('5.3', 'Oil and aqueous phase mass reconciliation', 'N/A',
     'AECN-only batch; no aqueous phase. [log:oil_type]',
     'N/A'),
    ('5.4', 'COBB values present and match Ops Notes', 'Flag',
     'COBB row 66 for "2-189" verified: cell A = "2-189" ✓, consumed_lot = "Rail_TIMX201060_2.27.26_redo_offload_3" ✓, BOL_number = "2314-R, 2316-R, 2317-R" matches Ops Notes BOLs Loaded ✓, offloaded_quantity_kg = 21,799.63152 = Ops Notes offload_3 Cargo (kg) ✓, injection_completion_date = 2026-04-03 = Ops Notes 4/3/2026 ✓. **Single COBB-side data-provenance issue for Garrett (NOT 2-189-impacting):** COBB row 66 origin_quantity_kg = 71,010 (stale pre-correction). Post-correction inbound aggregate = 80,610 kg per Ops Notes Total Loaded Mass — the 9.6 MT AECN BOL correction first surfaced on sister batch 2-186. Same flag carried by sister rows 63 (2-186) and 64 (2-187). Operational data-provenance for Garrett to clean up — verifier reads Ops Notes / Certify (both correct). [log:cobb_row_id_verified, cobb_cell_A_value, cobb_2-189_row, cobb_sister_rows, ops_notes.total_loaded_mass_mt, bol_billing_documents.total_bol_aggregate_kg]',
     'VR Sec 5.1'),
    ('5.5', 'pH measurement noted', 'Pass',
     'Ops Notes offload_3 row pH = 3.0. COBB row 66 pH = "3" (matches). [log:ops_notes.offload_table.2.ph, cobb_2-189_row.injection_ph]',
     'VR Sec 5.1'),
    ('5.6', 'Completion date present and matches COBB', 'Pass',
     'All three sources align: Ops Notes Completion Date 4/3/2026, COBB row 66 injection_completion_date 2026-04-03, Certify removal date 03 Apr 2026. [log:ops_notes.completion_date, cobb_2-189_row.injection_completion_date, certify_components]',
     'VR Sec 5.1'),
    ('5.7', 'No erroneous density measurements', 'N/A',
     'Density field expected blank for AECN per protocol. Ops Notes Batch Composition does not include density entry — expected behavior. [log:oil_type]',
     'N/A'),

    ('SECTION', '6. AECN Feedstock Emissions & Allocation'),
    ('6.1', 'AECN process emissions LCA CI spreadsheet present', 'Pass',
     '"Bioenergy AE Cote-Nord Canada - LCA CI calculation - 1Jul2025 to 31Dec2025 S.xlsx" present in Certify Sources, linked to Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail) datapoint. Vintage covers H2 2025 (correct for batch sampled in Apr 2026 since EFs are H1/H2 — H2 2025 carries through Q1 2026 reporting). [log:certify_sources_full_enumeration]',
     'VR Appendix 1'),
    ('6.2', 'LCA with stack emissions file present', 'Pass',
     '"BH 2025 LCA w Stack Emissions.xlsx" present in Certify Sources, linked to Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail). [log:certify_sources_full_enumeration]',
     'VR Appendix 1'),
    ('6.3', 'Supporting GHG data file present', 'Pass',
     '"Supporting GHG data 31Dec2025 S.xlsx" present in Certify Sources, linked to Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail). [log:certify_sources_full_enumeration]',
     'VR Appendix 1'),
    ('6.4', 'AECN pyrolysis process EF matches current value', 'Pass',
     'Modal value 0.104 tCO2e/t (display) for "Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail)". KB current native value 0.10402 tCO2e/t (H2 2025) — within rounding. Independent calc: 21.79963 t × 0.10402 = 2.2676 ≈ 2.267 tCO2e ✓ matches Certify "AECN pyrolysis (from rail loads) - process emissions" 2.267 t. [log:certify_components.AECN_pyrolysis_from_rail_loads_process, ef_verification.aecn_pyrolysis_process]',
     'Standard EFs 2026 tab'),
    ('6.5', 'AECN pyrolysis embodied EF applied', 'Pass',
     'Modal value 0.02116 tCO2e/t for "Carbon Intensity - AECN Pyrolysis Embodied Emissions (via Rail)". KB current value 0.0212 tCO2e/t (precise per Certify modal at 0.02116). Independent calc: 21.79963 t × 0.02116 = 0.4613 ≈ 0.461 tCO2e ✓ matches Certify "AECN pyrolysis (from rail loads) - embodied emissions" 0.461 t. [log:certify_components.AECN_pyrolysis_from_rail_loads_embodied, ef_verification.aecn_pyrolysis_embodied]',
     'Standard EFs 2026 tab'),
    ('6.6', 'Mass basis for pyrolysis EFs is Origin Bio-Oil Mass (pre-sparging)', 'Pass',
     'Pyrolysis components apply 0.104 / 0.02116 to "Allocated Bio-Oil Mass - AECN via Rail" = 21,799.63 kg = pre-sparging Cargo (kg). Sparging is applied at injection (Section 2.6) and only affects Mass of Injectant for sequestration. Methodology correct. [log:certify_datapoints.list, certify_components.AECN_pyrolysis_from_rail_loads_process]',
     'Bio-oil Geol Storage v1.1'),

    ('SECTION', '7. Gross-to-Net Calculation'),
    ('7.1', 'Gross CO2e = injectate mass x carbon wt% (decimal) x 3.667', 'Pass',
     'Independent calc: 21.51842 t (Mass of Injectant, post-sparging) × 0.413 (Carbon Content from CHN page 4) × 3.667 = 32.590 tCO2e. Certify Sequestrations 32.589 tCO2e — match within rounding. [log:certify_components.Injection_sequestration, lab_results.carbon_pct, ops_notes.offload_table.2.mass_of_bio_oil_injected_kg]',
     'GHG Stmt Sec E'),
    ('7.2', 'Process emissions correctly deducted', 'Pass',
     'All 8 activity components deducted: 3 transport process (AECN->SOPOR truck 0.268, SOPOR->Opelousas rail 1.242, Opelousas->Basco truck 0.107) + 3 transport embodied (0.021 + 0.019 + 0.010) + AECN pyrolysis process 2.267 + AECN pyrolysis embodied 0.461. Sum = 4.394 tCO2e ✓ matches Certify Activities total. Each component math reconciles independently. [log:certify_components, certify_calculation_view.activities_t]',
     'GHG Stmt Sec E'),
    ('7.3', 'Uncertainty discount correctly applied', 'Pass',
     'Discount 171 kg present in Calculation View. Per Isometric Standard S3.7.3 Option B (Variance Propagation), Certify computes the aggregate from per-component ± inputs — Isometric\'s aggregation is authoritative; QA scope is verifying the INPUTS (covered in Section 7a Layer 1+2: ±72.57 kg per discrete-weight datapoint, ±0.4% carbon). 0.171 t is ~0.5% of gross sequestration (32.589 t), consistent with typical AECN rail batch ratios (sister 2-186 was 0.5%, 2-187 was 0.5%). [log:certify_calculation_view.uncertainty_discount_kg, certify_calculation_view.sequestrations_t]',
     'Isometric Standard v2.0 S3.7'),
    ('7.4', 'Net removals consistent with internal calc and BCU offset', 'Pass',
     'Net = 32.589 (sequestrations) + 1.617 (BCU adjustments) - 4.394 (activities) - 0.171 (uncertainty) = 29.641 tCO2e. Certify shows Net CO2e Removed = 29.641 t (displayed as 29.64). Reconciles exactly. [log:certify_calculation_view.net_tco2e, certify_calculation_view.independent_check]',
     'GHG Stmt Sec E'),

    ('SECTION', '7a. Per-Component Uncertainty Verification'),
    ('7a.1', 'Layer 1 - Completeness: every required component has uncertainty input', 'Pass',
     'All mass-input components show ±72.57 kg in INPUTS section (verified by opening each component modal individually): Injection Sequestration (Mass of product 21,518.42 ±72.57), AECN->SOPOR Truck Process (Mass of load 21,799.63 ±72.57), SOPOR->Opelousas Rail Process (Mass of load 21,799.63 ±72.57), Opelousas->Basco Truck Process (Mass of load 21,799.63 ±72.57), AECN Pyrolysis Process (Mass 21,799.63 ±72.57), AECN Pyrolysis Embodied (Mass 21,799.63 ±72.57). Embodied transport components correctly have NO Mass-of-load datapoint (per-mile EFs, not per-T-mi). Carbon Content uncertainty ±0.1652% (= 41.3 × 0.004 = 0.1652, matches 0.4% relative). [log:certify_components.Injection_sequestration, certify_components.AECN_to_SOPOR_Tanker_Truck_Process, certify_components.SOPOR_to_Opelousas_Rail_Process, certify_components.Opelousas_to_Basco_Tanker_Truck_Process, certify_components.AECN_pyrolysis_from_rail_loads_process, certify_components.AECN_pyrolysis_from_rail_loads_embodied]',
     'Isometric Standard S3.7.3'),
    ('7a.2', 'Layer 2 - Correctness: reported ± matches expected value', 'Pass',
     'Per FC9 per-removal framing: a per-batch offload removal is fixed by ONE outbound discrete weight at Basco offload, so ±72.57 kg per discrete weight is the correct value on every leg (truck process + rail process + truck process). Verified by opening each component modal and reading the INPUTS section directly. Carbon content ±0.1652% = 41.3 × 0.004 confirms standard 0.4% relative uncertainty. Conservatism precondition (N_offloads ≥ N_inbound_BOLs) for railcar TIMX201060 will be formally validated at the railcar\'s last batch QA (2-188 = consumed_lot offload_4, the railcar\'s physical last pull). [log:certify_components, ef_verification]',
     'KB Section 7a / FC9'),
    ('7a.3', 'Layer 3 - Discount presence in net calc and consistency', 'Pass',
     'Aggregate uncertainty discount 171 kg in Calculation View. Visible in net formula: 32.589 sequestrations + 1.617 BCU - 4.394 activities - 0.171 uncertainty = 29.641 t Net. Magnitude reasonable given mass and carbon uncertainty inputs. [log:certify_calculation_view.uncertainty_discount_kg, certify_calculation_view.net_tco2e]',
     'Isometric Standard S3.7.3'),

    ('SECTION', '8. Data Integrity & Anomaly Check'),
    ('8.1', 'No unexplained outliers in mass, carbon content, or calculated values', 'Pass',
     'Carbon (41.3%), distances (107.5/2199/42.9 mi full-leg + 87.21/594.68 mi allocated), EFs (0.1143 truck process / 0.0259 rail process / 0.23756 truck embodied / 0.03113 rail embodied 2026 corrected value — same shared datapoint as sister batches 2-186/2-187; 2-182\'s stale 0.0000352240 was previously updated and remains correct here). Mass 21,799.63 kg pre-sparging / 21,518.42 kg post-sparging within expected per-truck offload range. Both scale tickets internally consistent (gross = axle sum). All math reconciles to Certify within rounding. [log:lab_results.carbon_pct, certify_components, ef_verification]',
     'GHG Stmt Sec C'),
    ('8.2', 'Spills/losses documented; unrecovered mass deducted', 'Pass',
     'No spills/losses for this offload. Ops Notes "Oil Left in Railcar (MT/kg)" cells correctly blank for 2-189 — railcar TIMX201060 still has remaining oil for sister batch 2-188 (consumed_lot offload_4, the physical last pull). Last-offload mass reconciliation belongs to 2-188\'s QA. [log:ops_notes.oil_left_in_railcar_mt, ops_notes.oil_left_in_railcar_kg]',
     'GHG Stmt Sec C'),
    ('8.3', 'Sister-batch sanity check (rail batch precondition for FC9 framing)', 'Pass',
     'Railcar TIMX201060 is mid-drawdown at this batch. Sister-batch state per COBB (rows 63-66, all consumed_lot Rail_TIMX201060_2.27.26_redo_offload_*): 2-186 row 63 = offload_1, Cargo 21,854.06 kg (1st pull, raw); 2-187 row 64 = offload_2, Cargo 21,736.13 kg (2nd pull, raw); 2-189 row 66 = offload_3, Cargo 21,799.63 kg (THIS BATCH, 3rd pull, raw); 2-188 row 65 = offload_4 (the railcar\'s physical LAST pull), Cargo 5,620.18 kg currently in COBB (still pre-correction policy back-calc; under post-correction BOL 80,610 kg, 2-188\'s Cargo will reset to ~15,220 kg — pending Garrett update at 2-188 reconciliation). **Section 8.3 last-offload reconciliation DEFERRED to 2-188\'s QA** (consumed_lot offload_4 = the railcar\'s physical last pull). For 2-189, the FC9 per-offload uncertainty framing (±72.57 kg per leg) is conservative iff N_offloads ≥ N_inbound_BOLs at last batch — that precondition check (and the railcar mass balance check) will be formally validated at 2-188\'s QA. 2-189 is independent of last-offload mechanics: this is a non-last raw-scale-ticket pull. [log:cobb_sister_rows, ops_notes.railcar_balance_check]',
     'KB Section 8 / FC9 / FC11'),

    ('SECTION', '9. Isometric Certify - Removal Component Mapping'),
    ('9.1', 'SEQUESTRATIONS: Gross injection sequestration matches internal gross', 'Pass',
     'Certify Sequestration component 32.589 tCO2e. Independent calc: 21.51842 t (Mass of Injectant) × 0.413 × 3.667 = 32.590 t. Match within rounding. Carbon content matches CHN PDF page 4 (41.3%). [log:certify_components.Injection_sequestration, lab_results.carbon_pct]',
     'Certify Components'),
    ('9.2', 'ACTIVITIES - Transport: 6 components correct for AECN rail route', 'Pass',
     '6 transport components present with correct EFs and distances. Rail embodied EF = 0.03113 kgCO2e/mi (2026 corrected value 0.0000311307 MTCO2e/mi in native units, NOT stale 0.0000352240 carried by sister batch 2-182). Rail-embodied stale-EF finding from 2-182 remains RESOLVED on 2-189. Allocated distances correctly use the post-correction BOL inbound (80,610 kg) as denominator (allocation fraction 0.27043334 = 21,799.63 / 80,610). All 6 transport components math-reconcile independently to Certify values within rounding. [log:certify_components, ef_verification.rail_embodied]',
     'Certify Components'),
    ('9.3', 'ACTIVITIES - Injection: LCS and sample transport', 'N/A',
     'AECN does not require LCS pre-treatment. Sample transport not separately accounted (typical for AECN rail batches). [log:oil_type]',
     'N/A'),
    ('9.4', 'ACTIVITIES - Pyrolysis: process + embodied entries present', 'Pass',
     'Both "AECN pyrolysis (from rail loads) - process emissions" (2.267 t) and "AECN pyrolysis (from rail loads) - embodied emissions" (0.461 t) present. Rail-batch naming convention used (deliberate signal to VVB to expect scale tickets as authoritative mass source). [log:certify_components.AECN_pyrolysis_from_rail_loads_process, certify_components.AECN_pyrolysis_from_rail_loads_embodied]',
     'Certify Components'),
    ('9.5', 'REDUCTIONS: applicable BCU categories populated', 'Pass',
     '"BCUs - Applied to Bio-Oil Transport Process Emissions" (1.617 t) is the only applicable category for typical AECN rail batches without site-emission burdens. Other categories (injection diesel, injection waste, pre-treat diesel, pre-treat waste) appear only for site-emission-burdened removals. [log:certify_components.BCUs_Applied_to_Bio_Oil_Transport_Process]',
     'Certify Components'),
    ('9.6', 'REDUCTIONS - BCU quantity >= corresponding activity emission', 'Pass',
     'BCU 1.617 t = exact sum of 3 transport process activities (0.268 + 1.242 + 0.107 = 1.617). BCU >= activity (here exactly equal). Embodied transport emissions (0.021 + 0.019 + 0.010 t) are NOT BCU-eligible per protocol — confirmed not double-claimed. [log:certify_components.BCUs_Applied_to_Bio_Oil_Transport_Process, certify_components.AECN_to_SOPOR_Tanker_Truck_Process, certify_components.SOPOR_to_Opelousas_Rail_Process, certify_components.Opelousas_to_Basco_Tanker_Truck_Process]',
     'Bio-oil Geol Storage v1.1'),
    ('9.7', 'COUNTERFACTUALS: zero per protocol', 'Pass',
     'Total removal counterfactuals amount = 0 tCO2e (Components tab > Removal counterfactuals section). Correct for bio-oil geological storage. [log:certify_components]',
     'Bio-oil Geol Storage v1.1'),
    ('9.8', 'Net removal matches internal calculation', 'Pass',
     'Certify Net = 29.641 t (displayed 29.64 t) = 32.589 - 4.394 + 1.617 - 0.171 ✓. Each input verified independently. [log:certify_calculation_view.net_tco2e, certify_calculation_view.independent_check]',
     'Certify Calculation View'),
    ('9.9', 'All Datapoints have descriptive source names', 'Flag',
     '4 non-descriptive filenames in Certify Sources (same shared rail-batch datapoints as sister batches 2-185, 2-186, 2-187): "Screenshot 2026-04-10 at 3.19.33 PM.png" (Truck Transport EF source), "Screenshot 2026-02-23 at 11.29.10 AM.png" (Pyrolysis Process EF support), "Screenshot 2026-01-27 125934.png" (Opelousas-Basco distance), "image (44).png" (SOPOR-Opelousas rail Watco distance email). Hygiene issue — content correct but filenames non-descriptive. **Additionally:** "Billing Truck 2314-T Corrections (1).pdf" has "-T" suffix where COBB lists "2314-R" — likely a filename typo (rail-load BOLs use "-R" suffix; this is the AECN BOL correction document referenced on sister batch 2-186 flag). **Additionally:** the rail PROCESS EF datapoint is sourced from a PDF named "Standard Emission Factors + Calculations - 2026 Railcar Embodied Emissions.pdf" — same filename also linked to the rail Embodied EF datapoint. The EF value itself is correct per 2026 GLEC V3.2, but the filename references the Embodied tab specifically; if the PDF was an Embodied-tab-only export, this is a wrong source link. Worth Garrett verifying. Same flag as 2-185/2-186/2-187 — propagates because these are SHARED rail-batch datapoint sources. [log:certify_sources_full_enumeration, certify_components.SOPOR_to_Opelousas_Rail_Process.ef_kg_per_t_mi, ef_verification.rail_process]',
     'Certify Sources'),
    ('9.10', 'All Sources match files in batch folder on Drive', 'Pass',
     'Sources tab shows 23 source-datapoint links from 18 unique files. All required Appendix 1 docs present in Certify Sources: 3 inbound BOLs (2314-R, 2316-R, 2317-R), AECN LCA CI spreadsheet, LCA with stack emissions, Supporting GHG data, AECN Process Emissions Calculation, Standard EFs Railcar Embodied, Charm-A3 BCU certificate, scale tickets, CHN PDF, Ops Notes, route maps, Truck Embodied EF source. Per Document Location Rule, presence in Certify Sources is sufficient. [log:certify_sources_tab_total, certify_sources_tab_unique_files, certify_sources_full_enumeration]',
     'Certify Sources'),
]


def build_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'QA Checklist'

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 100
    ws.column_dimensions['E'].width = 30

    row = 1

    for entry in ROWS:
        kind = entry[0]
        if kind == 'HEADER':
            ws.cell(row, 1, 'Charm Industrial — AECN Injection Batch QA Checklist')
            ws.cell(row, 1).font = Font(bold=True, size=14)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            ws.cell(row, 1).alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[row].height = 22
            row += 1
        elif kind == 'META':
            ws.cell(row, 1, entry[1]).font = Font(bold=True)
            ws.cell(row, 2, entry[2])
            row += 1
        elif kind == 'SPACER':
            row += 1
        elif kind == 'COLUMNS':
            headers = ['Item #', 'Check Item', 'Status', 'Notes / Evidence', 'Verifier Reference']
            for c, h in enumerate(headers, start=1):
                cell = ws.cell(row, c, h)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = WRAP_CENTER
                cell.border = BORDER
            ws.row_dimensions[row].height = 28
            row += 1
        elif kind == 'SECTION':
            ws.cell(row, 1, entry[1]).font = SECTION_FONT
            ws.cell(row, 1).fill = SECTION_FILL
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            ws.cell(row, 1).alignment = Alignment(horizontal='left', vertical='center', indent=1)
            ws.row_dimensions[row].height = 20
            row += 1
        else:
            item, check, status, evidence, verifier = entry
            ws.cell(row, 1, item)
            ws.cell(row, 2, check)
            ws.cell(row, 3, status)
            ws.cell(row, 4, evidence)
            ws.cell(row, 5, verifier)

            for c in range(1, 6):
                cell = ws.cell(row, c)
                cell.border = BORDER
                cell.alignment = WRAP

            if status == 'Fail':
                for c in range(1, 6):
                    ws.cell(row, c).fill = FAIL_FILL
                    ws.cell(row, c).font = FAIL_FONT
            elif status == 'Flag':
                for c in range(1, 6):
                    ws.cell(row, c).fill = FLAG_FILL
                    ws.cell(row, c).font = FLAG_FONT
            elif status == 'N/A':
                for c in range(1, 6):
                    ws.cell(row, c).font = NA_FONT

            ev_len = len(evidence or '')
            est_lines = max(1, ev_len // 90 + evidence.count('\n'))
            ws.row_dimensions[row].height = max(20, min(est_lines * 14, 240))

            row += 1

    row += 1
    pass_count = sum(1 for e in ROWS if len(e) >= 3 and e[0] not in ('HEADER','META','SPACER','COLUMNS','SECTION') and e[2] == 'Pass')
    fail_count = sum(1 for e in ROWS if len(e) >= 3 and e[0] not in ('HEADER','META','SPACER','COLUMNS','SECTION') and e[2] == 'Fail')
    flag_count = sum(1 for e in ROWS if len(e) >= 3 and e[0] not in ('HEADER','META','SPACER','COLUMNS','SECTION') and e[2] == 'Flag')
    na_count = sum(1 for e in ROWS if len(e) >= 3 and e[0] not in ('HEADER','META','SPACER','COLUMNS','SECTION') and e[2] == 'N/A')
    ws.cell(row, 1, 'SUMMARY').font = Font(bold=True, size=12)
    row += 1
    ws.cell(row, 1, 'Pass'); ws.cell(row, 2, pass_count)
    row += 1
    ws.cell(row, 1, 'Fail'); ws.cell(row, 2, fail_count)
    ws.cell(row, 1).fill = FAIL_FILL; ws.cell(row, 2).fill = FAIL_FILL
    row += 1
    ws.cell(row, 1, 'Flag'); ws.cell(row, 2, flag_count)
    ws.cell(row, 1).fill = FLAG_FILL; ws.cell(row, 2).fill = FLAG_FILL
    row += 1
    ws.cell(row, 1, 'N/A'); ws.cell(row, 2, na_count)
    row += 2
    ws.cell(row, 1, 'Net CDR (Certify, displayed)').font = Font(bold=True); ws.cell(row, 2, '29.64 tCO2e')
    row += 1
    ws.cell(row, 1, 'Material findings').font = Font(bold=True); ws.cell(row, 2, 'NONE — Net CDR 29.64 tCO2e is correctly calculated. 2-189 is the 3rd offload of railcar TIMX201060 (consumed_lot offload_3 of 4); raw-scale-ticket-driven and reconciles cleanly across scale tickets, Ops Notes, COBB offloaded_quantity_kg, and Certify Allocated Bio-Oil Mass (all = 21,799.63 kg). Independent gross-to-net calc reproduces Certify exactly: 21.51842 t × 0.413 × 3.667 = 32.590 t Sequestrations; 4.394 t Activities; 1.617 t BCU offset; 0.171 t uncertainty discount → Net 29.641 t (displayed 29.64). The railcar last-offload reconciliation belongs to 2-188 QA (consumed_lot offload_4, the railcar’s physical last pull). Section 7a per-offload uncertainty conservatism precondition check (N_offloads ≥ N_inbound_BOLs) is DEFERRED to 2-188’s QA per FC9.')
    row += 1
    ws.cell(row, 1, 'Hygiene flags').font = Font(bold=True); ws.cell(row, 2, '5.4 (COBB-side: row 66 origin_quantity_kg=71,010 stale [pre-correction; corrected = 80,610 per Ops Notes] — Garrett-side data-provenance, verifier-irrelevant, same flag as sister rows 63 [2-186] and 64 [2-187]); 9.9 (4 non-descriptive Certify filenames [image (44).png, Screenshot 2026-04-10 PM.png, Screenshot 2026-02-23 AM.png, Screenshot 2026-01-27 125934.png] propagating across rail batches, plus "Billing Truck 2314-T Corrections (1).pdf" with "-T" suffix where COBB lists "2314-R", plus rail PROCESS EF sourced from "Standard Emission Factors + Calculations - 2026 Railcar Embodied Emissions.pdf" — same filename also linked to rail Embodied EF, possible single-tab vs multi-tab PDF inconsistency)')
    print(f'Pass={pass_count}, Fail={fail_count}, Flag={flag_count}, N/A={na_count}')

    out_path = '/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/qa/batch-2-189/Batch_2-189_QA_Checklist.xlsx'
    wb.save(out_path)
    print(f'Saved {out_path}')


if __name__ == '__main__':
    build_workbook()
