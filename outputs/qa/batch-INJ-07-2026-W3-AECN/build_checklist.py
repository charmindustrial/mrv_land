#!/usr/bin/env python3
"""Build INJ-07-2026-W3 AECN weekly batch QA checklist."""
import sys, json
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2] / "skills" / "qa" / "batch-qa-aecn" / "scripts"))
from checklist_lib import build_with_validation  # noqa: E402

HEADER_FILL = PatternFill("solid", fgColor="D6E4F0")
FAIL_FILL   = PatternFill("solid", fgColor="FFC7CE")
FLAG_FILL   = PatternFill("solid", fgColor="FFF2CC")
FLAG_FONT   = Font(color="996600")
NA_FONT     = Font(color="808080", italic=True)
BOLD        = Font(bold=True)
WRAP        = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

D = json.loads((HERE / "_derived_preview.json").read_text())
D["carbon_pct"] = 40.2
D["carbon_2sd_low_pct"] = 27.2337
D["carbon_2sd_high_pct"] = 51.6070
D["carbon_z_score"] = 0.1279
D["hydrogen_pct"] = 7.4
D["hydrogen_2sd_low_pct"] = 5.7487
D["hydrogen_2sd_high_pct"] = 8.0001
D["density_g_cm3"] = 1.210
D["density_floor_g_cm3"] = 1.01
D["ph"] = 3.2
D["ph_target_low"] = 2.0
D["ph_target_high"] = 4.0
D["days_ticket_to_injection_max"] = 10
D["days_ticket_to_injection_min"] = 8
D["certify_sources_total"] = 31
D["n_bols_total"] = 6
D["n_railcars"] = 2
D["aecn_sopor_full_mi"] = 107.49722
D["aecn_sopor_km"] = 173
D["rail_full_mi"] = 2199
D["opel_basco_full_mi"] = 42.9
D["opel_basco_alt_mi"] = 40.9
D["ef_truck_fuel"] = 0.0001143
D["ef_rail_fuel"] = 2.5910374e-05
D["ef_truck_embodied"] = 0.0002375605888
D["ef_railcar_embodied"] = 3.113066508e-05
D["ef_railcar_embodied_stale"] = 3.5224e-05
D["tercs_retired"] = 2795
D["jun_map_distance_mi"] = 2275.569
D["jun_map_distance_km"] = 3662.173
D["charm_two_long_legs_mi"] = 2306.49722
D["bol_page_count_short"] = 3
D["bol_page_count_full"] = 5
D["shpx215079_prior_claimed_kg"] = 19513.52784
D["shpx215079_claimed_this_week_kg"] = 58749.23584
D["shpx215079_total_offloaded_kg"] = 78262.76368
D["gross_certify_t"] = 113.739
D["activities_certify_t"] = 15.838
D["adjustments_certify_t"] = 5.808

DERIVED = D

HDR = ("Reviewer: Claude QA Agent  |  Date: 2026-08-06  |  Injection Site: Basco 6  |  Oil Type: AECN (rail)  |  "
       "Ops week: Sat 2026-07-11 to Sat 2026-07-18  |  Injected: 2026-07-18  |  "
       "Certify GHG entry rmv_1KYV2NABZ1S0CZHJ (period 12-18 Jul 2026)  |  Net CDR 103.082 tCO2e")

SUMMARY = (
 "NO FAILS. The mass chain verifies end to end from the eight CAT Scale tickets through the Ops Notes and COBB to Certify, and "
 "every emission component, both allocated distances, the aggregate distance, every uncertainty input and the net arithmetic "
 "reconstruct to the reported figures from full-precision factors read off the 2026 emission factor sheet this session. "
 "All four offloads carry a matching tractor AND trailer across their full/empty ticket pair, so the reporting period's "
 "Trailer Standard Empty Weight substitution was not triggered and none was applied - consistent with the July VVB guidance, "
 "whose only application list covers Week 1. Neither railcar overdrew its inbound BOL aggregate, so the closing-offload "
 "back-calculation did not apply and all four Uncertainty Adjustment cells are correctly zero. Railcar SHPX215079 was fully "
 "drawn down this week across four offloads against three inbound BOLs, so the per-offload uncertainty framing remains "
 "conservative; its 67.24 kg shortfall against the BOL aggregate is a documented 0.0858 % loss that is not claimed. The BCU "
 "offset equals the transport process sum exactly, which is the error mode corrected on the Week 2 entry and is not repeated "
 "here. "
 "ONE ITEM WAS RAISED AND HAS BEEN RESOLVED: the Bio-Oil Mass uncertainty was entered as 290.98 kg where four discrete "
 "weighings give 290.28, and the Mass of Injectant datapoint on the same entry already carried the correct 290.28. The "
 "datapoint now reads 290.28 and both mass datapoints agree; re-read in Certify this session. Net and discount are unchanged "
 "at 103.082 t and 627 kg, the original gap being worth 0.005 kg of propagated uncertainty. "
 "FIVE FLAGS REMAIN, none of which changes a quantity. Two of the six BOL packages are short two pages - 2395-R and 2396-R are filed as "
 "three-page documents with no TSCA import certification and no purchase order, where the other four carry both. COBB labels "
 "this week INJ-20260712_bio_oil rather than INJ-07-2026-W3, so a verifier tracing by batch ID finds nothing. The pyrolysis "
 "process factor is stored at 0.104 against the sheet's final quantity of 0.1040167408. "
 "Two source files carry generic names, one of them the sole evidence for the 2,199-mile rail leg. And the "
 "purchase order bound into the four complete BOL packages records the order two ways and carries a Kansas ship-to on a "
 "Louisiana injection."
 )

ROWS = [
 ("INJ-07-2026-W3 - AECN Weekly Batch QA", HDR, "SECTION", "", ""),
 ("Scope",
  "Time-block (weekly) batch = all AECN oil injected in the week, aggregated into ONE Certify GHG entry. AECN is all rail this "
  "week: 4 truck-out offloads from 2 railcars - SHPX215079 (offloads 2, 3 and 4, closing the railcar) and GPRX5102 (offload 1 of "
  "4) - fed by 6 inbound BOLs. No WODO, Aqueous/QOWV, Kerry or AECN-truck oil was injected in this week: the Ops Notes Summary "
  "rows for those four modes are zero throughout and the only non-zero numbers on their tabs are unit-conversion constants. "
  "[log:batch_id] [log:railcars] [log:offload_count] [log:ops_notes.other_oil_types_this_week]", "SECTION", "", ""),
 ("SUMMARY OF FINDINGS", SUMMARY, "SECTION", "", ""),

 ("1. Batch Folder Completeness", "", "SECTION", "", ""),
 ("1.1", "Batch folder exists with correct naming", "Pass",
  "Drive folder id 1lXC2uzkp9AbOhv6O6OMGNWnXTPEZRQ38, holding the Ops Notes plus three subfolders (Scale Tickets, Testing, "
  "Completed QA Checklists). The Ops Notes Summary tab cell B2 and the AECN Rail tab cell B2 both read 'INJ-07-2026-W3', the "
  "Certify source file is named 'INJ-07-2026-W3 Ops Notes.xlsx', and the lab certificate Client ID is '07-2026-W3 AECN'. "
  "[log:ops_notes.batch_id] [log:drive_folder_url] [log:lab_results.client_id]", "VR Appendix 1"),
 ("1.2", "Scale Tickets subfolder present and populated", "Pass",
  "Fourteen files covering the four offloads. Offloads GPRX5102_1, SHPX215079_3 and SHPX215079_4 are filed twice each - once as "
  "a CAT-numbered PDF and once as a phone photograph of the same ticket - and SHPX215079_2 is filed as a photograph only. That "
  "gives one full and one empty weighing for every offload, eight distinct tickets, all eight opened and read this session. On "
  "every ticket the three printed axle weights sum exactly to the printed GROSS WEIGHT. "
  "[log:scale_tickets] [log:scale_tickets.axle_sum_check] [log:_derived_values.n_offloads]", "VR Appendix 1"),
 ("1.3", "Testing subfolder present with required lab results", "Pass",
  "Testing subfolder holds one file, 'CHARM INDUSTRIAL - 07-2026-W3 AECN (SR26-04478.001).pdf', opened and read this session. "
  "One certificate is the correct count for this entry - the week's other oil types were not injected, so there are no WODO or "
  "QOWV certificates to expect here. [log:lab_results.certificate_no] [log:ops_notes.other_oil_types_this_week]", "VR Appendix 1"),
 ("1.4", "Ops Notes spreadsheet present and complete", "Pass",
  "'INJ-07-2026-W3 Ops Notes' (Google Sheet, modified 2026-08-04) present in the batch folder and attached to the Certify "
  "removal. Eight tabs. The AECN Rail tab carries the batch header, a 2-row railcar register, a 6-row BOL origins ledger, the "
  "4-row offload table, the EE totals row and the Section 4 rollup, all populated. Loaded both with and without value-resolution "
  "so cell annotations were read as well as cell values; the workbook carries one comment, on AECN Rail H11. "
  "[log:ops_notes] [log:ops_notes.tabs] [log:ops_notes.cell_comments]", "VR Appendix 1"),
 ("1.5", "Injection site confirmed as Basco 6 - no Vaulted references", "Flag",
  "Injection site is unambiguously Basco 6. The SGS certificate records LOCATION 'BASCO LOUISIANA'; the offload route map runs "
  "to 4687 LA-115, St Landry LA 71367; Certify names three components 'Opelousas to Basco'; all eight scale tickets were weighed "
  "at CAT Scale 3122, Pilot Travel Center, I-49 Exit 40, Bunkie LA. A term-by-term scan of every machine-readable source opened "
  "this session - all 501 non-empty Ops Notes cells across all eight tabs, the one cell comment, and the Certify components, "
  "calculation view, datapoints, sources and GHG statement dumps - is saved as the artifact for this check. It returned zero "
  "hits for Vaulted, Hutchinson, Kansas, salt cavern, KDHE, Class III and Advantek, against 7 hits for Basco, 5 for SOPOR and 1 "
  "for Opelousas. One Kansas address does appear, in the image-based evidence rather than the machine-readable set: the standing "
  "Charm purchase order bound as page 5 of the four complete BOL packages, all four read this session, carries Ship To 'Charm "
  "Industrial, 838 SW Purity Springs Road, El Dorado, KS 67042'. That is a procurement address on a blanket order dated "
  "10/2/2025, not an injection or storage site, and it is not Vaulted Deep. It affects no quantity. Flagged only because a "
  "verifier reading that page in isolation sees a Kansas ship-to on a Louisiana injection. "
  "[log:ops_notes.site_reference_scan] [log:lab_results.location] [log:bols.kansas_address_note] "
  "[log:bols.purchase_order_page5]", "VR Appendix 1"),

 ("2. Scale Tickets & Mass Determination", "", "SECTION", "", ""),
 ("2.1", "Empty and full scale ticket images present for each truckload", "Pass",
  "Four full and four empty CAT Scale certificates, one pair per offload, all eight opened and read. GPRX5102 offload_1: full "
  "1312226190422 / empty 1312226191437. SHPX215079 offload_2: full 1312226189397 / empty 1312226190401; offload_3: full "
  "1312226190407 / empty 1312226190411; offload_4: full 1312226190408 / empty 1312226190413. All eight were weighed at the same "
  "scale, CAT 3122 at the Pilot Travel Center on I-49 Exit 40 in Bunkie LA. Six of the eight are $5.25 reweighs carrying a "
  "TICKET # OF FULL $ WEIGH cross-reference, and on several that reference points to a ticket from a different offload of the "
  "same trailer - a CAT fleet-billing artifact that does not alter the certified GROSS WEIGHT printed on each ticket, which is "
  "the value used. [log:scale_tickets.pairs] [log:scale_tickets.scale_all]", "VR Sec 5.1"),
 ("2.2", "Dates within 2 weeks of injection", "Pass",
  "Ticket dates run 7-08-26 to 7-10-26 against an injection completion date of 2026-07-18, so the widest gap is 10 days and the "
  "narrowest 8, both inside the 14-day window. The offloads physically preceded the injection week: the oil was trucked from "
  "Opelousas to Basco on 8-10 July and injected during the 12-18 July block, which is the expected shape for time-block batching "
  "where the batch is defined by injection week rather than offload date. "
  "[log:scale_tickets.all_ticket_dates] [log:_derived_values.days_ticket_to_injection_max] "
  "[log:_derived_values.days_ticket_to_injection_min] [log:injection_date_all_offloads]", "VR Sec 3.5"),
 ("2.3", "Tractor and trailer numbers match between empty and full tickets", "Pass",
  "All four pairs match on BOTH identifiers. GPRX5102 offload_1 CH1/KL192 on both tickets; SHPX215079 offload_2 CH-2/KL188 full "
  "and CH2/KL188 empty; offload_3 CH2/KL188 on both; offload_4 CH1/KL192 on both. Because no pair shows a rig mismatch and no "
  "empty ticket is missing, the reporting period's Trailer Standard Empty Weight substitution is not triggered on any offload - "
  "the guidance states the trigger as substituting 'when either (a) the empty scale ticket is missing, or (b) the empty ticket's "
  "tractor and/or trailer do not match the Full ticket'. Consistent with that, the July guidance doc's only application list "
  "covers Week 1, the Ops Notes Empty Scale Ticket cells carry no substitution annotation, and none of the four empty weights "
  "equals a published standard tare. [log:scale_tickets.all_rig_identifiers_match] "
  "[log:vvb_guidance.trailer_standard_empty_weight_trigger_quote] [log:vvb_guidance.week3_substitution_needed] "
  "[log:scale_tickets.standard_tare_substitutions_applied]", "VR Sec 5.1"),
 ("2.4", "Net mass correctly calculated in Ops Notes", "Pass",
  "Each offload recomputed from the ticket faces as (full - empty) x 0.453592 and compared with the Ops Notes Net Scaled Mass "
  "cell. GPRX5102_1: 72,320 - 29,180 = 43,140 lb = 19,567.95888 kg. SHPX215079_2: 72,540 - 28,600 = 43,940 lb = 19,930.83248 kg. "
  "SHPX215079_3: 72,680 - 28,820 = 43,860 lb = 19,894.54512 kg. SHPX215079_4: 70,780 - 29,060 = 41,720 lb = 18,923.85824 kg. "
  "All four agree with the Ops Notes to the fifth decimal. The four sum to 78,317.19472 kg, which is the Section 4 rollup figure. "
  "[log:_derived_values.scale_net_lb_o1] [log:_derived_values.scale_net_lb_o2] [log:_derived_values.scale_net_lb_o3] "
  "[log:_derived_values.scale_net_lb_o4] [log:_derived_values.scale_net_kg_o1] [log:_derived_values.scale_net_kg_o2] "
  "[log:_derived_values.scale_net_kg_o3] [log:_derived_values.scale_net_kg_o4] "
  "[log:_derived_values.scale_ticket_net_kg_total] [log:ops_notes.offload_table] "
  "[log:ops_notes.rollup_section4.deliveries_raw_downstream_scale_tickets_kg]", "VR Sec 5.1"),
 ("2.5", "Injectate mass uses conservative value: MIN(scale ticket net, BOL mass)", "Pass",
  "Railcar-balance check run before any mass conclusion, per the closing-offload policy. SHPX215079: the three inbound BOL "
  "weighbridge nets are 25,940 + 26,200 + 26,190 = 78,330 kg; the four offloads drawn from it across Weeks 2 and 3 total "
  "78,262.76368 kg, which is 67.23632 kg SHORT of the aggregate. GPRX5102: BOLs 25,600 + 26,050 + 26,040 = 77,690 kg against "
  "19,567.95888 kg drawn so far. Neither railcar overdraws its inbound aggregate, so the closing-offload back-calculation is not "
  "triggered on either, and correctly all four Uncertainty Adjustment cells read 0. The raw scale-ticket nets therefore stand as "
  "the claimed masses and are the conservative figures. "
  "[log:_derived_values.railcar_balance_SHPX215079_kg] [log:_derived_values.railcar_balance_GPRX5102_kg] "
  "[log:_derived_values.bol_aggregate_SHPX215079_kg_recalc] [log:_derived_values.bol_aggregate_GPRX5102_kg_recalc] "
  "[log:ops_notes.offload_table]", "GHG Stmt Sec C"),
 ("2.6", "1.29% sparging deduction applied", "Pass",
  "Sparging is applied per offload at the Ops Notes constant 0.0129, then filtration at 0.00195 on the post-sparging mass, and "
  "both were recomputed offload by offload. Totals: sparging 1,010.291812 kg and filtration 150.748461 kg, matching the Section 4 "
  "rollup exactly. Worked on the largest offload: 19,930.83248 x 0.0129 = 257.107739 sparging; (19,930.83248 - 257.107739) x "
  "0.00195 = 38.363763 filtration; 19,930.83248 - 257.107739 - 38.363763 = 19,635.36098 kg MOBI. "
  "[log:_derived_values.sparging_loss_recalc_kg] [log:_derived_values.filtration_loss_recalc_kg] "
  "[log:_derived_values.sparging_rate] [log:_derived_values.filtration_rate] [log:ops_notes.constants] "
  "[log:_derived_values.cargo_o2_kg] [log:_derived_values.sparging_loss_o2_kg] "
  "[log:_derived_values.filtration_loss_o2_kg] [log:_derived_values.mobi_o2_kg]", "GHG Stmt Sec C"),
 ("2.7", "All mass values consistent across scale tickets, Ops Notes and Certify", "Pass",
  "One chain, four sources, no divergence anywhere. Scale tickets sum to 78,317.19472 kg; the Ops Notes rollup carries the same "
  "figure on both the upstream (BOL basis) and downstream (scale ticket) lines because the Uncertainty Adjustment is zero; COBB "
  "rows 387-390 carry the same four per-offload figures to the fifth decimal and sum to the same total; and Certify's Bio-Oil "
  "Mass - AECN via Rail reads 78,317.2 kg. After sparging and filtration the Ops Notes MOBI is 77,156.15445 kg, which is "
  "Certify's Mass of Injectant to the fifth decimal. Because no back-calculation applied this week there is none of the expected "
  "COBB-versus-Ops-Notes divergence that a closing offload would produce. "
  "[log:_derived_values.mobi_recalc_kg] [log:_derived_values.cobb_offloaded_total_kg] "
  "[log:certify.uncertainty_modals_read] [log:cobb.offloaded_quantity_kg]", "VR Sec 5.1"),

 ("3. Bills of Lading & Transportation", "", "SECTION", "", ""),
 ("3.1", "Truck BOL(s) present for each shipment", "Pass",
  "All six inbound BOLs present and read: 2389-R, 2390-R and 2391-R into SHPX215079, and 2395-R, 2396-R and 2397-R into "
  "GPRX5102. They are not in the per-batch Drive folder; they sit in the reporting-period AECN BOLs folder and all six are also "
  "attached to the Certify removal on the Sources tab, which satisfies the document-location rule. "
  "[log:bols.bols_into_railcar] [log:certify.certify_sources_enumeration] [log:certify.certify_sources_tab_total] "
  "[log:drive_enumeration] [log:bols]", "VR Appendix 1"),
 ("3.2", "Rail delivery: 3 truck BOLs per railcar, destination = SOPOR railyard", "Pass",
  "Three BOLs per railcar, six in total, and every one carries the same Ship To: 'CHARM Industrial Inc. /WATCO, Societe du port "
  "ferroviaire de Baie-Comeau, 18 Route Maritime, Baie-Comeau, Quebec, Canada G4Z 2L6, Att: Cedric Berube'. That is the SOPOR "
  "railyard, not Basco, which is the correct inbound destination for a rail movement. "
  "[log:bols.destination_all] [log:bols.destination_is_sopor_not_basco] [log:_derived_values.n_bols_total]", "VR Appendix 1"),
 ("3.3", "BOL mass matches or exceeds reported injectate mass", "Pass",
  "Applied at railcar level, which is the right level for a rail batch. SHPX215079: BOL aggregate 78,330 kg against 78,262.76368 "
  "kg drawn across all four of its offloads - the aggregate exceeds the draw by 67.23632 kg. GPRX5102: BOL aggregate 77,690 kg "
  "against 19,567.95888 kg drawn this week, with the remaining three offloads recorded in COBB against the following week. "
  "Each BOL's booked mass is the page 3 Port-Cartier weighbridge Masse nette, not the page 1 volumetric estimate; the estimate "
  "is higher on every one of the six, so taking the weighbridge figure is the conservative choice each time (2395-R for example "
  "prints 26.75 MT estimated against 25.60 MT weighed). "
  "[log:_derived_values.railcar_balance_SHPX215079_kg] [log:bols.page1_estimated_vs_weighbridge_note] "
  "[log:bols.bols_into_railcar]", "VR Sec 3.5"),
 ("3.4", "Origin and destination correct on BOL", "Pass",
  "All six run BioEnergie AE Cote-Nord Canada, 175 Portage des Mousses, Port-Cartier, Quebec to the SOPOR railyard at "
  "Baie-Comeau, carrier Transport Baie-Comeau, driver Santos Bolduc, product Renewable Fuel Oil RFO / Huile pyrolytique. The "
  "onward legs are documented separately: the railcar inbound record and arrival dates in the Ops Notes railcar register "
  "(SHPX215079 at Opelousas 2026-06-18, GPRX5102 2026-06-23), and the Opelousas-to-Basco truck leg by the offload scale tickets "
  "weighed in Louisiana. [log:bols.origin_all] [log:bols.destination_all] [log:bols.carrier_all] "
  "[log:ops_notes.railcars_received]", "PDD Sec A"),
 ("3.5", "Transportation distance consistent with known AECN rail route", "Pass",
  "All three legs verified against their own evidence, and both allocated distances plus the aggregate distance rebuilt from the "
  "offload masses. AECN to SOPOR: the Google Maps card shows 173 km via QC-138 O, and 173 / 1.609344 = 107.49722 mi, which is "
  "the Certify datapoint to five decimals. SOPOR to Opelousas: 2,199 mi, evidenced by the rail carrier in writing - Watco's "
  "rail logistics specialist wrote 'The total milage for this move was 2199.' - which is the correct evidence type for a rail "
  "leg, there being no Google Maps equivalent. Opelousas to Basco: the map offers three routes at 42.9, 42.9 and 40.9 miles and "
  "Charm books 42.9, the longest. The allocations then reconstruct exactly: aggregate AECN-to-SOPOR = 107.5 x 3 trucks x each "
  "offload's share of its railcar, summed = 323.110963 mi against Certify's 323.11096; allocated rail = 2,199 x the same shares, "
  "summed = 2,203.165916 mi against Certify's 2,203.16606; allocated Opelousas-to-Basco = 42.9 x 2 for the round trip x 4 "
  "offloads = 343.2 mi, matching Certify exactly and correctly NOT mass-allocated, since each offload truck runs the leg in full. "
  "[log:_derived_values.aecn_sopor_km_to_mi] [log:_derived_values.allocated_truck_distance_mi_calc] "
  "[log:_derived_values.allocated_rail_distance_mi_calc] [log:_derived_values.allocated_opel_basco_mi_calc] "
  "[log:routes.aecn_to_sopor] [log:routes.sopor_to_opelousas_rail] [log:routes.opelousas_to_basco]",
  "Transport Emissions Acctg v1.1"),
 ("3.6", "Transport mode = tanker truck or railcar (AECN modes)", "Pass",
  "Six transport components, which is the correct count for a three-leg rail movement: tanker truck process and embodied for "
  "AECN to SOPOR, railcar process and embodied for SOPOR to Opelousas, and tanker truck process and embodied for Opelousas to "
  "Basco. Each process leg carries a per-ton-mile factor and each embodied leg a per-mile factor, which is the correct pairing - "
  "fuel burn scales with cargo and distance, vehicle manufacture does not. "
  "[log:certify.components] [log:certify.uncertainty_modals_read]", "Certify Datapoints"),
 ("3.7", "Billing documents present and consistent with BOL", "Flag",
  "The complete package for these shipments is five pages: the Straight Bill of Lading face, the AECN Shipping Details sheet, "
  "the Port-Cartier weighbridge transaction slip, the TSCA import certification and the Charm purchase order. Page counts were "
  "checked on all six files and every page of every file was read. Four packages are complete at five pages - 2389-R, 2390-R, "
  "2391-R and 2397-R. Two are filed at three pages: 2395-R and 2396-R carry no TSCA import certification and no purchase order. "
  "Both are GPRX5102 inbounds, and the third BOL into that same railcar, 2397-R on the same purchase order, does carry both "
  "pages. Neither absent page feeds a mass, a carbon intensity or any Certify datapoint, so the reported Net CDR is untouched; "
  "what is missing is the import-compliance and commercial backing a verifier would expect to find alongside the other four. "
  "Three smaller print issues on the pages that are present. First, the purchase order records the order two ways - header PO # "
  "Q42025, matching the Customer PO on all six BOL faces, but line item Q32025 beneath, struck through and hand-amended to "
  "Q42025 on 2390-R only and left as Q32025 on 2389-R, 2391-R and 2397-R; quantity 6,000.0 MT and value $1,740,000.00 are "
  "identical on every copy, so this is one blanket order printed inconsistently rather than two orders. Second, 2396-R prints "
  "its production period as '#0X-2026' on both page 1 and the page 2 Run # field where its two sister BOLs into the same "
  "railcar, loaded the same days, print '#06-2026'. Third, 2395-R carries a red banner reading 'Correction page 2: densite a 60 "
  "et litres finaux' with the density at 60F struck from 1.2034 to 1.185 and the litres at 60F struck from 22,575 to 21,273 in "
  "pen; the delivered mass on that BOL is weighbridge-derived and is unaffected by either correction. "
  "[log:bols.page_counts] [log:bols.package_completeness] [log:bols.tsca_page4] [log:bols.purchase_order_page5] "
  "[log:bols.prod_period_typo] [log:bols.bols_into_railcar]", "VR Appendix 1"),

 ("4. Bio-Oil Composition & Testing", "", "SECTION", "", ""),
 ("4.1", "CHN analysis PDF present in Testing subfolder", "Pass",
  "SGS Certificate of Analysis SR26-04478.001, SGS order 5470884, reported 24-Jul-2026 and authorised by Christopher Bartlett, "
  "Lab Manager, SGS North America Inc., St. Rose LA. Present in the Drive Testing subfolder and attached to the Certify removal "
  "as the source for Carbon Content. [log:lab_results.certificate_no] [log:lab_results.authorised_signatory] "
  "[log:certify.attachment_map]", "VR Appendix 1"),
 ("4.2", "Carbon content within expected AECN range", "Pass",
  "Carbon 40.2 % (m/m) by ASTM D5291 Method C. Against the AECN reference population read this session from the MRV analysis "
  "workbook - mean 39.420350877193 %, standard deviation 6.09332409018271 %, n = 285 - the 2-sigma band is 27.2337 % to 51.607 "
  "%, so 40.2 sits well inside it, about 0.13 standard deviations above the mean. Carbon and hydrogen both carry the SGS '#' "
  "annotation, defined on the certificate itself as 'Result is outside of test method limits and/or analytical range used in "
  "method precision study' - a note that the ASTM D5291 precision study was built on petroleum matrices rather than bio-oil, not "
  "a rejection of the result. The same value, 0.402 as a decimal, is what the Ops Notes Summary carries and what Certify reports "
  "as 40.20 %. [log:lab_results.carbon_pct] [log:aecn_carbon_reference.carbon_mean_pct] "
  "[log:aecn_carbon_reference.carbon_sd_pct] [log:aecn_carbon_reference.carbon_n] "
  "[log:_derived_values.carbon_2sd_low_pct] [log:_derived_values.carbon_2sd_high_pct] "
  "[log:ops_notes.summary_tab_aecn_row.carbon_pct_decimal] [log:aecn_carbon_reference] [log:lab_results]", "Biomass Feedstock Acctg v1.2"),
 ("4.3", "Hydrogen and nitrogen values reported and reasonable", "Pass",
  "Hydrogen 7.4 % (m/m) by ASTM D5291 Method C, against the same reference population's hydrogen mean of 6.87438596491228 % and "
  "standard deviation 0.562848694875852 %, giving a 2-sigma band of 5.7487 % to 8.0001 % - so 7.4 is inside it. Nitrogen is "
  "reported as '<0.75 %', below the method's detection limit, which is the normal result for this oil: the reference workbook "
  "notes 'Most N (%) values are below detection limit' and could compute nitrogen statistics from only 9 numeric entries against "
  "285 for carbon and hydrogen. [log:lab_results.hydrogen_pct] [log:lab_results.nitrogen_pct] "
  "[log:aecn_carbon_reference.hydrogen_mean_pct] [log:aecn_carbon_reference.hydrogen_sd_pct] "
  "[log:_derived_values.hydrogen_2sd_low_pct] [log:_derived_values.hydrogen_2sd_high_pct] "
  "[log:aecn_carbon_reference] [log:lab_results]", "Biomass Feedstock Acctg v1.2"),
 ("4.4", "Lab report references correct batch/sample ID", "Pass",
  "The certificate's CLIENT ID reads '07-2026-W3 AECN' and its LOCATION reads 'BASCO LOUISIANA'. That matches the Ops Notes "
  "Summary composite sample ID '07-2026-W3 AECN' and the Ops Notes lab reference 'COA SR26-04478.001', and the certificate is "
  "the file Certify links to the Carbon Content datapoint. Sampled 22-Jul-2026 07:00 by the client, received 24-Jul-2026 10:00, "
  "analysed the same day. The sampling date falls four days after the 18-Jul injection completion, which the certificate itself "
  "explains: SAMPLE SOURCE reads 'Submitted Sample' and SAMPLED BY reads 'Client', so this is a client-made-up composite rather "
  "than a lab-drawn sample. Charm's sampling protocol takes a jar sample at every transfer event - here the four offloads "
  "weighed 7-08 to 7-10 - and composites them for submission, so the 22-Jul stamp is the date the weekly composite was made up "
  "and sent, not a date on which oil was drawn from an already-injected stream. Sister batch W1 shows the same four-day interval "
  "(injected 04-Jul, sampled 08-Jul). [log:lab_results.client_id] [log:lab_results.location] [log:lab_results.sampled] "
  "[log:lab_results.sample_source] [log:lab_results.sampled_by] [log:lab_results.sampling_mechanism] "
  "[log:lab_results.sampling_date_vs_injection_note] [log:ops_notes.summary_tab_aecn_row.composite_sample_id]", "VR Sec 5.1"),
 ("4.5", "pH / conductivity (aqueous phase) results present", "N/A",
  "AECN is not an aqueous-phase oil, so the aqueous pH and conductivity checks do not apply here. The batch's own injection pH "
  "is checked separately at item 5.5. [log:oil_type]", "UIC Permit / LDENR"),
 ("4.6", "Density above the formation fluid threshold for safe injection", "Pass",
  "Density at 15 C = 1.210 g/cm3 by ASTM D4052, against the 1.01 g/cm3 formation-fluid floor - comfortably denser than the "
  "formation fluid, which is the condition for safe injection at a permeable reservoir. "
  "[log:lab_results.density_15C_g_cm3] [log:lab_results.density_method] [log:_derived_values.density_floor_g_cm3]",
  "VR Sec 5.1"),

 ("5. Ops Notes Accuracy & Internal Consistency", "", "SECTION", "", ""),
 ("5.1", "Correct batch ID and injection date", "Pass",
  "Batch ID 'INJ-07-2026-W3' on both the Summary and AECN Rail tabs. Week Start (Sat) 2026-07-11 and Week End (Sat) 2026-07-18. "
  "All four offload rows carry the date 2026-07-18, which is the injection completion date and matches COBB on all four rows and "
  "the Certify period end. Certify's period reads 12 Jul 2026 to 18 Jul 2026. "
  "[log:ops_notes.batch_id] [log:week_start_sat] [log:week_end_sat] [log:injection_date_all_offloads] [log:certify.period]", ""),
 ("5.2", "All scale ticket values correctly transcribed", "Pass",
  "Eight ticket faces compared cell by cell against the Ops Notes offload table. Full weights 72,320 / 72,540 / 72,680 / 70,780 "
  "lb and empty weights 29,180 / 28,600 / 28,820 / 29,060 lb are transcribed exactly, in the same offload order, and the derived "
  "Net Scaled Mass in each row reproduces from those two numbers. No transposition, no substituted value, no cell carrying a "
  "figure that does not appear on a ticket. [log:scale_tickets.pairs] [log:ops_notes.offload_table]", ""),
 ("5.3", "Oil and aqueous phase mass reconciliation", "N/A",
  "AECN-only entry - no aqueous phase was injected in this week, so there is no oil-versus-aqueous reconciliation to perform. "
  "[log:oil_type] [log:ops_notes.other_oil_types_this_week]", ""),
 ("5.4", "COBB values present and match", "Flag",
  "All four offloads are present in the 'Basco Injection -- COBB' tab at rows 387-390 and every value agrees with the Ops Notes: "
  "offloaded quantities 19,567.95888 / 19,930.83248 / 19,894.54512 / 18,923.85824 kg, origin quantities 77,690 and 78,330 kg, "
  "BOL strings '2395-R, 2396-R, 2397-R' and '2389-R, 2390-R, 2391-R', completion date 2026-07-18 and pH 3.2 on all four. "
  "quantity_consumed_in_injection equals offloaded_quantity_kg on every row. The flag is the row label: COBB records "
  "injection_batch as 'INJ-20260712_bio_oil' while the Ops Notes, the Certify source filename and the lab Client ID all use "
  "'INJ-07-2026-W3'. A verifier tracing from Certify to COBB by batch ID finds nothing and has to re-search by railcar or BOL, "
  "as this review did - the join key that does work is consumed_lot, which matches the four Ops Notes offload IDs exactly. This "
  "is the same labelling break flagged on Week 2, and it is a change of convention rather than a one-off: the Week 1 rows in the "
  "same tab are labelled 'INJ-07-2026-W1_bio_oil', which does tie. Values are correct; the label is not traceable. "
  "[log:cobb.rows] [log:cobb.batch_id_col_a] [log:cobb.batch_id_col_a_note] [log:cobb.row_id_at_qa] "
  "[log:_derived_values.cobb_offloaded_total_kg]", ""),
 ("5.5", "pH measurement noted", "Pass",
  "pH 3.2 recorded on all four offload rows and carried into the Section 4 rollup, inside the Ops Notes target range of 2.0 to "
  "4.0 held in the constants block. COBB independently records injection_ph 3.2 on all four rows. "
  "[log:ops_notes.offload_table] [log:ops_notes.rollup_section4.ph] [log:ops_notes.constants] [log:cobb.rows]", ""),
 ("5.6", "Completion date present and matches COBB", "Pass",
  "Ops Notes offload date 2026-07-18 on all four rows; COBB injection_completion_date 2026-07-18 on all four rows; Certify "
  "period ends 18 Jul 2026. COBB's order-created timestamps are 2026-07-08 19:04:42 for the SHPX215079 offloads and 2026-07-09 "
  "19:35:28 for the GPRX5102 offload, both consistent with the 7-08 to 7-10 scale ticket dates. "
  "[log:injection_date_all_offloads] [log:cobb.rows] [log:certify.period]", ""),
 ("5.7", "No erroneous density measurements in Batch Composition", "N/A",
  "The density field in the Ops Notes Batch Composition block is expected to be blank for AECN and is blank. Density for this "
  "oil is evidenced by the SGS certificate and is checked at item 4.6. [log:oil_type] [log:lab_results.density_15C_g_cm3]", ""),

 ("6. AECN Feedstock Emissions & Allocation", "", "SECTION", "", ""),
 ("6.1", "AECN process emissions LCA CI spreadsheet present", "Pass",
  "'Bioenergy AE Cote-Nord Canada - LCA CI calculation - 1Jul2025 to 31Dec2025 S.xlsx' is on the Certify Sources tab, linked to "
  "the Carbon Intensity - AECN Pyrolysis Process Emissions (via Rail) datapoint. The file's period is a half-year behind the "
  "#06-2026 production period the BOLs stamp on this oil, and the emission factor sheet's own derivation tab shows why that is "
  "correct rather than a mismatch. The 'H1 2026 AECN Process Emissions' tab builds the factor in two parts: a stack-emissions "
  "calculation off the GHGenius Ensyn RTP model, giving 0.00412683 MTCO2e per MT of bio-oil, plus a line labelled 'Total Other "
  "Process Emissions MT CO2e/MT Bio Oil (From Vendor LCA)' at 0.0997. The two sum to 0.10382683, which is the 0.1038 baseline "
  "carried on the 2026 tab. A vendor LCA can only report a period that has closed, so the most recent one available when the H1 "
  "2026 factor was set is precisely this 1Jul2025-31Dec2025 file. The attached document is the input the factor was actually "
  "built from. [log:certify.certify_sources_enumeration] [log:certify.attachment_map] [log:bols.bols_into_railcar] "
  "[log:ef_live_source.h1_2026_derivation_tab] [log:ef_live_source]", "VR Appendix 1"),
 ("6.2", "LCA with stack emissions file present", "Pass",
  "'BH 2025 LCA w Stack Emissions.xlsx' is on the Certify Sources tab, linked to the same pyrolysis process carbon-intensity "
  "datapoint. It backs the stack half of the factor: the emission factor sheet's H1 2026 derivation tab takes biomass input "
  "12,473 MT and bio-oil output 6,305,682 L at 17.67 MJ/L, applies IPCC AR6 100-year warming potentials of 27.0 for non-fossil "
  "methane and 273.0 for nitrous oxide, and returns total stack emissions of 0.00412683 MTCO2e per MT of bio-oil. "
  "[log:certify.certify_sources_enumeration] [log:certify.attachment_map] [log:ef_live_source.h1_2026_derivation_tab] "
  "[log:ef_live_source]", "VR Appendix 1"),
 ("6.3", "Supporting GHG data file present for the AECN production period", "Pass",
  "'Supporting GHG data 31Dec2025 S.xlsx' is on the Certify Sources tab, linked to the pyrolysis process carbon-intensity "
  "datapoint alongside the LCA CI calculation, the stack-emissions workbook and a supporting screenshot - four sources on that "
  "one datapoint. [log:certify.certify_sources_enumeration] [log:certify.attachment_map]", "VR Appendix 1"),
 ("6.4", "AECN pyrolysis process EF matches the current value", "Flag",
  "Certify's stored carbon intensity is 0.104 tCO2e/t. The emission factor sheet 2026 tab row 3, 'Production Emissions for AECN "
  "Oil--H1 2026', gives baseline 0.1038 and Final Quantity 0.1040167408 MTCO2e/MT after the 0.002088061302 conservatism uplift, "
  "and that Final Quantity is the authoritative figure. The 0.1038 baseline is itself built on the sheet's H1 2026 derivation "
  "tab as stack emissions 0.00412683 plus vendor-LCA other-process emissions 0.0997, summing to 0.10382683 MTCO2e per MT. "
  "The stored value is demonstrably the rounded one: 78.31719472 t x 0.104 "
  "= 8.144988 t, which is what Certify reports as 8.145, whereas the sheet's final factor gives 8.146299 t. The rounding "
  "understates pyrolysis process emissions by 0.001311 tCO2e - a non-conservative direction, but 1.3 kg on a 103 tonne removal, "
  "roughly 0.0013 % of net and far below materiality. Week 1 and Week 2 store the same 0.104, so this is a standing entry "
  "convention rather than a keying slip on this week; raised so the full-precision factor can be entered going forward. "
  "[log:_derived_values.ef_aecn_process_certify_stored] [log:_derived_values.ef_aecn_process_sheet_final] "
  "[log:_derived_values.ef_aecn_process_sheet_baseline] [log:_derived_values.ef_aecn_process_uncertainty_factor] "
  "[log:_derived_values.pyrolysis_process_recalc_t] [log:_derived_values.pyrolysis_process_at_sheet_final_ef_t] "
  "[log:_derived_values.pyrolysis_process_ef_rounding_delta_t] [log:ef_live_source.aecn_process_row3] "
  "[log:ef_live_source.h1_2026_derivation_tab] [log:ef_live_source]", "Certify Datapoints"),
 ("6.5", "AECN pyrolysis embodied EF applied", "Pass",
  "Certify stores 0.02116 tCO2e/t, exactly the emission factor sheet 2026 tab row 4 Final Quantity of 0.02116 MTCO2e/MT (the "
  "uplift on that row is 0.0, so baseline and final coincide). Applied: 78.31719472 t x 0.02116 = 1.657192 t against Certify's "
  "1.657 t. Sourced to 'AECN Embodied Emissions Calc Sheet (1).pdf' on the Sources tab. "
  "[log:_derived_values.ef_aecn_embodied] [log:_derived_values.pyrolysis_embodied_recalc_t] "
  "[log:ef_live_source.aecn_embodied_row4] [log:ef_live_source] [log:certify.attachment_map]", "Certify Datapoints"),
 ("6.6", "Mass basis for pyrolysis EFs is Origin Bio-Oil Mass (pre-sparging)", "Pass",
  "Both pyrolysis components take Bio-Oil Mass - AECN via Rail at 78,317.2 kg, which is the pre-treatment mass, not the "
  "post-treatment Mass of Injectant of 77,156.15445 kg. The 1,161.04 kg difference is the sparging and filtration losses that "
  "occur at the injection site, after the oil was produced. Charging pyrolysis on the full produced quantity is correct: those "
  "emissions were incurred to make all of it. [log:certify.uncertainty_modals_read] "
  "[log:_derived_values.sparging_loss_recalc_kg] [log:_derived_values.filtration_loss_recalc_kg] "
  "[log:_derived_values.mobi_recalc_kg]", "Certify Components"),

 ("7. Gross-to-Net Calculation", "", "SECTION", "", ""),
 ("7.1", "Gross CO2e = injectate mass x carbon wt% x 3.667", "Pass",
  "Rebuilt from source, not from Certify: 77.15615445 t x 0.402 x 3.667 = 113.738511 tCO2e against Certify's 113.739 t, agreeing "
  "to the third decimal. The mass is the Ops Notes MOBI, itself rebuilt from the eight ticket faces through sparging and "
  "filtration, and the carbon fraction is the SGS certificate value entered as a decimal. "
  "[log:_derived_values.gross_co2e_recalc_t] [log:_derived_values.mobi_recalc_kg] [log:_derived_values.carbon_decimal] "
  "[log:_derived_values.co2_c_ratio] [log:_derived_values.gross_certify_t]", "GHG Stmt Sec E"),
 ("7.2", "Process emissions correctly deducted", "Pass",
  "All eight emission components rebuilt from full-precision factors and summed independently: pyrolysis process 8.144988 + "
  "pyrolysis embodied 1.657192 = 9.802180 t, and transport process 5.808576 + transport embodied 0.226875 = 6.035451 t, for "
  "total activities of 15.837631 t against Certify's 15.838 t. Nothing is missing from the deduction and nothing is "
  "double-counted: the two pyrolysis components take the pre-treatment mass once, the three process legs take it once each "
  "against their own leg distance, and the three embodied legs take distance only. "
  "[log:_derived_values.activities_recalc_t] [log:_derived_values.pyrolysis_subtotal_t] "
  "[log:_derived_values.transport_process_sum_t] [log:_derived_values.transport_embodied_sum_t] "
  "[log:_derived_values.transport_total_t] [log:_derived_values.pyrolysis_process_recalc_t] "
  "[log:_derived_values.pyrolysis_embodied_recalc_t] [log:_derived_values.activities_certify_t]", "GHG Stmt Sec E"),
 ("7.3", "Uncertainty discount correctly applied", "Pass",
  "The discount is present in the calculation view at 627 kg and is 0.5513 % of gross. Rebuilt by variance propagation from the "
  "per-component inputs read out of the modals: the carbon and mass terms on the sequestration give a relative uncertainty of "
  "0.0054913, so sigma on gross is 624.57 kg, and adding the mass-driven terms on the two pyrolysis components and the three "
  "transport process legs in quadrature gives 625.56 kg at one sigma. That is 1.44 kg, or 0.23 %, below Certify's 627 kg, which "
  "is within the rounding of the component values the propagation runs on. Inputs and output are mutually consistent. "
  "[log:_derived_values.uncertainty_1sigma_reconstructed_kg] [log:_derived_values.uncertainty_discount_reported_kg] "
  "[log:_derived_values.uncertainty_discount_pct_of_gross] [log:certify.calculation_view]",
  "Isometric Standard v2.0 S3.7"),
 ("7.4", "Net removals consistent with the GHG Statement", "Pass",
  "Certify's calculation view reads 113.739 + 5.808 - 15.838 - 0.627 = 103.082 t. Rebuilt fully independently from source values "
  "- gross 113.738511 + BCU 5.808 - activities 15.837631 - discount 0.627 = 103.0819 t - the two agree. The July GHG statement "
  "ggs_1KZBX6AX41S029JS carries this entry at 103.08 tCO2e for 12 Jul 2026 to 18 Jul 2026, and its nine entries sum to 874.53 t, "
  "the statement total. The projected buffer pool allocation of 43.728 credits is 5.000 % of that total, matching Basco 6's Low "
  "Risk classification. [log:_derived_values.net_calc_t] [log:_derived_values.net_certify_t] "
  "[log:_derived_values.ghg_statement_entry_sum_t] [log:_derived_values.ghg_statement_total_t] "
  "[log:_derived_values.buffer_pct] [log:certify.ghg_statement]", "GHG Stmt Sec E"),

 ("7a. Per-Component Uncertainty Verification", "", "SECTION", "", ""),
 ("7a.1", "Layer 1 - Completeness: every component has uncertainty on the correct datapoint", "Pass",
  "All ten component modals opened individually and their INPUTS sections read. Mass-of-load uncertainty is present on every "
  "component that carries a mass datapoint: the injection sequestration, both pyrolysis components, and all three transport "
  "PROCESS legs. The three transport EMBODIED components correctly have no mass datapoint at all - embodied emissions are "
  "quantified per mile, not per ton-mile - so no mass uncertainty is expected or missing there. Carbon content carries its own "
  "uncertainty on the sequestration component. Distance datapoints carry no uncertainty, correct where the longest published "
  "route is used. The BCU component is a constant and carries none, as expected. "
  "[log:certify.uncertainty_modals_read] [log:certify.component_count]", "MRV Uncertainty Cheat Sheet"),
 ("7a.2", "Layer 2 - Correctness: reported +/- matches the expected value", "Pass",
  "RESOLVED 2026-08-06. Read from the modals, not inferred. Carbon content is right: 40.2 % +/- 0.1608, and 40.2 x 0.004 = "
  "0.1608, so the 0.4 % rule is applied exactly. Mass was originally raised as a Flag: four offloads means four discrete "
  "outbound weighings, so the expected mass uncertainty is 4 x 72.57 = 290.28 kg, and while the Mass of Injectant datapoint "
  "carried exactly that, the Bio-Oil Mass - AECN via Rail datapoint carried 290.98 kg - a figure that is not 4 x 72.57, is not "
  "a whole multiple of 72.57 at any count, and matched no method published in the reporting period's guidance doc, which was "
  "read in full. Two mass datapoints on the same entry, derived from the same four weighings, must carry the same figure. The "
  "datapoint has since been corrected to 290.28 kg. Re-read in Certify this session: the AECN pyrolysis process modal and the "
  "SOPOR to Opelousas rail process modal both now show 'Bio-Oil Mass - AECN via Rail 78,317.2 +/-290.28 kg', and because "
  "Bio-Oil Mass is a shared datapoint the single edit propagates to all five components that consume it - both pyrolysis "
  "components and all three transport process legs. Both mass datapoints now agree at 290.28 kg. The calculation view is "
  "unchanged at 627 kg discount and 103.082 t net, which is expected: the original discrepancy was worth 0.005 kg of "
  "propagated uncertainty. [log:certify.uncertainty_modals_read] [log:certify.correction_applied] "
  "[log:_derived_values.mass_uncertainty_injectant_kg] [log:_derived_values.mass_uncertainty_biooil_kg] "
  "[log:_derived_values.mass_uncertainty_expected_kg] [log:_derived_values.uncertainty_biooil_pre_correction_kg] "
  "[log:_derived_values.per_weighing_uncertainty_kg] [log:_derived_values.n_offloads] "
  "[log:_derived_values.carbon_uncertainty_abs] [log:_derived_values.carbon_uncertainty_rate] "
  "[log:vvb_guidance.mass_uncertainty_section_present]", "Isometric Standard v2.0 S3.7.3"),
 ("7a.3", "Layer 3 - Discount presence and consistency", "Pass",
  "The aggregate discount of 627 kg is present in the Certify calculation view and is carried into the net as 0.627 t. It "
  "reconstructs to 625.56 kg from the per-component inputs above, a 0.23 % gap attributable to rounding in the component values, "
  "so the inputs and the output are mutually consistent. The GHG statement's figure for this entry, 103.08 t, is the "
  "post-discount net. [log:_derived_values.uncertainty_discount_reported_kg] "
  "[log:_derived_values.uncertainty_1sigma_reconstructed_kg] [log:certify.ghg_statement]",
  "Certify Calculation View"),

 ("8. Data Integrity & Anomaly Check", "", "SECTION", "", ""),
 ("8.1", "No unexplained outliers in mass, carbon content or calculated values", "Pass",
  "The four offload masses run 18,923.86 to 19,930.83 kg, a spread of about 5 % around the mean, normal for tanker loads. Carbon "
  "40.2 % and hydrogen 7.4 % both sit inside their 2-sigma reference bands. The sister-batch allocation check holds: each "
  "offload's upstream allocation fraction equals its mass over its own railcar's total, the four fractions are 0.251872, "
  "0.254447, 0.253984 and 0.241591, and the three SHPX215079 fractions sum to 0.75002 - the three-quarters of that railcar drawn "
  "this week - with the GPRX5102 fraction accounting for its first quarter. No offload number is used twice on either railcar. "
  "Every emission component reconstructs to Certify from independently sourced factors and distances. "
  "[log:_derived_values.allocation_fraction_o1] [log:_derived_values.allocation_fraction_o2] "
  "[log:_derived_values.allocation_fraction_o3] [log:_derived_values.allocation_fraction_o4] "
  "[log:_derived_values.allocation_fraction_shpx_three_sum] [log:_derived_values.cargo_o4_kg] "
  "[log:_derived_values.cargo_o2_kg] [log:ops_notes.offload_table] "
  "[log:cobb.no_offload_number_collision] [log:_derived_values.activities_recalc_t] "
  "[log:aecn_carbon_reference] [log:lab_results]", "GHG Stmt Sec C"),
 ("8.2", "Any spills/losses documented, unrecovered mass deducted, incident report filed", "Pass",
  "One loss on the batch and it is documented and excluded rather than claimed. Railcar SHPX215079 was loaded with 78,330 kg "
  "across its three inbound BOLs and discharged 78,262.76368 kg across its four offloads, leaving 67.23632 kg unaccounted - "
  "0.0858 % of the load, consistent with residual heel and line loss. The Ops Notes railcar register carries that figure "
  "explicitly in its Loss column, and none of it is claimed: the mass that flows to Certify is the sum of the measured offloads, "
  "not the BOL aggregate. GPRX5102 is not yet fully drawn down, so no loss figure is computable for it this week. Beyond that, "
  "an explicit term scan for spill, spilled, leak, leaked, incident, release, released, overflow, contamination, recovered, "
  "unrecovered, clean-up, cleanup and damage was run across all 501 non-empty Ops Notes cells on all eight tabs, the workbook's "
  "single cell comment, and COBB rows 387-390 across all 29 columns; it returned zero hits on every term, and the artifact is "
  "saved with this review. The only loss-related content in the workbook is that railcar Loss (kg) field, which is a "
  "mass-balance line rather than an incident record, and the single cell comment concerns prior-week claiming. "
  "[log:_derived_values.shpx215079_loss_kg] [log:_derived_values.shpx215079_loss_pct] "
  "[log:ops_notes.railcars_received] [log:ops_notes.cell_comments] [log:ops_notes.spill_incident_scan] "
  "[log:ops_notes]", "GHG Stmt Sec C"),
 ("8.3", "If inventory reconciliation: carryover oil documented", "Pass",
  "This week draws from two railcars that each straddle weekly entries, and both hand-offs reconcile. SHPX215079: 19,513.52784 "
  "kg was claimed in Week 2 as offload_1 and 58,749.23584 kg is claimed here across offloads 2, 3 and 4; those sum to "
  "78,262.76368 kg, the railcar's total discharge. The Ops Notes records the prior claim in the Previously Claimed column and "
  "annotates it in the cell comment on AECN Rail H11, and COBB row 379 independently carries the same 19,513.52784 kg against "
  "injection date 2026-07-10. GPRX5102: offload_1 is claimed here and COBB rows 391-393 carry offloads 2, 3 and 4 against "
  "injection date 2026-07-25, so the remainder moves to the following week with no gap and no overlap; those four offloads sum "
  "to 77,690.00 kg, exactly the railcar's BOL aggregate. Independently, none of this week's eight CAT ticket numbers appears "
  "among Week 2's ten, so no weighing is counted twice. The railcar draw-down precondition for the per-offload uncertainty "
  "framing is met on the railcar that closed this week: SHPX215079 discharged in 4 offloads against 3 inbound BOLs, and 4 is "
  "greater than 3, so the framing is conservative. The same check on GPRX5102 falls due when its closing offload is reviewed "
  "next week. [log:_derived_values.shpx215079_prior_claimed_kg] [log:_derived_values.shpx215079_claimed_this_week_kg] "
  "[log:_derived_values.shpx215079_total_offloaded_kg] [log:_derived_values.gprx5102_all_offloads_sum_kg] "
  "[log:_derived_values.n_offloads_shpx215079] [log:_derived_values.n_bols_shpx215079] "
  "[log:ops_notes.cell_comments] [log:cobb.sister_rows_same_railcars] [log:scale_tickets.no_overlap_with_W2_note]",
  "VR Sec 3.5"),

 ("9. Isometric Certify - Removal Component Mapping", "", "SECTION", "", ""),
 ("9.1", "SEQUESTRATIONS: gross injection sequestration matches internal gross", "Pass",
  "One sequestration component, 'Injection sequestration' at 113.739 tCO2e, sitting under CO2 stored. Its two inputs read "
  "77,156.15445 kg mass of product and 40.2 % carbon content, and its blueprint is product_mass x carbon_content x "
  "co2e_of_carbon. Rebuilt from source that is 77.15615445 x 0.402 x 3.667 = 113.738511 t. "
  "[log:certify.components] [log:certify.uncertainty_modals_read] [log:_derived_values.gross_co2e_recalc_t]",
  "Certify Components"),
 ("9.2", "ACTIVITIES - Transport: process and embodied correct for the AECN rail route", "Pass",
  "All six components rebuilt from full-precision factors read off the sheet this session, never from Certify's rounded display. "
  "PROCESS legs, mass x distance x factor: AECN to SOPOR 78.31719472 t x 107.49722 mi x 0.0001143 = 0.962278 t (Certify 0.962); "
  "rail 78.31719472 t x 2,199 mi x 0.000025910374 = 4.462272 t (Certify 4.462); Opelousas to Basco 78.31719472 t x 42.9 mi x "
  "0.0001143 = 0.384026 t (Certify 0.384). Each process leg correctly uses its own one-way leg distance, not an allocated one. "
  "EMBODIED legs, allocated or aggregate distance x per-mile factor: AECN to SOPOR 323.110963 mi x 0.0002375605888 = 0.076758 t "
  "(Certify 0.077); rail 2,203.165916 mi x 0.00003113066508 = 0.068586 t (Certify 0.069); Opelousas to Basco 343.2 mi x "
  "0.0002375605888 = 0.081531 t (Certify 0.082). The railcar embodied factor is the current 2026 value of 0.0000311307 "
  "MTCO2e/mi, displayed as 0.03113 kgCO2e/mi in the modal, not the deprecated 0.000035224. Transport process sum 0.962 + 4.462 + "
  "0.384 = 5.808576 t; transport embodied sum 0.077 + 0.069 + 0.082 = 0.226875 t; bio-oil transport total 5.808576 + 0.226875 = "
  "6.035451 t against the Certify subtotal of 6.036 t. "
  "[log:_derived_values.aecn_sopor_truck_process_recalc_t] [log:_derived_values.rail_process_recalc_t] "
  "[log:_derived_values.opel_basco_truck_process_recalc_t] [log:_derived_values.aecn_sopor_truck_embodied_recalc_t] "
  "[log:_derived_values.rail_embodied_recalc_t] [log:_derived_values.opel_basco_truck_embodied_recalc_t] "
  "[log:_derived_values.transport_process_sum_t] [log:_derived_values.transport_embodied_sum_t] "
  "[log:_derived_values.transport_total_t] [log:_derived_values.ef_railcar_embodied] "
  "[log:_derived_values.ef_railcar_embodied_stale] [log:ef_live_source]", "Certify Components"),
 ("9.3", "ACTIVITIES - Injection: LCS and sample transport emissions", "N/A",
  "No LCS pre-treatment applies to AECN oil - LCS is a pH adjustment used on the aqueous fraction, and no aqueous oil was "
  "injected in this week. Sample transport is carried on the 2026 emission factor sheet at row 23 as 'Sample Transport "
  "Emissions ... N/A Immaterial', so there is no component to expect. "
  "[log:oil_type] [log:ops_notes.other_oil_types_this_week] [log:ef_live_source]", "Certify Components"),
 ("9.4", "ACTIVITIES - Pyrolysis: process and embodied entries present", "Pass",
  "Both present, under the rail naming convention: 'AECN pyrolysis (from rail loads) - process emissions' at 8.145 t and 'AECN "
  "pyrolysis (from rail loads) - embodied emissions' at 1.657 t, grouped under Pyrolysis or other process with a subtotal of "
  "9.802 t. The rail-specific label is the deliberate signal to the verifier that scale tickets, not the BOL, are the "
  "authoritative mass source for this entry. [log:certify.components] [log:certify.component_subtotals] "
  "[log:_derived_values.pyrolysis_subtotal_t]", "Certify Components"),
 ("9.5", "REDUCTIONS: applicable BCU categories populated", "Pass",
  "One reduction component, 'BCU - Transport Process' at 5.808 tCO2e, entered as a Constant CO2 reduction of 5,808 kgCO2e. That "
  "is the only category that applies: the other four - injection diesel, injection waste, pre-treatment diesel and pre-treatment "
  "waste - appear only on removals burdened with site emissions, and this entry carries none. The supporting certificate was "
  "opened and read, not merely confirmed present on the Sources list: a Transport Emission Reduction Certificate retired 18 Dec "
  "2025 by Amazon.com Services LLC on behalf of Charm Industrial, Inc, vintage 2024, commodity Renewable Diesel TERCS, 2,795 "
  "TERCs retired, verification body Turner Mason & Company, certificate TERC1012_2024_DF_RENEWABLE_DIESEL-13-47206-50000, with "
  "the certificate stating that 'Each TERC represents one metric ton of carbon dioxide equivalent (\"tCO2e\") reduced.' The 5.808 "
  "t drawn here sits against 2,795 t retired. This satisfies eligibility criterion EC1, which requires that 'BCUs must be "
  "purchased and retired by the Project Proponent in a volume equivalent to the volume of fuel substitution claimed.' "
  "[log:bcu_certificate] [log:_derived_values.tercs_retired] [log:protocol.ec1_quote] [log:certify.components]",
  "Certify Components"),
 ("9.6", "REDUCTIONS: BCU quantity >= corresponding activity emission", "Pass",
  "The BCU is 5.808 t against a transport process sum of 0.962 (AECN to SOPOR) + 4.462 (SOPOR to Opelousas rail) + 0.384 "
  "(Opelousas to Basco) = 5.808 t at Certify's precision, so the offset equals the emission it covers exactly. Rebuilt at full "
  "precision the three legs total 5.808576 t, so the 5,808 kg entered is 0.58 kg short of the fully-precise sum - a shortfall, "
  "which is the conservative direction. The cap is respected: the module allows BCUs only to 'substitute for some, or all, of "
  "project fuel usage' and Equation 5 caps each journey at the point where fuel usage net of substitution reaches zero, so an "
  "offset equal to the process emission is the ceiling and this entry sits on it. Transport embodied of 0.226875 t is correctly "
  "NOT offset - vehicle manufacture is not fuel usage and falls outside what the module permits BCUs to substitute for. Worth "
  "recording that this is the exact failure corrected on the Week 2 entry, where the BCU had been set to the pyrolysis process "
  "figure; here the BCU is 5.808 and pyrolysis process is 8.145, so the two are properly distinct. "
  "[log:_derived_values.bcu_constant_t] [log:_derived_values.transport_process_sum_t] "
  "[log:_derived_values.bcu_minus_transport_process_t] [log:_derived_values.transport_embodied_sum_t] "
  "[log:protocol.equation5_cap_quote] [log:protocol.section_4_0_scope_quote] "
  "[log:protocol.embodied_not_bcu_eligible_basis]", "Transport Emissions Acctg v1.1"),
 ("9.7", "COUNTERFACTUALS: zero per protocol", "Pass",
  "No counterfactual component exists on this removal. The Components treegrid holds exactly two top-level groups, Stored with "
  "two children and Emitted with eight, ten components in total, and the calculation view formula is Sequestrations + "
  "Adjustments - Activities - Uncertainty Discount, with no counterfactual term in it. Zero is the correct treatment for bio-oil "
  "geological storage. No loss component group is present either. "
  "[log:certify.counterfactuals] [log:certify.losses] [log:certify.component_count] [log:certify.calculation_view]",
  "Certify Components"),
 ("9.8", "Net removal matches internal calculation", "Pass",
  "Certify reports 103.082 t and the headline reads 103.08 t. Two independent rebuilds agree: using Certify's own component "
  "totals, 113.739 + 5.808 - 15.838 - 0.627 = 103.082 t; and rebuilding every input from source - gross 113.738511, BCU 5.808, "
  "activities 15.837631, discount 0.627 - gives 103.0819 t. The July GHG statement carries the same 103.08 t for this entry. "
  "[log:_derived_values.net_calc_t] [log:_derived_values.net_certify_t] [log:_derived_values.net_headline_t] "
  "[log:certify.calculation_view] [log:certify.ghg_statement]", "Certify Calculation View"),
 ("9.9", "All datapoints have descriptive source names", "Flag",
  "Twenty-nine of the thirty-one sources are descriptively named - scale tickets carry railcar, date, offload number and CAT "
  "ticket number; BOLs carry their number; the lab certificate carries the batch and the SGS reference; the Ops Notes carries "
  "the batch ID. Two do not. 'image (44).png' is the sole evidence for the 2,199-mile rail leg, which at 4.462 t is the largest "
  "single transport emission on the entry; it is a screenshot of a Gmail thread whose body reads only 'The total milage for this "
  "move was 2199', with no subject line, no visible date and no railcar identifier, and it is relied on for both railcars. "
  "'Screenshot 2026-02-23 at 11.29.10 AM.png' supports the pyrolysis carbon intensity. Both are hygiene rather than substance - "
  "the values are correct and independently corroborated - but a verifier cannot tie the rail email to these two movements from "
  "the file itself. Three further filenames carry export suffixes rather than content descriptions: 'Standard Emission Factors + "
  "Calculations (8).xlsx', 'AECN Embodied Emissions Calc Sheet (1).pdf' and 'AECN to SOPOR (Baie Comeau) (2).png'. "
  "[log:certify.generic_filenames_present] [log:certify.certify_sources_enumeration] "
  "[log:routes.sopor_to_opelousas_rail] [log:routes] [log:ef_live_source]", "Certify Datapoints"),
 ("9.10", "All sources match files in the batch folder on Drive", "Pass",
  "Thirty-one source rows on the Sources tab, all enumerated. The batch-specific evidence ties both ways: the eight scale "
  "tickets, the lab certificate and the Ops Notes are in the batch Drive folder and on the Sources tab. The six BOLs are on the "
  "Sources tab and in the reporting-period AECN BOLs folder rather than the batch folder, which the document-location rule "
  "allows. The remaining sources are period-level or standing references - the emission factor sheet, the railcar embodied "
  "derivation, the GLEC framework, the AECN LCA set, the BCU certificate and the route evidence - and are expected to live "
  "outside the batch folder. The evidence-attachment requirement is met: the Ops Notes is linked to six datapoints including "
  "both mass datapoints and all three allocated distances, and the scale tickets and BOLs are each linked to both Mass of "
  "Injectant and Bio-Oil Mass - AECN via Rail, fifteen sources on each. "
  "[log:certify.certify_sources_tab_total] [log:certify.certify_sources_enumeration] [log:certify.attachment_map] "
  "[log:certify.attachment_gate] [log:_derived_values.certify_sources_total] [log:ef_live_source] "
  "[log:drive_enumeration]", "Certify Sources"),
]


def build_workbook(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QA Checklist"
    ws.append(["Item", "Check", "Status", "Notes / Evidence", "Verifier Reference"])
    for c in ws[1]:
        c.font = BOLD; c.fill = HEADER_FILL; c.alignment = WRAP; c.border = BORDER
    for r in rows:
        ws.append(list(r))
        i = ws.max_row
        status = r[2]
        for col in range(1, 6):
            cell = ws.cell(i, col)
            cell.alignment = WRAP; cell.border = BORDER
        if status == "SECTION":
            for col in range(1, 6):
                ws.cell(i, col).fill = HEADER_FILL; ws.cell(i, col).font = BOLD
        elif status == "Fail":
            for col in range(1, 6):
                ws.cell(i, col).fill = FAIL_FILL
        elif status == "Flag":
            for col in range(1, 6):
                ws.cell(i, col).fill = FLAG_FILL; ws.cell(i, col).font = FLAG_FONT
        elif status == "N/A":
            for col in range(1, 6):
                ws.cell(i, col).font = NA_FONT
    for cw, w in [("A", 10), ("B", 46), ("C", 9), ("D", 132), ("E", 30)]:
        ws.column_dimensions[cw].width = w
    ws.freeze_panes = "A2"

    SR = json.loads((HERE / "batch_INJ-07-2026-W3_source_reads.json").read_text())

    # ---- Sheet 2: mass chain ----
    ws2 = wb.create_sheet("Mass Chain")
    ws2.append(["Offload", "Railcar", "Full (lb)", "Empty (lb)", "Tractor/Trailer", "Net (kg)",
                "Unc. Adj (kg)", "Sparging (kg)", "Filtration (kg)", "MOBI (kg)", "COBB (kg)", "Match"])
    for c in ws2[1]:
        c.font = BOLD; c.fill = HEADER_FILL; c.alignment = WRAP; c.border = BORDER
    cobb = {r["consumed_lot"]: r["offloaded_quantity_kg"] for r in SR["cobb"]["rows"]}
    tix = {p["offload"]: p for p in SR["scale_tickets"]["pairs"]}
    for o in SR["ops_notes"]["offload_table"]:
        t = tix[o["id"]]
        ws2.append([o["id"].replace("Rail_", ""), o["railcar"],
                    t["full_ticket"]["gross_weight_lb"], t["empty_ticket"]["gross_weight_lb"],
                    f"{t['full_ticket']['tractor']}/{t['full_ticket']['trailer']} = "
                    f"{t['empty_ticket']['tractor']}/{t['empty_ticket']['trailer']}",
                    round(o["cargo_kg"], 5), o["uncertainty_adj_kg"],
                    round(o["sparging_loss_kg"], 5), round(o["filtration_loss_kg"], 5),
                    round(o["mobi_kg"], 5), round(cobb[o["id"]], 5), "ties"])
    r = SR["ops_notes"]["rollup_section4"]
    ws2.append(["TOTAL", "", "", "", "", round(r["deliveries_raw_downstream_scale_tickets_kg"], 5),
                r["uncertainty_adj_kg"], round(r["sparging_loss_kg"], 6), round(r["filtration_loss_kg"], 6),
                round(r["mobi_kg"], 5), round(D["cobb_offloaded_total_kg"], 5), "ties"])
    for c in ws2[ws2.max_row]:
        c.font = BOLD
    for i in range(1, ws2.max_row + 1):
        for col in range(1, 13):
            ws2.cell(i, col).alignment = WRAP; ws2.cell(i, col).border = BORDER
    for cw, w in [("A", 34), ("B", 13), ("C", 11), ("D", 11), ("E", 22), ("F", 14),
                  ("G", 12), ("H", 13), ("I", 13), ("J", 14), ("K", 14), ("L", 9)]:
        ws2.column_dimensions[cw].width = w

    # ---- Sheet 3: Certify reconciliation ----
    ws3 = wb.create_sheet("Certify Recon")
    ws3.append(["Component / datapoint", "Reconstructed from source (show the work)", "Certify", "Match"])
    for c in ws3[1]:
        c.font = BOLD; c.fill = HEADER_FILL; c.alignment = WRAP; c.border = BORDER
    MATRIX = [
     ("BOL aggregate SHPX215079", "25,940 + 26,200 + 26,190 (page 3 weighbridge Masse nette)", "78,330 kg", "ties"),
     ("BOL aggregate GPRX5102", "25,600 + 26,050 + 26,040 (page 3 weighbridge Masse nette)", "77,690 kg", "ties"),
     ("Railcar balance check (FC11)",
      f"SHPX215079 discharged {D['shpx215079_total_offloaded_kg']} kg vs 78,330 BOL = {D['railcar_balance_SHPX215079_kg']} kg; "
      f"GPRX5102 {D['railcar_balance_GPRX5102_kg']} kg. Neither overdraws.",
      "all 4 Unc. Adj cells = 0", "no back-calc triggered"),
     ("Bio-Oil Mass - AECN via Rail", f"sum of 4 ticket nets = {D['scale_ticket_net_kg_total']} kg", "78.317 t", "ties"),
     ("Mass of Injectant",
      f"{D['scale_ticket_net_kg_total']} - {D['sparging_loss_recalc_kg']} sparging - {D['filtration_loss_recalc_kg']} filtration "
      f"= {D['mobi_recalc_kg']} kg", "77.156 t", "ties"),
     ("Carbon Content", "SGS SR26-04478.001 carbon 40.2 % (m/m), ASTM D5291 Method C", "40.20 %", "ties"),
     ("Injection sequestration", f"77.15615445 x 0.402 x 3.667 = {D['gross_co2e_recalc_t']} t", "113.739 t", "ties"),
     ("Pyrolysis process",
      f"78.31719472 x 0.104 = {D['pyrolysis_process_recalc_t']} t (sheet final factor 0.1040167408 would give "
      f"{D['pyrolysis_process_at_sheet_final_ef_t']} t)", "8.145 t", "ties (factor rounded, see 6.4)"),
     ("Pyrolysis embodied", f"78.31719472 x 0.02116 = {D['pyrolysis_embodied_recalc_t']} t", "1.657 t", "ties"),
     ("AECN->SOPOR truck PROCESS",
      f"78.31719472 t x 107.49722 mi x 0.0001143 = {D['aecn_sopor_truck_process_recalc_t']} t", "0.962 t", "ties"),
     ("SOPOR->Opelousas rail PROCESS",
      f"78.31719472 t x 2199 mi x 0.000025910374 = {D['rail_process_recalc_t']} t", "4.462 t", "ties"),
     ("Opelousas->Basco truck PROCESS",
      f"78.31719472 t x 42.9 mi x 0.0001143 = {D['opel_basco_truck_process_recalc_t']} t", "0.384 t", "ties"),
     ("AECN->SOPOR truck EMBODIED",
      f"{D['allocated_truck_distance_mi_calc']} mi x 0.0002375605888 = {D['aecn_sopor_truck_embodied_recalc_t']} t",
      "0.077 t", "ties"),
     ("SOPOR->Opelousas rail EMBODIED",
      f"{D['allocated_rail_distance_mi_calc']} mi x 0.00003113066508 = {D['rail_embodied_recalc_t']} t", "0.069 t", "ties"),
     ("Opelousas->Basco truck EMBODIED",
      f"{D['allocated_opel_basco_mi_calc']} mi x 0.0002375605888 = {D['opel_basco_truck_embodied_recalc_t']} t",
      "0.082 t", "ties"),
     ("Aggregate distance AECN->SOPOR",
      f"107.5 mi x 3 trucks x each offload's mass share of its railcar, summed = {D['allocated_truck_distance_mi_calc']} mi",
      "323.11096 mi", "ties"),
     ("Allocated distance rail",
      f"2199 mi x the same mass shares, summed = {D['allocated_rail_distance_mi_calc']} mi", "2,203.16606 mi", "ties"),
     ("Allocated distance Opelousas->Basco",
      f"42.9 mi x 2 (round trip) x 4 offloads = {D['allocated_opel_basco_mi_calc']} mi", "343.2 mi",
      "ties (not mass-allocated)"),
     ("Transport process sum", f"0.962 + 4.462 + 0.384 = {D['transport_process_sum_t']} t", "-", "basis for the BCU"),
     ("Transport embodied sum", f"0.077 + 0.069 + 0.082 = {D['transport_embodied_sum_t']} t", "-", "not BCU-eligible"),
     ("Bio-oil transport total",
      f"{D['transport_process_sum_t']} + {D['transport_embodied_sum_t']} = {D['transport_total_t']} t", "6.036 t", "ties"),
     ("BCU - Transport Process",
      f"equals the transport process sum {D['transport_process_sum_t']} t; entered as {D['bcu_constant_t']} t "
      f"({D['bcu_minus_transport_process_t']} t vs full precision)", "5.808 t", "at the cap, conservative"),
     ("Activities total",
      f"pyrolysis {D['pyrolysis_subtotal_t']} + transport {D['transport_total_t']} = {D['activities_recalc_t']} t",
      "15.838 t", "ties"),
     ("Uncertainty discount",
      f"carbon 0.4% + mass 4 x 72.57 = {D['mass_uncertainty_expected_kg']} kg propagated = "
      f"{D['uncertainty_1sigma_reconstructed_kg']} kg at 1 sigma", "627 kg", "ties within rounding"),
     ("Net CO2e removed",
      f"{D['gross_co2e_recalc_t']} + {D['bcu_constant_t']} - {D['activities_recalc_t']} - 0.627 = {D['net_calc_t']} t",
      "103.082 t", "ties"),
     ("GHG statement row",
      f"9 entries summing to {D['ghg_statement_entry_sum_t']} t; buffer {D['buffer_pct']} %",
      "103.08 t / 874.53 t", "ties"),
    ]
    for m in MATRIX:
        ws3.append(list(m)); r2 = ws3.max_row
        for col in range(1, 5):
            ws3.cell(r2, col).alignment = WRAP; ws3.cell(r2, col).border = BORDER
        if "FAIL" in str(m[-1]):
            for col in range(1, 5):
                ws3.cell(r2, col).fill = FAIL_FILL
    for cw, w in [("A", 38), ("B", 92), ("C", 26), ("D", 30)]:
        ws3.column_dimensions[cw].width = w
    return wb


if __name__ == "__main__":
    build_with_validation(
        rows=ROWS,
        source_reads_path=HERE / "batch_INJ-07-2026-W3_source_reads.json",
        out_path=HERE / "Batch_INJ-07-2026-W3_QA_Checklist.xlsx",
        build_workbook_fn=build_workbook,
        derived_values=DERIVED,
    )
    from collections import Counter
    tally = Counter(r[2] for r in ROWS if r[2] != "SECTION")
    print("tally:", dict(tally))
