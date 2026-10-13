#!/usr/bin/env python3
"""Build the Charm WODO QA checklist for batch INJ-09-2026-W1."""
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parents[2] / "skills" / "qa" / "batch-qa-aecn" / "scripts"))
from checklist_lib import build_with_validation, BuildAborted  # noqa: E402

LOG = HERE / "batch_INJ-09-2026-W1-WODO_source_reads.json"
OUT = HERE / "Charm_WODO_QA_Checklist_INJ-09-2026-W1.xlsx"

DERIVED = {
    "sum_gross_kg": 14926.0,
    "sum_net_produced_kg": 14177.2,
    "sum_net_injected_kg": 14075.28,
    "sum_standard_loss_kg": 101.92,
    "sum_tare_kg": 748.8,
    "delta_gross_kg": 0.0,
    "delta_net_produced_kg": 0.0,
    "delta_net_injected_kg": 0.0,
    "n_totes": 13,
    "lb_per_weighing_event": 2.5,
    "uncertainty_per_tote_kg": 1.13398,
    "uncertainty_expected_kg": 14.74174,
    "uncertainty_delta_kg": 0.0,
    "c_uncertainty_expected_pct": 0.2292,
    "c_uncertainty_reported_pct": 0.2292,
    "c_uncertainty_delta_pct": 0.0,
    "weighted_ci_recomputed": 0.7208894223,
    "production_emissions_recomputed_kg": 10220.19,
    "pe_feb_kg": 941.17,
    "pe_mar_kg": 3401.70,
    "pe_may_kg": 4256.77,
    "pe_jul_kg": 1620.54,
    "seq_recomputed_t": 29.575,
    "process_recomputed_t": 2.481,
    "embodied_recomputed_t": 0.151,
    "pyrolysis_recomputed_t": 10.220,
    "activities_recomputed_t": 12.852,
    "net_recomputed_t": 19.081,
    "batch_share_of_truck": 0.65,
    "batch_mileage_allocation_mi": 780.65,
    "bcu_kg": 2481,
    "standard_loss_mean_plus_1sd_kg": 7.8379,
    "longest_google_maps_route_mi": 1201,
    "pp_feb_mass_kg": 1107.4,
    "pe_sum_of_rounded_terms_kg": 10220.18,
    "mar_step1_kg": 2130.8,
    "mar_step2_kg": 3223.2,
    "may_step1_kg": 2201.8,
    "may_step2_kg": 3364.2,
    "may_step3_kg": 4489.6,
    "may_step4_kg": 5626.0,
    "pe_step1_kg": 4342.87,
    "pe_step2_kg": 8599.64,
    "pe_jul_share_ratio": 0.159,
    "pyrolysis_share_of_activities_ratio": 0.795,
    "pyrolysis_share_of_net_ratio": 0.536,
    "pp_mar_mass_kg": 4305.6,
    "pp_may_mass_kg": 6726.4,
    "pp_jul_mass_kg": 2037.8,
    "pyrolysis_at_displayed_ci_kg": 10220.34,
    "pyrolysis_share_of_activities_pct": 79.5,
    "pyrolysis_share_of_net_pct": 53.6,
    "map_route_1_mi": 1132,
    "map_route_2_mi": 1129,
    "map_route_3_mi": 1201,
    "seq_carbon_kg": 8065.14,
    "net_step1_t": 32.056,
    "net_step2_t": 19.204,
    "pe_jul_share_of_pyrolysis_pct": 15.9,
    "seq_recomputed_kg": 29574.85,
    "process_recomputed_kg": 2480.97,
    "embodied_recomputed_kg": 150.5,
    "process_exact_kg": 2480.9758,
    "process_at_1120_kg": 2481.8069,
    "bcu_short_at_1120_kg": 0.8069,
    "gross_alt_301257_kg": 1120.0,
    "mass_of_load_alt_kg": 14931,
    "tare_assumed_kg": 57.6,
    "mean_measured_empty_kg": 62.3327,
    "tare_plus_standard_loss_kg": 65.44,
    "mean_empty_plus_sd_kg": 65.4379,
    "bcu_over_kg": 0.0242,
    "inv_tare_576_wodo": 100,
    "inv_tare_576_qowv": 565,
    "inv_tare_6293_qowv": 73,
}

P, F, FL, NA = "Pass", "Fail", "Flag", "N/A"

ROWS = [
    ("SECTION 1 — BATCH FOLDER",),
    ("1.1", "Is there testing in the testing folder for the batch?", P,
     "The Testing subfolder holds 'SR26-05444 09-2026-W1 WODO.pdf' (plus the week's AECN certificate). Read "
     "from the PDF: SGS North America certificate SR26-05444.001, client ID '9-2026-W1 WODO', product "
     "'Oil - WODO', Fort Lupton CO, received 14-Sep-2026 11:45, analysed 15-Sep-2026 08:28 to 14:33. Carbon "
     "57.3 % (m/m) by ASTM D5291 Method C, hydrogen 7.4 %, nitrogen <0.75 %, kinematic viscosity at 40 C "
     "280.3 cSt, acid number 77 mg KOH/g, density at 15 C 1.200 g/cm3. Carbon content is present and "
     "unambiguous. [log:testing]",
     "VR Appendix 1"),
    ("1.2", "Are Ops Notes in the folder and fully completed?", P,
     "'INJ-09-2026-W1 Ops Notes' is in the batch folder, ten tabs. The WODO tab carries a complete 13-tote "
     "ledger with no blank cells in any populated row: tag, gross mass, tote mass, net oil produced, "
     "standard loss, net oil injected, production period, carbon intensity, production emissions and scale "
     "uncertainty are filled for every tote. Every summary cell resolves to a number with no formula errors: "
     "14926, 14177.2, 101.92, 14075.28, 0.7208894223, 10220.19352, 14.74174. Header carries batch ID "
     "INJ-09-2026-W1 and week 2026-08-29 to 2026-09-05. No cell comments on any WODO-bearing tab. "
     "[log:ops_notes_totals, totes, ops_notes_formulas]",
     "VR Appendix 1"),
    ("1.3", "Is the Bio-Oil Shipment Tracker in the batch folder and fully completed?", P,
     "Under the weekly time-block schema the shipment tracker is the Ops Notes tab 'WODO Shipments & EE "
     "Allocation'. It covers this batch completely: shipping order COLA_35 shipped 2026-09-01, 20 totes on "
     "the truck, 13 batch-relevant totes, 65% batch share, Charm CO to Basco LA one-way distance 1,201 mi, "
     "allocated miles 780.65, and the sheet's own tote-allocation check resolves to OK. "
     "[log:shipment_allocation]",
     "VR Appendix 1"),

    ("SECTION 2 — OPS NOTES",),
    ("2.1", "Is the 'Total Mass Transported' value correctly computed as the sum of the gross weight of all totes?", F,
     "Ops Notes WODO!B25 reports 14926 kg via '=SUM(B11:B23)' over the 13-tote Gross Mass column. Summing "
     "the 13 ledger rows independently (1165 + 1135 + 1111 + 1150 + 1140 + 1138 + 1179 + 1220 + 1183 + 1194 "
     "+ 1158 + 1115 + 1038) reproduces 14926 kg exactly, a delta of 0.0 kg. The rollup cell A33 ('=B25') "
     "carries the same figure, and it is the value Certify's Transport Process Mass of Load uses (see 5.3). "
     "The sum is right, but one input conflicts between sources. Tote 301257 is 1,115 kg gross in the Ops Notes "
     "(WODO!B22) and on the inventory's Q4 tab (row 725), but 1,120 kg on the inventory workbook's Oil_inventory_data "
     "tab (row 5, a 'Microsoft SQL Server Import'); all three give net 1,057.4 kg, so the difference is the tare "
     "(57.6 vs 62.6 kg). The other 12 tags agree across all three. At 1,120 kg, Total Mass Transported is 14,931 kg "
     "and the process emission 2,481.8069 kg. Fix: confirm 301257's gross from its scale record and correct whichever "
     "record is wrong; if 1,120 kg, update WODO!B22, Certify's Mass of Load and the BCU (8.1). "
     "[log:_derived_values.sum_gross_kg, _derived_values.delta_gross_kg, _derived_values.gross_alt_301257_kg, _derived_values.mass_of_load_alt_kg, "
     "_derived_values.process_at_1120_kg, ops_notes_totals, ops_notes_formulas, totes, gross_conflict_301257]",
     "Ops Notes WODO TOTAL row"),
    ("2.2", "Is the 'Total Oil Produced' value correctly computed as the sum of the net weight of oil in tote for all totes (gross minus tare per tote)?", FL,
     "Ops Notes reports 14177.2 kg at D25 ('=SUM(D11:D23)'); each per-tote Net Oil Produced cell is a live "
     "'=B-C' formula, Gross minus the 57.6 kg tote tare. Independently: 14926 - (13 x 57.6 = 748.8) = "
     "14177.2 kg, a delta of 0.0 kg. The same 57.6 kg tare appears per tote in the production inventory "
     "master for all 13 tags. This figure is what Certify's Pyrolysis component uses as its mass (see 7.4). The "
     "flag is the tare input: 57.6 kg is a standard figure, not a per-tote weighing. The Tote Loss Dataset states it "
     "outright ('Tote tare assumed (kg)' = 57.6, tab 'Charm WODO 2026' I10), the inventory master's Tare Weight column "
     "carries 57.6 kg on 100 WODO, 565 QOWV and 37 Whole Oil totes (62.93 kg on 73 QOWV and 62.9 kg on 3 Whole Oil totes), "
     "and the Ops Notes ledger header reads 'tote tag mass is authoritative'. It reaches Total Oil Produced, the pyrolysis "
     "mass, only: the Standard Loss is measured against the same assumed tare (loss = measured empty weight minus 57.6), "
     "so in Total Oil in Batch it cancels (57.6 + 7.84 = 65.44 kg against a mean measured empty weight of 62.3327 kg + "
     "3.1052 kg = 65.4379 kg). Fix: enter a measured empty weight per tote, or attach the basis for the 57.6 kg figure. "
     "[log:_derived_values.sum_net_produced_kg, _derived_values.sum_tare_kg, _derived_values.delta_net_produced_kg, "
     "_derived_values.tare_plus_standard_loss_kg, _derived_values.mean_empty_plus_sd_kg, ops_notes_formulas, tote_tare_methodology, "
     "tote_tare_methodology.tare_distribution, tote_tare_methodology.tote_loss_dataset_basis]",
     "Ops Notes WODO TOTAL row"),
    ("2.3", "Is the 'Total Oil in Batch' value correctly computed (WODO: net-in-tote less Standard Loss per tote)?", P,
     "Reported 14075.28 kg at F25 ('=SUM(F11:F23)'). Each per-tote Net Oil Injected cell is '=D-E' where E "
     "is '=$L$8', the cell labelled 'WODO Standard Loss (kg/tote), 2026 = Mean + 1 SD' = 7.84 kg, "
     "hyperlinked to the Tote Loss Dataset. That dataset was opened this session: tab 'Charm WODO 2026', 43 "
     "totes, average loss 4.7327 kg, sample (n-1) standard deviation 3.1052 kg, stated factor 7.84 kg/tote; "
     "4.7327 + 3.1052 = 7.8379, which rounds to 7.84 (the population SD would give 7.80, so the sample SD is the basis). Standard Loss IS applied, correctly for the WODO branch. Summed "
     "independently: 14177.2 - (7.84 x 13 = 101.92) = 14075.28 kg, a delta of 0.0 kg, and the sheet's own "
     "Standard Loss column total E25 is 101.92 kg. The standard tare flagged in 2.2 cancels out of this figure. [log:_derived_values.sum_net_injected_kg, "
     "_derived_values.sum_standard_loss_kg, _derived_values.delta_net_injected_kg, "
     "_derived_values.standard_loss_mean_plus_1sd_kg, standard_loss_methodology, ops_notes_formulas]",
     "Ops Notes WODO TOTAL row / Tote Loss Dataset"),
    ("2.4", "Is uncertainty correctly computed as 2.5 lb per tote, summed across totes, converted to kg?", P,
     "Per-tote Scale Uncertainty J11:J23 is '=2.5*$L$4' with L4 = 0.453592 kg/lb, giving 1.13398 kg per "
     "tote; J25 '=SUM(J11:J23)' = 13 x 1.13398 = 14.74174 kg. This is the exact Mass of Product uncertainty "
     "Certify shows on the Sequestration component (plus or minus 14.74174 kg) and, rounded for display, on "
     "the Transport Process and Pyrolysis components (plus or minus 14.74 kg). Charm per-tote precision, not "
     "AECN's 72.57 kg per-truck figure. [log:_derived_values.uncertainty_per_tote_kg, "
     "_derived_values.uncertainty_expected_kg, _derived_values.uncertainty_delta_kg, "
     "_derived_values.lb_per_weighing_event, _derived_values.n_totes, ops_notes_formulas]",
     "Ops Notes WODO!J11:J25, L4"),

    ("SECTION 3 — BIO-OIL SHIPMENT TRACKER",),
    ("3.1", "Are totes for the RP correctly identified?", P,
     "All 13 WODO totes for the week are identified on the 'WODO Shipments & EE Allocation' tab as "
     "batch-relevant against shipping order COLA_35 (D6 '=COUNTA(WODO!A11:A23)' = 13), the same 13 tags on "
     "the WODO ledger. Independently confirmed in the production inventory master ('BGN & DC Production "
     "Bio-Oil Inventory', tab 'Q4 2025 -2026 Production Invent'): exactly 20 rows carry Ship Status COLA_35, "
     "date loaded and shipped 2026-09-01; 13 are WODO and are these 13 tags, the other 7 are June 2026 "
     "QOWV totes (301224, 301221, 301295, 301213, 301247, 301296, 301217) that are not in this batch. "
     "[log:shipment_allocation, shipment_allocation._primary_source_opened, totes]",
     "Ops Notes allocation tab / inventory master"),
    ("3.2", "Do all totes identified for RP match the Ops Notes for the applicable batch?", P,
     "The allocation tab's batch-relevant count is a live formula over the WODO ledger itself, so the two "
     "agree by construction (13 = 13, check cell B16 resolves to OK). Against the independent inventory "
     "master, every one of the 13 Ops Notes tags appears with Ship Status COLA_35 and fraction WODO. "
     "[log:shipment_allocation, totes]",
     "WODO Shipments & EE Allocation tab"),
    ("3.3", "Are computed outputs (RP Share of Truck, RP Allocated Miles, Batch Mileage Allocation) auditable formulas rather than hard-coded values?", P,
     "Every computed cell was opened and its formula read, not inferred from appearance. Batch-Relevant "
     "Totes D12 = '=$D$6' (D6 '=COUNTA(WODO!A11:A23)'). Batch Share of Truck E12 = '=IFERROR(D12/C12,\"\")'. "
     "Mileage F12 = '=$B$8'. Allocated Miles G12 = '=IFERROR(E12*F12,\"\")'. Total G14 = '=SUM(G12:G13)'. "
     "The only typed values in the chain are the inputs: 20 totes on truck (C12) and the 1,201 mi one-way "
     "distance (B8), both evidenced elsewhere (inventory master; routing map). [log:shipment_allocation]",
     "Auditability"),
    ("3.4", "Are computed outputs computed correctly?", P,
     "Reproduced against the spec. Batch Share of Truck = 13 batch-relevant totes / 20 totes on COLA_35 = "
     "0.65, matching E12. Allocated Miles = 0.65 x 1,201 = 780.65 mi, matching G12 and G14. One-way distance "
     "throughout, no round-trip doubling. The 20-tote denominator is confirmed by the inventory master's 20 "
     "COLA_35 rows. [log:_derived_values.batch_share_of_truck, _derived_values.batch_mileage_allocation_mi, "
     "shipment_allocation]",
     "WODO Shipments & EE Allocation tab"),
    ("3.5", "Do the totes assigned to each batch match the Ops Notes on (a) tag, (b) fraction, (c) gross mass, (d) production period?", P,
     "Per-tote, per-attribute comparison of the Ops Notes ledger against the inventory master's COLA_35 "
     "rows: all 13 tags present; all 13 fraction WODO; all 13 gross masses identical (300055 1165, 301072 "
     "1135, 301076 1111, 301095 1150, 301097 1140, 301140 1138, 301143 1179, 301165 1220, 301181 1183, "
     "301190 1194, 301197 1158, 301257 1115, 301302 1038); all 13 production periods identical (300055 "
     "February 2026; 301072/301076/301095/301097 March 2026; 301140/301143/301165/301181/301190/301197 May "
     "2026; 301257/301302 July 2026). Zero mismatches. [log:totes, shipment_allocation]",
     "Inventory master vs Ops Notes"),

    ("SECTION 4 — CERTIFY: INJECTION SEQUESTRATION",),
    ("4.1", "Does the C content in the testing match the C content in the Sequestration component?", P,
     "Certificate SR26-05444.001 reports carbon 57.3 % (m/m), read from the PDF. The Certify sequestration "
     "modal shows Carbon Content of Bio-oil Injected 57.3 %, and the Ops Notes Summary tab carries 0.573 "
     "for WODO against COA SR26-05444.001. Exact match. Recomputed sequestration: 14075.28 x 0.573 = 8065.14 kg carbon; "
     "8065.14 x 3.667 = 29574.85 kgCO2e = 29.575 t, Certify 29.575 t. [log:testing, certify.components.0, "
     "_derived_values.seq_recomputed_t, _derived_values.seq_carbon_kg, _derived_values.seq_recomputed_kg, ops_notes_totals]",
     "Certify cmp_1M2PJXHXC1S013HZ"),
    ("4.2", "Is the C content in the Sequestration component properly evidenced (test copy plus Ops Notes copy)?", P,
     "The carbon content datapoint shows 2 sources in the modal; the Sources tab and calculation-view "
     "datapoint list name them as 'INJ-09-2026-W1 Ops Notes - 2026-09-18' and 'SR26-05444 09-2026-W1 "
     "WODO.pdf', both linked to the specific datapoint. The attached Ops Notes copy was downloaded from "
     "Certify and matches the live sheet cell for cell on the Summary, WODO, allocation and Whole Oil tabs. "
     "[log:certify.attachment_map, certify.certify_sources_enumeration, certify.attached_ops_notes_copy_verification]",
     "Certify Sources tab"),
    ("4.3", "Does the Mass of Product measurement uncertainty match the uncertainty value in Ops Notes?", P,
     "Read from the modal INPUTS: Mass of Injectant 14,075.28 kg plus or minus 14.74174 kg. Ops Notes J25 = "
     "14.74174 kg (1.13398 x 13). Exact match. [log:certify.components.0, "
     "_derived_values.uncertainty_expected_kg, _derived_values.uncertainty_per_tote_kg]",
     "Certify cmp_1M2PJXHXC1S013HZ INPUTS"),
    ("4.4", "Does the C content have the appropriately computed 0.4 % relative uncertainty?", P,
     "Certify shows 57.3 plus or minus 0.2292 %. Expected 0.004 x 57.3 = 0.2292 percentage points. Exact "
     "match, delta 0.0. The Ops Notes Summary tab derives the same figure (G19 '=E19*H19' = 0.573 x 0.004 = "
     "0.002292). [log:certify.components.0, _derived_values.c_uncertainty_expected_pct, "
     "_derived_values.c_uncertainty_reported_pct, _derived_values.c_uncertainty_delta_pct, ops_notes_totals]",
     "Certify cmp_1M2PJXHXC1S013HZ INPUTS"),

    ("SECTION 5 — CERTIFY: TRANSPORT PROCESS EMISSIONS",),
    ("5.1", "Does the Transport Process emission use the correct and up-to-date Flatbed Truck emission factor?", P,
     "The modal shows 0.1384 kgCO2e per tonne-mile on the datapoint 'Emissions Factor - Truck Transport "
     "Fuel Use (Well to Wheel) - Flatbed Truck'. The live Standard Emission Factors 2026 tab, exported this "
     "session, carries 'Truck Transport Fuel Use (Well to Wheel) -- Flatbed Truck' = 0.0001384 MT CO2e/t-mi "
     "(GLEC V3.2 2025, status Updated) = 0.1384 kg/t-mi. Flatbed row, not tanker. [log:certify.components.1, standard_efs_2026]",
     "Standard EFs 2026 tab"),
    ("5.2", "Is the Transport Process distance value correct and evidenced by a Google Maps route where the longest given value is used?", P,
     "Distance in Certify is 1,201 mi ('Distance from Miniforge CO to Basco'), sourced to 'CO to Basco "
     "Routing Map.png'. That image was downloaded from Certify and read: three Google Maps options to 4675 "
     "LA-115, St Landry, LA 71367, read from the image pixels (screenshots/CO_to_Basco_Routing_Map.png, not text-extractable): "
     "via US-287 S 1132 mi, via US-287 S and I-20 E 1129 mi, via I-70 E 1201 mi. 1201 is the longest and is the value used. The map's origin field reads 'Your location' with "
     "the pin at the Denver/Fort Lupton area rather than a typed Fort Lupton address; this is the same reused "
     "upload (src_1KSTWP2AP1S0QQPF) that backs every 2026 WODO removal for this lane, whose Fort Lupton origin "
     "was accepted on INJ-07-2026-W1. [log:google_maps_options, google_maps_note, _derived_values.map_route_1_mi, _derived_values.map_route_2_mi, _derived_values.map_route_3_mi, "
     "_derived_values.longest_google_maps_route_mi, certify.components.1]",
     "CO to Basco Routing Map.png"),
    ("5.3", "Does the Mass of Load match the Total Mass Transported value in Ops Notes?", P,
     "Mass of Load in Certify is 14,926 kg on the datapoint 'Transported CO to Basco - WODO - Mass', exactly "
     "the Ops Notes Total Mass Transported (B25 = 14926) - the gross basis including tote tare, not Total Oil "
     "Produced (14,177.2) or Total Oil in Batch (14,075.28). The datapoint is now correctly named for oil "
     "type, basis and lane (the 'Trip Count Source' mislabel seen on the last two WODO batches is gone). "
     "Recomputed: 14.926 t x 1,201 mi x 0.1384 kg/t-mi = 2,480.9758 kgCO2e = 2.481 t, Certify 2.481 t. "
     "[log:certify.components.1, _derived_values.process_recomputed_t, _derived_values.process_recomputed_kg, _derived_values.process_exact_kg, ops_notes_totals]",
     "Certify cmp_1M2PJXHXC1S0MRDM"),
    ("5.4", "Is the Mass of Load evidenced by a copy of the batch Ops Notes?", P,
     "The Mass of Load datapoint shows 1 source: 'INJ-09-2026-W1 Ops Notes - 2026-09-18', linked to the "
     "specific datapoint per the Sources tab expansion. [log:certify.attachment_map, certify.certify_sources_enumeration]",
     "Certify Sources tab"),
    ("5.5", "Does the uncertainty reported for Mass of Load match the uncertainty value in Ops Notes?", P,
     "Read from the modal INPUTS: Mass of Load 14,926 kg plus or minus 14.74 kg, the display-rounded form of "
     "the same 14.74174 kg carried on the Sequestration component - correct, since it is the same 13 totes "
     "weighed once each. [log:certify.components.1, _derived_values.uncertainty_expected_kg]",
     "Certify cmp_1M2PJXHXC1S0MRDM INPUTS"),

    ("SECTION 6 — CERTIFY: TRANSPORT EMBODIED EMISSIONS",),
    ("6.1", "Does the Transport Embodied emission use the correct and up-to-date Non-Tanker Truck emission factor?", P,
     "The modal shows 0.19279 kgCO2e per mile on 'Emissions Factor - Vehicle Embodied Emissions - "
     "Non-Tanker Truck'. The live 2026 tab carries 'Vehicle Embodied Emissions --Non-Tanker Truck' = "
     "0.0001927908256 MT CO2e/mi (GLEC V3.2 2025) = 0.19279 kg/mi. Label and value agree; non-tanker row, "
     "per-mile basis. [log:certify.components.2, standard_efs_2026]",
     "Standard EFs 2026 tab"),
    ("6.2", "Does the aggregate distance transported match the Batch Mileage Allocation in the RP Bio-Oil Shipment Tracker?", P,
     "Certify aggregate distance ('Aggregate Distance Transported CO to LA - WODO') is 780.65 mi. The "
     "allocation tab's Allocated Miles G12/G14 resolve to 780.65 mi (0.65 x 1,201). Exact match. This batch "
     "shares its truck with 7 QOWV totes, so the allocation is a genuine 13/20 share rather than the full "
     "one-way distance. Recomputed embodied: 780.65 mi x 0.19279 kg/mi = 150.5 kgCO2e = 0.151 t, Certify "
     "0.151 t. [log:certify.components.2, _derived_values.batch_mileage_allocation_mi, "
     "_derived_values.embodied_recomputed_t, _derived_values.embodied_recomputed_kg, shipment_allocation]",
     "Certify cmp_1M2PJXHXC1S08D99 INPUTS"),
    ("6.3", "Is the aggregate distance evidenced with the RP Bio-Oil Shipment Tracker, batch Charm Contents, and applicable route Google Map?", P,
     "The aggregate distance datapoint carries 2 sources: 'INJ-09-2026-W1 Ops Notes - 2026-09-18' (which "
     "embeds the WODO Shipments & EE Allocation tracker that derives 780.65 mi) and 'CO to Basco Routing "
     "Map.png' (the 1,201 mi route). No standalone Charm Contents sheet exists for any 2026 weekly WODO "
     "batch - confirmed absent from both the Drive folder and Certify's 6 sources - its per-period content "
     "lives in the attached Ops Notes. Both practical evidence types are attached (the gap raised on "
     "INJ-08-2026-W5 does not recur). [log:certify.attachment_map, certify.charm_contents_present, "
     "certify.charm_contents_note]",
     "Certify Sources tab"),

    ("SECTION 7 — CERTIFY: PRODUCTION EMISSIONS",),
    ("7.1", "Are there Production Emissions components for ALL Production Periods in the batch?", P,
     "The batch draws on four production periods: February 2026 (1 tote, CI 0.8498961226), March 2026 (4 "
     "totes, CI 0.790064533), May 2026 (6 totes, CI 0.6328458891) and July 2026 (2 totes, H2 2026 blended "
     "CI 0.795241). All four are captured in one aggregated component, 'Pyrolysis - Charm WODO Emissions' "
     "at 10.220 tCO2e. Per-period reconciliation from the Net Oil Produced column: February 1107.4 kg x 0.8498961226 = "
     "941.17; March mass 1077.4 + 1053.4 = 2130.8, 2130.8 + 1092.4 = 3223.2, 3223.2 + 1082.4 = 4305.6 kg, and 4305.6 x "
     "0.790064533 = 3401.70; May mass 1080.4 + 1121.4 = 2201.8, 2201.8 + 1162.4 = 3364.2, 3364.2 + 1125.4 = "
     "4489.6, 4489.6 + 1136.4 = 5626.0, 5626.0 + 1100.4 = 6726.4 kg, and 6726.4 x 0.6328458891 = 4256.77; July "
     "mass 1057.4 + 980.4 = 2037.8 kg, and 2037.8 x 0.795241 = 1620.54. Sum: 941.17 + 3401.70 = 4342.87, "
     "4342.87 + 4256.77 = 8599.64, 8599.64 + 1620.54 = 10220.18 kgCO2e from rounded terms, 10220.19 unrounded = Ops Notes I25 (10220.19352) = Certify 10.220 t. No period omitted; "
     "aggregation into one weighted component is not a finding. [log:production_periods, "
     "_derived_values.pe_feb_kg, _derived_values.pe_mar_kg, _derived_values.pe_may_kg, _derived_values.pe_jul_kg, "
     "_derived_values.pe_sum_of_rounded_terms_kg, _derived_values.mar_step1_kg, _derived_values.mar_step2_kg, _derived_values.may_step1_kg, _derived_values.may_step2_kg, _derived_values.may_step3_kg, _derived_values.may_step4_kg, _derived_values.pe_step1_kg, _derived_values.pe_step2_kg, _derived_values.pp_feb_mass_kg, _derived_values.pp_mar_mass_kg, _derived_values.pp_may_mass_kg, _derived_values.pp_jul_mass_kg, totes, "
     "_derived_values.production_emissions_recomputed_kg, certify.components.3]",
     "Certify Components tab"),
    ("7.2", "Does the CI match the Ops Notes weighted-average CI across all totes in the injection?", FL,
     "The CI matches; the Flag is carried for datapoint-label hygiene only. Certify's modal shows 0.7209 "
     "kgCO2e/kg (calculation view 0.72); the Ops Notes weighted average H25 = I25/D25 = 10220.19352 / 14177.2 "
     "= 0.7208894223, which rounds to 0.7209. At full precision 14,177.2 kg x 0.7208894223 = 10,220.19 kg and "
     "at the displayed 0.7209 it is 14,177.2 x 0.7209 = 10,220.34 kg; both round to Certify's 10.220 t. Each period's CI is applied to the right totes, confirmed by reading the per-tote lookup "
     "formulas: 300055 '=$L$14' (February 0.8498961226), the four March totes '=$L$15' (0.790064533), the "
     "six May totes '=$L$17' (0.6328458891), the two July totes '=$L$19' -> '$L$7' (WODO 6-Mo CI Jul-Dec "
     "2026, 0.795241). Every monthly CI was verified against its own PP Summary (Feb T2 0.8498961226, Mar T2 "
     "0.790064533, May T2 0.6328458891) and the blended CI against the Jan-Jun 2026 mass-weighted CI "
     "workbook (D18 0.795241059). Pre-July totes correctly keep their monthly CI; only the July-produced "
     "totes take the blended figure. Flag because the two pyrolysis datapoints are titled just 'Carbon "
     "emission factor' and 'Mass' with no oil type, lane or period - every other datapoint on this removal, "
     "and the equivalent datapoints on INJ-08-2026-W5 ('WODO Tote - Carbon Intensity', 'WODO Tote - Mass "
     "(Gross-Tote) for pyrolysis'), are descriptively named. A verifier can still confirm the values from "
     "the attached Ops Notes, so verification is not prevented. Whether the two names were typed or are Certify's "
     "defaults for this component type was not determinable from the modal; if defaults, the fix sits with "
     "Isometric rather than the data entry. The same label gap is cross-referenced at 7.4. [log:_derived_values.weighted_ci_recomputed, "
     "ops_notes_totals, production_periods, pp_summary_sources, totes, certify.components.3, certify.datapoints, _derived_values.pyrolysis_at_displayed_ci_kg]",
     "Ops Notes WODO TOTAL row / PP Summaries"),
    ("7.3", "Is the CI evidenced (Ops Notes weighted-average across totes)?", FL,
     "The evidence standard is met for 11 of 13 totes and thinly for the other 2. The 'Carbon emission factor' "
     "datapoint carries 1 source, 'INJ-09-2026-W1 Ops Notes - 2026-09-18', whose WODO tab holds the per-tote "
     "CI column, the weighted-average cell H25, and a per-month CI lookup table in which every monthly row "
     "K11:L18 carries a 'PP Summary' hyperlink in column M. The blended row K19/L19 (0.795241, applied to the "
     "two July 2026 totes: mass 1057.4 + 980.4 = 2037.8 kg, 2037.8 x 0.795241 = 1620.54 kgCO2e, and 1620.54 / 10220.19 = 0.159, so 15.9% of the pyrolysis term) has M19 empty, and the Jan-Jun 2026 "
     "mass-weighted CI workbook that produces 0.795241059 is not among this removal's 6 Certify sources - "
     "unlike INJ-08-2026-W5, where the CI datapoint carried that workbook plus two PP Summaries alongside the "
     "Ops Notes. The figure was verified this session from the workbook itself (CI Calculation!D18), so the "
     "value is right and verification is not prevented; the trace for the blended CI just lives outside the "
     "attached evidence. Fix: hyperlink the workbook in Ops Notes M19 or attach it to the datapoint. "
     "[log:certify.attachment_map, certify.certify_sources_enumeration, ops_notes_totals, pp_summary_sources, "
     "production_periods, _derived_values.pe_jul_kg, _derived_values.pp_jul_mass_kg, _derived_values.production_emissions_recomputed_kg, _derived_values.pe_jul_share_of_pyrolysis_pct, _derived_values.pe_jul_share_ratio]",
     "Certify Sources tab"),
    ("7.4", "Does the mass match the sum of Net Weight of Bio-Oil in Tote values for the PP (Ops Notes)?", P,
     "Certify mass is 14,177.2 kg (datapoint 'Mass'). The Ops Notes sum of Net Oil Produced across all 13 "
     "totes, D25, is 14177.2 kg, a delta of 0.0 kg. The component correctly uses net-in-tote (produced "
     "mass), not net-in-injection-batch (14,075.28). The datapoint is titled only 'Mass'; that label gap is "
     "carried as the 7.2 Flag. This mass rests on the standard 57.6 kg tare flagged in 2.2. [log:_derived_values.sum_net_produced_kg, "
     "_derived_values.delta_net_produced_kg, certify.components.3]",
     "Certify cmp_1M2QZVM7P1S05PA8 INPUTS"),
    ("7.5", "Does the uncertainty match the sum of uncertainties for the PP (Ops Notes)?", P,
     "Read from the modal INPUTS: Mass 14,177.2 kg plus or minus 14.74 kg, the display-rounded 14.74174 kg "
     "(1.13398 x 13). Uncertainty is present on this dominant input - pyrolysis is 10.220 / 12.852 = 0.795 (79.5%) of Activities "
     "and 10.220 / 19.081 = 0.536 (53.6%) of net - so no v1.2 6.5 omission question arises. [log:certify.components.3, "
     "_derived_values.uncertainty_expected_kg, _derived_values.pyrolysis_recomputed_t, _derived_values.activities_recomputed_t, "
     "_derived_values.net_recomputed_t, _derived_values.pyrolysis_share_of_activities_pct, _derived_values.pyrolysis_share_of_net_pct, _derived_values.pyrolysis_share_of_activities_ratio, _derived_values.pyrolysis_share_of_net_ratio]",
     "v1.2 6.5.1"),
    ("7.6", "Is the mass evidenced with a copy of the Charm Contents sheet for the batch?", P,
     "No Charm Contents sheet exists for this batch or any 2026 weekly WODO batch (absent from the Drive "
     "folder's four items and from Certify's 6 sources). Its per-period mass breakdown lives in the Ops "
     "Notes WODO tab (Production Period column against Net Oil Produced), and that Ops Notes is attached to "
     "the 'Mass' datapoint. Same basis accepted on every prior 2026 weekly WODO/Aqueous batch. "
     "[log:certify.charm_contents_present, certify.charm_contents_note, certify.attachment_map]",
     "VR Appendix 1"),

    ("SECTION 8 — OTHER",),
    ("8.1", "Are all diesel transport-process emissions offset by BCUs, and does the BCU total equal (never exceed) the transport-process sum?", F,
     "One diesel transport-process emission on this removal, the flatbed leg at 2.481 tCO2e (14.926 t x "
     "1,201 mi x 0.1384 kg/t-mi = 2,480.9758 kg). Reductions carries 'BCU - Transport Process' with a Constant "
     "CO2 reduction of 2,481 kgCO2e. The BCU exceeds the emission it offsets: at full precision the process emission is "
     "14.926 t x 1,201 mi x 0.1384 kg/t-mi = 2,480.9758 kg, so the constant is 0.0242 kg over. The BCU QA Rule reads 'BCU "
     "quantity (tCO2e) in Reductions must NOT exceed the corresponding diesel/process emission (tCO2e) in Activities', so "
     "over-coverage of any size is a Fail. Fix: enter the constant as 2,480.9758 kg; if 301257's gross is 1,120 kg (2.1), the "
     "process emission is 2,481.8069 kg and the constant is 0.8069 kg short, so enter 2,481.8069 kg instead. Net stays "
     "19.081 t either way. Evidenced by 'Charm-RetirementCertificate-RD-EACs-"
     "20251218.pdf': TERC Certificate of Retirement, Amazon.com Services LLC retiring on behalf of Charm "
     "Industrial Inc, Dec 18 2025, vintage 2024, Renewable Diesel TERCS, 2,795 TERCs, verification body "
     "Turner Mason & Company, certificate TERC1012_2024_DF_RENEWABLE_DIESEL-13-47206-50000 - read from this "
     "removal's own attached copy, downloaded from Certify this session (source src_1KVRCCD7J1S087E3, "
     "byte-identical to the copy read on INJ-08-2026-W2) - the shared "
     "inventory certificate drawn down across 2026 WODO/QOWV batches. Embodied emissions correctly not "
     "offset. Net check: 29.575 + 2.481 = 32.056; 32.056 - 12.852 = 19.204; 19.204 - 0.123 = 19.081 t, Certify 19.081 t. "
     "[log:_derived_values.bcu_kg, _derived_values.process_recomputed_t, _derived_values.process_recomputed_kg, _derived_values.net_recomputed_t, _derived_values.net_step1_t, _derived_values.net_step2_t, "
     "_derived_values.process_exact_kg, _derived_values.bcu_over_kg, _derived_values.process_at_1120_kg, _derived_values.bcu_short_at_1120_kg, bcu_vs_process, gross_conflict_301257, standard_efs_2026, retirement_certificate, certify.components.4, certify.calculation_view]",
     "BCU QA Rule (batch-qa-charm SKILL.md)"),

    ("SECTION 9 — AQUEOUS BUFFERING (LCS)",),
    ("9.1", "If LCS use is reported in Ops Notes, does a matching Certify component exist?", NA,
     "Not applicable to a WODO batch - LCS is an aqueous pre-treatment. The WODO and allocation tabs carry no "
     "LCS/caustic/buffer column, and the 5 enumerated Certify components include no LCS component. "
     "[log:ops_notes_totals, certify.components]",
     "Aqueous-only item"),
]

BOLD = Font(bold=True, size=11)
BODY = Font(size=10)
TITLE = Font(bold=True, size=14)
HDR = PatternFill("solid", fgColor="D6E4F0")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
FLAG_FONT = Font(size=10, color="996600")
NA_FONT = Font(size=10, italic=True, color="808080")

HEADER_ROW = 12


def build_workbook(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WODO QA INJ-09-2026-W1"

    ws.cell(row=1, column=1, value="Charm WODO Injection Batch QA Checklist").font = TITLE
    meta = [
        ("Batch ID", "INJ-09-2026-W1 (WODO)"),
        ("Reporting Period", "Certify removal window 30 Aug 2026 to 04 Sep 2026; Ops Notes injection week 2026-08-29 to 2026-09-05"),
        ("Reviewer", "Claude (MRV QA agent)"),
        ("Date", "2026-09-21 (revised 2026-09-29: standard tare, BCU rule)"),
        ("Injection Site", "Basco 6, Louisiana (Class V UIC)"),
        ("Oil Type", "WODO (Charm-produced, Fort Lupton CO), 13 IBC totes on flatbed COLA_35 (13 of 20 totes)"),
        ("Production Periods", "February 2026 (1 tote, CI 0.8498961226); March 2026 (4 totes, CI 0.790064533); May 2026 (6 totes, CI 0.6328458891); July 2026 (2 totes, H2 2026 blended CI 0.795241)"),
        ("Certify Removal ID", "rmv_1M2PJXHXC1S04Y5N"),
        ("Certify Net", "19.081 tCO2e (displayed 19.08) = Sequestrations 29.575 + BCU 2.481 - Activities 12.852 - Uncertainty Discount 0.123"),
    ]
    r = 2
    for k, v in meta:
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v).font = BODY
        r += 1

    hdrs = ["#", "Check", "Status", "Evidence / Finding", "Reference"]
    for i, h in enumerate(hdrs, start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=h)
        c.font = BOLD
        c.fill = HDR
        c.alignment = Alignment(vertical="center")

    r = HEADER_ROW + 1
    for row in rows:
        if len(row) == 1:
            c = ws.cell(row=r, column=1, value=row[0])
            c.font = BOLD
            c.fill = HDR
            for col in range(2, 6):
                ws.cell(row=r, column=col).fill = HDR
            r += 1
            continue
        num, check, status, notes, ref = row
        vals = [num, check, status, notes, ref]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if status == "Fail":
                c.fill = FAIL_FILL
            elif status == "Flag":
                c.fill = FLAG_FILL
                c.font = FLAG_FONT
            elif status == "N/A":
                c.font = NA_FONT
        r += 1

    r += 1
    c = ws.cell(row=r, column=1, value="Sidebar Observations (no Pass/Fail status - flagged for attention)")
    c.font = BOLD
    c.fill = HDR
    for col in range(2, 6):
        ws.cell(row=r, column=col).fill = HDR
    r += 1

    sidebar = [
        "The SGS certificate marks the carbon result (57.3 %) and hydrogen result (7.4 %) with '#': "
        "'Result is outside of test method limits and/or analytical range used in method precision study' "
        "(ASTM D5291 Method C). The values are reported and used as-is; the mark is a method-precision "
        "caveat from the lab, not a rejection. Same pattern as prior WODO certificates at 55%+ carbon.",

        "The same Ops Notes workbook carries a 'Whole Oil' tab for one tote (301310, 1,173 kg gross, 1,104 "
        "kg injected, August 2026 lot) shipped on COLA_32 and using the WODO H2 2026 CI 'BORROWED for Whole "
        "Oil (no DC Whole Oil CI yet)'. That tote is not in this removal (13 WODO totes only) and no "
        "separate Certify entry for it was looked for this session; whoever QAs the Whole Oil entry should "
        "note the borrowed-CI basis.",

        "The September VVB Guidance doc exists (created 2026-09-21) but its only section is scoped to AECN "
        "Rail Week 3 (INJ-09-2026-W3) loaded-mass basis. It says nothing about WODO Week 1, the 7.84 kg/tote "
        "standard loss, the H2 2026 blended CI, or the shared-truck allocation. Not a finding for this "
        "removal, but a verifier reading only the guidance gets no pointer to any of those methods.",

        "The routing-map source (src_1KSTWP2AP1S0QQPF) is a reused upload whose ID predates this removal's "
        "components; its origin field reads 'Your location'. A regenerated screenshot with the Fort Lupton "
        "address typed as origin would remove the only ambiguity left on the distance evidence.",
    ]
    for s in sidebar:
        ws.cell(row=r, column=1, value="-").font = BODY
        c = ws.cell(row=r, column=4, value=s)
        c.font = BODY
        c.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1

    r += 1
    counts = {}
    for rr in range(HEADER_ROW + 1, r):
        v = ws.cell(row=rr, column=3).value
        if v in ("Pass", "Fail", "Flag", "N/A"):
            counts[v] = counts.get(v, 0) + 1
    total = sum(counts.values())

    summary = (
        "STATUS SUMMARY (2026-09-21 run; revised 2026-09-29): {} Pass - {} Flag - {} Fail - {} N/A  ({} items)  |  "
        "TWO FAILS. 2.1: tote 301257's gross is 1,115 kg in the Ops Notes and Q4 inventory but 1,120 kg on the Oil_inventory_data tab. "
        "8.1: the BCU (2,481 kg) is 0.0242 kg over the process emission of 2,480.9758 kg (0.8069 kg short if the gross is 1,120 kg). "
        "All five Certify components recompute to Certify's figures from primary sources: "
        "sequestration 14075.28 x 0.573 = 8065.14 kg carbon, 8065.14 x 3.667 = 29574.85 kgCO2e = 29.575 t; transport process 14.926 t x 1,201 mi x "
        "0.1384 kg/t-mi = 2.481 t; transport embodied 780.65 mi x 0.19279 kg/mi = 0.151 t; pyrolysis 14.1772 "
        "t x weighted-avg CI 0.7208894 = 10.220 t; net 29.575 + 2.481 = 32.056, less activities 12.852 = 19.204, less discount 0.123 = 19.081 t against "
        "Certify 19.081 t. Both emission factors match the live 2026 Standard EFs tab. All Ops Notes totals "
        "reconcile to a zero delta across the 13-tote ledger; the 7.84 kg/tote standard loss reproduces from "
        "the Tote Loss Dataset (mean 4.7327 + SD 3.1052). Tote roster verified per attribute against the "
        "production inventory master, which also confirms the 20-tote truck (13 WODO + 7 QOWV) behind the "
        "0.65 allocation share. The BCU (2,481 kg), evidenced by the verified Renewable Diesel TERC retirement "
        "certificate, is 0.0242 kg over the flatbed process emission of 2,480.9758 kg (8.1). Carbon uncertainty is exactly 0.4% relative "
        "and the aggregate-distance datapoint carries both Ops Notes and routing map, so neither of the "
        "INJ-08-2026-W5 Fails recurs; the Mass-of-Load datapoint is now correctly named. THREE FLAGS. Item 2.2: the 57.6 kg tote tare is a standard figure, not a per-tote weighing. Item 7.2, "
        "the pyrolysis component's two datapoints are titled just 'Carbon emission factor' and 'Mass' with "
        "no WODO/period identifier, unlike every other datapoint on the removal and unlike the prior batch's "
        "equivalents. Values correct and evidenced. Item 7.3, the blended H2 2026 CI (0.795241) "
        "applied to the two July-produced totes is a bare constant in the attached Ops Notes (lookup row M19 has "
        "no hyperlink) and the Jan-Jun 2026 mass-weighted CI workbook is not attached, unlike the prior batch; "
        "the value was verified from that workbook this session."
        .format(counts.get("Pass", 0), counts.get("Flag", 0), counts.get("Fail", 0), counts.get("N/A", 0), total)
    )
    c = ws.cell(row=r, column=1, value=summary)
    c.font = BOLD
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)   # Excel does not auto-size merged cells
    ws.row_dimensions[r].height = 15 * -(-len(summary) // 170)
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0

    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 9
    ws.column_dimensions["D"].width = 100
    ws.column_dimensions["E"].width = 32
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    return wb


try:
    build_with_validation(ROWS, LOG, OUT, build_workbook, derived_values=DERIVED)
except BuildAborted as e:
    print("BUILD ABORTED\n", e)
    raise SystemExit(1)
