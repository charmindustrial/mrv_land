import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "QA Checklist"

HEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FLAG_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FLAG_FONT = Font(color="996600")
NA_FONT = Font(italic=True, color="808080")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

for c, w in {"A": 19, "B": 60, "C": 10, "D": 95, "E": 32}.items():
    ws.column_dimensions[c].width = w

row = 1

def header_block(lines):
    global row
    for label, value in lines:
        ws.cell(row=row, column=1, value=label).font = BOLD
        ws.cell(row=row, column=2, value=value).alignment = WRAP
        row += 1
    row += 1

def section_header(title):
    global row
    ws.cell(row=row, column=1, value=title)
    for col in range(1, 6):
        ws.cell(row=row, column=col).fill = HEADER_FILL
        ws.cell(row=row, column=col).font = BOLD
    row += 1

def col_headers():
    global row
    for i, h in enumerate(["Item", "Question", "Status", "Evidence / Finding", "Reference"], start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.fill = HEADER_FILL
        cell.font = BOLD
    row += 1

def item_row(num, question, status, evidence, ref):
    global row
    ws.cell(row=row, column=1, value=num).alignment = WRAP
    ws.cell(row=row, column=2, value=question).alignment = WRAP
    ws.cell(row=row, column=3, value=status).alignment = WRAP
    ws.cell(row=row, column=4, value=evidence).alignment = WRAP
    ws.cell(row=row, column=5, value=ref).alignment = WRAP
    if status == "Fail":
        for col in range(1, 6):
            ws.cell(row=row, column=col).fill = FAIL_FILL
    elif status == "Flag":
        for col in range(1, 6):
            ws.cell(row=row, column=col).fill = FLAG_FILL
            ws.cell(row=row, column=col).font = FLAG_FONT
    elif status == "N/A":
        for col in range(1, 6):
            ws.cell(row=row, column=col).font = NA_FONT
    row += 1

header_block([
    ("Batch ID", "INJ-09-2026-W2 (QOWV / Aqueous)"),
    ("Reporting Period", "September 1 - September 30 2026; Certify removal window 05 Sep 2026 to 11 Sep 2026; Ops Notes injection week 2026-09-05 to 2026-09-12"),
    ("Reviewer", "Claude (QA agent)"),
    ("Date", "2026-09-22 (re-checked 2026-09-29 against the live Ops Notes, edited 2026-09-29 15:13 UTC, and live Certify)"),
    ("Injection Site", "Basco 6, Louisiana"),
    ("Oil Type", "Aqueous (QOWV) - Charm-produced, Fort Lupton CO, IBC totes"),
    ("Production Periods", "Ops Notes labels: July 2025 (PP8), September 2025, October 2025, December 2025, April 2026, May 2026. Inventory master: July 2025 (PP#8), August 2025 (PP#9), September 2025 (PP#10), April 2026, May 2026 (8 totes differ, see 3.5). 38 totes, one production-mass-weighted Pyrolysis component"),
    ("Certify Removal ID", "rmv_1M2TTAZYD1S0BCCS"),
    ("Net CO2e Removed", "Submitted 2026-09-22: 4.484 tCO2e (Sequestrations 9.512 + Adjustments 6.454 - Activities 11.440 - Uncertainty discount 0.041). Live 2026-09-29: 4.643 tCO2e (Activities 11.282; pyrolysis 4.650 -> 4.491 t)"),
])

col_headers()

section_header("SECTION 1 -- Batch Folder")
item_row("1.1", "Is there testing in the testing folder for the batch?", "Pass",
    "SGS Certificate of Analysis SR26-05494.001 is in the batch folder as 'SR26-05494 W2 QOWV AECN COA.pdf' and is attached in Certify. Page 1 covers this removal: Client ID 9-2026-W2-QOWV, Product Description 'Oil - QOWV', location FORT LUPTON CO, sampled 05-Sep-2026 12:00, Carbon 8.3 % (m/m) by ASTM D5291 (Method C), Hydrogen 9.9 %, Nitrogen <0.75 %, Kinematic Viscosity 0.9420 cSt, Acid Number 41 mg KOH/g, Density at 15 C 1.047 g/cm3 (above the 1.01 g/cm3 formation-fluid threshold). Page 2 is the AECN composite SR26-05494.002 (40.4 % C), which belongs to the AECN removal, not this one [log:testing.qowv_certificate, _derived_values_numeric].",
    "VR Appendix 1")
item_row("1.2", "Are Ops Notes in the folder and fully completed?", "Fail",
    "As submitted (2026-09-22 export) the QOWV tab was complete: 38 tote rows (A11:A48) with no blanks, and summary cells Total Mass Transported 38,830.4 kg, Total Oil Produced 36,641.6 kg, Total Oil Injected 31,250.9 kg, weighted CI 0.1269082907, production emissions 4,650.122824 kg, uncertainty 43.09124 kg. In the live sheet (export 2026-09-29, edited 15:13 UTC) the weighted-average CI cell H50 is blank (' '), so the sheet no longer states the CI Certify uses; I50 = 4,491.438649 kg over D50 = 36,641.6 kg gives 0.1225775798. Fix: restore H50 = SUMPRODUCT(D11:D48,H11:H48)/SUM(D11:D48) [log:ops_notes, ops_notes_totals, live_ops_notes_2026_09_29, _derived_values_numeric].",
    "VR Appendix 1")
item_row("1.3", "Is the Bio-Oil Shipment Tracker in the batch folder and fully completed?", "Pass",
    "No standalone tracker file exists for 2026 weekly batches; the allocation function is carried by the 'QOWV Shipments & EE Allocation' tab of the same Ops Notes workbook, which is the structure every 2026 weekly QOWV batch has used. The tab is complete for this batch: 2 shipping orders (Truck 1 / CO_AQUEOUS_12.01.25 and COLA_34), tote counts 40 and 20, batch-relevant totes 18 and 20 summing to the 38-tote ledger, one-way distance 1,201 mi, allocated-mile total 1,741.45 mi. Whether its Truck 1 row reflects the trucks that ran is 3.4 [log:shipment_tracker, _derived_values_numeric].",
    "VR Appendix 1")

section_header("SECTION 2 -- Ops Notes")
item_row("2.1", 'Is the "Total Mass Transported" value correctly computed as the sum of the gross weight of all totes?', "Pass",
    "QOWV!B50 = SUM(B10:B48) = 38,830.4 kg, a live sum formula over the gross-mass column, not a typed constant. Independently re-summing the 38 per-tote gross values gives 38,830.4 kg, delta 0.0 kg. The same figure is Certify's Transport Process Mass of Load [log:ops_notes_totals, _derived_values_numeric].",
    "VR Sec 5.1")
item_row("2.2", 'Is the "Total Oil Produced" value correctly computed as the sum of the net weight of oil in tote for all totes (gross - tare per tote)?', "Flag",
    "The arithmetic is right: per-tote D = B - C on every row; QOWV!D50 = SUM(D10:D48) = 36,641.6 kg, and an independent re-sum across all 38 totes gives the same, delta 0.0 kg, matching Certify's QOWV Pyrolysis Mass. The flag is the tare input. Column C (Tote Mass) is not a per-tote weighing: it reads 57.6 kg on all 38 rows, and the inventory master's Tare Weight column carries 57.6 kg on 566 QOWV totes and 62.93 kg on 73, the same figures on steel and plastic pallets, with no note on either header cell. Total Oil Produced, the pyrolysis mass, therefore rests on a standard tare; Total Oil in Batch is unaffected, since it uses the measured post-unload weight (column E). Fix: enter a measured empty weight per tote in column C, or attach the basis for the 57.6 kg standard figure [log:ops_notes_totals, inventory_master.tare_distribution, _derived_values_numeric].",
    "VR Sec 5.1")
item_row("2.3", 'Is the "Total Oil in Batch" value correctly computed (Aqueous: measured pre-unload minus measured post-unload per tote; Standard Loss does NOT apply)?', "Pass",
    "Correct for Aqueous. Per-tote F = IF(E=\"\",\"\",B-E): gross weight minus the measured post-unload empty tote weight (column E ranges 144.0-368.0 kg across the 38 rows, so it is a per-tote weighing, not the 57.6 kg nominal tote mass in column C). QOWV!F50 = SUM(F9:F48) = 31,250.9 kg; independent re-sum = 31,250.9 kg, delta 0.0 kg. No Standard Loss column, constant or deduction appears anywhere on the QOWV tab; the sheet's only loss constant (Sparging Mass Loss L5 = 0.0129) is referenced by no QOWV row, and Summary!E12/F12 report sparging 0.0 and filtration 0.0 for QOWV. The figure is Certify's Mass of Injectant [log:ops_notes_totals, _derived_values_numeric].",
    "GHG Stmt Sec C")
item_row("2.4", "Is uncertainty correctly computed as 2.5 lb per tote, summed across totes, converted to kg?", "Pass",
    "Per-tote J = 2.5*$L$4 where L4 = 0.453592 kg/lb, giving 1.13398 kg per tote on all 38 rows. QOWV!J50 = SUM(J9:J48) = 43.09124 kg. Expected 38 x 1.13398 = 43.09124 kg, delta 0.00000 kg. This is Charm's per-tote weighing precision, not AECN's 72.57 kg per-truck figure [log:ops_notes_totals, _derived_values_numeric].",
    "Isometric Standard S3.7.3")

section_header("SECTION 3 -- Bio-Oil Shipment Tracker (QOWV Shipments & EE Allocation tab)")
item_row("3.1", "Are totes for the RP correctly identified?", "Pass",
    "Batch-relevant totes across the two shipping orders are 18 (Truck 1 / CO_AQUEOUS_12.01.25) + 20 (COLA_34) = 38, matching the 38-row IBC ledger exactly. The tab's own tote count D6 = COUNTA(QOWV!A11:A48) resolves to 38, so the count is driven off the ledger rather than typed [log:shipment_tracker, _derived_values_numeric].",
    "VR Sec 3.5")
item_row("3.2", "Do all totes identified for RP match the Ops Notes for the applicable batch?", "Fail",
    "Checked per attribute, tag by tag, against the production inventory master and the receiving record (Oil_inventory_data): tag 38 of 38, fraction QOWV 38 of 38, gross mass 38 of 38 against the receiving record, production period 30 of 38 as submitted (see 3.5). Four totes also carry a Fort Lupton production weighing 0.4 kg above the Louisiana receiving weighing the ledger uses (015162 946.0 vs 945.6 kg; 245103, 245140, 245142 likewise); two weighings, not an error. The roster is complete: the Bio-oil Injection Tracker ('Basco Injection -- COBB', rows 760-797) holds 38 rows under 'INJ-09-2026-W2_aqueous_fraction', one per ledger tag. The allocation tab's counts, 18 + 20, tie to the ledger; the split of the 18 across trucks is 3.4 [log:totes, tote_inventory_crosscheck, work_orders, cobb_injection_record, shipment_tracker, _derived_values_numeric].",
    "VR Sec 3.5")
item_row("3.3", "Are computed outputs (Batch Share of Truck, Allocated Miles, Batch Mileage Allocation) auditable formulas rather than hard-coded values?", "Pass",
    "Every computed cell is a live formula referencing its inputs, read from the workbook exported from the live sheet this session: E12 '=D12/C12' and E13 '=D13/C13' (batch share), G12 '=E12*F12' and G13 '=E13*F13' (allocated miles), G15 '=SUM(G12:G13)' (batch mileage allocation), D6 '=COUNTA(QOWV!A11:A48)' (tote count), and both mileage cells F12/F13 '=$B$8' pointing at the single 1,201 mi distance input. No computed output is a typed constant [log:shipment_tracker_cells_are_formulas, shipment_tracker, _derived_values_numeric].",
    "Auditability - VR Sec 5.1")
item_row("3.4", "Are computed outputs computed correctly?", "Fail",
    "The formulas reproduce as written: Truck 1 18/40 = 0.45 share x 1,201 mi = 540.45 mi; COLA_34 20/20 x 1,201 = 1,201 mi; total 1,741.45 mi, the value Certify carries. The Truck 1 denominator is wrong. The receiving record (Oil_inventory_data, column shipping_order_number) shows 2025-12-01 as two orders of 20 totes: CO_AQUEOUS_12.01.25, carrying 9 of this batch's totes, and CO_AQUEOUS_12.1.25_T2, carrying the other 9; no order in that record holds more than 20 totes. By order the allocation is 9/20 + 9/20 + 20/20 of 1,201 mi = 2,281.9 mi, and embodied emission 0.440 t against 0.336 t: 0.104 t understated, so the net is overstated. Fix: split row 12 into CO_AQUEOUS_12.01.25 (20 totes, 9 batch totes) and CO_AQUEOUS_12.1.25_T2 (20 totes, 9 batch totes), change Certify's 'QOWV Embodied Distance (CO to B6)' to 2,281.9 mi, and re-attach the corrected Ops Notes [log:shipment_tracker, shipping_orders_2025_12, _derived_values_numeric].",
    "Transportation Emissions Acctg v1.1")
item_row("3.5", "Do the totes assigned to each batch match the Ops Notes on (a) tag, (b) fraction, (c) gross mass, (d) production period?", "Fail",
    "(a) tag 38 of 38; (b) fraction QOWV 38 of 38; (c) gross mass 38 of 38 against the receiving record (see 3.2); (d) production period 30 of 38 as submitted. Five totes were labelled 'December 2025', their Louisiana receiving date, while their production work orders date them July-September 2025 (015162 WO00001147 2025-07-28; 228490 WO00001371 2025-09-19) and the inventory places them in PP#8 (015162), PP#9 (245140) and PP#10 (228490, 245103, 245142). The live sheet (2026-09-29) now labels those five July, August and September 2025. Three remain in conflict between records: 228470 and 228485 are 'October 2025' in the Ops Notes and in their work orders (WO00001397 and WO00001399, 2025-10-08) but '2025 PP#10' in the inventory (rows 228, 230); 301108 is 'May 2026' in the Ops Notes and its work order (2026-05-01) but 'April 2026 PP' in the inventory (row 557) and the receiving record. Fix: confirm the period of 228470, 228485 and 301108 and correct whichever record is wrong [log:totes, tote_inventory_crosscheck, work_orders, live_ops_notes_2026_09_29, _derived_values_numeric].",
    "VR Sec 3.5")
section_header("SECTION 4 -- Certify: Injection Sequestration")
item_row("4.1", "Does the C content in the testing match the C content in the Sequestration component?", "Pass",
    "Modal INPUTS read directly: Carbon content of product 8.3 %. Lab certificate SR26-05494.001 reports Carbon 8.3 % (m/m). Delta 0.0 percentage points. Ops Notes Summary!E21 carries the same 8.3 % against composite SR26-05494.001 dated 9/5/2026 [log:certify.component_by_key.injection_sequestration, testing.qowv_certificate, _derived_values_numeric].",
    "VR Sec 5.1")
item_row("4.2", "Is the C content properly evidenced (test copy + Ops Notes copy)?", "Pass",
    "The Carbon Content datapoint carries 2 sources, and the Sources tab names both: 'SR26-05494 W2 QOWV AECN COA.pdf' (the lab certificate) and 'INJ-09-2026-W2 Ops Notes.xlsx'. Both evidence types are linked to the datapoint itself, not merely to the removal [log:certify.attachment_map, certify.certify_sources_enumeration].",
    "VR Appendix 1")
item_row("4.3", "Does the Mass of Product measurement uncertainty match the uncertainty value in Ops Notes?", "Pass",
    "Modal INPUTS read directly: Mass of product 31,250.9 kg plus or minus 43.09124 kg. Ops Notes uncertainty total J50 = 43.09124 kg. Exact match, carried at full precision rather than rounded [log:certify.component_by_key.injection_sequestration, ops_notes_totals, _derived_values_numeric].",
    "Isometric Standard S3.7.3")
item_row("4.4", "Does the C content have the appropriately computed 0.4% relative uncertainty?", "Pass",
    "Modal INPUTS read directly: 8.3 % plus or minus 0.0332. Expected 0.004 x 8.3 = 0.0332 percentage points absolute. Exact match, and it is entered as an absolute value rather than as '0.4'. Ops Notes Summary!G21 computes the same figure as E21*H21 with H21 = 0.004 [log:certify.component_by_key.injection_sequestration, _derived_values_numeric].",
    "Isometric Standard S3.7.3")

section_header("SECTION 5 -- Certify: Transport Process Emissions")
item_row("5.1", "Does the Transport Process emission use the correct and up-to-date Flatbed Truck emission factor?", "Pass",
    "Modal shows 'Flatbed, Northern America, WTW' at 0.086 kgCO2e/(km*t). The live Standard Emission Factors sheet, 2026 tab, row 7 'Truck Transport Fuel Use (Well to Wheel) -- Flatbed Truck' = 0.0001384 MT CO2e/t-mi (GLEC V3.2 2025), which is 0.086 kgCO2e/(km*t) in Certify's display units. The component recomputes at full precision: 38.8304 t x 1,201 mi x 0.0001384 = 6.454 t, matching Certify [log:standard_efs_2026, certify.component_by_key.transport_process, _derived_values_numeric].",
    "Standard EFs 2026 tab, row 7")
item_row("5.2", "Is the Transport Process distance value correct and evidenced by a Google Maps route where the longest given value is used?", "Pass",
    "Certify distance is 1,201 mi ('Distance from Miniforge CO to Basco'), sourced to 'CO to Basco Routing Map.png'. That image was downloaded from this removal's own Certify source and read from the pixels: three route options are displayed to destination '4675 LA-115, St Landry, LA 71367' -- via US-287 S 1,132 mi, via US-287 S and I-20 E 1,129 mi, and via I-70 E 1,201 mi. Certify uses 1,201 mi, the longest [log:google_maps_options, google_maps_note, _derived_values_numeric].",
    "Charm convention - longest route")
item_row("5.3", "Does the Mass of Load match the Total Mass Transported value in Ops Notes?", "Pass",
    "Modal Mass of load 38,830.4 kg = Ops Notes Total Mass Transported (QOWV!B50) 38,830.4 kg, delta 0.0 kg. This is the gross figure including tote tares, not Total Oil Produced (36,641.6 kg) and not Total Oil Injected (31,250.9 kg), which is the correct basis for a per-tonne-mile transport factor [log:certify.component_by_key.transport_process, ops_notes_totals, _derived_values_numeric].",
    "VR Sec 5.1")
item_row("5.4", "Is the Mass of Load evidenced by a copy of the batch Ops Notes?", "Pass",
    "The 'QOWV Transport Mass (CO to B6)' datapoint carries 1 source, 'INJ-09-2026-W2 Ops Notes.xlsx', confirmed on the Sources tab as linked to that datapoint [log:certify.attachment_map, certify.certify_sources_enumeration].",
    "VR Appendix 1")
item_row("5.5", "Does the uncertainty reported for Mass of Load match the uncertainty value in Ops Notes?", "Pass",
    "Modal shows 38,830.4 kg plus or minus 43.09 kg read from the modal, the display rounding of the same 38 x 1.13398 = 43.09124 kg the Ops Notes carries and the Sequestration component holds at full precision. Same 38 totes, same weighing events, so the two components agree as they should [log:certify.component_by_key.transport_process, ops_notes_totals, _derived_values_numeric].",
    "Isometric Standard S3.7.3")

section_header("SECTION 6 -- Certify: Transport Embodied Emissions")
item_row("6.1", "Does the Transport Embodied emission use the correct and up-to-date Non-Tanker Truck emission factor?", "Pass",
    "Modal shows 'Emissions Factor - Vehicle Embodied Emissions - Flatbed Truck' at 0.19279 kgCO2e/mi. The live Standard Emission Factors sheet, 2026 tab, row 12 'Vehicle Embodied Emissions --Non-Tanker Truck' = 0.0001927908256 MT CO2e/mi (GLEC V3.2 2025) -- the same value, and the per-mile (not per-tonne-mile) basis the embodied convention requires. Recompute: 1,741.45 mi x 0.0001927908256 = 0.336 t, matching Certify [log:standard_efs_2026, certify.component_by_key.transport_embodied, _derived_values_numeric].",
    "Standard EFs 2026 tab, row 12")
item_row("6.2", "Does the aggregate distance transported match the Batch Mileage Allocation in the RP Bio-Oil Shipment Tracker?", "Pass",
    "Certify 'QOWV Embodied Distance (CO to B6)' = 1,741.45 mi. The allocation tab's batch mileage total G15 = SUM(G12:G13) = 540.45 + 1,201 = 1,741.45 mi. Exact match [log:certify.component_by_key.transport_embodied, shipment_tracker, _derived_values_numeric].",
    "Transportation Emissions Acctg v1.1")
item_row("6.3", "Is the aggregate distance evidenced with the Shipment Tracker, batch Charm Contents, and applicable route Google Map?", "Pass",
    "The datapoint carries 2 sources and the Sources tab confirms both: 'INJ-09-2026-W2 Ops Notes.xlsx' (which embeds the QOWV Shipments & EE Allocation tab deriving 1,741.45 mi) and 'CO to Basco Routing Map.png' (the 1,201 mi route). No standalone Charm Contents sheet exists for this batch -- absent from both the four-item Drive folder and the six Certify sources -- and its per-period mass breakdown lives in the Ops Notes Production Period column, the basis accepted on every 2026 weekly batch [log:certify.attachment_map, batch_folder_inventory, certify.certify_sources_enumeration].",
    "VR Appendix 1")

section_header("SECTION 7 -- Certify: Production (Pyrolysis) Emissions")
item_row("7.1", "Are Production Emissions captured for ALL Production Periods in the batch?", "Pass",
    "Six production periods feed this batch -- July 2025 (1 tote), September 2025 (10), October 2025 (2), December 2025 (5), April 2026 (6), May 2026 (14) -- and all 38 totes carry a period and a period-specific CI, so no period's emissions are omitted. They are captured in one production-mass-weighted component, 'Pyrolysis - Charm QOWV PP Emissions' (4.650 tCO2e), whose blended CI reconciles exactly to the per-tote sum: sum(net produced x per-tote CI) = 4,650.122824 kg. On the inventory master the batch also draws on August 2025 (PP#9, tote 245140), whose emissions are carried at the December 2025 CI -- see 7.2 [log:production_period_tote_counts, ops_notes_totals, tote_inventory_crosscheck, _derived_values_numeric].",
    "Bio-oil Geological Storage v1.2")
item_row("7.2", "Does the CI match the Ops Notes weighted-average CI across all totes in the injection?", "Fail",
    "As submitted, Certify's 0.126908 tCO2e/t matched H50 = 0.1269082907, but the five December-labelled totes carried the December 2025 CI ($L$12 = 0.1630616644) instead of their production periods' CIs, overstating production emissions by 158.6842 kg (net 0.159 t low). The live sheet has relabelled them: I50 = 4,491.438649 kg, CI 0.1225775798, and Certify's pyrolysis now reads 4.491 t (read 2026-09-29). Still open: H50 is blank (1.2), the new CIs in H11:H15 are typed numbers (7.3), and the corrected Ops Notes must be re-attached to the 'QOWV Pyrolysis Emissions - Carbon Intensity (Weighted)' datapoint. If the inventory is right for 228470, 228485 and 301108 (3.5), production emissions rise a further 55.2061 kg, to 4,546.6447 kg (CI 0.1240842287) [log:tote_inventory_crosscheck, live_ops_notes_2026_09_29, certify.live_change_since_qa, pp9_summary, pp_summary_sources, _derived_values_numeric].",
    "GHG Stmt Sec C")
item_row("7.3", "Is the CI evidenced (Ops Notes weighted-average across totes)?", "Flag",
    "The CI datapoint carries the batch Ops Notes as its source, which holds the per-tote CI column and the weighted-average cell, so the evidence standard is met. The flag is for traceability of two of the six periods. The QOWV tab's constants block (K10:M20) hyperlinks a Production Period Summary for PP#8 and for each month November 2025 through June 2026, and the rows for those periods reference it by cell ($L$12, $L$16, $L$17). The September 2025 CI (0.12868956697570233, 10 totes) and the October 2025 CI (0.1110176228550779, 2 totes) are instead typed constants in column H with no entry in the constants block and no link -- 12 of 38 totes. Summing those rows' production-emissions cells gives September 1,254.5946 kg + October 217.7944 kg = 1,472.389 kg of the 4,650.122824 kg component. Both were traced and confirmed correct against their own Production Period Summaries, but a verifier working only from the submitted evidence cannot make that trace. In the live sheet (2026-09-29) H11:H15 are also typed numbers, and the August 2025 CI (PP#9, 0.11056982394624702, tote 245140) has no entry in the constants block. Adding September 2025, October 2025 and August 2025 to the constants block with their PP Summary links, and referencing them from column H, clears this. A related point for the same fix: none of the six Production Period Summaries states a calendar date range on its face, so the period-to-month mapping (for instance PP#8 carrying the July 2025 lot) rests on the Ops Notes label alone [log:pp_summary_sources, _derived_values_numeric].",
    "VR Appendix 1")
item_row("7.4", "Does the mass match the sum of Net Weight of Bio-Oil in Tote values for the periods that apply to the component?", "Pass",
    "Certify Mass 36,641.6 kg. Ops Notes sum of Net Oil Produced across all 38 totes (D50) = 36,641.6 kg, delta 0.0 kg. The component correctly uses net-in-tote produced mass, not the net-injected mass (31,250.9 kg), since pyrolysis emissions were incurred to produce the full quantity [log:certify.component_by_key.pyrolysis, ops_notes_totals, _derived_values_numeric].",
    "GHG Stmt Sec C")
item_row("7.5", "Does the uncertainty match the sum of uncertainties for the periods that apply to the component?", "Pass",
    "Modal INPUTS read directly: 36,641.6 kg plus or minus 43.09 kg read from the modal, the display rounding of 38 x 1.13398 = 43.09124 kg. Uncertainty is present on this input, which matters because pyrolysis at 4.650 tCO2e is 104 % of the 4.484 t net removal -- a dominant input that does not qualify for the under-1 % omission allowance [log:certify.component_by_key.pyrolysis, _derived_values_numeric].",
    "Bio-oil Geological Storage v1.2 Sec 6.5.1")
item_row("7.6", "Is the mass evidenced with a copy of the Charm Contents sheet for the batch?", "Pass",
    "No Charm Contents sheet exists for this batch or any 2026 weekly QOWV batch -- confirmed absent from both the Drive folder (four items: Kerry Oil, Scale Tickets, Ops Notes, COA) and the six Certify sources. The per-period mass breakdown lives in the Ops Notes QOWV tab Production Period column, and the Ops Notes is attached and linked to the 'QOWV Pyrolysis Mass' datapoint [log:certify.attachment_map, batch_folder_inventory, certify.certify_sources_enumeration].",
    "VR Appendix 1")

section_header("SECTION 8 -- Other")
item_row("8.1", "Are all diesel transport-process emissions offset by BCUs, and does the BCU total equal (never exceed) the transport-process sum?", "Fail",
    "The batch has one diesel transport-process emission, the flatbed leg, and it is offset: 'BCU - Transport Process' carries a Constant CO2 reduction of 6,454 kgCO2e, evidenced by the attached TERC Certificate of Retirement (2,795 TERCs retired 18 Dec 2025, vintage 2024, Renewable Diesel, Amazon.com Services LLC on behalf of Charm Industrial, Inc, verified by Turner, Mason & Company). The embodied leg is correctly not offset. The BCU does not equal the emission it offsets: Certify's process component, on its 0.086 kgCO2e/(km*t) factor, is 6,454.494 kg (6,454.327 kg on the EF sheet's 0.0001384 MT/t-mi), so the 6,454 kg constant is 0.494 kg short. The checklist's fail mode for this item is 'BCU < transport process activity (under-reduction)'. Fix: enter the constant as 6,454.494 kg, equal to Certify's own process component [log:certify.component_by_key.bcu, certify.component_by_key.transport_process, bcu_vs_process, retirement_certificate, standard_efs_2026, _derived_values_numeric].",
    "BCU QA Rule (batch-qa-charm SKILL.md)")
section_header("SECTION 9 -- Aqueous Buffering (LCS)")
item_row("9.1", "If LCS use is reported in Ops Notes, does a matching Certify component exist?", "N/A",
    "No LCS use is reported. Every cell string across all ten tabs of the workbook was scanned for 'LCS', 'calcium' and 'buffer' with zero hits; the QOWV tab's only loss constant is Sparging Mass Loss L5 = 0.0129, referenced by no QOWV row, and the Summary tab reports sparging 0.0 kg and filtration 0.0 kg for QOWV. With no LCS use there is no component to require [log:lcs_use_kg, lcs_search_method, ops_notes].",
    "VR Sec 3.1")

section_header("CALCULATION CHECK -- independent recompute")
item_row("C.1", "Does the net removal recompute from source values?", "Pass",
    "Recomputed from the Ops Notes export and the live 2026 EF sheet, independent of Certify's own arithmetic: Sequestration 31,250.9 kg x 8.3 % x 3.667 = 9.512 t; Transport process 38.8304 t x 1,201 mi x 0.0001384 = 6.454 t; Transport embodied 1,741.45 mi x 0.0001927908256 = 0.336 t; Pyrolysis 36.6416 t x 0.126908 = 4.650 t; Net = 9.512 + 6.454 - (6.454 + 0.336 + 4.650) - 0.041 = 4.484 t, matching Certify's 4.484 t. All 18 checks in recompute_qowv.py return OK (exit 0). The 41 kg uncertainty discount is consistent with variance propagation on the carbon content (0.004 relative) and the three uncertain masses (43.09124 kg each on sequestration, transport process and pyrolysis): estimate 41.24 kg. This reproduces the removal as submitted. 4.484 t carries two offsetting errors: production emissions overstated by 158.6842 kg (7.2, the December labels) and embodied emissions understated by 104.194 kg (3.4), plus the 0.494 kg BCU shortfall (8.1). The live 4.643 t reflects the December relabel only. Counterfactuals are zero. Ledger integrity checks also clean: no duplicate tote tags, no tote whose injected mass exceeds its gross, gross masses 944.0-1,130.0 kg and post-unload tote weights 144.0-368.0 kg [log:_derived_values, _derived_values_numeric, standard_efs_2026, discount_check, shipping_orders_2025_12, live_ops_notes_2026_09_29].",
    "Bio-oil Geological Storage v1.2")

section_header("SIDEBAR OBSERVATIONS (no status - for Garrett's attention)")
for note in [
    "Allocation tab rows 56-71 hold a per-tote block headed '--- Shipment 2025-12-10 --- (14 of 35 totes on this truck are W4-relevant)', listing 14 tags all assigned to batch INJ-05-2026-W4. The allocation totals reference rows 12-13 only, so no figure in this batch draws on those rows. Stale content in a submitted evidence file; recommend deleting it from the W2 workbook.",
    "The October 2025 Production Period Summary (Drive 1AZymoR1vHkfpH4ZKO5J6CklWiVlCfhVS872y-2LipDw) still names its tab 'September PP Summary'. Its numbers are genuinely October's: row 2 Mass of Output reads 14,227 kg for WODO against 12,630 kg in the same cell of the September file, both read this session, so the data is right and the tab name is stale. Same observation was logged on INJ-07-2026-W1-QOWV; it has not been fixed.",
    "On page 1 of the lab certificate (SR26-05494.001, QOWV) the '#' footnote marks Carbon 8.3 % (the Hydrogen row carries none; on the AECN page 2 both Carbon and Hydrogen are marked): 'Result is outside of test method limits and/or analytical range used in method precision study'. Noted so the verifier is not surprised by it.",
    "The Bio-oil Injection Tracker's quantity_consumed_in_injection column is the Ops Notes net injected mass minus 57.6 kg on 34 of the 38 rows and twice that on rows 776, 777, 780 and 784 (301132 reads 1,420.8 = 2 x 710.4 kg). Certify uses the ledger figures.",
    "Bio-oil Injection Tracker rows 760-764 (015162, 228490, 245103, 245140, 245142) carry the tag in consumed_lot with security_tag blank; the other 33 rows carry it in security_tag.",
    "Truck 1 (CO_AQUEOUS_12.01.25) shipped on 2025-12-01 and COLA_34 on 2026-08-26, so 18 of the 38 totes were injected on 2026-09-05, about nine months after shipping. Both are costed over the direct 1,201 mi Colorado-to-Louisiana route, with no Kansas repositioning of the kind the August RP guidance documented for that month's totes.",
]:
    ws.cell(row=row, column=2, value=note).alignment = WRAP
    row += 1

row += 1
ws.cell(row=row, column=1, value="STATUS COUNTS").font = BOLD
cnt = {k: sum(1 for r in range(1, row) if ws.cell(r, 3).value == k) for k in ("Pass", "Flag", "Fail", "N/A")}
ws.cell(row=row, column=2, value=f"Pass {cnt['Pass']}  |  Flag {cnt['Flag']}  |  Fail {cnt['Fail']}  |  N/A {cnt['N/A']}  ({sum(cnt.values())} items)")
ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
print(cnt)
wb.save("Charm_Aqueous_QOWV_QA_Checklist_INJ-09-2026-W2.xlsx")
print("saved")
