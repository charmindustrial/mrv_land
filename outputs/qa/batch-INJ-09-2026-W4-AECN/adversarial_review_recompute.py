"""Gate 3 adversarial review: independent arithmetic for INJ-09-2026-W4-AECN.

Built by the reviewer from primary files only:
  - scale-ticket gross weights typed from the 32 ticket images (28 W4 + 4 W2 SHPX215079)
  - BOL figures typed from the 12 BOL page images
  - emission factors read at full precision from source_exports/Standard_Emission_Factors_export_run3_2026-10-07T0213Z.xlsx
  - Certify inputs and outputs read from source_exports/certify_graphql_GetGhgEntryQAData_run3recheck_2026-10-07T0826Z.json
  - Ops Notes read from source_exports/INJ-09-2026-W4_Ops_Notes_live_export_run3b_2026-10-07T0825Z.xlsx
Does not read the QA's recompute script or its output.
"""
import json, math, datetime as dt
import openpyxl

B = "/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/qa/batch-INJ-09-2026-W4-AECN/"
out = []
def p(*a):
    s = " ".join(str(x) for x in a); print(s); out.append(s)

KG_PER_LB = 0.453592
KM_PER_MI = 1.609344

# ---------------------------------------------------------------- 1. tickets
# (railcar, offload, full_lb, empty_lb, full_datetime, empty_datetime, full_axles, empty_axles)
T = [
 ("SHPX215053",1,73800,29400,"2026-09-22 16:19","2026-09-23 07:00",(10240,31840,31720),(9500,11320,8580)),
 ("SHPX215053",2,73440,28640,"2026-09-22 16:16","2026-09-23 07:08",(11900,28480,33060),(10480,9740,8420)),
 ("SHPX215053",3,75200,28680,"2026-09-23 11:43","2026-09-24 05:08",(12100,29460,33640),(9580,11260,7840)),
 ("SHPX215053",4,64180,29640,"2026-09-23 12:18","2026-09-24 05:09",(10120,27220,26840),(0,21140,8500)),
 ("SHPX215039",1,74500,28680,"2026-09-18 13:18","2026-09-18 16:37",(9940,32360,32200),(9360,11300,8020)),
 ("SHPX215039",2,78620,29480,"2026-09-18 13:16","2026-09-18 16:00",(0,23800,54820),(0,20900,8580)),
 ("SHPX215039",3,70620,29280,"2026-09-19 12:34","2026-09-19 16:17",(0,10040,60580),(0,17920,11360)),
 ("SHPX215039",4,63340,28620,"2026-09-19 12:30","2026-09-19 15:39",(9940,27040,26360),(9440,11320,7860)),
 ("SHPX215091",1,74200,28680,"2026-09-17 11:22","2026-09-17 13:23",(10000,32280,31920),(9420,11380,7880)),
 ("SHPX215091",2,72200,29240,"2026-09-17 11:57","2026-09-17 14:00",(0,23780,48420),(0,18300,10940)),
 ("SHPX215091",3,76060,28560,"2026-09-17 16:25","2026-09-18 06:44",(9940,33060,33060),(9400,11280,7880)),
 ("SHPX215091",4,64560,29580,"2026-09-17 16:11","2026-09-18 07:02",(0,9940,54620),(9680,11400,8500)),
 ("SHPX215079",3,75800,29200,"2026-09-05 09:29","2026-09-05 12:35",(10320,32960,32520),(9580,11420,8200)),
 ("SHPX215079",4,60700,28260,"2026-09-05 14:20","2026-09-06 07:50",(12480,21420,26800),(9340,11060,7860)),
]
W2_079 = [("SHPX215079",1,76720,28660,"2026-09-04 14:07","2026-09-05 07:23"),
          ("SHPX215079",2,73860,28420,"2026-09-05 09:24","2026-09-05 12:03")]

wb = openpyxl.load_workbook(B+"source_exports/INJ-09-2026-W4_Ops_Notes_live_export_run3b_2026-10-07T0825Z.xlsx", data_only=True)
ws = wb["AECN Rail"]
p("=== 1. Scale tickets vs Ops Notes F28:F41 (net kg = (full-empty) lb x 0.453592)")
nets = {}
axle_ok = True
for i,(car,o,f,e,fd,ed,fa,ea) in enumerate(T):
    r = 28+i
    net = (f-e)*KG_PER_LB
    nets[(car,o)] = net
    ops = ws[f"F{r}"].value
    if sum(fa)!=f or sum(ea)!=e: axle_ok=False
    p(f"{car} o{o}: ({f}-{e}) lb = {f-e} lb -> {net:.5f} kg | Ops F{r} {ops:.5f} | diff {net-ops:+.6f} | Ops D/E {ws[f'D{r}'].value}/{ws[f'E{r}'].value}")
tot = sum(nets.values())
p(f"sum of 14 nets {tot:.4f} kg | Ops A46 {ws['A46'].value:.4f} | axle readings sum to gross on all 28: {axle_ok}")
lbs = [f-e for (_,_,f,e,*_) in T]; fulls=[f for (_,_,f,*_) in T]
p(f"net range {min(lbs)}..{max(lbs)} lb; full gross range {min(fulls)}..{max(fulls)} lb")
inj = dt.datetime(2026,9,25)
p("ticket age at 2026-09-25 (days, full weigh):")
for (car,o,f,e,fd,ed,*_) in T:
    d = (inj-dt.datetime.strptime(fd,"%Y-%m-%d %H:%M")).total_seconds()/86400
    p(f"  {car} o{o}: full {fd} -> {d:.2f} d; empty {ed}")
p("W2 SHPX215079 offloads 1-2:")
w2 = 0
for (car,o,f,e,fd,ed) in W2_079:
    n=(f-e)*KG_PER_LB; w2+=n
    p(f"  o{o}: ({f}-{e}) lb -> {n:.5f} kg ({fd} / {ed})")
p(f"  W2 total {w2:.5f} kg | Ops H13 {ws['H13'].value}")

# ---------------------------------------------------------------- 2. BOLs
p("\n=== 2. BOLs (old format TR95-012-R00 / Shipping Details TR95-010-R01 / weighbridge)")
# bol: (railcar, tare, full, delivered, qty_MT, litres60F, density60F, nette)
BOL = {
 "2497-R":("SHPX215053",19210,45120,25910,25.91,21648,1.1969,25910),
 "2498-R":("SHPX215053",19280,45010,25730,25.73,21497,1.1969,25730),
 "2500-R":("SHPX215053",19050,44980,25930,25.93,21664,1.1969,25930),
 "2459-R":("SHPX215039",19120,44980,25860,25.86,21566,1.1991,25860),
 "2460-R":("SHPX215039",19540,45460,25920,25.92,21652,1.1971,25920),
 "2461-R":("SHPX215039",19270,45130,25860,25.86,21602,1.1971,25860),
 "2465-R":("SHPX215091",19370,45390,26020,26.02,21720,1.1980,26020),
 "2466-R":("SHPX215091",19290,45150,25860,25.86,21586,1.1980,25860),
 "2467-R":("SHPX215091",19330,45220,25890,25.89,21634,1.1967,25890),
 "2453-R":("SHPX215079",19440,45430,25990,25.99,21595,1.2035,25990),
 "2454-R":("SHPX215079",19280,45190,25910,25.91,21608,1.1991,25910),
 "2455-R":("SHPX215079",19130,45100,25970,25.97,21658,1.1991,25970),
}
loaded = {}
maxgap=0
for b,(car,ta,fu,de,mt,L,rho,net) in BOL.items():
    g = L*rho-de; allow=L*0.00005
    maxgap=max(maxgap,abs(g))
    p(f"{b} {car}: full-tare {fu-ta} vs delivered {de} [{fu-ta==de}]; kg/1000 {de/1000} vs Qty {mt} [{abs(de/1000-mt)<1e-9}]; nette {net} [{net==de}]; L x rho {L*rho:.2f} gap {g:+.2f} kg (allow {allow:.2f})")
    loaded[car]=loaded.get(car,0)+de
p(f"max |litres x density - delivered| = {maxgap:.2f} kg")
for car,v in loaded.items(): p(f"loaded {car} = {v} kg")

# ---------------------------------------------------------------- 3. railcar balance, loss, cut
p("\n=== 3. Railcar balance")
deliv = {}
for (car,o),n in nets.items(): deliv[car]=deliv.get(car,0)+n
deliv["SHPX215079"] += w2
cut = deliv["SHPX215079"]-loaded["SHPX215079"]
loss = {}
for car in ["SHPX215053","SHPX215039","SHPX215091","SHPX215079"]:
    l = loaded[car]-deliv[car]
    loss[car]=max(0,l)
    p(f"{car}: loaded {loaded[car]} | delivered {deliv[car]:.5f} | loaded-delivered {l:+.5f} kg ({l/loaded[car]*100:.3f} %) | Ops G {ws['G'+str({'SHPX215053':10,'SHPX215039':11,'SHPX215091':12,'SHPX215079':13}[car])].value}")
o4_cut = nets[("SHPX215079",4)]-cut
p(f"SHPX215079 over-draw {cut:.5f} kg -> offload 4 carries {o4_cut:.5f} kg (= {loaded['SHPX215079']} - {w2+nets[('SHPX215079',3)]:.5f}) | Ops J41 {ws['J41'].value}")
lost_total = sum(loss.values())
p(f"total lost mass {lost_total:.5f} kg")
p(f"SHPX215079 cross-week mass held (offloads 3+4 after cut) {nets[('SHPX215079',3)]+o4_cut:.5f} kg")

# ---------------------------------------------------------------- 4. mass used, sparging, filtration, MOBI
p("\n=== 4. Rollup")
mass_used = tot - cut
spar = mass_used*0.0129
filt = (mass_used-spar)*0.00196
mobi = mass_used-spar-filt
p(f"mass used {mass_used:.5f} (Ops C46 {ws['C46'].value:.5f}); sparging {spar:.6f} (D46 {ws['D46'].value:.6f}); filtration {filt:.6f} (E46 {ws['E46'].value:.6f}); MOBI {mobi:.4f} (F46 {ws['F46'].value:.4f})")
BR = sum(n for (c,o),n in nets.items() if c=="SHPX215053")
OP = mass_used-BR
p(f"Baton Rouge delivered {BR:.5f} kg (C49 {ws['C49'].value}); Opelousas delivered {OP:.5f} kg (C50 {ws['C50'].value})")

# ---------------------------------------------------------------- 5. factors
p("\n=== 5. Standard EF sheet (2026 tab) full precision")
ef = openpyxl.load_workbook(B+"source_exports/Standard_Emission_Factors_export_run3_2026-10-07T0213Z.xlsx", data_only=True)["2026"]
CI_P = ef["F4"].value; CI_E = ef["F5"].value
TT_P = ef["F6"].value; RAIL_P = ef["F10"].value; TT_E = ef["F11"].value; RAIL_E = ef["F13"].value
for k,v in [("F4 AECN process H2 2026",CI_P),("F5 AECN embodied",CI_E),("F6 tanker process MT/(t mi)",TT_P),("F10 rail process MT/(t mi)",RAIL_P),("F11 tanker embodied MT/mi",TT_E),("F13 railcar embodied MT/mi",RAIL_E)]:
    p(f"  {k}: {v!r}")

cq = json.load(open(B+"source_exports/certify_graphql_GetGhgEntryQAData_run3recheck_2026-10-07T0826Z.json"))
E = json.loads(cq[1]["text"])["ghgEntry"]
comp = {}
for g in E["componentGroups"]:
    for ca in g["componentAttributions"]:
        c = ca["component"]
        ins = {s["blueprintInput"]["inputKey"]:s["dataPoint"] for s in (c.get("inputData") or {}).get("scalarInputs",[])}
        comp[c["displayName"]] = (ca["carbonFlux"]["co2eTransferred"], ins, c["componentType"], g["displayName"])
def q(name,key): return comp[name][1][key]["quantity"]["magnitude"]
def sd(name,key): return comp[name][1][key]["quantity"]["standardDeviation"]
def flux_t(name):
    f=comp[name][0]; return f["magnitude"]/(1000 if f["unit"].startswith("kg") else 1)
p("Certify factor values vs sheet:")
p(f"  tanker process Certify {q('AECN to SOPOR Tanker Truck Transport Process Emissions','carbon_intensity')} kg/(km t) x 1.609344 = {q('AECN to SOPOR Tanker Truck Transport Process Emissions','carbon_intensity')*KM_PER_MI:.10f} kg/(t mi) vs sheet {TT_P*1000}")
p(f"  rail process (Opelousas) Certify {q('SOPOR to Opelousas Rail Transport Process Emissions','carbon_intensity')} -> {q('SOPOR to Opelousas Rail Transport Process Emissions','carbon_intensity')*KM_PER_MI:.10f} kg/(t mi) vs sheet {RAIL_P*1000}")
p(f"  rail process (Baton Rouge) Certify GLEC library {q('SOPOR to Baton Rouge Rail Process Emissions','carbon_intensity')} kg/(km t)")
p(f"  tanker process (Baton Rouge) Certify GLEC library {q('Baton Rouge to Basco Tanker Truck Process Emissions','carbon_intensity')} kg/(km t) vs sheet {TT_P*1000/KM_PER_MI:.10f}")
re_c = q('SOPOR to Opelousas/Baton Rouge Rail Transport Embodied Emissions','carbon_intensity')
te_c = q('AECN to SOPOR Tanker Truck Transport Embodied Emissions','carbon_intensity')
p(f"  rail embodied Certify {re_c} kg/km x 1.609344 = {re_c*KM_PER_MI:.10f} kg/mi vs sheet {RAIL_E*1000:.10f}")
p(f"  tanker embodied Certify {te_c} kg/km x 1.609344 = {te_c*KM_PER_MI:.10f} kg/mi vs sheet {TT_E*1000:.10f}")
p(f"  pyrolysis CI Certify {q('AECN pyrolysis (from rail loads) - process emissions','carbon_intensity')} vs sheet {CI_P}; embodied {q('AECN pyrolysis (from rail loads) - embodied emissions','carbon_intensity')} vs {CI_E}")

# ---------------------------------------------------------------- 6. components on Certify inputs
p("\n=== 6. Components recomputed on Certify's own inputs (t)")
M = q("AECN to SOPOR Tanker Truck Transport Process Emissions","mass")      # t
def tr(name, mass_t, dist, dist_unit, ci_kgkmt):
    km = dist*KM_PER_MI if dist_unit=="mi" else dist
    return mass_t*km*ci_kgkmt/1000
rows = {}
for name,mk in [("AECN to SOPOR Tanker Truck Transport Process Emissions","mass"),("SOPOR to Opelousas Rail Transport Process Emissions","mass"),
                ("Opelousas to Basco - Tanker Truck Transport Process Emissions","mass"),("SOPOR to Baton Rouge Rail Process Emissions","mass"),
                ("Baton Rouge to Basco Tanker Truck Process Emissions","mass")]:
    ins = comp[name][1]
    d = ins["distance"]["quantity"]
    v = tr(name, ins["mass"]["quantity"]["magnitude"], d["magnitude"], d["unit"], ins["carbon_intensity"]["quantity"]["magnitude"])
    rows[name]=v
    p(f"{name}: {ins['mass']['quantity']['magnitude']} t x {d['magnitude']} {d['unit']} x {ins['carbon_intensity']['quantity']['magnitude']} = {v:.6f} | Certify {flux_t(name):.6f}")
for name in ["SOPOR to Opelousas/Baton Rouge Rail Transport Embodied Emissions","Opelousas/Baton Rouge to Basco Tanker Truck Transport Embodied Emissions","AECN to SOPOR Tanker Truck Transport Embodied Emissions"]:
    ins = comp[name][1]; d=ins["distance"]["quantity"]
    v = d["magnitude"]*KM_PER_MI*ins["carbon_intensity"]["quantity"]["magnitude"]/1000
    rows[name]=v
    p(f"{name}: {d['magnitude']} mi x {ins['carbon_intensity']['quantity']['magnitude']} kg/km = {v:.6f} | Certify {flux_t(name):.6f}")
pp = q("AECN pyrolysis (from rail loads) - process emissions","mass")*q("AECN pyrolysis (from rail loads) - process emissions","carbon_intensity")
pe = q("AECN pyrolysis (from rail loads) - embodied emissions","mass")*q("AECN pyrolysis (from rail loads) - embodied emissions","carbon_intensity")
p(f"pyrolysis process {pp:.6f} | Certify {flux_t('AECN pyrolysis (from rail loads) - process emissions'):.6f}; embodied {pe:.6f} | Certify {flux_t('AECN pyrolysis (from rail loads) - embodied emissions'):.6f}")
proc_names = ["AECN to SOPOR Tanker Truck Transport Process Emissions","SOPOR to Opelousas Rail Transport Process Emissions","Opelousas to Basco - Tanker Truck Transport Process Emissions","SOPOR to Baton Rouge Rail Process Emissions","Baton Rouge to Basco Tanker Truck Process Emissions"]
emb_names = ["SOPOR to Opelousas/Baton Rouge Rail Transport Embodied Emissions","Opelousas/Baton Rouge to Basco Tanker Truck Transport Embodied Emissions","AECN to SOPOR Tanker Truck Transport Embodied Emissions"]
proc_c = sum(flux_t(n) for n in proc_names); emb_c = sum(flux_t(n) for n in emb_names)
loss_c = flux_t("AECN Bio-oil Loss"); bcu = flux_t("BCU - Transport Process")
pyro_c = flux_t('AECN pyrolysis (from rail loads) - process emissions')+flux_t('AECN pyrolysis (from rail loads) - embodied emissions')
aecn_act = proc_c+emb_c+loss_c+pyro_c
site_act = sum(flux_t(n) for n,(f,i,t,g) in comp.items() if t=="ACTIVITY" and g=="Bio-oil injection")
site_red = sum(flux_t(n) for n,(f,i,t,g) in comp.items() if t=="ADJUSTMENT" and g=="Bio-oil injection")
seq = flux_t("Injection sequestration")
p(f"transport process sum {proc_c:.7f}; embodied sum {emb_c:.6f}; pyrolysis {pyro_c:.6f}; loss {loss_c}; AECN activities {aecn_act:.6f}")
p(f"BCU {bcu} vs process sum {proc_c:.7f}: BCU - sum = {bcu-proc_c:+.7f} t")
p(f"site activities {site_act:.6f}; site BCUs {site_red:.5f}; total activities {aecn_act+site_act:.6f}; adjustments {bcu+site_red:.5f}")
disc = E["uncertaintyDiscount"]["magnitude"]
net_c = seq-(aecn_act+site_act)+(bcu+site_red)-disc
p(f"sequestration {seq:.6f} = {q('Injection sequestration','product_mass')} t x {q('Injection sequestration','carbon_content')/100} x 3.667 = {q('Injection sequestration','product_mass')*q('Injection sequestration','carbon_content')/100*3.667:.6f}")
p(f"net = {seq:.6f} - {aecn_act+site_act:.6f} + {bcu+site_red:.5f} - {disc:.10f} = {net_c:.6f} (Certify co2eRemoved {E['co2eRemoved']['magnitude']})")
p(f"MOBI check: Certify injectant {q('Injection sequestration','product_mass')} t vs my MOBI {mobi/1000:.7f} t; Bio-oil mass {M} t vs my mass used {mass_used/1000:.7f} t")
p(f"gross on my MOBI: {mobi/1000*0.44*3.667:.6f} t")

# ---------------------------------------------------------------- 7. uncertainty discount (one combined one-sigma)
p("\n=== 7. Uncertainty discount re-derived")
C = q("Injection sequestration","carbon_content")/100; Mi = q("Injection sequestration","product_mass")
terms = {}
terms["carbon"] = Mi*sd("Injection sequestration","carbon_content")/100*3.667
terms["injectant mass"] = sd("Injection sequestration","product_mass")*C*3.667
gci = q("Basco Site Electricity Use (RECs)","grid_carbon_intensity")
terms["grid electricity"] = sd("Basco Site Electricity Use (RECs)","grid_electricity_use")*gci/1000
sens_bio = CI_P*0 + q('AECN pyrolysis (from rail loads) - process emissions','carbon_intensity') + q('AECN pyrolysis (from rail loads) - embodied emissions','carbon_intensity')
sens_bio += 107.49722*KM_PER_MI*q("AECN to SOPOR Tanker Truck Transport Process Emissions","carbon_intensity")/1000
sens_bio += 42.9*KM_PER_MI*q("Opelousas to Basco - Tanker Truck Transport Process Emissions","carbon_intensity")/1000
terms["bio-oil mass (4 components)"] = sd("AECN to SOPOR Tanker Truck Transport Process Emissions","mass")*sens_bio
terms["Opelousas rail mass"] = sd("SOPOR to Opelousas Rail Transport Process Emissions","mass")*2199*KM_PER_MI*q("SOPOR to Opelousas Rail Transport Process Emissions","carbon_intensity")/1000
terms["Baton Rouge truck mass"] = sd("Baton Rouge to Basco Tanker Truck Process Emissions","mass")*136*KM_PER_MI*0.071/1000
terms["site diesel"] = sd("B6 Injection site diesel emissions","mass")*3.87/1000
terms["site gasoline"] = sd("B6 Injection site gasoline emissions","mass_of_fuel")*3.74/1000
for k,v in terms.items(): p(f"  {k}: {v:.6f} t")
d0 = math.sqrt(sum(v*v for v in terms.values()))
p(f"root-sum-square {d0:.10f} t | Certify {disc:.10f} | diff {d0-disc:+.2e} | {d0/seq*100:.3f} % of gross")
br_rail_term_now = 0.29028*2153*0.0161/1000
d1 = math.sqrt(d0**2+br_rail_term_now**2)
p(f"with Baton Rouge rail mass +/-0.29028 t on today's inputs: term {br_rail_term_now:.6f} -> discount {d1:.10f} (+{d1-d0:.7f})")

# ---------------------------------------------------------------- 8. corrections on the evidence
p("\n=== 8. Evidence-based corrections")
br_rail_mi = 2152
br_rail_ev = BR/1000*br_rail_mi*KM_PER_MI*0.0161/1000
br_rail_ev_sheet = BR/1000*br_rail_mi*RAIL_P
p(f"Baton Rouge rail on 2,152 mi: {BR/1000:.5f} t x 2152 mi x 1.609344 x 0.0161/1000 = {br_rail_ev:.6f} t (sheet factor {br_rail_ev_sheet:.6f}); Certify {flux_t('SOPOR to Baton Rouge Rail Process Emissions'):.6f}; understated {br_rail_ev-flux_t('SOPOR to Baton Rouge Rail Process Emissions'):.6f}")
p(f"  2,153 km = {2153/KM_PER_MI:.1f} mi = {2153/KM_PER_MI/2199*100:.1f} % of 2,199 mi")
op_tr_ev = OP/1000*42.9*KM_PER_MI*q("Opelousas to Basco - Tanker Truck Transport Process Emissions","carbon_intensity")/1000
p(f"Opelousas truck on Opelousas mass {OP/1000:.5f} t: {op_tr_ev:.6f} t vs Certify {flux_t('Opelousas to Basco - Tanker Truck Transport Process Emissions'):.6f}; overstated {flux_t('Opelousas to Basco - Tanker Truck Transport Process Emissions')-op_tr_ev:.6f}")
proc_ev = proc_c - flux_t('SOPOR to Baton Rouge Rail Process Emissions') + br_rail_ev - flux_t('Opelousas to Basco - Tanker Truck Transport Process Emissions') + op_tr_ev
p(f"process sum on evidence {proc_ev:.7f} t (vs Certify {proc_c:.7f}; +{proc_ev-proc_c:.7f}); BCU would draw {proc_ev-bcu:.7f} t more")
p(f"Opelousas-mass fix alone: sum {proc_c-flux_t('Opelousas to Basco - Tanker Truck Transport Process Emissions')+op_tr_ev:.7f}; BCU 19.339 exceeds it by {bcu-(proc_c-flux_t('Opelousas to Basco - Tanker Truck Transport Process Emissions')+op_tr_ev):.7f}")
act_ev = aecn_act+site_act + (proc_ev-proc_c)
p(f"activities on evidence {act_ev:.6f} t (Certify {aecn_act+site_act:.6f}; delta {act_ev-(aecn_act+site_act):+.6f})")
# discount with corrected inputs and BR rail +/-
terms2 = dict(terms)
terms2["bio-oil mass (4 components)"] = sd("AECN to SOPOR Tanker Truck Transport Process Emissions","mass")*(sens_bio - 42.9*KM_PER_MI*q("Opelousas to Basco - Tanker Truck Transport Process Emissions","carbon_intensity")/1000)
terms2["Opelousas truck mass (on 190.58 t, +/-0.7257)"] = 0.7257*42.9*KM_PER_MI*q("Opelousas to Basco - Tanker Truck Transport Process Emissions","carbon_intensity")/1000
terms2["Baton Rouge rail mass"] = 0.29028*br_rail_mi*KM_PER_MI*0.0161/1000
d2 = math.sqrt(sum(v*v for v in terms2.values()))
p(f"discount on corrected inputs (Opelousas truck on its own mass +/-725.7 kg, BR rail 2,152 mi +/-290.28 kg): {d2:.7f} t")
net_ev_bcu_fixed = seq - act_ev + (proc_ev+site_red) - d2
net_ev_bcu_kept = seq - act_ev + (bcu+site_red) - d2
p(f"net on evidence, BCU = corrected process sum: {net_ev_bcu_fixed:.6f} t (vs Certify {net_c:.6f}; {net_ev_bcu_fixed-net_c:+.6f})")
p(f"net on evidence, BCU left at 19.339: {net_ev_bcu_kept:.6f} t ({net_ev_bcu_kept-net_c:+.6f})")
p(f"BCU exact-figure gap today: process sum {proc_c:.7f} - BCU {bcu} = {proc_c-bcu:.7f} t ({(proc_c-bcu)*1000:.4f} kg) below")

# ---------------------------------------------------------------- 9. distances
p("\n=== 9. Allocated distances")
L42 = sum(( (nets[k]-(cut if k==('SHPX215079',4) else 0)) / loaded[k[0]])*107.5*3 for k in nets)
L42c = L42*107.49722/107.5
p(f"AECN-SOPOR allocated on 107.5: {L42:.6f} mi (Ops L42 {ws['L42'].value}); on 107.49722: {L42c:.5f} (Certify {q('AECN to SOPOR Tanker Truck Embodied Emissions' if False else 'AECN to SOPOR Tanker Truck Transport Embodied Emissions','distance')})")
share79 = (nets[('SHPX215079',3)]+o4_cut)/loaded['SHPX215079']
rail_rt = 2*(2152+2199+2199+2199*share79)
rail_ow = rail_rt/2
p(f"rail allocated round trip {rail_rt:.5f} (Ops M42 {ws['M42'].value}); one way {rail_ow:.5f}; difference x sheet railcar factor {(rail_rt-rail_ow)*RAIL_E:.6f} t, x Certify factor {(rail_rt-rail_ow)*KM_PER_MI*re_c/1000:.6f} t")
w2_share = w2/loaded['SHPX215079']*2199
p(f"SHPX215079 rail miles: W2 one way {w2_share:.5f} + W4 round trip {2*2199*share79:.5f} = {w2_share+2*2199*share79:.5f} (full round trip would be {2*2199}, one way {2199})")
tr_emb = (10*42.9*2+4*136*2)
p(f"truck embodied distance 10 x 42.9 x 2 + 4 x 136 x 2 = {tr_emb} mi (Ops N42 {ws['N42'].value})")
p(f"tanker embodied factor on sheet carries deadhead uplift: base 0.0001652020785 x 1.438 = {0.0001652020785*1.438:.10f} (sheet F11 {TT_E})")
p(f"  truck embodied with round trip AND uplift {tr_emb*TT_E:.6f} t; round trip without uplift {tr_emb*0.0001652020785:.6f} t; difference {tr_emb*(TT_E-0.0001652020785):.6f} t")

# ---------------------------------------------------------------- 10. loss component
p("\n=== 10. AECN Bio-oil Loss")
def upstream_per_t(rail_mi):
    proc = 107.49722*TT_P + rail_mi*RAIL_P            # t CO2e per t
    return proc
def loss_car(car, rail_mi):
    lm = loss[car]/1000
    share = loss[car]/loaded[car]
    pyro = lm*(CI_P_c+CI_E_c)
    proc = share*loaded[car]/1000*upstream_per_t(rail_mi)
    emb = share*(107.5*3*TT_E + rail_mi*RAIL_E)
    return pyro, proc, emb
CI_P_c = 0.13712; CI_E_c = 0.02116
for car in ["SHPX215053"]:
    py,pr,em = loss_car(car,2199)
    p(f"example-image basis (SHPX215053 on 2,199 mi, Certify CIs): pyrolysis process {loss[car]/1000*CI_P_c:.5f} embodied {loss[car]/1000*CI_E_c:.5f} transport process {pr:.5f} embodied {em:.5f} total {py+pr+em:.5f} (image 0.07833)")
    p(f"  full-railcar upstream process 77.57 t on 2,199 mi = {77.57*upstream_per_t(2199):.4f} t (image 5.3728); upstream embodied 322.5 x {TT_E} + 2199 x {RAIL_E} = {322.5*TT_E+2199*RAIL_E:.5f} (image 0.14507)")
tot_op = sum(sum(loss_car(c,2199)) for c in ["SHPX215053","SHPX215039","SHPX215091"])
tot_br = sum(loss_car("SHPX215053",2152)) + sum(sum(loss_car(c,2199)) for c in ["SHPX215039","SHPX215091"])
p(f"loss on Opelousas distances for all three {tot_op:.6f} t; with SHPX215053 on 2,152 mi {tot_br:.6f} t; Certify {loss_c}; Certify - guidance basis {loss_c-tot_br:+.6f}")
emb_double = sum(loss[c]/loaded[c]*({'SHPX215053':2152}.get(c,2199))*RAIL_E for c in ["SHPX215053","SHPX215039","SHPX215091"])
p(f"rail embodied on lost share also charged in full-car rail allocation: {emb_double:.6f} t (one-way basis)")
p(f"guidance '~0.229 t per t': loss/lost mass = {loss_c/(lost_total/1000):.4f} t/t")

# ---------------------------------------------------------------- 11. carbon statistics
p("\n=== 11. Carbon range")
mrv = openpyxl.load_workbook(B+"source_exports/MRV_AECN_Analysis.xlsx_run3.xlsx", data_only=True)["Summary Statistics"]
mean, sdv, n = mrv["B5"].value, mrv["B6"].value, mrv["B7"].value
p(f"mean {mean} sd {sdv} n {n}: band {mean-2*sdv:.2f}..{mean+2*sdv:.2f}; 44.0 is {(44.0-mean)/sdv:.3f} SD above mean")
p(f"carbon +/-: 0.004 x 44.0 = {0.004*44.0}; mass +/-: 14 x 72.57 = {14*72.57:.2f}; 10 x 72.57 = {10*72.57:.1f}; 4 x 72.57 = {4*72.57:.2f}")

# ---------------------------------------------------------------- 12. September drafts
p("\n=== 12. September results workbook vs report v1.7")
rw = openpyxl.load_workbook(B+"source_exports/September1-September30_Bio-Oil_GHG_Statement_Results_run3b_2026-10-07T0826Z.xlsx", data_only=True)["Sheet1"]
H=J=K=L=Mx=G=0
for r in range(2,13):
    h,j,k,l,m,c = [rw.cell(r,col).value or 0 for col in (8,10,11,12,13,18)]
    H+=h; J+=j; K+=k; L+=l; Mx+=m; G+=h*c*3.667/1000
p(f"workbook: injected {H:.2f} kg; gross {G:.2f}; process {J:.4f}; site {K:.4f}; reductions {L:.4f}; discount {Mx:.4f}; net {G-J-K+L-Mx:.2f}")
p("report v1.7 Section E: gross 1988.19, process 384.65, site 20.48, discount 11.00, reductions 171.95, net 1744.01")
p(f"W4 AECN row: H {rw['H10'].value} J {rw['J10'].value} K {rw['K10'].value} L {rw['L10'].value} M {rw['M10'].value} -> net {rw['H10'].value*0.44*3.667/1000-rw['J10'].value-rw['K10'].value+rw['L10'].value-rw['M10'].value:.4f}")
p(f"EAC tracker September usage 171.9503 vs workbook reductions {L:.4f}: short by {L-171.9503:.4f}")
p(f"BCU certificate 2,795 TERCs vs tracker total 2765.361153: {2795-2765.361153:.6f}")

# ---------------------------------------------------------------- 13. extra checks for challenges / observations
p("\n=== 13. Extra checks")
att = openpyxl.load_workbook(B+"certify_run3/files/src_1M3W1NWED1S0K0MP_INJ-09-2026-W4_Ops_Notes.xlsx", data_only=True)["AECN Rail"]
g = [att[f"G{r}"].value for r in (10,11,12,13)]
p(f"attached Ops copy (2026-10-01, linked to the loss) G10:G13 = {g}; sum {sum(g):.5f} kg vs {lost_total:.5f} kg behind Certify's loss")
p(f"  loss emission at Certify's own rate {loss_c/(lost_total/1000):.6f} t/t on {sum(g):.5f} kg = {loss_c/(lost_total/1000)*sum(g)/1000:.6f} t (Certify 0.145)")
ag = q("Argon Electricity Usage (RECs)","grid_electricity_use")
p(f"Argon grid electricity {ag} kWh, +/- shown: {sd('Argon Electricity Usage (RECs)','grid_electricity_use')}; KB 2% = {0.02*ag:.4f} kWh -> term {0.02*ag*gci/1000:.6f} t")
d3 = math.sqrt(d0**2 + (0.02*ag*gci/1000)**2)
p(f"  discount with Argon grid +/-2% added: {d3:.10f} (+{d3-d0:.7f} t)")
for nm in ["Basco Site Electricity Use (RECs)","Argon Electricity Usage (RECs)"]:
    pu = q(nm,"procured_power_electricity_use"); rc = q(nm,"procured_power_carbon_intensity")
    p(f"  {nm}: procured {pu} kWh, +/- shown {sd(nm,'procured_power_electricity_use')}; REC CI {rc}")
p(f"SHPX215091 full weighs 2026-09-17 and SHPX215039 o1-o2 full weighs 2026-09-18 precede the Certify period start {E['started']}")
p(f"Certify period {E['started']} to {E['completed']}")
loss_guid = tot_br
net_ev_full = seq - (act_ev - (loss_c - loss_guid)) + (proc_ev+site_red) - d2
p(f"net on evidence, BCU = corrected process sum AND loss at guidance basis {loss_guid:.6f}: {net_ev_full:.6f} t; QA checklist 9.8 states 362.069035 -> gap {net_ev_full-362.069035:+.6f} t")
p(f"share of week's mass held across weeks: {(nets[('SHPX215079',3)]+o4_cut)/mass_used*100:.2f} %")
import hashlib, glob
hs = {hashlib.sha256(open(f,'rb').read()).hexdigest()[:16] for f in glob.glob(B+"certify_run3/files/*Standard_Emission_Factors___Calculations__8_.xlsx")}
p(f"Standard EF workbook copies in Certify files: {len(glob.glob(B+'certify_run3/files/*Standard_Emission_Factors___Calculations__8_.xlsx'))}, distinct hashes {hs}")
p(f"2026-09-18 is a {dt.date(2026,9,18).strftime('%A')}; 2026-09-25 is a {dt.date(2026,9,25).strftime('%A')}")

open(B+"adversarial_review_recompute_output.txt","w").write("\n".join(out)+"\n")
