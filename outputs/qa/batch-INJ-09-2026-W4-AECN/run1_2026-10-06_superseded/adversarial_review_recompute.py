#!/usr/bin/env python3
"""Independent Gate 3 recompute for INJ-09-2026-W4-AECN (rmv_1M3Z5YJN51S0SN1Q).

Written by the adversarial reviewer. Does NOT import or reuse recompute_INJ-09-2026-W4-AECN.py.
Inputs read directly:
  - source_exports/certify_graphql_GetGhgEntryQAData_2026-10-06.json  (unrounded Certify pull)
  - source_exports/INJ-09-2026-W4_Ops_Notes_live_export_2026-10-06.xlsx (live Ops Notes)
  - source_exports/Standard_Emission_Factors_export_2026-10-06.xlsx     (2026 tab)
  - certify_sources/files/src_1M3VZCDQF1..., src_1M3W1NWED1...           (Ops copies attached in Certify)
  - drive_files/MRV_AECN_Analysis.xlsx                                  (carbon statistics)
  - ticket and BOL values typed below from the reviewer's own page-image reads this session
Every printed figure shows its formula.
"""
import json, math, os, datetime
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
P = lambda *a: os.path.join(HERE, *a)
MI_KM = 1.609344
LB_KG = 0.453592

def line(label, formula, value, unit=""):
    print(f"{label}: {formula} = {value}{(' ' + unit) if unit else ''}")

def hdr(t):
    print("\n== " + t)

# ---------------------------------------------------------------- reviewer's own reads
# (full gross lb, empty gross lb, full ticket, empty ticket, full date, empty date, full tractor/trailer, empty tractor/trailer)
TICKETS = {
 "SHPX215053_1": (73800, 29400, "1186426265990", "1124426266608", "2026-09-22", "2026-09-23", "CH1/KL720", "CH1/KL720 RL"),
 "SHPX215053_2": (73440, 28640, "1186426265989", "1186426266998", "2026-09-22", "2026-09-23", "CH2/KL188 RL", "CH2/KL188 RL"),
 "SHPX215053_3": (75200, 28680, "1124426266610", "1186426267015", "2026-09-23", "2026-09-24", "CH2/KL188 RL", "CH2/KL188"),
 "SHPX215053_4": (64180, 29640, "1186426266003", "1186426267016", "2026-09-23", "2026-09-24", "CH1/KL720 RL", "CH1/KL720"),
 "SHPX215039_1": (74500, 28680, "1124426261599", "1312226261666", "2026-09-18", "2026-09-18", "CH2/KL188", "CH2/KL188"),
 "SHPX215039_2": (78620, 29480, "1124426261598", "1312226261663", "2026-09-18", "2026-09-18", "CH1/KL720", "CH1/KL720"),
 "SHPX215039_3": (70620, 29280, "1124426262601", "1312226262682", "2026-09-19", "2026-09-19", "CH1/KL720", "CH1/KL720"),
 "SHPX215039_4": (63340, 28620, "1124426262600", "1312226262679", "2026-09-19", "2026-09-19", "CH2/KL188", "CH2/KL188"),
 "SHPX215091_1": (74200, 28680, "1312226260630", "1312226260636", "2026-09-17", "2026-09-17", "CH2/KL188", "CH2/KL188"),
 "SHPX215091_2": (72200, 29240, "1124426260591", "1312226260638", "2026-09-17", "2026-09-17", "CH1/KL720", "CH1/KL720"),
 "SHPX215091_3": (76060, 28560, "1312226260641", "1312226261651", "2026-09-17", "2026-09-18", "CH2/KL188", "CH2/KL188"),
 "SHPX215091_4": (64560, 29580, "1124426260594", "1312226261653", "2026-09-17", "2026-09-18", "CH1/KL720", "CH1/KL720"),
 "SHPX215079_3": (75800, 29200, "1124426248529", "1312226248408", "2026-09-05", "2026-09-05", "CH1/KL192", "CH1/KL192"),
 "SHPX215079_4": (60700, 28260, "1124426248533", "1124426249537", "2026-09-05", "2026-09-06", "CH2/KL188", "CH2/KL188"),
}
OPS_ROW = {"SHPX215053_1": 28, "SHPX215053_2": 29, "SHPX215053_3": 30, "SHPX215053_4": 31,
           "SHPX215039_1": 32, "SHPX215039_2": 33, "SHPX215039_3": 34, "SHPX215039_4": 35,
           "SHPX215091_1": 36, "SHPX215091_2": 37, "SHPX215091_3": 38, "SHPX215091_4": 39,
           "SHPX215079_3": 40, "SHPX215079_4": 41}
# BOL Shipping Details "Product Shipment Summary": (tare kg, full kg, delivered kg, Qty shipped MT, litres 60F, density 60F)
BOLS = {
 "SHPX215053": {"2497-R": (19210, 45120, 25910, 25.91, 21648, 1.1969), "2498-R": (19280, 45010, 25730, 25.73, 21497, 1.1969),
                "2500-R": (19050, 44980, 25930, 25.93, 21664, 1.1969)},
 "SHPX215039": {"2459-R": (19120, 44980, 25860, 25.86, 21566, 1.1991), "2460-R": (19540, 45460, 25920, 25.92, 21652, 1.1971),
                "2461-R": (19270, 45130, 25860, 25.86, 21602, 1.1971)},
 "SHPX215091": {"2465-R": (19370, 45390, 26020, 26.02, 21720, 1.1980), "2466-R": (19290, 45150, 25860, 25.86, 21586, 1.1980),
                "2467-R": (19330, 45220, 25890, 25.89, 21634, 1.1967)},
 "SHPX215079": {"2453-R": (19440, 45430, 25990, 25.99, 21595, 1.2035), "2454-R": (19280, 45190, 25910, 25.91, 21608, 1.1991),
                "2455-R": (19130, 45100, 25970, 25.97, 21658, 1.1991)},
}
# W2 offloads 1-2 of SHPX215079: reviewer's reads of ../batch-INJ-09-2026-W2-AECN/tickets/*SHPX215079*offload_1/2.pdf
# (1124426247525 9-04-26 14:07 76720 lb / 1124426248527 9-05-26 07:23 28660 lb; 1124426248528 9-05-26 09:24 73860 lb /
#  1124426248531 9-05-26 12:03 28420 lb reweigh; all CH2 / KL188)
W2_079 = {"offload_1": (76720, 28660), "offload_2": (73860, 28420)}
W2_079_RAIL_ALLOC_MI = (615.6079326, 582.0479496)   # W4 log prior_weeks.w2_SHPX215079.w2_rail_alloc_mi
LOSS_IMAGE = {"mass_kg": 341.43, "alloc": 0.0044, "transport_process_railcar_t": 5.3728, "transport_embodied_railcar_t": 0.14507,
              "pyro_process_t": 0.04682, "pyro_embodied_t": 0.00722, "total_t": 0.07833}

# ---------------------------------------------------------------- Certify pull
raw = json.load(open(P("source_exports", "certify_graphql_GetGhgEntryQAData_2026-10-06.json")))
g = json.loads(raw[2]["text"])["ghgEntry"]
UNIT_T = {"t": 1.0, "tCO₂e": 1.0, "kgCO₂e": 0.001, "kg": 0.001}
comps = {}
for grp in g["componentGroups"]:
    for att in grp["componentAttributions"]:
        c = att["component"]
        flux = att["carbonFlux"]["co2eTransferred"]
        inputs = {}
        for si in c["inputData"]["scalarInputs"]:
            dp = si["dataPoint"]; q = dp["quantity"]
            inputs[si["blueprintInput"]["inputKey"]] = {"dp": dp["id"], "name": dp["displayName"], "v": q["magnitude"],
                                                         "sd": q["standardDeviation"], "u": q["unit"], "nsrc": len(dp["references"])}
        comps[c["displayName"]] = {"type": c["componentType"], "bp": c["componentBlueprint"]["key"], "group": grp["displayName"],
                                   "dir": att["carbonFlux"]["direction"], "t": flux["magnitude"] * UNIT_T[flux["unit"]], "in": inputs}
DISC = g["uncertaintyDiscount"]["magnitude"]

def to_t(v, u):
    return v * {"t": 1.0, "kg": 0.001}[u]

def to_km(v, u):
    return v * {"km": 1.0, "mi": MI_KM}[u]

def comp_value(c, override=None):
    """Recompute one component in tCO2e from its inputs (override: {inputKey: value})."""
    ins = {k: dict(v) for k, v in c["in"].items()}
    if override:
        for k, v in override.items():
            ins[k]["v"] = v
    bp = c["bp"]
    if bp == "transport":
        ef = ins["carbon_intensity"]; assert ef["u"] == "kgCO₂e/(km⋅t)"
        return ef["v"] * to_km(ins["distance"]["v"], ins["distance"]["u"]) * to_t(ins["mass"]["v"], ins["mass"]["u"]) / 1000
    if bp == "distance_based_ci_emissions":
        return ins["carbon_intensity"]["v"] * to_km(ins["distance"]["v"], ins["distance"]["u"]) / 1000
    if bp == "mass_based_ci_emissions":
        ci = ins["carbon_intensity"]
        if ci["u"] == "tCO₂e/t":
            return ci["v"] * to_t(ins["mass"]["v"], ins["mass"]["u"])
        return ci["v"] * ins["mass"]["v"] / 1000          # kgCO2e/kg x kg
    if bp == "fuel_usage_by_mass":
        return ins["fuel_combustion_carbon_intensity"]["v"] * ins["mass_of_fuel"]["v"] / 1000
    if bp == "grid_electricity_use_with_recs":
        gi, gu = ins["grid_carbon_intensity"]["v"], ins["grid_electricity_use"]["v"]
        ri, pu = ins["procured_power_carbon_intensity"]["v"], ins["procured_power_electricity_use"]["v"]
        return (gu * gi - pu * gi + pu * ri) / 1000
    if bp in ("constant_activity_emissions", "constant_reduction"):
        k = list(ins)[0]
        return ins[k]["v"] * {"tCO₂e": 1.0, "kgCO₂e": 0.001}[ins[k]["u"]]
    if bp == "carbon_rich_substance_sequestration":
        return to_t(ins["product_mass"]["v"], ins["product_mass"]["u"]) * ins["carbon_content"]["v"] / 100 * 3.667
    raise ValueError(bp)

def net_terms(cs):
    seq = sum(comp_value(c) for c in cs.values() if c["type"] == "SEQUESTRATION")
    act = sum(comp_value(c) for c in cs.values() if c["type"] == "ACTIVITY")
    adj = sum(comp_value(c) for c in cs.values() if c["type"] == "ADJUSTMENT")
    return seq, act, adj

def discount(cs, show=False):
    """One combined one-sigma: each datapoint with a +/- moved by one sigma in every component that uses it; terms in quadrature."""
    base = sum(v for v in [net_terms(cs)[0] - net_terms(cs)[1] + net_terms(cs)[2]])
    dps = {}
    for cn, c in cs.items():
        for k, i in c["in"].items():
            if i["sd"]:
                dps.setdefault(i["dp"], {"name": i["name"], "sd": i["sd"], "uses": []})["uses"].append((cn, k))
    tot = 0.0; terms = []
    for dp, info in dps.items():
        moved = {cn: dict(c, **{"in": {k: dict(v) for k, v in c["in"].items()}}) for cn, c in cs.items()}
        for cn, k in info["uses"]:
            moved[cn]["in"][k]["v"] += info["sd"]
        s, a, d = net_terms(moved)
        delta = (s - a + d) - base
        terms.append((info["name"], dp, info["sd"], delta, [u[0] for u in info["uses"]]))
        tot += delta ** 2
    if show:
        for n, dp, sd, delta, uses in terms:
            print(f"   term {n} ({dp}) +/-{sd}: change in net {delta:+.6f} t  [used by: {', '.join(uses)}]")
    return math.sqrt(tot)

# ---------------------------------------------------------------- Ops Notes live
ops_path = P("source_exports", "INJ-09-2026-W4_Ops_Notes_live_export_2026-10-06.xlsx")
wv = openpyxl.load_workbook(ops_path, data_only=True)["AECN Rail"]
wf = openpyxl.load_workbook(ops_path)["AECN Rail"]

# ---------------------------------------------------------------- Standard EF sheet 2026 tab
ef = openpyxl.load_workbook(P("source_exports", "Standard_Emission_Factors_export_2026-10-06.xlsx"), data_only=True)["2026"]
EF_PYRO = ef["F4"].value; EF_EMB = ef["F5"].value; EF_TRUCK = ef["F6"].value
EF_RAIL = ef["F10"].value; EF_TRUCK_EE = ef["F11"].value; EF_RAIL_EE = ef["F13"].value

print("INDEPENDENT RECOMPUTE - INJ-09-2026-W4-AECN - adversarial reviewer -",
      datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%MZ"))

hdr("0. Standard EF sheet 2026 tab (export 2026-10-06), cells read")
for cell, lab in (("F4", "AECN process H2 2026 (MTCO2e/MT)"), ("F5", "AECN embodied (MTCO2e/MT)"), ("F6", "tanker process (MT/t-mi)"),
                  ("F10", "rail process (MT/t-mi)"), ("F11", "tanker embodied (MT/mi)"), ("F13", "railcar embodied (MT/mi)")):
    print(f"{lab}: 2026!{cell} = {ef[cell].value}")

# ---------------------------------------------------------------- 1. tickets
hdr("1. Scale tickets (reviewer's image reads) vs Ops Notes D/E/F")
nets = {}
for k, (full, empty, tf, te, df, de, rf, re_) in TICKETS.items():
    r = OPS_ROW[k]
    net = (full - empty) * LB_KG
    nets[k] = net
    okD = wv[f"D{r}"].value == full; okE = wv[f"E{r}"].value == empty
    print(f"{k}: ({full} - {empty}) lb x {LB_KG} = {net:.5f} kg | Ops D{r}={wv[f'D{r}'].value} E{r}={wv[f'E{r}'].value} F{r}={wv[f'F{r}'].value:.5f} | "
          f"D/E match={okD and okE} | rig full {rf} / empty {re_} | dates {df} / {de}")
sum_nets = sum(nets.values())
line("Sum of 14 nets", "sum((full-empty) x 0.453592)", f"{sum_nets:.5f}", "kg")
line("Ops A46", "read", wv["A46"].value, "kg")

hdr("2. BOLs (reviewer's reads) - internal ties and railcar totals")
bol_tot = {}
for rc, bs in BOLS.items():
    s = 0
    for b, (tare, full, dlv, qty, lit, dens) in bs.items():
        s += qty * 1000
        print(f"{b}: full {full} - tare {tare} = {full - tare} (delivered {dlv}); {dlv}/1000 = {dlv/1000} (Qty shipped {qty}); "
              f"{lit} L x {dens} = {lit*dens:.2f} kg (gap {lit*dens-dlv:+.2f})")
    bol_tot[rc] = s
    line(f"{rc} loaded", "sum(Qty shipped) x 1000", f"{s:.0f}", f"kg | Ops E = {wv['E' + str({'SHPX215053':10,'SHPX215039':11,'SHPX215091':12,'SHPX215079':13}[rc])].value}")

hdr("3. Railcar balances, losses and the SHPX215079 last-offload cut")
w2 = {k: (f - e) * LB_KG for k, (f, e) in W2_079.items()}
line("W2 SHPX215079 offloads 1-2", f"({W2_079['offload_1'][0]}-{W2_079['offload_1'][1]})x0.453592 + ({W2_079['offload_2'][0]}-{W2_079['offload_2'][1]})x0.453592",
     f"{w2['offload_1']:.5f} + {w2['offload_2']:.5f} = {sum(w2.values()):.5f}", f"kg (Ops H13 {wv['H13'].value})")
raw079 = sum(w2.values()) + nets["SHPX215079_3"] + nets["SHPX215079_4"]
over = raw079 - bol_tot["SHPX215079"]
line("SHPX215079 raw total", f"{sum(w2.values()):.5f} + {nets['SHPX215079_3']:.5f} + {nets['SHPX215079_4']:.5f}", f"{raw079:.5f}", "kg")
line("Over-draw", f"{raw079:.5f} - {bol_tot['SHPX215079']:.0f}", f"{over:.5f}", f"kg (Ops J41 {wv['J41'].value})")
cut4 = bol_tot["SHPX215079"] - sum(w2.values()) - nets["SHPX215079_3"]
line("Offload 4 after cut", f"{bol_tot['SHPX215079']:.0f} - {sum(w2.values()):.5f} - {nets['SHPX215079_3']:.5f}", f"{cut4:.5f}", "kg")
print("Chronology (reviewer reads of W2 and W4 tickets): off1 full 09-04 14:07; off2 full 09-05 09:24; off3 full 09-05 09:29; off4 full 09-05 14:20, empty 09-06 07:50 -> off4 is the last weighing")
loss = {}
for rc, rows in (("SHPX215053", range(1, 5)), ("SHPX215039", range(1, 5)), ("SHPX215091", range(1, 5))):
    d = sum(nets[f"{rc}_{i}"] for i in rows)
    loss[rc] = bol_tot[rc] - d
    line(f"{rc} loss", f"{bol_tot[rc]:.0f} - {d:.5f}", f"{loss[rc]:.5f}", f"kg ({loss[rc]/bol_tot[rc]*100:.3f} %)")
delivered = sum_nets - over
line("Delivered (post-cut)", f"{sum_nets:.5f} - {over:.5f}", f"{delivered:.5f}", f"kg | Certify Bio-Oil Mass {comps['AECN pyrolysis (from rail loads) - process emissions']['in']['mass']['v']} t")
br_mass = sum(nets[f"SHPX215053_{i}"] for i in range(1, 5))
op_mass = delivered - br_mass
line("Baton Rouge mass (SHPX215053)", "sum of its 4 nets", f"{br_mass:.5f}", "kg")
line("Opelousas mass", f"{delivered:.5f} - {br_mass:.5f}", f"{op_mass:.5f}", "kg")

hdr("4. Sparging, filtration, MOBI, gross")
sp = delivered * 0.0129; fi = delivered * (1 - 0.0129) * 0.00196; mobi = delivered - sp - fi
line("Sparging", f"{delivered:.5f} x 0.0129", f"{sp:.6f}", f"kg (Ops D46 {wv['D46'].value})")
line("Filtration", f"{delivered:.5f} x (1-0.0129) x 0.00196", f"{fi:.6f}", f"kg (Ops E46 {wv['E46'].value})")
line("MOBI", f"{delivered:.5f} - {sp:.6f} - {fi:.6f}", f"{mobi:.6f}", f"kg (Ops F46 {wv['F46'].value}; Certify {comps['Injection sequestration']['in']['product_mass']['v']} t)")
gross = mobi / 1000 * 0.44 * 3.667
line("Gross", f"{mobi/1000:.7f} t x 0.440 x 3.667", f"{gross:.6f}", f"t (Certify {comps['Injection sequestration']['t']:.6f})")
cs = openpyxl.load_workbook(P("drive_files", "MRV_AECN_Analysis.xlsx"), data_only=True)["Summary Statistics"]
mu, sdv = cs["B5"].value, cs["B6"].value
line("Carbon z-score", f"(44.0 - {mu}) / {sdv}", f"{(44.0-mu)/sdv:.4f}", "SD")
line("Two-SD band", f"{mu} -/+ 2 x {sdv}", f"{mu-2*sdv:.2f} to {mu+2*sdv:.2f}", "%")

hdr("5. Every AECN component recomputed from Certify's own unrounded inputs")
AECN = ["AECN to SOPOR Tanker Truck Transport Process Emissions", "SOPOR to Opelousas Rail Transport Process Emissions",
        "Opelousas to Basco - Tanker Truck Transport Process Emissions", "SOPOR to Baton Rouge Rail Process Emissions",
        "Baton Rouge to Basco Tanker Truck Process Emissions", "AECN to SOPOR Tanker Truck Transport Embodied Emissions",
        "SOPOR to Opelousas/Baton Rouge Rail Transport Embodied Emissions", "Opelousas/Baton Rouge to Basco Tanker Truck Transport Embodied Emissions",
        "AECN pyrolysis (from rail loads) - process emissions", "AECN pyrolysis (from rail loads) - embodied emissions",
        "AECN Bio-oil Loss", "Injection sequestration", "BCU - Transport Process"]
for n in AECN:
    c = comps[n]
    ins = "; ".join(f"{i['name']} {i['v']} {i['u']}" + (f" +/-{i['sd']}" if i['sd'] else "") + f" [{i['nsrc']} src]" for i in c["in"].values())
    print(f"{n}: recomputed {comp_value(c):.6f} t vs Certify {c['t']:.6f} t | inputs: {ins}")

hdr("6. Transport legs on the evidence basis (Standard EF sheet factors, Ops/guidance distances)")
aecn_sopor_mi = 173 / MI_KM
ev = {}
ev["aecn_sopor"] = delivered / 1000 * aecn_sopor_mi * EF_TRUCK
line("AECN-SOPOR truck process", f"{delivered/1000:.5f} t x (173 km / 1.609344 = {aecn_sopor_mi:.5f} mi) x {EF_TRUCK}", f"{ev['aecn_sopor']:.6f}", "t")
ev["opel_rail"] = op_mass / 1000 * 2199 * EF_RAIL
line("Opelousas rail process", f"{op_mass/1000:.5f} t x 2199 mi x {EF_RAIL}", f"{ev['opel_rail']:.6f}", "t")
ev["opel_truck"] = op_mass / 1000 * 42.9 * EF_TRUCK
line("Opelousas-Basco truck process (Opelousas mass)", f"{op_mass/1000:.5f} t x 42.9 mi x {EF_TRUCK}", f"{ev['opel_truck']:.6f}", "t")
cert_opel_truck = comps["Opelousas to Basco - Tanker Truck Transport Process Emissions"]["t"]
line("  Certify Opelousas truck overcharge", f"{cert_opel_truck:.6f} - {ev['opel_truck']:.6f}", f"{cert_opel_truck-ev['opel_truck']:.6f}", "t")
ev["br_rail"] = br_mass / 1000 * 2152 * EF_RAIL
line("Baton Rouge rail process on 2,152 mi", f"{br_mass/1000:.5f} t x 2152 mi x {EF_RAIL}", f"{ev['br_rail']:.6f}", "t")
cert_br_rail = comps["SOPOR to Baton Rouge Rail Process Emissions"]["t"]
line("  Certify Baton Rouge rail shortfall", f"{ev['br_rail']:.6f} - {cert_br_rail:.6f}", f"{ev['br_rail']-cert_br_rail:.6f}", "t")
line("  Certify '2153 km' in miles", "2153 / 1.609344", f"{2153/MI_KM:.2f}", "mi")
line("  2153 km as share of Watco 2,199 mi (3,538.9 km)", f"2153 / (2199 x 1.609344)", f"{2153/(2199*MI_KM):.3f}", "")
# ASSUMPTION: approximate public coordinates, not read from a session source; used only as a sanity check
def gc(a, b):
    R = 6371.0
    la1, lo1, la2, lo2 = map(math.radians, (*a, *b))
    h = math.sin((la2-la1)/2)**2 + math.cos(la1)*math.cos(la2)*math.sin((lo2-lo1)/2)**2
    return 2*R*math.asin(math.sqrt(h))
bc, brg = (49.22, -68.15), (30.45, -91.15)
line("  ASSUMPTION straight-line Baie-Comeau to Baton Rouge", "haversine((49.22N,68.15W),(30.45N,91.15W)), coordinates are approximate public values, not session-sourced", f"{gc(bc, brg):.0f}", "km")
ev["br_truck"] = br_mass / 1000 * 136 * EF_TRUCK
line("Baton Rouge-Basco truck process", f"{br_mass/1000:.5f} t x 136 mi x {EF_TRUCK}", f"{ev['br_truck']:.6f}", f"t (Certify {comps['Baton Rouge to Basco Tanker Truck Process Emissions']['t']:.6f}, GLEC library 0.071 kg/(km t))")
proc_cert = sum(comps[n]["t"] for n in AECN[:5])
proc_ev = sum(ev.values())
line("Transport process sum, Certify", " + ".join(f"{comps[n]['t']:.6f}" for n in AECN[:5]), f"{proc_cert:.7f}", "t")
line("Transport process sum, evidence", " + ".join(f"{v:.6f}" for v in ev.values()), f"{proc_ev:.7f}", "t")

hdr("7. Allocated distances and embodied legs")
L = sum(((wv[f"F{r}"].value - wv[f"J{r}"].value) / wv[f"E{ {'SHPX215053':10,'SHPX215039':11,'SHPX215091':12,'SHPX215079':13}[wv[f'C{r}'].value] }"].value) * 107.5 * 3 for r in range(28, 42))
line("AECN-SOPOR allocated (Ops method)", "sum((F-J)/E_railcar x 107.5 x 3)", f"{L:.6f}", f"mi (Ops L42 {wv['L42'].value}; Certify {comps['AECN to SOPOR Tanker Truck Transport Embodied Emissions']['in']['distance']['v']})")
line("  same on 107.49722 mi", f"{L:.6f} x 107.49722/107.5", f"{L*107.49722/107.5:.6f}", "mi")
print(f"Ops M28 formula (live): {getattr(wf['M28'].value, 'text', wf['M28'].value)}")
fr079 = (wv["F40"].value - wv["J40"].value + wv["F41"].value - wv["J41"].value) / 77870
one_way = 2152 + 2199 + 2199 + 2199 * fr079
line("Rail allocated, one way", f"2152 + 2199 + 2199 + 2199 x {fr079:.6f}", f"{one_way:.6f}", "mi")
line("Rail allocated, round trip (live M42)", f"2 x {one_way:.6f}", f"{2*one_way:.6f}", f"mi (Ops M42 {wv['M42'].value}; Certify {comps['SOPOR to Opelousas/Baton Rouge Rail Transport Embodied Emissions']['in']['distance']['v']})")
line("Round-trip extra rail embodied", f"{one_way:.6f} mi x {EF_RAIL_EE}", f"{one_way*EF_RAIL_EE:.6f}", "t")
w4_079 = wv["M40"].value + wv["M41"].value
line("SHPX215079 rail miles booked across weeks", f"W2 {W2_079_RAIL_ALLOC_MI[0]} + {W2_079_RAIL_ALLOC_MI[1]} (one way) + W4 {wv['M40'].value:.6f} + {wv['M41'].value:.6f} (round trip)",
     f"{sum(W2_079_RAIL_ALLOC_MI)+w4_079:.4f}", "mi (one-way railcar = 2,199; round trip = 4,398)")
line("Truck embodied allocated", "2 x (42.9 x 10 + 136 x 4)", f"{2*(42.9*10+136*4):.1f}", f"mi (Ops N42 {wv['N42'].value})")
for n, d_mi, e in (("AECN to SOPOR Tanker Truck Transport Embodied Emissions", wv["L42"].value, EF_TRUCK_EE),
                   ("SOPOR to Opelousas/Baton Rouge Rail Transport Embodied Emissions", wv["M42"].value, EF_RAIL_EE),
                   ("Opelousas/Baton Rouge to Basco Tanker Truck Transport Embodied Emissions", wv["N42"].value, EF_TRUCK_EE)):
    line(f"{n} (Ops miles x sheet EF)", f"{d_mi} x {e}", f"{d_mi*e:.6f}", f"t (Certify {comps[n]['t']:.6f})")

hdr("8. Pyrolysis")
for n, e in (("AECN pyrolysis (from rail loads) - process emissions", EF_PYRO), ("AECN pyrolysis (from rail loads) - embodied emissions", EF_EMB)):
    line(n + " (sheet CI)", f"{delivered/1000:.5f} t x {e}", f"{delivered/1000*e:.6f}", f"t (Certify {comps[n]['t']:.6f})")

hdr("9. AECN Bio-oil Loss")
per_t_truck = aecn_sopor_mi * EF_TRUCK
im = LOSS_IMAGE
line("Image: transport process per railcar", f"77.57 t x ({aecn_sopor_mi:.5f} x {EF_TRUCK} + 2199 x {EF_RAIL})", f"{77.57*(per_t_truck+2199*EF_RAIL):.4f}", f"t (image {im['transport_process_railcar_t']})")
line("Image: transport embodied per railcar", f"{aecn_sopor_mi:.5f} x 3 x {EF_TRUCK_EE} + 2199 x {EF_RAIL_EE}", f"{aecn_sopor_mi*3*EF_TRUCK_EE + 2199*EF_RAIL_EE:.5f}", f"t (image {im['transport_embodied_railcar_t']})")
line("Image total SHPX215053", f"{im['alloc']} x {im['transport_process_railcar_t']} + {im['alloc']} x {im['transport_embodied_railcar_t']} + {im['pyro_process_t']} + {im['pyro_embodied_t']}",
     f"{im['alloc']*im['transport_process_railcar_t'] + im['alloc']*im['transport_embodied_railcar_t'] + im['pyro_process_t'] + im['pyro_embodied_t']:.5f}", f"t (image {im['total_t']})")
def loss_t(rc, rail_mi):
    m = loss[rc]; L_ = bol_tot[rc]
    tp = L_/1000 * (per_t_truck + rail_mi * EF_RAIL); te = aecn_sopor_mi*3*EF_TRUCK_EE + rail_mi*EF_RAIL_EE
    return (m / L_) * tp + (m / L_) * te + m/1000 * (0.13712 + 0.02116)
la = sum(loss_t(rc, 2199) for rc in loss)
lb = loss_t("SHPX215053", 2152) + loss_t("SHPX215039", 2199) + loss_t("SHPX215091", 2199)
line("Loss, all three on Opelousas 2,199 mi", " + ".join(f"{rc} {loss_t(rc,2199):.6f}" for rc in loss), f"{la:.6f}", "t (Certify 0.145)")
line("Loss, SHPX215053 on its 2,152 mi leg", f"{loss_t('SHPX215053',2152):.6f} + {loss_t('SHPX215039',2199):.6f} + {loss_t('SHPX215091',2199):.6f}", f"{lb:.6f}", "t")
line("Loss per tonne lost", f"{la:.6f} / {sum(loss.values())/1000:.6f}", f"{la/(sum(loss.values())/1000):.4f}", "t/t (guidance 'approximately 0.229')")
rail_ee_in_loss = sum((loss[rc]/bol_tot[rc]) * 2199 * EF_RAIL_EE for rc in loss)
line("Rail embodied of lost oil also inside the loss figure", "sum(loss/BOL x 2199 x railcar EE EF)", f"{rail_ee_in_loss:.6f}", "t (main rail embodied already charges FULL railcars the whole trip)")

hdr("10. BCU against transport process")
bcu = comps["BCU - Transport Process"]["t"]
line("BCU minus Certify process sum", f"{bcu} - {proc_cert:.7f}", f"{bcu-proc_cert:+.7f}", "t")
line("BCU minus evidence process sum", f"{bcu} - {proc_ev:.7f}", f"{bcu-proc_ev:+.7f}", "t")
only_opel = proc_cert - cert_opel_truck + ev["opel_truck"]
line("If only the Opelousas truck mass is fixed", f"{proc_cert:.7f} - {cert_opel_truck:.6f} + {ev['opel_truck']:.6f}", f"{only_opel:.7f}", f"t (BCU would exceed by {bcu-only_opel:.7f})")

hdr("11. Net and uncertainty discount, Certify as entered")
s, a, d = net_terms(comps)
dis = discount(comps, show=True)
line("Discount rebuilt", "sqrt(sum of squared terms)", f"{dis:.10f}", f"t (Certify {DISC})")
line("Net", f"{s:.6f} + {d:.6f} - {a:.6f} - {dis:.6f}", f"{s+d-a-dis:.6f}", "t (Calculation view 362.068)")
line("Discount share of gross", f"{dis:.6f} / {s:.6f}", f"{dis/s*100:.3f}", "%")

hdr("12. Corrected scenarios (cumulative)")
import copy
cc = copy.deepcopy(comps)
cc["SOPOR to Baton Rouge Rail Process Emissions"]["in"]["distance"].update(v=2152 * MI_KM, u="km")
cc["SOPOR to Baton Rouge Rail Process Emissions"]["in"]["carbon_intensity"]["v"] = EF_RAIL * 1000 / MI_KM
def report(tag, cs_, bcu_follow):
    cs2 = copy.deepcopy(cs_)
    if bcu_follow:
        p = sum(comp_value(cs2[n]) for n in AECN[:5])
        cs2["BCU - Transport Process"]["in"]["constant_reduction"].update(v=p, u="tCO₂e")
    s_, a_, d_ = net_terms(cs2); di = discount(cs2)
    print(f"{tag} | BCU {'= corrected process sum' if bcu_follow else 'held at 19.339'} {comp_value(cs2['BCU - Transport Process']):.7f} t | "
          f"activities {a_:.6f} | adjustments {d_:.6f} | discount {di:.6f} | net {s_+d_-a_-di:.6f} t (Certify-as-entered {s+d-a-dis:.6f}; change {(s_+d_-a_-di)-(s+d-a-dis):+.6f})")
report("A  Baton Rouge rail on 2,152 mi, sheet rail EF", cc, False)
report("A  Baton Rouge rail on 2,152 mi, sheet rail EF", cc, True)
ot = cc["Opelousas to Basco - Tanker Truck Transport Process Emissions"]["in"]
op_dp = cc["SOPOR to Opelousas Rail Transport Process Emissions"]["in"]["mass"]
ot["mass"] = dict(op_dp)
report("B  A + Opelousas truck on Opelousas mass (its +/-0.7257 t)", cc, False)
report("B  A + Opelousas truck on Opelousas mass (its +/-0.7257 t)", cc, True)
cc["SOPOR to Baton Rouge Rail Process Emissions"]["in"]["mass"]["sd"] = 0.29028
cc["SOPOR to Baton Rouge Rail Process Emissions"]["in"]["mass"]["dp"] = cc["Baton Rouge to Basco Tanker Truck Process Emissions"]["in"]["mass"]["dp"]
report("C  B + Baton Rouge rail mass +/-0.29028 t (same datapoint as truck leg)", cc, False)
report("C  B + Baton Rouge rail mass +/-0.29028 t (same datapoint as truck leg)", cc, True)
cc["AECN Bio-oil Loss"]["in"]["constant_activity_emissions"].update(v=lb*1000, u="kgCO₂e")
report("D  C + loss with SHPX215053 on 2,152 mi", cc, False)
report("D  C + loss with SHPX215053 on 2,152 mi", cc, True)
cc2 = copy.deepcopy(comps)
cc2["SOPOR to Baton Rouge Rail Process Emissions"]["in"]["mass"]["sd"] = 0.29028
report("7a.1 alone: Baton Rouge rail mass +/-0.29028 t (separate datapoint)", cc2, False)
cc3 = copy.deepcopy(cc2); cc3["SOPOR to Baton Rouge Rail Process Emissions"]["in"]["mass"]["dp"] = cc3["Baton Rouge to Basco Tanker Truck Process Emissions"]["in"]["mass"]["dp"]
report("7a.1 alone: Baton Rouge rail mass +/-0.29028 t (shared datapoint)", cc3, False)

hdr("13. Ops copies attached in Certify")
for fn in ("src_1M3VZCDQF1S0FQ0D_INJ-09-2026-W4_Ops_Notes.xlsx", "src_1M3W1NWED1S0K0MP_INJ-09-2026-W4_Ops_Notes.xlsx"):
    w = openpyxl.load_workbook(P("certify_sources", "files", fn), data_only=True)["AECN Rail"]
    print(f"{fn}: K10={w['K10'].value} L10={w['L10'].value} G13={w['G13'].value} M42={w['M42'].value} N42={w['N42'].value} A49={w['A49'].value}")
print(f"live: K10={wv['K10'].value} L10={wv['L10'].value} G13={wv['G13'].value} M42={wv['M42'].value} N42={wv['N42'].value} A49={wv['A49'].value}")
print(f"live headers: P9='{wv['P9'].value}' Q9='{wv['Q9'].value}'; P13={wv['P13'].value} Q13={wv['Q13'].value}")

hdr("14. Ticket timing and rail-uncertainty precondition")
for k in ("SHPX215079_3", "SHPX215079_4"):
    f_ = datetime.date.fromisoformat(TICKETS[k][4])
    line(f"{k} days full weigh to injection", "2026-09-25 - " + TICKETS[k][4], (datetime.date(2026, 9, 25) - f_).days, "days")
for rc in BOLS:
    n_off = 4
    print(f"{rc}: offloads {n_off} vs inbound BOLs {len(BOLS[rc])} -> {'conservative' if n_off >= len(BOLS[rc]) else 'deficit'}")
