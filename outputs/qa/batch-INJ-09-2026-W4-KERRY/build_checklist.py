#!/usr/bin/env python3
"""Builds Kerry_Batch_INJ-09-2026-W4_QA_Checklist.xlsx from the canonical Kerry tanker template.

Findings cite the source-read log (batch_INJ-09-2026-W4-KERRY_source_reads.json). After the template
writes the sheet, every section / sidebar header row is filled #D6E4F0 across A-E (the template fills
only the merged A cell), and row heights are set from the evidence length so Excel shows the full note.
"""
import json, math, sys
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Alignment

WS = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(WS / "skills/qa/batch-qa-kerry/templates"))
sys.path.insert(0, str(WS / ".claude/skills/batch-qa-aecn/scripts"))
from build_kerry_checklist import build_kerry_checklist  # noqa: E402
from checklist_lib import audit_evidence  # noqa: E402

B = Path(__file__).resolve().parent
LOG = json.loads((B / "batch_INJ-09-2026-W4-KERRY_source_reads.json").read_text())
OUT = B / "Kerry_Batch_INJ-09-2026-W4_QA_Checklist.xlsx"

F = {}

# ---------------- Section 1: Injection Sequestration
F["1.1"] = dict(status="Pass", ref="SR26-05719.001 p.1; Certify Injection sequestration modal; Ops Notes Summary E19", evidence=(
    "Carbon 57.9 % (m/m) on SGS Certificate of Analysis SR26-05719.001 (client ID '9-2026-W4-KERRY', sampled 25-Sep-2026 07:00, "
    "ASTM D5291 Method C). The Injection sequestration modal reads Carbon Content of Bio-oil Injected 57.9 %, and Ops Notes "
    "Summary E19 holds 0.579 against composite sample SR26-05719.001, composited 2026-09-25. One tanker, one composite sample. "
    "[log:lab_results.carbon_pct, lab_results.certificate, lab_results.client_id, "
    "certify.component_by_key.injection_sequestration.carbon_content_pct, ops_notes.summary.E19_c_fraction, "
    "ops_notes.summary.C19_composite_sample_id]"))
F["1.2"] = dict(status="Pass", ref="Certify Sources tab; Drive Testing folder", evidence=(
    "The carbon datapoint links two sources on the Sources tab: 'SR26-05719 W4 Kerry AECN QOWV COA.pdf' and "
    "'INJ-09-2026-W4 Ops Notes - 2026-10-01'. The CoA attached in Certify (src_1M3VZ8S8W1S013HS) and the Drive Testing copy "
    "are the same file (identical sha256). Page 1 of 3 is the Kerry certificate and its carbon result is the value reported. "
    "[log:certify.attachment_map, lab_results.certify_copy_sha256, lab_results.drive_copy_sha256, lab_results.certificate]"))
F["1.3"] = dict(status="Pass", ref="Certify Injection sequestration modal (INPUTS)", evidence=(
    "Injection sequestration modal INPUTS: carbon content 57.9 with uncertainty 0.2316 %. 57.9 x 0.004 gives 0.2316, the "
    "0.4 % convention; Ops Notes Summary G19 carries the same 0.002316 as a fraction. "
    "[log:certify.component_by_key.injection_sequestration.carbon_content_pm, ops_notes.summary.G19_c_uncertainty_fraction]"))
F["1.4"] = dict(status="Pass", ref="Certify Injection sequestration modal; Ops Notes Kerry C18, Summary G11", evidence=(
    "Certify Mass of Injectant (per injection) reads 17,994.817 kg in the Injection sequestration modal (17.995 t on the "
    "Datapoints tab). Ops Notes Kerry C18 'MOBI (kg)' holds 17994.817; Summary G11 'MOBI (kg)' and B19 carry the same figure. Depends on the ticket attribution open under 1.7. "
    "[log:certify.component_by_key.injection_sequestration.mass_of_product_kg, ops_notes.rollup.C18_mobi_kg, "
    "ops_notes.summary.G11_mobi_kg, ops_notes.summary.B19_kerry_mobi_kg]"))
F["1.5"] = dict(status="Flag", ref="Ops Notes Kerry E10, F10, A18:C18, M4:M5; AECN Rail N7:O7; Aug W5 Kerry O5 note; VVB Guidance Jul 2026; CAR-474", evidence=(
    "Formulas read from the live sheet: E10 '=C10-D10' (74,000 - 34,300 lb, 39,700 lb); F10 '=E10*M4' with M4 0.453592 "
    "kg/lb gives 18,007.6024 kg; B18 '=A18*$M$5' with M5 0.00071 gives 12.7853977 kg filtration loss; C18 '=A18-B18' gives "
    "17,994.817 kg. The recompute reproduces each figure, and Summary E11 sparging loss for Kerry is 0, as the July 2026 VVB "
    "guidance requires ('Kerry takes a filtration deduction only. No sparging loss is applied to Kerry oil'). "
    "The Kerry solids constant itself is in question. M5 (0.00071) matches the CAR-474 SGS row 'MANITOWOC CALCINER PHASE', "
    "0.071 % solids, but it sits below the AECN constant on the same workbook (AECN Rail O7, 'Filtration Solids Content', "
    "0.00196), while the July guidance tells the verifier 'Kerry calciner oil measures materially higher in solids than AECN "
    "oil'. The August W5 Kerry O5 cell note records a lab discrepancy: the CAR-474 Slack announcement "
    "\"quotes 'Kerry Manitowoc Calciner Phase Separated Tar - 0.71% solids'\"; the note holds that \"the data file is "
    "authoritative and the Slack text is treated as a transcription slip\" and adds \"Confirm with Steven Rowland before this "
    "constant is ever changed.\" Only the 0.71 % figure is consistent with the July wording. On this week's sheet M5 carries no cell note and no link, though the July "
    "guidance states 'The Kerry solids value and its source analysis are cited on the constants block of the Kerry tab in "
    "the weekly ops notes', and CAR-474 is not one of the Certify sources. At 0.71 % the injected mass would be 17,879.748 kg "
    "and sequestration 37.9622 t, 0.2443 t (0.64 %) lower: below materiality, but over-crediting if the lab confirms 0.71 %. "
    "Depends on the ticket attribution open under 1.7. "
    "[log:ops_notes.kerry_row, ops_notes.rollup, ops_notes.constants, ops_notes.summary.E11_sparging_loss_kg, "
    "vvb_guidance.kerry_method_july_verbatim, filtration_constant_source, solids_constant_check]"))
F["1.6"] = dict(status="Pass", ref="Certify Injection sequestration modal (INPUTS); both CAT Scale tickets", evidence=(
    "Injection sequestration modal INPUTS: mass of product 17,994.817 kg with uncertainty 72.57 kg. Both tickets are CAT "
    "Scale (Certified Automated Truck Scale) weighings, a commercial truck scale, and the tanker is one discrete weighing "
    "event (one full + empty pair), so 72.57 kg is the expected precision. Depends on the ticket attribution open under 1.7. "
    "[log:certify.component_by_key.injection_sequestration.mass_pm_kg, scale_tickets.scale_type, "
    "scale_tickets]"))
F["1.7"] = dict(status="Fail", ref="Kerry Oil scale tickets; COBB 'Basco Injection -- COBB' W2-W4 calciner rows; W3 Kerry Certify removal rmv_1M3SWNGK51S0GH00; W3 Ops Notes; VVB Guidance Jul-Sep 2026; KB Section 10A", evidence=(
    "The credited ticket pair is genuine and its arithmetic ties (1.5), but the evidence does not consistently place this "
    "truckload in this removal (19-25 Sep 2026). "
    "(a) Full ticket #1189526252027 is dated 9-09-26 (Wilson Township WI) and the post-offload Empty ticket #1319926254821 "
    "9-11-26 (Lake Village AR); Drive shows the empty-ticket image filed 2026-09-14. The Ops Notes Kerry row is dated "
    "2026-09-25, 16 days after the Full ticket, against the knowledge-base check 'For each pair: dates within 2 weeks of "
    "injection [VR Sec 3.5]'; COBB records injection completion on 2026-09-23 (14 days). "
    "(b) The live COBB row (cell A 'INJ-09-2026-W4_calciner', lot WILA004_offload_1) records the WILA004 offload order as "
    "completed 2026-09-24 18:04:36, 13 days after this pair's post-offload empty weighing and one day after the row's own "
    "injection completion date. "
    "(c) COBB's offload completion dates track the empty tickets elsewhere: WILA002 completed 2026-09-02 19:41:31, the date of "
    "W2's empty ticket (9-02-26); WILA003 completed 2026-09-12 14:17:56, the day after this pair's empty ticket. The COBB "
    "quantity cannot tell the loads apart (WILA003 and WILA004 both carry 18,007.6024 kg). The W3 Kerry removal in Certify "
    "(rmv_1M3SWNGK51S0GH00, 12-18 Sep 2026) credits WILA003 with the later pair: Sept_W3 Full 9-16-26 (76,680 lb) / Empty "
    "9-18-26 (34,420 lb), Mass of Injectant 19.155 t. W3's composite carbon sample is dated 2026-09-16 (W3 Summary D19), the "
    "morning that truck was weighed full in Wisconsin. By ticket date the tanker IDs run WILA002 (8-31), WILA004 (9-09), "
    "WILA003 (9-16). "
    "(d) The July 2026 VVB guidance covers a tank hold ('held in a dedicated tank and progressively injected in combination "
    "with AECN injections'; 'the associated removal is not considered complete until the tank is empty'). Neither it nor the "
    "August or September guidance, nor any Ops Notes cell note, addresses one tank completing a later load before an earlier "
    "one, or the COBB dates above. "
    "If this removal should carry the 16-18 Sep pair instead, its sequestration would be 40.6702 t rather than 38.2065 t "
    "(2.4637 t, 6.45 % of this removal, above the 5 % materiality threshold), and its transport process emission 2.38818 t "
    "(0.14467 t more, with the BCU rising by the same amount). Rows 1.4, 1.5, 1.6, 2A.5, 2A.6, 2A.7 and 2C.1 rest on this "
    "ticket pair. "
    "[log:scale_tickets.full_ticket, scale_tickets.empty_ticket, ops_notes.kerry_row.A10_date, cobb_tracker.w4_row, "
    "cobb_tracker.w3_row, cobb_tracker.w2_row, cross_week_kerry.w2_tickets, cross_week_kerry.w3_tickets, "
    "cross_week_kerry.w3_certify, cross_week_kerry.w3_summary_row19, cross_week_kerry.attribution_sensitivity, "
    "cross_week_kerry.trip_order_by_ticket_date, vvb_guidance.kerry_method_july_verbatim, vvb_guidance.bearing_on_this_batch, "
    "reference_rules]"))

# ---------------- Section 2A: Transport, process
F["2A.1"] = dict(status="Pass", ref="Certify process modal; Standard EFs live, 2026 tab row 6", evidence=(
    "Process modal INPUTS: 'Emissions Factor - Truck Transport Fuel Use (Well to Wheel) - Tanker Truck' 0.1143 "
    "kgCO2e/(t.mi). Live Standard Emission Factors sheet, 2026 tab, row 6 'Truck Transport Fuel Use (Well to Wheel) -- "
    "Tanker Truck' 0.0001143 MT CO2e/t-mi, GLEC V3.2 2025: the same value, current tab. The tanker class matches the vehicle: "
    "both tickets record commodity 'TANKER' on trailer 1013, and the July 2026 VVB guidance states Kerry oil 'is delivered to "
    "the Basco injection site in bulk tanker trailers'. "
    "[log:certify.component_by_key.transport_process, standard_efs_2026, scale_tickets.full_ticket.commodity, "
    "vvb_guidance.kerry_method_july_verbatim.delivery_vehicle]"))
F["2A.2"] = dict(status="Pass", ref="Certify Sources tab; attached EF workbook src_1KWZ2NVQM1S08AWN, 2026 tab row 5", evidence=(
    "The process EF datapoint links 'Standard Emission Factors + Calculations (8).xlsx' (src_1KWZ2NVQM1S08AWN). Opened the "
    "attached workbook: 2026 tab row 5 'Truck Transport Fuel Use (Well to Wheel) -- Tanker Truck' 0.0001143 MT CO2e/t-mi, "
    "matching the reported 0.1143 kgCO2e/(t.mi) and the live sheet. "
    "[log:certify.attachment_map, standard_efs_2026.certify_attached_copy, standard_efs_2026]"))
F["2A.3"] = dict(status="Pass", ref="Certify process modal; Kerry to Basco Distance.png", evidence=(
    "Process modal distance 1,090 mi ('Route Distance Kerry to Basco'). The route map runs 4343 Expo Dr, Manitowoc WI to "
    "4687 LA-115, St Landry LA and offers 1,059, 1,071 and 1,090 mi; 1,090 mi is the longest, as the convention requires. "
    "The Full ticket was weighed at Wilson Township WI on I-43, consistent with this origin. "
    "[log:certify.component_by_key.transport_process.distance_mi, route_map, scale_tickets.full_ticket.location]"))
F["2A.4"] = dict(status="Pass", ref="Certify Sources tab; attached src_1KT1W9XED1S08J0H", evidence=(
    "'Route Distance Kerry to Basco' links 'Kerry to Basco Distance.png' (src_1KT1W9XED1S08J0H). Opened the attached image: "
    "Google Maps, Manitowoc WI to St Landry LA, three routes, longest 'via I-55 S', 1,090 miles. "
    "[log:route_map, certify.attachment_map]"))
F["2A.5"] = dict(status="Pass", ref="Certify process modal; Ops Notes Summary B11, Kerry A18", evidence=(
    "Process modal mass of load 'Allocated Kerry Oil Gross Mass' 18.0076 t. Ops Notes Summary B11 'Σ Deliveries / Cargo (kg)' "
    "and Kerry A18 hold 18007.6024 kg: the pre-filtration ticket net, which is the oil-only mass hauled for a tanker load "
    "(FC-K2), not the post-filtration MOBI. Depends on the ticket attribution open under 1.7. "
    "[log:certify.component_by_key.transport_process.mass_of_load_t, ops_notes.summary.B11_deliveries_cargo_kg, "
    "ops_notes.rollup.A18_sum_net_kg]"))
F["2A.6"] = dict(status="Pass", ref="Both CAT Scale ticket images; Ops Notes Kerry E10:F10, A18", evidence=(
    "Full 74,000 lb less Empty 34,300 lb gives 39,700 lb; at 0.453592 kg/lb that is 18,007.6024 kg (Kerry F10 '=E10*M4', "
    "A18 sums F10:F15). Recomputed from the two ticket images: the same figure. Tractor 07 / trailer 1013 on both tickets. "
    "Depends on the ticket attribution open under 1.7. "
    "[log:scale_tickets.net_kg, scale_tickets.tractor_trailer_match, ops_notes.kerry_row]"))
F["2A.7"] = dict(status="Pass", ref="Certify process modal (INPUTS)", evidence=(
    "Process modal INPUTS: mass of load 18.0076 t with uncertainty 0.07257 t (72.57 kg), one discrete CAT Scale weighing "
    "for the tanker. Depends on the ticket attribution open under 1.7. "
    "[log:certify.component_by_key.transport_process.mass_of_load_pm_t, scale_tickets.scale_type, "
    "scale_tickets]"))

# ---------------- Section 2B: Transport, embodied
F["2B.1"] = dict(status="Pass", ref="Certify embodied modal; Standard EFs live, 2026 tab row 11", evidence=(
    "Embodied modal: 'Emissions Factor - Vehicle Embodied Emissions - Tanker Truck' 0.23756 kgCO2e/mi (0.1476 kgCO2e/km on "
    "the Datapoints tab). Live 2026 tab row 11 'Vehicle Embodied Emissions --Tanker Truck' 0.0002375605888 MT CO2e/mi, "
    "derived on the 'Truck and Trailer EE 2026' tab (B17); per mile, not per ton-mile (FC8). "
    "[log:certify.component_by_key.transport_embodied, standard_efs_2026]"))
F["2B.2"] = dict(status="Pass", ref="Certify Sources tab; attached EF workbook src_1KVJRSYT91S09G2E, 2026 tab row 10", evidence=(
    "The embodied EF datapoint links 'Standard Emission Factors + Calculations (8).xlsx' (src_1KVJRSYT91S09G2E, the same "
    "file as the process-side copy). Attached workbook, 2026 tab row 10 'Vehicle Embodied Emissions --Tanker Truck' "
    "0.0002375605888 MT CO2e/mi, matching the reported value. "
    "[log:certify.attachment_map, standard_efs_2026.certify_attached_copy, standard_efs_2026]"))
F["2B.3"] = dict(status="Pass", ref="Certify embodied modal", evidence=(
    "Embodied modal distance 1,090 mi, the same 'Route Distance Kerry to Basco' datapoint the process component uses "
    "(1 source); the longest map option. "
    "[log:certify.component_by_key.transport_embodied.distance_mi, route_map.longest_mi]"))
F["2B.4"] = dict(status="Pass", ref="Certify Sources tab; attached src_1KT1W9XED1S08J0H", evidence=(
    "Same datapoint and source as 2A.4: 'Kerry to Basco Distance.png', opened; longest route 1,090 miles. "
    "[log:route_map, certify.attachment_map]"))

# ---------------- Section 2C: BCU
F["2C.1"] = dict(status="Pass", ref="Certify BCU modal; retirement certificate src_1M1MFJPW11S08KY2", evidence=(
    "BCU component 'BCUs - Applied to Bio-Oil Transport Process Emissions': constant reduction 2,243.51 kgCO2e (2.244 t). "
    "The recomputed process emission is 18.0076024 t x 1,090 mi x 0.0001143 t/(t.mi), 2,243.513 kg; the BCU matches it and "
    "does not exceed it. Evidence: 'Charm-RetirementCertificate-RD-EACs-20251218.pdf', 2,795 Renewable Diesel TERCs retired "
    "Dec 18, 2025 on behalf of Charm Industrial, Inc (vintage 2024, Turner, Mason & Company). Embodied emissions carry no "
    "BCU. Depends on the ticket attribution open under 1.7. [log:certify.component_by_key.bcu, certify.component_by_key.transport_process, retirement_certificate, "
    "standard_efs_2026]"))

# ---------------- Section 3: MRV parameters
F["3.1"] = dict(status="Pass", ref="SR26-05719.001", evidence=(
    "Density at 15 °C 1.201 g/cm³ (ASTM D4052), above the 1.01 g/cm³ formation-fluid floor. One tanker load, one sample. "
    "[log:lab_results.density_15C_g_cm3]"))
F["3.2"] = dict(status="Pass", ref="SR26-05719.001; SR26-05719.002", evidence=(
    "Hydrogen 6.4 % (m/m) (ASTM D5291 Method C), reported. No Kerry-specific range is established yet; the AECN sample on "
    "the same report reads 6.6 %. [log:lab_results.hydrogen_pct, lab_results.aecn_sample_same_report]"))
F["3.3"] = dict(status="Pass", ref="SR26-05719.001", evidence=(
    "Nitrogen <0.75 % (m/m) (ASTM D5291 Method C), reported below the method reporting level. [log:lab_results.nitrogen]"))
F["3.4"] = dict(status="Pass", ref="SR26-05719.001", evidence=(
    "Kinematic viscosity at 40 °C 545.2 cSt (ASTM D445), reported. [log:lab_results.kinematic_viscosity_cst_40C]"))
F["3.5"] = dict(status="Pass", ref="SR26-05719.001", evidence=(
    "Acid number 154 mg KOH/g (ASTM D664 Method A), reported; SGS marks it '#' (outside the method precision study range), "
    "as it does the carbon and hydrogen results. [log:lab_results.acid_number_mgKOH_g, lab_results.lab_flags]"))
F["3.6"] = dict(status="Pass", ref="Ops Notes Kerry D18 (cell note), G10, M6; COBB W4 row; VVB Guidance Jul 2026", evidence=(
    "Ops Notes Kerry D18 pH 2.97, cell note 'Wellhead pH = weekly median of Ignition PH-1041, gated to StateRun=1; provided "
    "by G. Lutz 2026-09-30.', carried to G10; COBB WILA004 injection_ph 2.97. Within the 2.0 - 4.0 target in M6. The July "
    "2026 VVB guidance: 'we do not have an independent pH value that is specific to Kerry. Rather, the pH measured over the "
    "course of the month’s AECN injections is net of Kerry oil as well.' "
    "[log:ops_notes.rollup.D18_ph, vvb_guidance.ops_notes_cell_comments, cobb_tracker.w4_row.injection_ph, "
    "ops_notes.constants.M6_ph_target, vvb_guidance.kerry_method_july_verbatim.ph]"))

SIDEBAR = [
    dict(title="System boundary (waste product)", ref="Certify Components tab; Ops Notes Summary E11", detail=(
        "Components tab lists 4 components: Injection sequestration 38.206 t; Kerry to Basco Tanker Truck Transport Process "
        "Emissions 2.244 t; Kerry to Basco Tanker Truck Transport Embodied Emissions 0.259 t; BCUs - Applied to Bio-Oil "
        "Transport Process Emissions 2.244 t. There is no Kerry production, pyrolysis, biomass-sourcing or counterfactual "
        "component, and Kerry sparging loss is 0 kg (Summary E11). Matches the waste boundary.")),
    dict(title="W2 / W3 Kerry records that bear on 1.7", ref="COBB calciner rows; W2 and W3 Ops Notes Kerry tabs", detail=(
        "These belong to the W2 and W3 QAs. COBB WILA003 quantity 18,007.6024 kg vs W3 Ops Notes WILA003 net 19,168.79792 kg; "
        "COBB WILA002 quantity 17,000 kg vs W2 Ops Notes 18,234.3984 kg; the W2 Ops Notes Kerry row is dated 2026-09-09 while "
        "its tickets are dated 8-31-26 (full) and 9-02-26 (empty). The W4 ticket pair appears on no other week's Kerry tab "
        "(Aug W5, Sep W1-W3 and Oct W1 checked).")),
    dict(title="CoA sample location label", ref="SR26-05719.001, .002", detail=(
        "SR26-05719.001 gives LOCATION 'FORT LUPTON, CO' for the Kerry sample (the AECN sample, .002, reads the same). Kerry "
        "oil is delivered to Basco (tickets: 'DELIVERY SITE: CHARM'); the client ID '9-2026-W4-KERRY', SGS order 5506143 and "
        "sample date tie the certificate to this batch. Label only.")),
    dict(title="Kerry tab column labels", ref="Ops Notes Kerry C9:F9, C10, F10, A17:C18, I10:K10", detail=(
        "C10, under the header 'Gross Mass (lb) — transport', holds the full-vehicle ticket weight (74,000 lb), not the oil "
        "transported; F10 ('Net Injected (kg) — MOBI') and A18 ('Σ Net Injected (kg)') hold the pre-filtration 18,007.60 kg, "
        "while C18 ('MOBI (kg)') holds 17,994.817 kg, the figure in Certify. I10:K10 (C content and its uncertainties) are blank on the tanker row; "
        "the composite carbon sits on Summary row 19. Values trace; the labels could mislead a reader.")),
    dict(title="No Bill of Lading; COBB origin quantity", ref="Drive W4 folder; Certify Sources; COBB W4 row", detail=(
        "No Bill of Lading among the Drive folder files or the 8 Certify sources, as on earlier Kerry tanker batches; per the "
        "July guidance the CAT ticket pair is the mass evidence. COBB BOL_number holds the tanker UID 'WILA004', and "
        "origin_quantity_kg reads 20,000 kg, a round figure that matches no weighing and enters no Certify datapoint.")),
    dict(title="Transport EF vehicle class", ref="Standard EFs 2026 tab rows 6, 11; both W4 tickets; VVB Guidance Jul 2026", detail=(
        "This removal uses tanker-truck factors (process 0.0001143 MT CO2e/t-mi, embodied 0.0002375605888 MT CO2e/mi), matching the "
        "vehicle: both tickets record commodity 'TANKER' and the July guidance describes delivery in bulk tanker trailers. Kerry moved "
        "from totes on flatbeds to tankers in July 2026; tanker deliveries take the tanker-truck factors, tote deliveries took the "
        "flatbed and non-tanker factors. The first Kerry tanker batch (INJ-07-2026-W4-KERRY) was reported with the flatbed process "
        "(0.0001384) and non-tanker embodied (0.0001927908256) factors.")),
    dict(title="COBB lot naming and repeated blocks", ref="COBB 'Basco Injection -- COBB'", detail=(
        "Kerry tanker rows use consumed_lot 'WILA00N_offload_1' (W4: 'WILA004_offload_1'); the kc_NNNN pattern in the Kerry QA "
        "reference applies to the tote batches K-001 / K-002. The COBB tab repeats its dataset in three blocks (W4 calciner at "
        "rows 1790, 2718, 3681; pH blank in the third).")),
]

meta = dict(batch_id="INJ-09-2026-W4-KERRY", reporting_period="2026-09-01 to 2026-09-30 (removal 19-25 Sep 2026)",
            reviewer="Claude QA Agent", date="2026-10-06", injection_site="Basco 6", delivery_mode="Tanker",
            certify_removal="rmv_1M3T2J3V51S0NCHH")

# pre-build audit of every row against the log (same rules as audit_checklist.py)
bad = []
for k, v in F.items():
    ok, gaps = audit_evidence(k, v["status"], v["evidence"], LOG)
    if not ok:
        bad.append((k, gaps))
if bad:
    for k, g in bad:
        print("AUDIT", k, g)
    sys.exit(1)

res = build_kerry_checklist(str(OUT), meta, F, SIDEBAR, mode="tanker")

# post-process: headers #D6E4F0 across A-E, wrap + heights
wb = openpyxl.load_workbook(OUT)
ws = wb.active
HDR = PatternFill("solid", fgColor="D6E4F0")
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if isinstance(a, str) and (a.startswith("Section") or a.startswith("Sidebar")) and ws.cell(r, 3).value in (None, ""):
        rng = f"A{r}:E{r}"
        if rng in [str(m) for m in ws.merged_cells.ranges]:
            ws.unmerge_cells(rng)   # merged B:E cells do not keep their own fill
        for c in range(1, 6):
            ws.cell(r, c).fill = HDR
    d = ws.cell(r, 4).value
    if isinstance(d, str) and len(d) > 90 and ws.cell(r, 1).value not in (None, "Item #", "#"):
        lines = sum(math.ceil(max(len(p), 1) / 95) for p in d.split("\n"))
        b_lines = math.ceil(len(str(ws.cell(r, 2).value or "")) / 52)
        ws.row_dimensions[r].height = max(30, 15 * max(lines, b_lines) + 6)
        for c in range(1, 6):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
# print layout: one page wide, landscape; column A wide enough for the title-block labels
ws.column_dimensions["A"].width = 17
ws.page_setup.orientation = "landscape"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_options.gridLines = False
# status counts row
cnt = res["counts"]
r = ws.max_row + 2
ws.cell(r, 1, "STATUS COUNTS")
ws.cell(r, 2, f"Pass {cnt['Pass']} | Fail {cnt['Fail']} | Flag {cnt['Flag']} | N/A {cnt['N/A']} (25 items)")
wb.save(OUT)
print(json.dumps(res["counts"]), OUT.name)
