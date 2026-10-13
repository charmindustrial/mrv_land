#!/usr/bin/env python3
"""Assemble batch_INJ-09-2026-W4-QOWV_source_reads.json from the exports taken this session
(source_exports/), the tag analyses (inventory_tag_hits.json, cobb_w4_aqueous_rows.json,
cross_week_tag_scan.json) and the Certify views read in Playwright on 2026-10-06.
Values are computed from the files; Certify observations are recorded as read on screen."""
import json, re, collections, datetime
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
OPS = HERE / "source_exports/INJ-09-2026-W4_Ops_Notes.xlsx"
ATT = HERE / "source_exports/Certify_attached_Ops_Notes_src_1M3VZCDQF1S0FQ0D.xlsx"
EFX = HERE / "source_exports/Standard_EFs_live_export.xlsx"
INV = HERE / "source_exports/inventory_master_export.xlsx"
OUT = HERE / "batch_INJ-09-2026-W4-QOWV_source_reads.json"
CERTIFY_URL = "https://registry.isometric.com/account/certify/project/prj_1HZSSWBQM1S08H83/ghg-entry/rmv_1M3SG7D701S0HZ1Z/edit"
OPS_URL = "https://docs.google.com/spreadsheets/d/1zgrkVKQsGlzwjiumtn1vkYSaCt5_XTr7UWs3lR3RORQ/edit"
EF_URL = "https://docs.google.com/spreadsheets/d/1RPm-t6EyKIk_MQicbTitx1JLjj2rBkH7kz7Nx0N1ug4/edit?gid=684080705"
INV_URL = "https://docs.google.com/spreadsheets/d/1cD9cd_y_SI8GA1Yv8TFtDPqevKKF1Jh42WFRWSSMJ6g/edit"

wv = openpyxl.load_workbook(OPS, data_only=True); wf = openpyxl.load_workbook(OPS)
q, qf = wv["QOWV"], wf["QOWV"]; sh, shf = wv["QOWV Shipments & EE Allocation"], wf["QOWV Shipments & EE Allocation"]
sm, smf = wv["Summary"], wf["Summary"]

# ---------------- totes ----------------
totes = []
for r in range(11, 53):
    g = q.cell(r, 7).value
    totes.append(dict(row=r, tag=str(int(q.cell(r, 1).value)), gross_kg=q.cell(r, 2).value, tote_mass_kg=q.cell(r, 3).value,
                      net_oil_produced_kg=q.cell(r, 4).value, empty_tote_weight_kg=q.cell(r, 5).value,
                      net_oil_injected_kg=q.cell(r, 6).value, production_period_label=g,
                      production_period=g.split(", ")[-1], ci_kgco2e_per_kg=q.cell(r, 8).value,
                      ci_formula=qf.cell(r, 8).value, production_emissions_kg=q.cell(r, 9).value,
                      scale_uncertainty_kg=q.cell(r, 10).value, fraction="QOWV", source_doc="INJ-09-2026-W4 Ops Notes, QOWV tab"))
n = len(totes)
pp_counts = collections.Counter(t["production_period"] for t in totes)
pp_mass = collections.defaultdict(float); pp_pe = collections.defaultdict(float)
for t in totes:
    pp_mass[t["production_period"]] += t["net_oil_produced_kg"]; pp_pe[t["production_period"]] += t["production_emissions_kg"]

# ---------------- inventory record per tote ----------------
hits = {h["tag"]: h for h in json.load(open(HERE / "inventory_tag_hits.json"))}
for t in totes:
    inv = hits[t["tag"]]["inventory"][0]
    t.update(inventory_row=inv["inv_row"], inventory_pp=inv["pp"], inventory_fraction=inv["fraction"],
             inventory_gross_kg=inv["gross"], inventory_tare_kg=inv["tare"], ship_status=inv["ship_status"],
             date_shipped=inv["date_shipped"][:10], spray_paint=inv["paint"])
wr = {w["tag"]: w for w in json.load(open(HERE / "tote_weighing_records.json"))}
# tote 300022 has no record under its own tag; its lot 4746 is weighed under tag 300020 (Historical Tote Record row 840)
htr = openpyxl.load_workbook(HERE / "source_exports/bio_oil_injection_tracker_live_2026-10-06.xlsx", read_only=True, data_only=True)["Historical Tote Record"]
r840 = next(htr.iter_rows(min_row=840, max_row=840, values_only=True))
assert r840[6] == 300020 and r840[7] == 4746 and r840[5] == 1012.0 and r840[3] == "WO00001655", r840
lot4746 = {"row": 840, "reported_on": f"{r840[1]:%Y-%m-%d %H:%M:%S}", "order": r840[3], "gross": r840[5], "net": r840[0],
           "lot": "4746", "filed_under_tag": "300020"}
for t in totes:
    w = wr[t["tag"]]
    t["weighing_site"] = w["q4_site"]
    t["weighing_record"] = w["htr"][0] if w["htr"] else (lot4746 if t["tag"] == "300022" else None)
    t["weighing_date"] = t["weighing_record"]["reported_on"][:19] if t["weighing_record"] else None
    t["weighing_location"] = ("not recorded in the weighing record (the Historical Tote Record has no location column); the Q4 production tab's "
                              f"column A (header 'M') reads '{w['q4_site']}'; inferred: Charm's Fort Lupton, CO pyrolysis facility [log:tote_weighing_records.location_basis]")
wdates = sorted(t["weighing_date"] for t in totes if t["weighing_date"])
late_july = [dict(tag=t["tag"], weighing_date=t["weighing_date"], order=t["weighing_record"]["order"]) for t in totes
             if t["production_period"] == "July 2026" and t["weighing_date"][:7] != "2026-07"]
late_june = [dict(tag=t["tag"], weighing_date=t["weighing_date"], order=t["weighing_record"]["order"], inventory_row=t["inventory_row"],
                  net_oil_produced_kg=t["net_oil_produced_kg"], ci_kgco2e_per_kg=t["ci_kgco2e_per_kg"]) for t in totes
             if t["production_period"] == "June 2026" and t["weighing_date"][:7] != "2026-06"]
late_june_kg = round(sum(x["net_oil_produced_kg"] for x in late_june), 4)
match = dict(tag=n, fraction=sum(t["inventory_fraction"] == "QOWV" for t in totes),
             gross=sum(abs(t["inventory_gross_kg"] - t["gross_kg"]) < 1e-9 for t in totes),
             production_period=sum(t["inventory_pp"].replace(" PP", "") == t["production_period"] for t in totes))
by_truck = collections.Counter(t["ship_status"] for t in totes)

# all totes on each truck, every fraction
invws = openpyxl.load_workbook(INV, data_only=True, read_only=True)["Q4 2025 -2026 Production Invent"]
truck_all = collections.defaultdict(lambda: collections.Counter())
for row in invws.iter_rows(min_row=3, values_only=True):
    s = row[11]
    if s and re.search(r"COLA[\s_-]*(3[7-9]|4[01])\b", str(s)):
        truck_all[re.sub(r"[\s_-]+", "_", str(s).strip()).upper()][row[7]] += 1
tare_dist = collections.Counter(t["inventory_tare_kg"] for t in totes)

# ---------------- allocation tab ----------------
trucks = []
for r in range(12, 16):
    trucks.append(dict(row=r, truck=sh.cell(r, 1).value, ship_date=str(sh.cell(r, 2).value)[:10], totes_on_truck=sh.cell(r, 3).value,
                       batch_totes=sh.cell(r, 4).value, share=sh.cell(r, 5).value, share_formula=shf.cell(r, 5).value,
                       mileage=sh.cell(r, 6).value, mileage_formula=shf.cell(r, 6).value,
                       allocated_miles=sh.cell(r, 7).value, allocated_miles_formula=shf.cell(r, 7).value))
inv_den = {"COLA-37": sum(truck_all["COLA_37"].values()), "COLA-38": sum(truck_all["COLA_38"].values()),
           "COLA-39": sum(truck_all["COLA_39"].values()), "COLA-40": sum(truck_all["COLA_40"].values())}
alloc_inv = sum(t["batch_totes"] / inv_den[t["truck"]] * 1201 for t in trucks)
alloc_tab_c40 = sum(t["batch_totes"] / (18 if t["truck"] == "COLA-40" else inv_den[t["truck"]]) * 1201 for t in trucks)

# ---------------- EF sheet ----------------
efw = openpyxl.load_workbook(EFX, data_only=True)["2026"]
def efrow(r): return dict(row=r, label=efw.cell(r, 2).value, unit=efw.cell(r, 3).value, value=efw.cell(r, 6).value, source_cited=efw.cell(r, 7).value)
EF = dict(tanker_process=efrow(6), flatbed_process=efrow(7), tanker_embodied=efrow(11), non_tanker_embodied=efrow(12))

# ---------------- PP summaries (live) + V2 blended CI workbook ----------------
pps = {}
for m in ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]:
    ws = openpyxl.load_workbook(HERE / f"source_exports/pp_summaries/{m}2026_PP_Summary.xlsx", data_only=True).worksheets[0]
    rows = {ws.cell(r, 13).value: [ws.cell(r, c).value for c in range(13, 21)] for r in range(2, 5)}
    pps[m] = dict(tab=ws.title, aqueous_mass_kg=rows["Aqueous"][1], aqueous_emissions_mt=rows["Aqueous"][6], aqueous_ci=rows["Aqueous"][7],
                  wodo_mass_kg=rows["WODO"][1], wodo_emissions_mt=rows["WODO"][6], wodo_ci=rows["WODO"][7],
                  days_in_period=next((ws.cell(r, 3).value for r in range(50, 60) if ws.cell(r, 2).value == "Days in Production Period"), None))
v2 = openpyxl.load_workbook(HERE / "source_exports/pp_summaries/V2_Charm_BioOil_Mass_Weighted_CI_Jan-Jun_2026.xlsx", data_only=True)["CI Calculation"]
v2rows = {v2.cell(r, 1).value: [v2.cell(r, c).value for c in range(1, 8)] for r in range(7, 13)}
live_aq_m = sum(p["aqueous_mass_kg"] for p in pps.values()); live_aq_e = sum(p["aqueous_emissions_mt"] for p in pps.values())
live_wo_m = sum(p["wodo_mass_kg"] for p in pps.values()); live_wo_e = sum(p["wodo_emissions_mt"] for p in pps.values())
july = [t for t in totes if t["production_period"] == "July 2026"]
july_mass = sum(t["net_oil_produced_kg"] for t in july)

# ---------------- cross-week + tracker ----------------
xw = json.load(open(HERE / "cross_week_tag_scan.json")); cobb = json.load(open(HERE / "cobb_w4_aqueous_rows.json")); cobbL = json.load(open(HERE / "cobb_w4_aqueous_rows_live.json"))

L = {
 "batch_id": "INJ-09-2026-W4-QOWV",
 "oil_type": "Aqueous",
 "oil_type_label_in_ops_notes": "QOWV",
 "reporting_period": "September 1 - September 30 2026 RP; Ops Notes injection week 2026-09-18 to 2026-09-25 (QOWV!B3/B4); Certify removal window 19 Sep 2026 to 25 Sep 2026",
 "certify_removal_url": CERTIFY_URL, "certify_removal_id": "rmv_1M3SG7D701S0HZ1Z",
 "drive_folder_url": "https://drive.google.com/drive/u/0/folders/1Ed6zGEgZO69vZPbvmDo0l4Sw4zzZiZl0",
 "T0": "2026-10-06T14:34:05Z",
 "routing": {"_primary_source_opened": [{"file": "Certify GHG entry rmv_1M3SG7D701S0HZ1Z, Components tab", "action": "modal_opened", "live_source_url": CERTIFY_URL, "artifact_path": "certify_initial_snapshot.yml"}],
             "evidence": "Removal 4.08 tCO2e, 19 Sep 2026 to 25 Sep 2026; components 'CO to LA Flatbed Truck Transport Embodied Emissions - QOWV', 'CO to LA Flatbed Truck Transport Process Emissions - QOWV', 'Pyrolysis - Charm QOWV PP Emissions', 'Injection sequestration', 'BCU - Transport Process' -> Aqueous (QOWV) batch, batch-qa-charm skill"},
 "preflight": {
   "knowledge_base": "knowledge-base/Charm_Isometric_Knowledge_Base.md read in full (lines 1-1299) 2026-10-06 ~14:34Z: Sections 0-16 incl. BCU QA Rule (never exceed), uncertainty Layer 1-3, Document Location Rule",
   "lessons_learned": "skills/qa/batch-qa-charm/references/lessons_learned.md read in full: FC-C1 to FC-C12 plus cross-cutting FCs",
   "checklist_structure": "skills/qa/batch-qa-charm/references/checklist_structure.md read in full: Sections 1-9, sidebar, citation discipline, custody rule (2026-09-29) under 2.3",
   "failure_categories_printed": ["FC-C1 Aqueous has no Standard Loss", "FC-C2 a PP omitted entirely is the fail, aggregation passes", "FC-C3 allocation cells must be formulas", "FC-C4 LCS reported needs a component", "FC-C5 roster per-attribute match", "FC-C6 longest Google Maps route", "FC-C7 Mass of Load = gross", "FC-C8 1.134 kg per tote", "FC-C9 PP/PB slash and n=3 are settled", "FC-C10 fraction CI to the right totes", "FC-C11 Max-stated vs invented rules", "FC-C12 dominant input needs uncertainty", "cross-cutting: no cop-outs, no Per Max, missing evidence = Fail, Sources tab is inventory"],
   "registry_checked": "skills/qa/batch-qa-charm/qa_active_charm_batches.json read in full (700 lines): no INJ-09-2026-W4 entry -> fresh QA"},
 "vvb_guidance": {
   "_primary_source_opened": [{"file": "VVB Guidance September 1 - September 30", "action": "doc_read_in_full", "drive_id": "1pib26xToLPFxaOla2p3pv4ummq5ekPpAn65FFuqrlPc",
                               "live_source_url": "https://docs.google.com/document/d/1pib26xToLPFxaOla2p3pv4ummq5ekPpAn65FFuqrlPc/edit", "artifact_path": "source_exports/VVB_Guidance_Sep_2026.txt", "modified": "2026-10-05T07:16:43Z"}],
   "read_in_full": True,
   "sections": ["AECN Rail - Oil Withdrawn for Research and Development (Week 3, INJ-09-2026-W3)", "AECN Rail - New Rail Destination at Baton Rouge (Railcar SHPX215053, Week 4, INJ-09-2026-W4)", "AECN Rail - Bio-oil Loss", "Site Gasoline - E-85 Fills Quantified at an E-85 Emission Factor", "Site Diesel - Rail-Yard Air Compressor Fuel", "Immateriality of Argon Transport and removal from calculations"],
   "verbatim_week4_window": "SHPX215053, offloaded in Week 4 (INJ-09-2026-W4, 18 – 25 September 2026)",
   "verbatim_loss_principle": "Delivered oil and lost oil together therefore carry the pyrolysis and upstream transport emissions of the full loaded mass.",
   "bearing_on_this_batch": "No section covers QOWV, totes, tote tare, the QOWV allocation tab or Charm production CI. The Week 4 section is AECN rail routing only. The loss section states the custody principle applied in 2.3/7.4.",
   "ops_notes_cell_comments": "Formula-mode load of the live export: 2 cell notes in the workbook, AECN Truck!H28 and Kerry!D18, both 'Wellhead pH = weekly median of Ignition PH-1041, gated to StateRun=1; provided by G. Lutz 2026-09-30.' None on the QOWV, QOWV Shipments & EE Allocation or Summary tabs."},
 "batch_folder_inventory": [
   {"name": "Testing", "type": "folder", "drive_id": "1sbuy0NI_yqQmXd4e326PCZ5VmaSUHnsM", "present": True, "contents": "SR26-05719 W4 Kerry AECN QOWV COA.pdf (48,654 bytes)"},
   {"name": "INJ-09-2026-W4 Ops Notes", "type": "google_sheet", "drive_id": "1zgrkVKQsGlzwjiumtn1vkYSaCt5_XTr7UWs3lR3RORQ", "present": True},
   {"name": "QOWV Shipments & EE Allocation (tab of the Ops Notes)", "type": "sheet_tab", "present": True},
   {"name": "Scale Tickets", "type": "folder", "present": True, "contents": "4 AECN railcar subfolders, not QOWV"},
   {"name": "Kerry Oil", "type": "folder", "present": True, "contents": "Kerry scale tickets, not QOWV"},
   {"name": "RailLogix (Carville LA) to Basco B6 Truck Route - 136 mi.png", "type": "image", "present": True, "contents": "AECN rail leg map, not QOWV"}],
 "batch_folder_listing": {"_primary_source_opened": [{"file": "Drive folder 1Ed6zGEgZO69vZPbvmDo0l4Sw4zzZiZl0 and subfolders (Drive API parentId listing, all pages)", "action": "folder_listed", "artifact_path": "batch_folder_listing.txt", "negative_check_tokens": ["Charm Contents.xlsx", "Shipment Tracker.xlsx"], "items_found": []}],
                          "note": "No standalone Bio-Oil Shipment Tracker file and no Charm Contents sheet in the folder; the allocation tracker is the 'QOWV Shipments & EE Allocation' tab of the Ops Notes workbook."},
 "n_totes": n,
 "production_periods": sorted(pp_counts),
 "production_period_tote_counts": dict(pp_counts),
 "production_period_net_produced_kg": {k: round(v, 4) for k, v in pp_mass.items()},
 "production_period_emissions_kg": {k: round(v, 6) for k, v in pp_pe.items()},
 "ops_notes": {
   "_primary_source_opened": [{"file": "INJ-09-2026-W4 Ops Notes", "action": "sheet_grid_read", "live_source_url": OPS_URL, "artifact_path": "source_exports/INJ-09-2026-W4_Ops_Notes.xlsx", "modified": "2026-10-05T15:09:43Z"}],
   "tabs": wv.sheetnames,
   "batch_id_cell": f"QOWV!B2 = {q['B2'].value}",
   "week": f"QOWV!B3 = {q['B3'].value:%Y-%m-%d} ({q['B3'].value:%A}), B4 = {q['B4'].value:%Y-%m-%d} ({q['B4'].value:%A}); labels 'Week Start (Sat)' / 'Week End (Sat)'",
   "section_header": q["A8"].value,
   "ledger_rows": f"QOWV!A11:A52 - {n} tote rows, every row populated for tag, gross, tote mass, net produced, empty tote weight, net injected, production period, CI, production emissions, scale uncertainty",
   "constants": {"L3_lbs_per_kg": q["L3"].value, "L4_kg_per_lb": q["L4"].value, "L5_sparging_mass_loss": q["L5"].value, "L6_ph_target": q["L6"].value,
                 "L7_qowv_6mo_ci_jul_dec_2026": q["L7"].value, "L14_Feb2026": q["L14"].value, "L18_Jun2026": q["L18"].value, "L19_July2026_and_later": q["L19"].value,
                 "L11_Nov2025": q["L11"].value, "L12_Dec2025": q["L12"].value, "L13_Jan2026": q["L13"].value, "L15_Mar2026": q["L15"].value, "L16_Apr2026": q["L16"].value, "L17_May2026": q["L17"].value},
   "constants_links": {f"M{r}": qf.cell(r, 13).value for r in range(10, 22) if qf.cell(r, 13).value},
   "L7_link": "K7 'QOWV 6-Mo CI (kg CO2e/kg), Jul-Dec 2026', L7 typed 0.120947, M7 empty: no hyperlink to a source; row 19 K19 'July 2026 and later (Jan-Jun 2026 rolling)', L19 '=$L$7', M19 empty",
   "per_tote_formulas": {"D": qf["D11"].value, "F": qf["F11"].value, "H_june": qf["H11"].value, "H_july": qf["H34"].value, "H_feb": qf["H51"].value, "I": qf["I11"].value, "J": qf["J11"].value},
   "stray_formulas": "Row 53 and rows 55-74 carry template formulas G=IF(AND(D<>\"\",F<>\"\"),(D-F)*$L$4,\"\") and H=IF(G=\"\",\"\",G*(1-$L$5)) that evaluate blank; no total references them.",
   "whole_oil_tabs": "'Whole Oil' B2 and 'Whole Oil Shipments & EE Alloca' A1/A6 still read INJ-09-2026-W3 with zero totes (stale template from the W3 workbook)."},
 "ops_notes_totals": {
   "total_mass_transported_kg": q["B54"].value, "total_mass_transported_formula": f"QOWV!B54 = {qf['B54'].value}",
   "total_oil_produced_kg": q["D54"].value, "total_oil_produced_formula": f"QOWV!D54 = {qf['D54'].value}; per-tote D = B - C",
   "total_oil_in_batch_kg": q["F54"].value, "total_oil_in_batch_formula": f"QOWV!F54 = {qf['F54'].value}; per-tote F = {qf['F11'].value}",
   "standard_loss_applied": False,
   "weighted_avg_ci_kgco2e_per_kg": q["H54"].value, "weighted_avg_ci_formula": f"QOWV!H54 = {qf['H54'].value}",
   "production_emissions_kg": q["I54"].value, "production_emissions_formula": f"QOWV!I54 = {qf['I54'].value}; per-tote I = H*D",
   "uncertainty_kg": q["J54"].value, "uncertainty_formula": f"QOWV!J54 = {qf['J54'].value}; per-tote J = 2.5*$L$4",
   "residue_left_in_totes_kg": round(q["D54"].value - q["F54"].value, 4)},
 "summary_tab": {"B12_qowv_cargo": sm["B12"].value, "E12_sparging": sm["E12"].value, "F12_filtration": sm["F12"].value, "G12_mobi": sm["G12"].value,
                 "C20_composite": sm["C20"].value, "D20_date": f"{sm['D20'].value:%Y-%m-%d}", "E20_qowv_c": sm["E20"].value, "G20_c_unc": sm["G20"].value, "H20_testing_unc": sm["H20"].value,
                 "G20_formula": smf["G20"].value, "F20_link": smf["F20"].value},
 "lcs_use_kg": None,
 "lcs_search_method": {"_primary_source_opened": [{"file": "INJ-09-2026-W4 Ops Notes (all 10 tabs)", "action": "sheet_grid_read", "artifact_path": "lcs_scan.txt", "negative_check_tokens": ["HIT:"], "items_found": []}],
                       "note": "Case-insensitive scan of every string cell for LCS, calcium, caustic, NaOH, buffer: zero hits across 10 tabs. Summary!E12 sparging 0.0 and F12 filtration 0.0 for QOWV."},
 "totes": totes,
 "tote_weighing_records": {
   "_primary_source_opened": [{"file": "Bio-oil Injection Tracker, tab 'Historical Tote Record' (live export 2026-10-06), production weighing record per tag: reported_quantity, reported_on, order_number, gross_weight, sectag, lot_number", "action": "sheet_grid_read",
                               "live_source_url": "https://docs.google.com/spreadsheets/d/116ZyeotERBTpPrHnWmxLfEjXpqPguTpHNAQKohy5j1Q/edit", "artifact_path": "tote_weighing_records.json"}],
   "result": f"41 of 42 tags carry a production weighing record under their own tag; gross matches the ledger on 40 of the 41. 301274's record (row 1181, OIL260721143336, 2026-07-21) reads 1080.0 kg against 1036.0 in the ledger (checklist 2.3). Tote 300022 has no record under its own tag: its lot 4746 is weighed under tag 300020 (row 840, WO00001655, {lot4746['reported_on']}, gross 1012.0, matching the ledger) (checklist 3.1); that record is its weighing record and date. The 41 own-tag records run {min(w['htr'][0]['reported_on'] for w in wr.values() if w['htr'])[:10]} to {wdates[-1][:10]}; with lot 4746's record, all 42 run {wdates[0][:10]} to {wdates[-1][:10]}. No record read gives a weighing location (see location_basis).",
   "with_record": sum(1 for w in wr.values() if w["htr"]), "gross_match": sum(1 for w in wr.values() if w["htr"] and all(abs(h["gross"] - w["ledger_gross"]) < 1e-9 for h in w["htr"])),
   "weighing_date_range": [wdates[0], wdates[-1]], "with_weighing_date": len(wdates),
   "location_basis": {
     "recorded": "None. The Historical Tote Record's columns (row 2) are reported_quantity, reported_on, production_year, order_number, code, gross_weight, sectag, lot_number: no location. The Q4 production tab's column A, headed 'M', reads 'BGN' on all 42 rows; the tab title reads 'BGN Tote Inventory as of October 2025 to August 2026, DC Tote Inventory August to September 2026'.",
     "what_BGN_names": ["Slack #pyro-ops-handover-notes, Leo, 2026-01-08: 'Swapped out quench oil tote on BGN 1 & helped swap quench oil tote on BGN 5' (BGN names pyrolyzer units)",
                        "Slack #injection-ops-la, rachel, 2026-08-07: 'remaining BGN oil coming from CO'"],
     "production_site": "knowledge-base/pdd_full.txt line 450: 'Charm produces bio-oil at our pyrolysis facility in Fort Lupton, CO'",
     "conclusion": "Inferred, not recorded: the totes were weighed at Charm's Fort Lupton, CO pyrolysis facility, where the BGN units produce. No weighing record states it. The production system's own location fields (inventory master 'Oil_inventory_data' site_location / location, see production_system_location) record where a tote sat when the data was pulled, not where it was weighed."},
   "july_totes_weighed_after_july": {"totes": late_july,
     "note": "The Q4 production tab labels these four totes 'July 2026 PP' ('7/1-7/31'; Form/tote Date blank), but their production weighing records are dated 2026-08-31. If they are August production, nothing in this batch changes: July and August 2026 product both take the Jan-Jun 2026 blended CI, and production emissions sit in one aggregated component (7.1). Checklist 3.5."},
   "june_totes_weighed_in_july": {"totes": late_june, "net_oil_produced_kg": late_june_kg, "june_ci_L18": q["L18"].value,
     "note": "The Q4 production tab and the Ops Notes label these four totes 'June 2026 PP' ('6/1-6/30') and they carry the June CI; their production weighing records, order numbers (OIL260701...) and the production system's production_date are dated 2026-07-01. Those dates are the weigh-in, not the fill (production_date sits seconds to minutes before reported_on). The records place the totes in June (see corroboration): the June summary's Aqueous mass and the Jan-Jun blend's June mass both count them, and they carry June's spray paint. No change to the removal. Checklist 3.5 (documentation flag: no 2026 tote has a fill date)."}},
 "inventory_master": {
   "_primary_source_opened": [{"file": "BGN & DC Production Bio-Oil Inventory, tab 'Q4 2025 -2026 Production Invent'", "action": "sheet_grid_read", "drive_id": "1cD9cd_y_SI8GA1Yv8TFtDPqevKKF1Jh42WFRWSSMJ6g", "live_source_url": INV_URL, "artifact_path": "source_exports/inventory_master_export.xlsx"}],
   "attribute_match": match, "batch_totes_by_ship_status": dict(by_truck),
   "truck_totals_all_fractions": {k: dict(v) for k, v in truck_all.items()},
   "truck_totals_count": {k: sum(v.values()) for k, v in truck_all.items()},
   "tare_distribution_batch": {str(k): v for k, v in tare_dist.items()},
   "pp_date_ranges": {t["inventory_pp"]: hits[t["tag"]]["inventory"][0]["date_range"] for t in totes},
   "tote_300022":"Q4 row 410: February 2026 PP, tag 300022, MFO Lot 4746, QOWV, gross 1012.0, tare 57.6, net 954.4, COLA_37, shipped 2026-09-09, spray paint Blue2602",
   "tote_300020": "Q4 row 420: February 2026 PP, tag 300020, MFO Lot 4770, QOWV, gross 1085.0, tare 57.6, net 1027.4, COLA_25, shipped 2026-07-16",
   "oil_inventory_data_tab": "Tab 'Oil_inventory_data' A1 reads 'Microsoft SQL Server Import / Last updated 7 weeks ago'; predates the September shipments, so truck membership is read from the Production Inventory tab's Ship Status column (the basis the October Whole Oil allocation tabs cite as 'Tier-1 resolution').",
   "cola40_correction_note": sh["A19"].value},
 "shipment_tracker": {
   "_primary_source_opened": [{"file": "'QOWV Shipments & EE Allocation' tab, INJ-09-2026-W4 Ops Notes", "action": "sheet_grid_read", "live_source_url": OPS_URL, "artifact_path": "source_exports/INJ-09-2026-W4_Ops_Notes.xlsx"}],
   "one_way_distance_mi": sh["B8"].value, "one_way_distance_cell": f"B8 = {sh['B8'].value}, '{sh['A8'].value}', '{sh['C8'].value}'",
   "tote_count_cell": f"D6 = {shf['D6'].value} = {sh['D6'].value}", "E6_text": sh["E6"].value,
   "trucks": trucks,
   "totals": {"totes_on_trucks": sh["C17"].value, "batch_totes": sh["D17"].value, "allocated_miles_mi": sh["G17"].value, "formulas": f"C17 {shf['C17'].value}, D17 {shf['D17'].value}, G17 {shf['G17'].value}"},
   "tote_level_rows": "None: rows 1-19 hold the summary, one-way distance, the four truck rows, totals and the A19 correction note; rows 20-74 are empty. No row ties a tag to a truck, and there is no 'every batch tote is allocated to a truck' check row (the Whole Oil allocation tab carries one at A16).",
   "inventory_denominators": inv_den, "allocated_miles_on_inventory_denominators": round(alloc_inv, 6), "allocated_miles_tab_c40_18": round(alloc_tab_c40, 6)},
 "shipment_tracker_cells_are_formulas": {
   **{f"share_E{t['row']}": t["share_formula"] for t in trucks}, **{f"allocated_miles_G{t['row']}": t["allocated_miles_formula"] for t in trucks},
   **{f"mileage_F{t['row']}": t["mileage_formula"] for t in trucks}, "batch_mileage_allocation_G17": shf["G17"].value, "tote_count_D6": shf["D6"].value,
   "all_are_formulas": all(str(x).startswith("=") for x in [*(t["share_formula"] for t in trucks), *(t["allocated_miles_formula"] for t in trucks), shf["G17"].value]),
   "typed_inputs": "C12:C15 (totes on truck 20, 6, 20, 18) and D12:D15 (batch totes 15, 5, 20, 2) are typed counts; B8 1201 typed distance",
   "method": "Formula-mode load of the workbook exported from the live Google Sheet this session (openpyxl without data_only)."},
 "other_batches_on_shared_trucks": {
   "_primary_source_opened": [{"file": "INJ-10-2026-W2 Ops Notes, 'Whole Oil Shipments & EE Alloca' rows 6 and 12", "action": "sheet_grid_read", "live_source_url": "https://docs.google.com/spreadsheets/d/1PAyyKyKrEWMucw6y8Fs5KV9TW4q8AvshhWnC1hCLe68/edit", "artifact_path": "source_exports/prior_weeks/INJ-10-2026-W2_Ops_Notes.xlsx"},
                              {"file": "INJ-10-2026-W3 Ops Notes, 'Whole Oil Shipments & EE Alloca' rows 6 and 12", "action": "sheet_grid_read", "live_source_url": "https://docs.google.com/spreadsheets/d/18OucJBh4kSXfKWt8h3ir6098oLjTfC2NDnHRBrOWDVY/edit", "artifact_path": "source_exports/prior_weeks/INJ-10-2026-W3_Ops_Notes.xlsx"}],
   "inj10_w2_whole_oil": "COLA_38 | 2026-09-10 | 18.0 | 3 | 0.1666666667 | 1201 | 200.1666667; E6 'COLA_38 = 3 of 18 totes on the truck (12 DC Whole Oil + 6 BGN QOWV; the other 15 are not in this batch). Tier-1 resolution: Ship Status on Q4 2025-2026 Production Inventory Master.'",
   "inj10_w3_whole_oil": "COLA_38 | 2026-09-10 | 18.0 | 5 | 0.2777777778 | 1201 | 333.6111111; E6 'COLA_38 = 5 of 18 totes on the truck (12 DC Whole Oil + 6 BGN QOWV; the other 13 are not in this batch)'",
   "cola38_shares_claimed": round(5 / 6 + 3 / 18 + 5 / 18, 10)},
 "cross_week_double_count": {
   "_primary_source_opened": [{"file": "Weekly Ops Notes exported 2026-10-06: INJ-08-2026-W5, INJ-09-2026-W1, W2, W3, INJ-10-2026-W1, W2, W3 (every tab, every cell)", "action": "sheet_grid_read", "artifact_path": "cross_week_tag_scan.json"}],
   "result": "None of the 42 tags appears in any other week's workbook (0 hits in each of 7 workbooks). Control: tag 301009 is found 4 times in the W3 workbook (QOWV!A11, G11, allocation A21/B21).",
   "hits_per_workbook": {k: len(v) for k, v in xw.items()}},
 "cobb_injection_record": {
   "_primary_source_opened": [{"file": "Bio-oil Injection Tracker, tab 'Basco Injection -- COBB', live export 2026-10-06 19:26 BST (after the 2026-10-06 Google sign-in), rows 861-3723", "action": "sheet_grid_read", "drive_id": "116ZyeotERBTpPrHnWmxLfEjXpqPguTpHNAQKohy5j1Q",
                               "live_source_url": "https://docs.google.com/spreadsheets/d/116ZyeotERBTpPrHnWmxLfEjXpqPguTpHNAQKohy5j1Q/edit", "artifact_path": "cobb_w4_aqueous_rows_live.json"},
                              {"file": "Same tab, 2026-09-28 export (history)", "action": "sheet_grid_read", "artifact_path": "cobb_w4_aqueous_rows.json"}],
   "live_read_attempts_2026_10_06": "Before the sign-in: (1) export?format=xlsx -> Google 'Verify it's you' page; (2) export retried -> same; (3) gviz/tq CSV of the COBB tab -> same; (4) htmlview -> same; (5) Drive full-text search -> no result; (6) Drive metadata -> modifiedTime 2026-10-05T19:24:28Z and a 'Naming Conventions' snippet; (7) Drive connector download as text/csv -> 'Internal error encountered.'. After the MRV lead's sign-in (login window showed the tracker open), the export was taken at 19:26 BST through the headless browser profile playwright-profile-mrv-3, the main profile being held by another Claude session.",
   "header_A1": cobbL["A1"], "header": cobbL["header"], "w4_counts": cobbL["w4_counts"],
   "result": "168 rows 'INJ-09-2026-W4_aqueous_fraction' in four stacked imports (Import Time column W: 2026-09-29 08:37:20 from row 861, 2026-09-29 11:54:56 from row 1791, 2026-09-30 11:14:57 from row 2719, 2026-10-05 12:24:22 from row 3682): each of the 42 totes appears once per import, and every September batch (W1-W4) is repeated the same way. injection_completion_date 2026-09-25 on all 168. Security tags match the ledger on 41 of 42; every copy of the 42nd reads security_tag 300020 with consumed_lot 4746, quantity 666.4 kg and tote_gross_weight 1012.0 kg, which is the ledger's tote 300022 (lot 4746, gross 1012.0, net injected 724 less 57.6 kg). tote_gross_weight matches the ledger on every other row except 301274.",
   "quantity_pattern": "quantity_consumed_in_injection = Ops Notes net injected - 57.6 kg on 155 of 168 rows; 301204 reads 1,540.8 (= 2 x 770.4) in the first block and 770.4 in the other three; 301215 821.4 kg and 301274 852.9 kg in all four blocks; the four 300020 rows carry 666.4 kg.",
   "ph_by_week": {"block_861_1790": {"W1": 2.89, "W2": 2.92, "W3": 2.7, "W4": 2.97}, "block_1791_2718": {"W1": 2.89, "W2": 2.92, "W3": 2.7, "W4": 2.97},
                  "block_2719_3681": {"W1": None, "W2": None, "W3": None, "W4": "2.97 on 64 rows, blank on 1"}, "block_3682_3805": {"W4": None},
                  "note": "One pH per week, stamped on every product row of that week (aqueous, bio-oil, whole oil, calciner); the W4 value 2.97 equals the Ops Notes AECN Truck!H28 and Kerry!D18 weekly wellhead pH (cell note 'Wellhead pH = weekly median of Ignition PH-1041, gated to StateRun=1'). The newest block carries no pH."},
   "tag_300020": {"production_system_Oil_inventory_data": ["row 360: lot 4746, security_tag 300020, gross 1012.0, net 954.4, WO00001655, 2026-02-17", "row 350: lot LOT00004770, security_tag 300020, gross 1085.0, net 144.4, WO00001665, 2026-02-24, shipping order COLA25"],
                  "historical_tote_record": ["row 840: 300020 / lot 4746, gross 1012.0", "row 850: 300020 / lot LOT00004770, gross 1085.0"],
                  "claude_c_production_tab": ["row 410: 300020 / 4746 / 1012.0", "row 420: 300020 / 4770 / 1085.0"],
                  "q4_production_inventory_tab": "row 410: tag 300022 / lot 4746 / 1012.0; row 420: tag 300020 / lot 4770 / 1085.0",
                  "injections": "Tracker: tag 300020 / lot LOT00004770, 883.0 kg, batch INJ-20260726_aqueous_fraction, completed 2026-07-31 (rows 455, 1380, 2308, 3271); tag 300020 / lot 4746, 666.4 kg, INJ-09-2026-W4 (rows 901, 1831, 2759, 3722). Lots and weights agree across records, so no mass is counted twice; the metal tag number is shared by two lots in the production system."},
   "post_unload_conflicts": {"301215": {"ops_notes_E26_kg": 442.5, "tracker_gross_kg": 1041.0, "tracker_implied_kg": round(1041 - 57.6 - 821.4, 1), "ops_net_injected_kg": 598.5, "tracker_basis_net_injected_kg": round(1041 - (1041 - 57.6 - 821.4), 1)},
                             "301166_same_442_5": "Ops Notes E17 (301166) also reads 442.5 kg; the tracker agrees for 301166 (quantity 593.9 = 651.5 - 57.6)"},
   "gross_conflicts": {"301274": {"ops_notes_B40_kg": 1036.0, "q4_inventory_tab_row694_kg": 1036.0, "tracker_gross_kg": 1080.0,
                                  "production_system_Oil_inventory_data_row39": "aqueous_fraction, lot 301274, tag 301274, produced 2026-07-21 20:34, work order OIL260721143336, net 1022.4, gross 1080.0",
                                  "historical_tote_record_row1181": "reported 1022.4, 2026-07-21 20:35, OIL260721143336, gross 1080.0",
                                  "tracker_implied_post_unload_kg": round(1080 - 57.6 - 852.9, 1), "ops_notes_E40_kg": 169.5,
                                  "if_1080": {"mass_gain_kg": 44.0, "seq_gain_kg": round(44.0 * 0.079 * 3.667, 6), "pe_gain_kg": round(44.0 * 0.120947, 6), "net_gain_kg": round(44.0 * 0.079 * 3.667 - 44.0 * 0.120947, 6)}}},
   "changes_since_0928": "Compared by column name over the columns both exports share: injection_ph added on 126 of the W4 aqueous rows (2.97); 301204's quantity corrected to 770.4 in blocks 2-4. The live tab also carries tote_gross_weight, weighing_order and weighing_date, which the 28 Sep export did not have (it had Proportion of Origin Consumed, Origin Mass, Consumed Origin Mass in those positions), so 301274's tracker gross of 1080.0 is visible for the first time. The W4 rows went from 42 to 168 through the repeated blocks.",
   "history_0928": {"result": "42 rows (856-897), completion 2026-09-25, injection_ph blank, tote_gross_weight blank; same 300020/lot 4746 row; 301204 1,540.8", "artifact_path": "cobb_w4_aqueous_rows.json"},
   "ops_notes_unchanged_since_0928": "A 2026-09-28 export of this workbook (outputs/qa/batch-INJ-09-2026-W3-QOWV/source_exports/prior_weeks/INJ-09-2026-W4_Ops_Notes_1zgrkV.xlsx) has 0 cell differences on the QOWV and allocation tabs against today's export."},
 "testing": {
   "_primary_source_opened": [{"file": "Certify source src_1M3VZ8S8W1S013HS 'SR26-05719 W4 Kerry AECN QOWV COA.pdf'", "action": "pdf_read", "artifact_path": "Certify_attached_COA_src_1M3VZ8S8W1S013HS.pdf",
                               "sha256": "276ccd65301889cf4e037a09ee2f397d6e2224f336dbfc0895947baf766943f6", "bytes": 48654, "drive_testing_copy": "Testing/SR26-05719 W4 Kerry AECN QOWV COA.pdf, Drive id 1Viar0c3xdOh-pwkpbsNY0buJPDH5AGEJ"},
                              {"file": "Drive Testing folder copy, 'SR26-05719 W4 Kerry AECN QOWV COA.pdf' (Drive id 1Viar0c3xdOh-pwkpbsNY0buJPDH5AGEJ), downloaded 2026-10-06 after the sign-in", "action": "pdf_read", "artifact_path": "Drive_Testing_COA_1Viar0c3xdOh.pdf",
                               "sha256": "276ccd65301889cf4e037a09ee2f397d6e2224f336dbfc0895947baf766943f6", "evidence": "page 3 read: SR26-05719.003, Client ID 9-2026-W4-QOWV, Carbon 7.9 % (m/m); sha256 identical to the Certify-attached copy"}],
   "qowv_certificate": {"page": 3, "certificate_no": "SR26-05719.003", "client_id": "9-2026-W4-QOWV", "location": "FORT LUPTON, CO", "product_description": "Oil",
                        "sampled": "25-Sep-2026 10:00", "received": "28-Sep-2026 11:00", "analysed": "28-Sep-2026 14:29 - 28-Sep-2026 23:23", "carbon_pct": 7.9, "carbon_method": "ASTM D5291 (Method C)",
                        "carbon_out_of_range_flag": "# - Result is outside of test method limits and/or analytical range used in method precision study",
                        "hydrogen_pct": 7.9, "nitrogen_pct": "<0.75", "kinematic_viscosity_40C_cSt": 0.902, "acid_number_mgKOH_g": 36, "density_15C_g_cm3": 1.039},
   "formation_fluid_density": {"value_g_cm3": 1.01, "source": "knowledge-base/Charm_Isometric_Knowledge_Base.md line 313, section 'Density Requirements'",
                               "quote": "All streams exceed formation fluid density (1.01 g/cm\u00b3) \u2014 confirmed via density analysis for safe injection.", "lab_density_15C_g_cm3": 1.039},
   "other_pages": "Page 1 SR26-05719.001 Client ID 9-2026-W4-KERRY, C 57.9 %; page 2 SR26-05719.002 Client ID 9-2026-W4-AECN, C 44.0 % - other removals."},
 "standard_efs_2026": {
   "_primary_source_opened": [{"file": "Standard Emission Factors + Calculations, 2026 tab", "action": "sheet_grid_read", "live_source_url": EF_URL, "artifact_path": "source_exports/Standard_EFs_live_export.xlsx"}],
   "tab": "2026", **EF,
   "flatbed_km_t_equivalent_kgco2e": round(EF["flatbed_process"]["value"] * 1000 / 1.609344, 7),
   "tanker_km_t_equivalent_kgco2e": round(EF["tanker_process"]["value"] * 1000 / 1.609344, 7),
   "non_tanker_per_km_kgco2e": round(EF["non_tanker_embodied"]["value"] * 1000 / 1.609344, 7),
   "tanker_embodied_per_km_kgco2e": round(EF["tanker_embodied"]["value"] * 1000 / 1.609344, 7)},
 "certify_attached_ef_sheet": {"_primary_source_opened": [{"file": "Certify sources src_1KWZ2NVQM1S08AWN and src_1KVJRSYT91S09G2E 'Standard Emission Factors + Calculations (8).xlsx' (byte-identical, sha256 d49871d1...)", "action": "sheet_grid_read", "artifact_path": "source_exports/Certify_attached_StdEFs_src_1KWZ2NVQM1S08AWN.xlsx"}],
                               "rows": "2026 tab: row 5 'Truck Transport Fuel Use (Well to Wheel) -- Tanker Truck' 0.0001143; row 6 '-- Flatbed Truck' 0.0001384; row 10 'Vehicle Embodied Emissions --Tanker Truck' 0.0002375605888; row 11 '--Non-Tanker Truck' 0.0001927908256 (MT CO2e/t-mi and MT CO2e/mi, GLEC V3.2 2025)"},
 "google_maps_options": [{"route": "via US-287 S and US-287 S", "time": "17 hr 22 min", "distance_mi": 1132, "label": "Fastest route now, avoids road closures"},
                         {"route": "via US-287 S, US-287 S and I-20 E", "time": "17 hr 32 min", "distance_mi": 1129, "label": "Avoids road closures"},
                         {"route": "via I-70 E", "time": "17 hr 36 min", "distance_mi": 1201, "label": "Avoids road closures; This route has tolls"}],
 "google_maps_note": {"_primary_source_opened": [{"file": "Certify source src_1KSTWP2AP1S0QQPF 'CO to Basco Routing Map.png'", "action": "image_read", "artifact_path": "screenshots/CO_to_Basco_Routing_Map_src_1KSTWP2AP1S0QQPF.png", "sha256": "9168453b581f7236bdb7503a6a95c9ac5590009ca5de4dbc8e6f409f07dc5aec"}],
                      "note": "Read from the pixels: destination '4675 LA-115, St Landry, LA 71367', origin field 'Your location' (pin just north of Denver). Byte-identical to the map attached on INJ-09-2026-W3-QOWV.",
                      "max_ruling": "skills/qa/batch-qa-charm/qa_active_charm_batches.json, INJ-07-2026-W1 max_rulings_2026_08_06: '5.2 -> Pass: routing-map origin confirmed as Fort Lupton, unnamed-origin flag dropped.'"},
 "retirement_certificate": {"_primary_source_opened": [{"file": "Certify source src_1M1MFJPW11S08KY2 'Charm-RetirementCertificate-RD-EACs-20251218.pdf'", "action": "pdf_read", "artifact_path": "Certify_attached_RetirementCertificate_src_1M1MFJPW11S08KY2.pdf", "sha256": "bd185a8a63236ab0c79165fbd964347603b4dc8880d7230e9453b45de45721c6"}],
                            "read": "TERC Certificate of Retirement: Retiring Organization Amazon.com Services LLC; On Behalf Of Charm Industrial, Inc; Date of Retirement Dec 18, 2025; Vintage Year 2024; Commodity Type Renewable Diesel TERCS; Number of TERCs Retired 2,795; Verification Body Turner, Mason & Company; Certificate Number TERC1012_2024_DF_RENEWABLE_DIESEL-13-47206-50000; 'Each TERC represents one metric ton of carbon dioxide equivalent (tCO2e) reduced.'",
                            "tercs": 2795},
 "pp_summary_sources": {
   "_primary_source_opened": [{"pp": m, "file": f"{m} 26 Production Period Summary (live Google Sheet)", "action": "sheet_grid_read",
                               "live_source_url": f"https://docs.google.com/spreadsheets/d/{fid}/edit", "artifact_path": f"source_exports/pp_summaries/{m}2026_PP_Summary.xlsx"}
                              for m, fid in [("Jan", "1bPZ7pNsR5rv-om-BypF4lK3nczjn4HJNO2dzimwnx7E"), ("Feb", "1vH7bborVnc-TCby3wEkpzi7GGMOI5jMJooEW0V-rOQ4"), ("Mar", "1O7y30opa_u1_deqFm_w2qgVyxi_75X00xnzKuyVyh6o"),
                                             ("Apr", "10Qwx7K6sUcOBrCebGUGVaVeLQnhzvcT8sHQ0Sy3VBBM"), ("May", "1m66t1jsLGc4e6AWnCFMsFbkSceZ9IYC_-bfTsbP_OYc"), ("Jun", "1MUQajDI2a_Cn9xWGPP7z4PhhphZnfl2_CSTnpB07Wgk")]],
   "live": pps,
   "batch_uses": "June 2026 (L18, M18 link 1MUQajDI2a...) on 32 totes; February 2026 (L14, M14 link 1vH7bbor...) on 2 totes; July 2026 totes use L7 (Jan-Jun 2026 blended).",
   "july_2026": {
     "_primary_source_opened": [
       {"pp": "Jul", "file": "Jul 26 Production Period Summary (Google Sheet, created 2026-09-22)", "action": "sheet_grid_read", "drive_id": "124Yhg_sw5bpdkBSq6gt8IlMj13xNZAAJV5rm9rdTIWE",
        "live_source_url": "https://docs.google.com/spreadsheets/d/124Yhg_sw5bpdkBSq6gt8IlMj13xNZAAJV5rm9rdTIWE/edit", "artifact_path": "source_exports/pp_summaries/Jul2026_PP_Summary.xlsx",
        "sha256": "e0424eb38e30028438807af2e535948442bb12eb5898247b998cd124cbeb1bcc"},
       {"pp": "Jul", "file": "Jul 26 Production Period Summary (Google Sheet, created 2026-09-11)", "action": "sheet_grid_read", "drive_id": "1nIz_QcIeNbNiUu4UyiatlZENkAhIgA8wm0r3mE7plgg",
        "live_source_url": "https://docs.google.com/spreadsheets/d/1nIz_QcIeNbNiUu4UyiatlZENkAhIgA8wm0r3mE7plgg/edit", "artifact_path": "source_exports/pp_summaries/Jul2026_PP_Summary_1nIz.xlsx",
        "sha256": "8e2d69767cc67336f606b669c17363a90a7546a5389b697b12725b1550b2c754"},
       {"pp": "Jul", "file": "July 26 Production Period Output Summary (Google Sheet, owner garrett.lutz@, created 2026-09-11)", "action": "sheet_grid_read", "drive_id": "1k5le7x-EYfdqiSEmad1bA1NliBJm7ryGRw0lE6CrN_I",
        "live_source_url": "https://docs.google.com/spreadsheets/d/1k5le7x-EYfdqiSEmad1bA1NliBJm7ryGRw0lE6CrN_I/edit", "artifact_path": "source_exports/pp_summaries/Jul2026_PP_Output_Summary_1k5le.csv",
        "sha256": "65720e9145f3475b79fd6dacd71f9ba83aa5b2ddb5c540d90206e960e0310a02"}],
     "read": "Both 'Jul 26 Production Period Summary' files have the same four tabs as June, the summary tab still named 'Jun 26 PP Summary', and match the June summary (source_exports/pp_summaries/Jun2026_PP_Summary.xlsx) on every non-empty cell: 425 stored values and 426 formula cells (the extra one, C89 '=C84', has no stored value), with no differences (Aqueous N4 46616.2 kg, T4 0.1236896674; WODO N2 10778.0 kg, T2 1.110535564): they are copies of June, not July figures. 'July 26 Production Period Output Summary' has one tab, 'Sheet1' (xlsx export, source_exports/pp_summaries/Jul2026_PP_Output_Summary_1k5le.xlsx, sha256 b8e9578df2f94840ad6d977af538315d64bef69386c1dad905c069dc4a3a8f65), and it is an unfilled template: the WODO and Aqueous mass cells are blank and every calculated cell on those rows reads #DIV/0!; only five biochar carbon readings are entered.",
     "conclusion": "No July 2026 production-period figures (masses, CIs) exist in any summary found. The July totes' CI does not depend on one: July-December 2026 product takes the Jan-Jun 2026 blended CI (7.1-7.3). July date range from the Q4 production tab: '7/1-7/31'."},
   "constants_match_live": {"Feb": q["L14"].value == pps["Feb"]["aqueous_ci"] or abs(q["L14"].value - pps["Feb"]["aqueous_ci"]) < 1e-9,
                            "Jun": abs(q["L18"].value - pps["Jun"]["aqueous_ci"]) < 1e-9,
                            "Jan": abs(q["L13"].value - pps["Jan"]["aqueous_ci"]) < 1e-9, "Mar": abs(q["L15"].value - pps["Mar"]["aqueous_ci"]) < 1e-9,
                            "Apr": abs(q["L16"].value - pps["Apr"]["aqueous_ci"]) < 1e-9, "May": abs(q["L17"].value - pps["May"]["aqueous_ci"]) < 1e-9}},
 "blended_ci_source": {
   "_primary_source_opened": [{"file": "V2_Charm_BioOil_Mass_Weighted_CI_Jan-Jun_2026.xlsx (Drive id 1J1_DxRBIbYTa3o65q_zEv_GSB6Ci6Pt8, owner max@, modified 2026-09-08)", "action": "sheet_grid_read",
                               "live_source_url": "https://drive.google.com/file/d/1J1_DxRBIbYTa3o65q_zEv_GSB6Ci6Pt8/view", "artifact_path": "source_exports/pp_summaries/V2_Charm_BioOil_Mass_Weighted_CI_Jan-Jun_2026.xlsx"},
                              {"file": "July RP reporting package copies of the six monthly summaries (local, saved 2026-08-11/12)", "action": "sheet_grid_read", "artifact_path": "source_exports/pp_summaries/V2_Charm_BioOil_Mass_Weighted_CI_Jan-Jun_2026.xlsx"}],
   "method_text": v2["A2"].value, "read_note": v2["A3"].value,
   "qowv_result": {"mass_kg": v2["B19"].value, "emissions_mt": v2["C19"].value, "ci": v2["D19"].value},
   "wodo_result": {"mass_kg": v2["B18"].value, "emissions_mt": v2["C18"].value, "ci": v2["D18"].value},
   "v2_qowv_emissions_by_month_mt": {k: v[4] for k, v in v2rows.items()},
   "sources_tab": "Sources tab: Jan 'January 2026 Production Period Summary.xlsx' 1D07oKMClJBm0ZEgdsuN_APdGrhzoLw9L; Feb 'Feb 26 Production Period Summary.xlsx' 1wFC1c4DWZZqLXeS422lgThPu23quOJTT; Mar 160LXWi92KCgonrEbY4omp8QmXpQAGXlX; Apr 1pmsqSW2toU5KAxYin82KfaF_BzcKYQ0y; May 1XU3CTs0HKjAOXmo7H6LzJQhPzyYeT1N6; Jun 1c9objRn_ahdCwpHWX8kIYL8PWe1l77Ci - xlsx copies, not the Google Sheets the Ops Notes links",
   "live_recompute": {"qowv_mass_kg": round(live_aq_m, 2), "qowv_emissions_mt": round(live_aq_e, 9), "qowv_ci": round(live_aq_e * 1000 / live_aq_m, 10),
                      "wodo_mass_kg": round(live_wo_m, 2), "wodo_emissions_mt": round(live_wo_e, 9), "wodo_ci": round(live_wo_e * 1000 / live_wo_m, 10)},
   "earlier_use_of_live_monthly_ci": "qa_active_charm_batches.json INJ-07-2026-W1 (T0 2026-08-06): 'Jan 2026 (6 totes, CI 0.931994)' - the live Jan WODO CI (0.9319944878), against 0.8258470104 in the copy V2 read on 2026-08-11",
   "july_totes": len(july), "july_net_produced_kg": round(july_mass, 4)},
 "pp_summary_basis": {
   "_primary_source_opened": [{"file": "Six monthly summaries: the July RP reporting-package copies V2 was built from (Drive mount My Drive/MRV/Removals Reporting/2026/Verified Periods/July 1 - July 31 RP/.../H1 2026 Charm CI/, local copy in Downloads/Final Reporting Package July 1 - July 31 RP) against the live Google Sheets exported 2026-10-06",
                               "action": "sheet_grid_read", "artifact_path": "pp_summary_basis_comparison.json"}],
   "by_month": {m: {"copy_I19_formula": v["copy"]["I19f"], "live_I19_formula": v["live"]["I19f"], "copy_H18": v["copy"]["H18"], "live_H18": v["live"]["H18"],
                    "live_bcus_applied_C10_mt": v["live"]["C10"], "live_aqueous_share": v["live"]["aq_share"], "live_aqueous_mt": v["live"]["aq_mt"],
                    "live_aqueous_after_bcu_mt": v["live_aq_after_bcu_mt"], "copy_aqueous_mt": v["copy"]["aq_mt"], "C85_live": v["live"]["C85"], "C85_copy": v["copy"]["C85"]}
               for m, v in json.load(open(HERE / "pp_summary_basis_comparison.json")).items()},
   "result": "Copies: Biomass Transport I19 = C11*F7 ('Remaining Emissions to Allocate', after BCUs) for Jan-Apr and Jun; May I19 = C9*F7 with C9 = 0 (no biomass transport). Live sheets: I19 = C9*F7 ('Total Emissions', before BCUs) under the heading 'Consolidated Emissions Inventory (No GHG Neutralization)'. The inputs differ too (June: tonnes 109.03 in the package copy vs 129.1156462585034 on the linked sheet, miles 354.5 vs 405.5); the package copies are the later versions (pp_summary_basis.ruling_2026_10_07.version_history).",
   "after_bcu_live_blend": {"qowv": None, "wodo": None},
   "may_C85": "May summary C85 'Transport Emissions (Embodied)': 0.73646 live vs 0.073646 in the copy (factor of 10)."},
 "certify": {
   "_primary_source_opened": [{"file": "Certify GHG entry rmv_1M3SG7D701S0HZ1Z (Components, each component detail view, Calculation view, Datapoints, Sources, source detail pages)", "action": "modal_opened", "live_source_url": CERTIFY_URL, "artifact_path": "certify_initial_snapshot.yml"}],
   "removal_window": "19 Sep 2026 to 25 Sep 2026",
   "net_co2e_removed_t": 4.076, "net_header_display_t": 4.08,
   "calculation_view": "Net CO2e Removed 4.076 t = Sequestrations 9.902 t + Adjustments 5.895 tCO2e - Activities 11.674 tCO2e - Uncertainty Discount 47 kg",
   "components_tab_totals": "Mitigated 15.797 tCO2e (Bio-oil transport 5.895 = BCU - Transport Process; CO2 stored 9.902 = Injection sequestration); Emitted 11.674 tCO2e (Bio-oil transport 6.664 = embodied 0.769 + process 5.895; Pyrolysis or other process 5.010); 5 components",
   "uncertainty_discount_kg": 47, "components_tab_total": 5,
   "component_by_key": {
     "injection_sequestration": {"name": "Injection sequestration", "id": "cmp_1M3SG7D701S0E4E9", "value_t": 9.902,
                                 "inputs": {"carbon_content_label": "Carbon Content of Bio-oil Injected (per injection)", "carbon_content_pct": 7.9, "carbon_content_uncertainty_pct_abs": 0.03523, "carbon_content_sources": 2,
                                            "mass_of_product_label": "Mass of Injectant (per injection)", "mass_of_product_kg": 34179.5, "mass_of_product_uncertainty_kg": 47.62716, "mass_sources": 1},
                                 "blueprint": "Carbon rich substance sequestration: output = product_mass * carbon_content * {co2e_of_carbon}"},
     "transport_process": {"name": "CO to LA Flatbed Truck Transport Process Emissions - QOWV", "id": "cmp_1M3SG7D701S0MXD4", "value_t": 5.895, "share_of_net_pct_display": 144.49,
                           "inputs": {"ef_label": "Emissions Factor - Truck Transport Fuel Use (Well to Wheel) - Tanker Truck", "ef_kgco2e_per_t_mi": 0.1143, "ef_datapoints_tab_display": "0.07 kgCO2e/(km*t)", "ef_source": "Standard Emission Factors + Calculations (8).xlsx",
                                      "distance_label": "Distance from Miniforge CO to Basco", "distance_mi": 1201, "distance_source": "CO to Basco Routing Map.png",
                                      "mass_of_load_label": "QOWV Transport Mass (CO to B6) [Certify input]", "mass_of_load_kg": 42945, "mass_of_load_uncertainty_kg_displayed": 47.627, "mass_source": "INJ-09-2026-W4 Ops Notes.xlsx"},
                           "blueprint": "Transport emissions: output = mass * distance * carbon_intensity"},
     "transport_embodied": {"name": "CO to LA Flatbed Truck Transport Embodied Emissions - QOWV", "id": "cmp_1M3SG7D701S0D2XB", "value_t": 0.769, "share_of_net_pct_display": 18.85,
                            "inputs": {"ef_label": "Emissions Factor - Vehicle Embodied Emissions - Tanker Truck", "ef_kgco2e_per_mi": 0.23756, "ef_datapoints_tab_display": "0.1476 kgCO2e/km", "ef_source": "Standard Emission Factors + Calculations (8).xlsx",
                                       "distance_label": "QOWV Embodied Distance (CO to B6) [Certify input]", "distance_mi": 3236.03, "distance_sources": 2},
                            "blueprint": "Distance-based emissions: output = distance * carbon_intensity"},
     "pyrolysis": {"name": "Pyrolysis - Charm QOWV PP Emissions", "id": "cmp_1M3SG7D701S09318", "value_t": 5.01, "share_of_net_pct_display": 122.79,
                   "inputs": {"ci_label": "QOWV Pyrolysis Emissions - Carbon Intensity (Weighted)", "ci_tco2e_per_t": 0.123623, "ci_sources": 1,
                              "mass_label": "QOWV Pyrolysis Mass [Certify input]", "mass_kg": 40525.8, "mass_uncertainty_kg_displayed": 47.627, "mass_sources": 1},
                   "blueprint": "Mass-based CI emissions: output = mass * carbon_intensity"},
     "bcu": {"name": "BCU - Transport Process", "id": "cmp_1M3SGFS3K1S0BBQP", "value_t": 5.895, "inputs": {"constant_co2_reduction_kg": 5895, "source": "Charm-RetirementCertificate-RD-EACs-20251218.pdf"},
             "blueprint": "Constant CO2 reduction"}},
   "counterfactuals": "none present (expected zero for bio-oil geological storage)",
   "datapoints_tab_total": 10,
   "datapoints": ["QOWV Embodied Distance (CO to B6) [Certify input] 3,236 mi REPORTED", "Emissions Factor - Vehicle Embodied Emissions - Tanker Truck 0.1476 kgCO2e/km DERIVED", "QOWV Transport Mass (CO to B6) [Certify input] 42.945 t REPORTED",
                  "Distance from Miniforge CO to Basco 1,201 mi DERIVED", "Emissions Factor - Truck Transport Fuel Use (Well to Wheel) - Tanker Truck 0.07 kgCO2e/(km*t) DERIVED", "QOWV Pyrolysis Mass [Certify input] 40.526 t REPORTED",
                  "QOWV Pyrolysis Emissions - Carbon Intensity (Weighted) 0.12 tCO2e/t REPORTED", "Mass of Injectant (per injection) 34.179 t DERIVED", "Carbon Content of Bio-oil Injected (per injection) 7.90 % DERIVED", "Constant CO2 reduction 5.895 tCO2e REPORTED"],
   "certify_sources_tab_total": 6,
   "certify_sources_enumeration": [
     {"name": "INJ-09-2026-W4 Ops Notes.xlsx", "source_id": "src_1M3VZCDQF1S0FQ0D", "visibility": "PRIVATE", "datapoints": ["QOWV Embodied Distance (CO to B6) [Certify input]", "QOWV Transport Mass (CO to B6) [Certify input]", "QOWV Pyrolysis Mass [Certify input]", "QOWV Pyrolysis Emissions - Carbon Intensity (Weighted)", "Mass of Injectant (per injection)", "Carbon Content of Bio-oil Injected (per injection)"]},
     {"name": "SR26-05719 W4 Kerry AECN QOWV COA.pdf", "source_id": "src_1M3VZ8S8W1S013HS", "visibility": "PRIVATE", "datapoints": ["Carbon Content of Bio-oil Injected (per injection)"]},
     {"name": "Charm-RetirementCertificate-RD-EACs-20251218.pdf", "source_id": "src_1M1MFJPW11S08KY2", "visibility": "PUBLIC", "datapoints": ["Constant CO2 reduction"]},
     {"name": "Standard Emission Factors + Calculations (8).xlsx", "source_id": "src_1KWZ2NVQM1S08AWN", "visibility": "PUBLIC", "datapoints": ["Emissions Factor - Truck Transport Fuel Use (Well to Wheel) - Tanker Truck"]},
     {"name": "Standard Emission Factors + Calculations (8).xlsx", "source_id": "src_1KVJRSYT91S09G2E", "visibility": "PUBLIC", "datapoints": ["Emissions Factor - Vehicle Embodied Emissions - Tanker Truck"]},
     {"name": "CO to Basco Routing Map.png", "source_id": "src_1KSTWP2AP1S0QQPF", "visibility": "PUBLIC", "datapoints": ["QOWV Embodied Distance (CO to B6) [Certify input]", "Distance from Miniforge CO to Basco"]}],
   "attachment_map": {
     "Mass of Injectant (per injection)": ["INJ-09-2026-W4 Ops Notes.xlsx"],
     "Carbon Content of Bio-oil Injected (per injection)": ["SR26-05719 W4 Kerry AECN QOWV COA.pdf", "INJ-09-2026-W4 Ops Notes.xlsx"],
     "QOWV Transport Mass (CO to B6) [Certify input]": ["INJ-09-2026-W4 Ops Notes.xlsx"],
     "QOWV Pyrolysis Mass [Certify input]": ["INJ-09-2026-W4 Ops Notes.xlsx"],
     "QOWV Pyrolysis Emissions - Carbon Intensity (Weighted)": ["INJ-09-2026-W4 Ops Notes.xlsx"],
     "QOWV Embodied Distance (CO to B6) [Certify input]": ["INJ-09-2026-W4 Ops Notes.xlsx", "CO to Basco Routing Map.png"],
     "Distance from Miniforge CO to Basco": ["CO to Basco Routing Map.png"],
     "Emissions Factor - Truck Transport Fuel Use (Well to Wheel) - Tanker Truck": ["Standard Emission Factors + Calculations (8).xlsx"],
     "Emissions Factor - Vehicle Embodied Emissions - Tanker Truck": ["Standard Emission Factors + Calculations (8).xlsx"],
     "Constant CO2 reduction": ["Charm-RetirementCertificate-RD-EACs-20251218.pdf"]},
   "attachment_map_note": "Read from the Sources tab; the Ops Notes row's '+ 5 more' popover expanded and read; the routing map's second datapoint read from the Datapoints tab.",
   "attached_ops_notes_copy": {"_primary_source_opened": [{"file": "Certify source src_1M3VZCDQF1S0FQ0D 'INJ-09-2026-W4 Ops Notes.xlsx'", "action": "sheet_grid_read", "artifact_path": "source_exports/Certify_attached_Ops_Notes_src_1M3VZCDQF1S0FQ0D.xlsx"}],
                               "comparison": "Cell-for-cell identical to the live export on the QOWV, QOWV Shipments & EE Allocation and Summary tabs in formula mode (0 differences each). In value mode the Summary tab differs only in I8/J8/I13/J13 (AECN rail mileage, 15,102.69 vs 7,598.34 mi and 1,946 vs 1,201.2 mi), outside this batch."}},
 "charm_contents": {"_primary_source_opened": [{"file": "Drive batch folder listing + Certify Sources tab (6 sources)", "action": "folder_listed", "artifact_path": "batch_folder_listing.txt", "negative_check_tokens": ["Charm Contents.xlsx"], "items_found": []}],
                    "status": "No standalone Charm Contents sheet in the Drive batch folder or among the 6 Certify sources. The per-period mass breakdown is the Ops Notes QOWV tab Production Period column, which is attached to the pyrolysis mass datapoint."},
 "week_window": {"ops_B3": "2026-09-18 (Friday)", "ops_B4": "2026-09-25 (Friday)", "certify": "19 Sep 2026 to 25 Sep 2026", "vvb_guidance": "18 - 25 September 2026", "tracker_completion": "2026-09-25 for all 42 rows (2026-09-28 export)", "lab_sampled": "25-Sep-2026 10:00"},
}
bc = json.load(open(HERE / "pp_summary_basis_comparison.json"))
aqA = sum(v["live_aq_after_bcu_mt"] for v in bc.values()); woA = sum(v["live_wo_after_bcu_mt"] for v in bc.values())
L["pp_summary_basis"]["after_bcu_live_blend"] = {"qowv": round(aqA * 1000 / live_aq_m, 10), "wodo": round(woA * 1000 / live_wo_m, 10),
                                                 "qowv_mt": round(aqA, 9), "wodo_mt": round(woA, 9)}
# per-month after-BCU Aqueous CI (live data) for the June and February totes
L["pp_summary_basis"]["after_bcu_ci"] = {m: round(bc[m]["live_aq_after_bcu_mt"] * 1000 / bc[m]["live"]["aq_mass"], 10) for m in bc}
L["shipment_tracker"]["cola40_documented_count"] = {"tab_C15": 18, "inventory_master": inv_den["COLA-40"],
    "annotation": "A19: 'physical counts are COLA-40=18/COLA-41=20' ... 'Live master now shows COLA-40=17/COLA-41=20 -- one short of the stated physical 18/20; flagged for Garrett to reconcile upstream.'"}
L["shipment_tracker"]["allocated_miles_corrected_c38_only"] = round(alloc_tab_c40, 6)
# MRV lead rulings, 2026-10-07 (chat), and the 7.2 figures on the ruled basis
_cmp = json.load(open(HERE / "pp_summary_basis_comparison.json"))
_after_tot = lambda m, v: _cmp[m][v]["I32"] - _cmp[m][v]["I19"] + _cmp[m][v]["C11"]   # month total with biomass transport after its BCUs
_ac = L["pp_summary_basis"]["after_bcu_ci"]; _L7 = q["L7"].value
_ci_rule = {"June 2026": _ac["Jun"], "February 2026": _ac["Feb"], "July 2026": _L7}
_pe_sub = sum(t["net_oil_produced_kg"] * t["ci_kgco2e_per_kg"] for t in totes)
_pe_rule = sum(t["net_oil_produced_kg"] * _ci_rule[t["production_period"]] for t in totes)
_pe_rule_live_blend = sum(t["net_oil_produced_kg"] * (L["pp_summary_basis"]["after_bcu_live_blend"]["qowv"] if t["production_period"] == "July 2026" else _ci_rule[t["production_period"]]) for t in totes)
_ci_pkg = {"June 2026": _cmp["Jun"]["copy"]["aq_ci"], "February 2026": _cmp["Feb"]["copy"]["aq_ci"], "July 2026": _L7}   # package versions (V2's sources), all I19 = C11*F7
_pe_pkg = sum(t["net_oil_produced_kg"] * _ci_pkg[t["production_period"]] for t in totes)
_prod = sum(t["net_oil_produced_kg"] for t in totes)
_v2m = L["blended_ci_source"]["v2_qowv_emissions_by_month_mt"]; _jm = sum(t["net_oil_produced_kg"] for t in totes if t["production_period"] == "July 2026")
_blend_diff_by_month = {m: round(_jm * (_cmp[m]["live_aq_after_bcu_mt"] - _v2m[f"{m} 2026"]) * 1000 / L["blended_ci_source"]["qowv_result"]["mass_kg"], 4) for m in _cmp}
L["max_rulings"] = {"2026-10-07": {
  "bcu_scope": {"quote": "BCUs should be applied to diesel transport, including biomass transport process emissions and handling (not embodied emissions).",
                "effect": "Production CIs count biomass transport after its BCUs: process and log handling offset, embodied not (summary I19 = C11*F7, 'Remaining Emissions to Allocate'). The live monthly summaries count it before (I19 = C9*F7); the Jan-Jun blend (V2) was built on the after-BCU basis. Checklist 7.2."},
  "cobb": {"quote": "On COBB, I defer to Garrett. He works directly with the team that uses COBB and I do not, so if they have changed up he would know that first.",
           "garrett_positions_read": ["Slack #team-mrv-alerts, Garrett Lutz, 2026-10-05 08:24 BST, Week 2 QOWV: 'COBB is not source of truth heree'; Week 3 QOWV: 'cobb not source of truth here'",
                                      "same message, Week 1 WODO: 'These are all acceptable discrepancies-BGN Inv takes precedent over oil inv data, rounding error here, red herring' (source hierarchy; routed to the MRV lead for a ruling the same day, #team-mrv-alerts 2026-10-05 11:16 BST)"],
           "effect": "The injection tracker's COBB tabs are not used as a cross-check: 2.3 drops tote 301215 (a COBB-only conflict) and 3.1 drops the COBB injection entries. Tote 301274's gross conflict rests on the production system's own weighing (Oil_inventory_data row 39, Historical Tote Record row 1181), not COBB, and stays in 2.3."}}}
L["pp_summary_basis"]["ruling_2026_10_07"] = {
  "ci_by_period": _ci_rule, "pe_submitted_kg": round(_pe_sub, 4), "pe_ruling_kg": round(_pe_rule, 4), "pe_change_kg": round(_pe_rule - _pe_sub, 4),
  "weighted_ci_ruling": round(_pe_rule / _prod, 10), "june_change_kg": round(sum(t["net_oil_produced_kg"] * (_ac["Jun"] - t["ci_kgco2e_per_kg"]) for t in totes if t["production_period"] == "June 2026"), 4),
  "feb_change_kg": round(sum(t["net_oil_produced_kg"] * (_ac["Feb"] - t["ci_kgco2e_per_kg"]) for t in totes if t["production_period"] == "February 2026"), 4),
  "blend_v2_after_bcu": L["blended_ci_source"]["qowv_result"]["ci"], "blend_live_after_bcu": L["pp_summary_basis"]["after_bcu_live_blend"]["qowv"],
  "july_change_if_blend_rebuilt_from_live_kg": round(_pe_rule_live_blend - _pe_rule, 4),
  "month_totals_after_bcu_mt_copy_vs_live": {m: [round(_after_tot(m, "copy"), 6), round(_after_tot(m, "live"), 6)] for m in _cmp},
  "package_version": {"ci_by_period": _ci_pkg, "pe_kg": round(_pe_pkg, 4), "pe_change_kg": round(_pe_pkg - _pe_sub, 4), "weighted_ci": round(_pe_pkg / _prod, 10)},
  "linked_sheets_version": {"ci_by_period": {"June 2026": _ac["Jun"], "February 2026": _ac["Feb"], "July 2026": L["pp_summary_basis"]["after_bcu_live_blend"]["qowv"]},
                            "pe_kg": round(_pe_rule_live_blend, 4), "pe_change_kg": round(_pe_rule_live_blend - _pe_sub, 4), "weighted_ci": round(_pe_rule_live_blend / _prod, 10),
                            "blend_difference_on_july_totes_by_month_kg": _blend_diff_by_month},
  "version_history": {
    "_primary_source_opened": [{"file": "dated local downloads of the linked sheets (~/Downloads/Charm Production Docs/{January,March,April} 2026 Production Period/, ~/Downloads/{Feb,May,Jun} 26 Production Period Summary.xlsx) and the July RP package versions (~/Downloads/July 1 - July 31 RP/Charm BO MRV/H1 2026 Charm CI/*/), plus 'Biomass Transport Emissions.xlsx' in its June folder", "action": "sheet_grid_read"}],
    "linked_sheet_downloads": "Jan 07-11 03:15, Mar 07-11 03:14, Apr 07-16 21:47, Feb 08-10 21:17, May 08-11 02:02, Jun 08-11 02:05 (local time): all I19 = C9*F7 and already holding today's biomass-transport values (Jan C9 2.495338746, Feb C9 1.797501254, Jun C13 129.1156462585034 / C14 405.5)",
    "package_versions": "saved 2026-08-11 11:28 (Feb) to 22:39 (Jun) local, after every download: I19 = C11*F7 in every month with biomass transport (May I19 = C9*F7 with C9 = 0); Jan C9 2.365369654, Feb C9 1.701467837, Jun C13 109.03 / C14 354.5",
    "evidence_for_package": ["Feb linked sheet C5 1.4172200170521543 and C6 0.09439038819643568 are typed values at the 2026 factors, while its own 'Emission Factors' tab lists E3 0.000135156 and E6 6.44927e-05 and the package's 'Biomass Transport Emissions.xlsx' EmissionFactors D9 reads 'First day the 2026 EFs apply. Feb 2026 and earlier = 2025 EFs.'",
                             "June package C13 109.03 and C14 354.5 match 'Biomass Transport Emissions.xlsx' 'Monthly Totals' row 14 (5 loads, 109.0340136 t, 354.5 mi); no file read derives the linked sheet's 129.1156462585034 t and 405.5 mi"],
    "conclusion": "The package versions are the later, reworked ones; the linked sheets were not edited after them (except June J6, which moved its T4 from 0.123429378 in the 08-11 02:05 download to 0.1236896674 today). Which version governs is for the blend owner to confirm."},
  "note": "On the ruled basis the July blend in use (V2) stands only with the package versions; mixing package July with linked-sheet June/February (pe_ruling_kg) is not a consistent basis. See package_version and linked_sheets_version."}
# production-period corroboration and the production system's location field (second addendum review)
_inv = openpyxl.load_workbook(INV, read_only=True, data_only=True)
_q4 = list(_inv["Q4 2025 -2026 Production Invent"].iter_rows(min_row=3, values_only=True))
_oi = list(_inv["Oil_inventory_data"].iter_rows(min_row=1, values_only=True))
assert _oi[1][1:7] == ("lot_number", "security_tag", "production_date", "inventory_itag", "site_location", "location"), _oi[1]
_norm = lambda v: str(v).replace(".0", "") if v is not None else None
_oi_by_tag = {}
for i, r in enumerate(_oi[2:], start=3):
    if r and r[2] is not None:
        _oi_by_tag.setdefault(_norm(r[2]), []).append((i, r))
_q4_labels = collections.defaultdict(set)          # a tag can sit on more than one Q4 row (e.g. 301025: February and March 2026 PP)
for r in _q4:
    if r[5] and r[1]:
        _q4_labels[_norm(r[5])].add(r[1])
_mon = {m: k for k, m in enumerate(["January", "February", "March", "April", "May", "June", "July", "August", "September"], start=1)}
_next = collections.defaultdict(list)
for r in htr.iter_rows(min_row=3, values_only=True):
    if r[6] is None or not isinstance(r[1], datetime.datetime) or r[1].year != 2026:
        continue
    for pp_ in _q4_labels.get(_norm(int(r[6])), ()):
        if str(pp_).endswith(" 2026 PP") and _mon.get(str(pp_).split(" ")[0]) and r[1].month == _mon[str(pp_).split(" ")[0]] + 1:
            _next[pp_].append(f"{r[1]:%Y-%m-%d}")
jun_q = [r for r in _q4 if r[1] == "June 2026 PP" and r[7] == "QOWV"]
lab26 = [r for r in _q4 if r[1] and re.fullmatch(r"(January|February|March|April|May|June|July|August|September) 2026 PP", str(r[1]))]
paint = collections.Counter(f"{r[1]} / {r[18]}" for r in _q4 if r[1] in ("June 2026 PP", "July 2026 PP"))
L["tote_weighing_records"]["june_totes_weighed_in_july"]["corroboration"] = {
  "_primary_source_opened": [{"file": "inventory master tabs 'Q4 2025 -2026 Production Invent' and 'Oil_inventory_data'; Jun 26 PP Summary; V2 Jan-Jun blend workbook", "action": "sheet_grid_read",
                              "artifact_path": "source_exports/inventory_master_export.xlsx"}],
  "june_summary_aqueous_mass_N4_kg": pps["Jun"]["aqueous_mass_kg"],
  "q4_june_qowv_rows": len(jun_q), "q4_june_qowv_net_kg": round(sum(r[10] for r in jun_q), 4),
  "four_totes_inside_june_qowv_rows": sorted(_norm(r[5]) for r in jun_q if _norm(r[5]) in {x["tag"] for x in late_june}),
  "v2_blend_june_qowv_mass_D12_kg": v2["D12"].value,
  "spray_paint_column_S_counts": dict(paint),
  "spray_paint_eight_totes": {_norm(r[5]): r[18] for r in _q4 if _norm(r[5]) in {x["tag"] for x in late_june} | {x["tag"] for x in late_july}},
  "production_date_four_june_totes": {t: {"row": i, "production_date": f"{r[3]:%Y-%m-%d %H:%M:%S}", "produced_on_work_order": r[7], "production_period_col_M": r[12] if len(r) > 12 else None}
                                      for t in sorted(x["tag"] for x in late_june) for i, r in _oi_by_tag.get(t, [])},
  "weighed_in_month_after_label_2026": {k: {"count": len(v), "first": min(v), "last": max(v)} for k, v in _next.items()},
  "tags_on_two_q4_labels": {t: sorted(v) for t, v in _q4_labels.items() if len(v) > 1},
  "q4_form_tote_date_blank_2026": f"{sum(1 for r in lab26 if r[3] in (None, ''))} of {len(lab26)} rows labelled January-September 2026",
  "q4_august_2026_rows_by_fraction": dict(collections.Counter(r[7] for r in _q4 if r[1] == "August 2026 PP"))}
_q4_rows = {i: r for i, r in enumerate(_q4, start=3)}
_shift = {}
for tg in ("301274", "301277", "301227"):
    bgn = [(i, r[8], r[10]) for i, r in _q4_rows.items() if _norm(r[5]) == tg]
    prod = [(i, r[7], r[10], r[11]) for i, r in _oi_by_tag.get(tg, [])]
    hrec = [(i, r[3], r[0], r[5]) for i, r in enumerate(htr.iter_rows(min_row=1, values_only=True), start=1) if i > 2 and r[6] is not None and _norm(int(r[6])) == tg]
    _shift[tg] = {"bgn_q4_row_gross_net": bgn, "oil_inventory_data_row_order_net_gross": prod, "historical_tote_record_row_order_net_gross": hrec}
assert _shift["301274"]["bgn_q4_row_gross_net"][0][1] == _shift["301277"]["oil_inventory_data_row_order_net_gross"][0][3] == 1036.0
assert _shift["301277"]["bgn_q4_row_gross_net"][0][1] == _shift["301227"]["oil_inventory_data_row_order_net_gross"][0][3] == 1070.0
# production-system records (Manufacturo SQL imports), kept apart from the injection tracker's COBB tabs (batch-qa-charm FC-C14)
_oi39 = _oi[38]                                   # Oil_inventory_data row 39
assert _norm(_oi39[2]) == "301274" and _oi39[7] == "OIL260721143336", _oi39
_t274 = next(t for t in totes if t["tag"] == "301274")
_gc_old = L["cobb_injection_record"].pop("gross_conflicts")["301274"]
_dg = round(_oi39[11] - _t274["gross_kg"], 4)
_t300 = L["cobb_injection_record"].pop("tag_300020")
L["cobb_injection_record"]["tag_300020_injections"] = _t300.pop("injections")
L["production_records"] = {
  "_primary_source_opened": [{"file": "inventory master tabs 'Oil_inventory_data' and 'Q4 2025 -2026 Production Invent'; injection tracker tab 'Historical Tote Record' (SQL import, not COBB)", "action": "sheet_grid_read",
                              "artifact_path": "source_exports/inventory_master_export.xlsx"}],
  "gross_conflict_301274": {"ops_notes_B40_kg": _t274["gross_kg"], "q4_inventory_tab_row694_kg": _gc_old["q4_inventory_tab_row694_kg"],
                            "production_system_gross_kg": _oi39[11], "production_system_net_kg": _oi39[10], "production_system_work_order": _oi39[7],
                            "production_system_Oil_inventory_data_row39": _gc_old["production_system_Oil_inventory_data_row39"],
                            "historical_tote_record_row1181": _gc_old["historical_tote_record_row1181"], "ops_notes_E40_kg": _t274["empty_tote_weight_kg"],
                            "if_production_gross": {"mass_gain_kg": _dg, "seq_gain_kg": round(_dg * L["testing"]["qowv_certificate"]["carbon_pct"] / 100 * 3.667, 6),
                                                    "pe_gain_kg": round(_dg * _t274["ci_kgco2e_per_kg"], 6),
                                                    "net_gain_kg": round(_dg * L["testing"]["qowv_certificate"]["carbon_pct"] / 100 * 3.667 - _dg * _t274["ci_kgco2e_per_kg"], 6)}},
  "tag_300020": _t300,
  "bgn_row_shift_301274": dict(_shift, read="BGN (Q4 tab) gives 301274 1036.0/978.4 kg, which is the production system's weighing of 301277 (OIL260721142141); BGN's 301277 row gives 1070.0/1012.4 kg, which is 301227's weighing (OIL260721133028). Each production record carries its own work order. The BGN rows look shifted, so the production system's 1080.0 kg (OIL260721143336) fits 301274."),
  "source_hierarchy": "The order among the Ops Notes, the BGN inventory and the production system's records is not yet ruled (batch-qa-charm FC-C14; operations' view, #team-mrv-alerts 2026-10-05: 'BGN Inv takes precedent over oil inv data')."}
_batch_tags = {t["tag"] for t in totes}
_loc = collections.Counter((r[5], r[6]) for t in _batch_tags for i, r in _oi_by_tag.get(t, []))
_lot4746 = [(i, r) for i, r in enumerate(_oi[2:], start=3) if _norm(r[1]) == "4746"]
L["tote_weighing_records"]["location_basis"]["production_system_location"] = {
  "batch_tags_found": len([t for t in _batch_tags if t in _oi_by_tag]), "site_location_location_counts": {f"{a} / {b}": c for (a, b), c in _loc.items()},
  "lot_4746": [{"row": i, "site_location": r[5], "location": r[6]} for i, r in _lot4746],
  "tab_site_location_counts": dict(collections.Counter(r[5] for r in _oi[2:] if r and r[5])),
  "moves_with_tote_example": {"row": 350, "lot": _oi[349][1], "security_tag": _oi[349][2], "site_location": _oi[349][5], "location": _oi[349][6]},
  "read": "'Oil_inventory_data' (A1 'Microsoft SQL Server Import / Last updated 7 weeks ago') has site_location and location columns. The batch's tags it carries, plus lot 4746, read 'CHARM' / 'OPA'; already-shipped lot LOT00004770 reads 'CHARM_LOUISIANA' / 'Bio Oil Incoming', so the field follows the tote: it gives where a tote sat at the data pull, not where it was weighed."}
for t in totes:                                   # name the production system's location field on each tote row
    hit = _oi_by_tag.get(t["tag"]) or ([(i, r) for i, r in _lot4746] if t["tag"] == "300022" else [])
    t["weighing_location"] += ("; production system site_location / location at the data pull: "
                               + (f"{hit[0][1][5]} / {hit[0][1][6]} (Oil_inventory_data row {hit[0][0]}{', lot 4746' if t['tag'] == '300022' else ''})" if hit
                                  else "none (the Oil_inventory_data pull, 'Last updated 7 weeks ago', has no row for this tag)"))
L["adversarial_review"] = {
  "artifact_path": "adversarial_review_INJ-09-2026-W4-QOWV.md", "reviewer": "independent subagent, files on disk only",
  "completed": datetime.datetime.fromtimestamp((HERE / "adversarial_review_INJ-09-2026-W4-QOWV.md").stat().st_mtime, datetime.UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
  "accepted": ["7.2 kept Fail, ground rewritten: the ledger mixes two CI bases (June/February live summaries count biomass transport before BCUs; the July blended CI was built from copies that count it after BCUs). Quantified both ways; basis is a method decision.",
               "2.3 Flag -> Fail: conflicting post-unload weights between two records (conflicting evidence = Fail). Reason for weighing the two rows differently from the tracker's general offset added to the note: on 39 rows plus halved row 867 the tracker sits exactly 57.6 kg (the standard tote tare) below the ledger, a constant definitional difference, while 301215 breaks it by 280.5 kg (tracker 222.9 kg above the ledger's net injected). The live tracker read (2026-10-06) showed 301274's difference is its gross weight (tracker 1080.0 kg vs 1036.0 kg in the Ops Notes and inventory master), with the post-unload weight agreeing at 169.5 kg; 2.3 was rewritten accordingly. The operations position on the tracker as a record (registry INJ-09-2026-W2-QOWV notes) goes to the MRV lead in the delivery message, not the checklist.",
               "3.4 COLA-40 18 is the documented physical count (A19 annotation), not an error; COLA-38 alone carries the Fail; corrected allocation 2,568.8056 mi",
               "C.1 corrected net given as a range across the CI basis choice",
               "5.1/8.1 fix wording: BCU equals whatever the corrected process component outputs (7,138.434 kg if Certify keeps GLEC 0.086 kg/(km*t) as on W3)",
               "7.3 wording: the V2 workbook is not among the six Certify sources; link target named",
               "1.1 wording: file-size wording dropped; Drive id and formation-fluid density threshold cited",
               "7.4/7.5 per-period figures added; column B questions set to the skill's exact text; sidebar adds the WODO allocation tab's stale W2 header, the basis note on the WODO blended CI, and the May C85 factor-of-ten",
               "Addendum review after the live tracker read (adversarial_review_addendum_INJ-09-2026-W4-QOWV.md), each point re-checked in the files: 2.3's 301274 reasoning corrected (the production system's own weighing, Oil_inventory_data row 39 and Historical Tote Record row 1181, gives 1080.0 kg, so two production records disagree) and its fix widened to the dependent items; 3.1 Pass -> Flag (security tag 300020 assigned to lots 4746 and LOT00004770 in the production system, recorded as injected in July and in W4); sidebar tracker notes reworded (every September batch repeats across blocks; pH is the weekly wellhead value stamped on all products; residue shares 0.3714 and 0.3914 of each tote's net oil); 1.1 agreed. The operations position on the tracker as a record goes to the MRV lead in the delivery message.", "Final audit points applied: sidebar describes the four blocks as stacked imports with their Import Time stamps and puts 301204's doubled figure in the oldest; 2.1 and 2.2 point to 2.3 for 301274's disputed gross; 2.3's fix names 4.3; 2.3's offset wording is limited to the newest import; 3.4 gives 2,576.6552 mi as the allocation if the master's 17 for COLA-40 proves right.",
               "Second addendum review (adversarial_review_addendum2_INJ-09-2026-W4-QOWV.md) of 2.1, 3.5, 7.1, 7.3 and C.1, each point re-checked in the files by the QA agent: 3.5 kept Flag but its ground rewritten. The June summary's Aqueous mass (N4 46,616.2 kg) is the sum of the 48 QOWV rows labelled June 2026, these four included, and the Jan-Jun blend's June mass (D12) is the same figure. The totes carry their labelled period's spray paint. Next-month weigh-ins are routine. The production system's production_date is the weigh-in. So the -11.008 kg counterfactual and 'conservative' are dropped, and the flag is the missing fill date (blank on all 421 rows for 2026), with a new fix. C.1's 3.5 phrase replaced with 'does not change the net'. 2.1 date range splits lot 4746's record from the 41 own-tag records. The location basis now cites the production system's site_location/location field. 7.3 says the July copies show June's CI, and the output summary is confirmed one tab. 7.1 upheld."],
  "addendum3": {"artifact_path": "adversarial_review_addendum3_INJ-09-2026-W4-QOWV.md", "reviewer": "independent subagent, files on disk only",
                "completed": datetime.datetime.fromtimestamp((HERE / "adversarial_review_addendum3_INJ-09-2026-W4-QOWV.md").stat().st_mtime, datetime.UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
                "accepted": "Review of the rows changed by the 2026-10-07 rulings (2.3, 3.1, 7.2, 7.3, C.1, sidebar, recompute, drafts). Each point was re-checked in the files by the QA agent and accepted. 7.2: the mixed figure (linked June/February with the package July blend) was replaced by the two consistent versions, package 4,790.2737 kg and linked 4,786.3316 kg; the version direction was corrected (the package copies are the later, reworked ones, shown by dated downloads); the evidence favoring the package versions was added (February 2025 factors, June per-load workbook); the blend difference was attributed to March and May C85; the fix was limited to Ops Notes L18/L14 and Certify. 2.3: BGN row-shift evidence added (1,036.0 kg is 301277's weighing; 301277's row carries 301227's 1,070.0 kg); conversation citations replaced by FC-C14; 'similar conflict' dropped. Sidebar: 2025 summaries added, the May overstatement removed, the May C85 note corrected. Scan file: scope header added. C.1: nets on both versions. Drafts: stale and backwards wording fixed, Garrett 7.2 held. Skill FC-C13: the Fail condition is BCUs applied (C10 > 0), the version note and the 2025 summaries added; FC-C14 quote given in full. Recompute: the 7.2 CIs are computed from the summary workbooks (source_exports/pp_summaries/ and package/) and the 2.3 effect from Oil_inventory_data row 39 and the Ops Notes row; the production-system entries moved from cobb_injection_record to production_records. 2.3 carries the sentence that the order among the three records is not yet ruled. 3.1 and 7.3 upheld."},
  "addendum2": {"artifact_path": "adversarial_review_addendum2_INJ-09-2026-W4-QOWV.md", "reviewer": "independent subagent, files on disk only",
                "completed": datetime.datetime.fromtimestamp((HERE / "adversarial_review_addendum2_INJ-09-2026-W4-QOWV.md").stat().st_mtime, datetime.UTC).strftime("%Y-%m-%dT%H:%M:%SZ")},
  "rejected": {}, "verified_by_qa_agent": "Copy vs live I19 formulas and C10 BCUs re-read for all six months (pp_summary_basis_comparison.json); after-BCU blends recomputed; WODO allocation tab A1 re-read."}
OUT.write_text(json.dumps(L, indent=1, default=str))
print("wrote", OUT.name, "totes", n, "pp", dict(pp_counts), "match", match, "trucks", dict(by_truck))
print("inv denominators", inv_den, "alloc inv", round(alloc_inv, 4), "alloc tab C40=18", round(alloc_tab_c40, 4))
print("live QOWV blended", L["blended_ci_source"]["live_recompute"])
print("constants match live", L["pp_summary_sources"]["constants_match_live"])
