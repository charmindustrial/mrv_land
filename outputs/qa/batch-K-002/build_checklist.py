"""Build the Kerry K-002 QA checklist xlsx.

Per skill spec — sections 1, 2, 3, 4, 5, 6, 6a, 7, 8.
Columns: A item#, B description, C status, D notes/evidence, E verifier ref.
Formatting: FAIL red across A-E; FLAG yellow #FFF2CC + dark gold #996600 font; PASS no fill;
N/A gray italic; section headers preserved.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = "/Users/maxlavine/Desktop/Claude/Skills/MRV Expert/outputs/qa/batch-K-002/Kerry_Batch_K-002_QA_Checklist.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "QA Checklist"

# Column widths
widths = {"A": 8, "B": 56, "C": 9, "D": 100, "E": 28}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# Style helpers
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
SECTION_FILL = PatternFill("solid", fgColor="9DC3E6")  # blue
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
FLAG_FONT = Font(color="996600", bold=False)
NA_FONT = Font(italic=True, color="808080")
SECTION_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(border_style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP = Alignment(wrap_text=True, vertical="top")

def write_title_block(ws):
    ws["A1"] = "Kerry Injection Batch QA Checklist"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:E1")
    rows = [
        ("Batch ID", "K-002"),
        ("Reporting Period", "2026-04-21"),
        ("Reviewer", "Claude QA Agent"),
        ("Date", "2026-05-05"),
        ("Injection Site", "Basco 6"),
        ("Oil Type", "Kerry"),
        ("Delivery Mode", "Tote (4 totes: KC_0005, KC_0006, KC_0007, KC_0008)"),
        ("Certify Removal", "rmv_1KQ0QGNBV1S0B315 (Net 9.378 tCO₂e)"),
    ]
    for i, (k, v) in enumerate(rows, start=2):
        ws[f"A{i}"] = k
        ws[f"B{i}"] = v
        ws[f"A{i}"].font = HEADER_FONT
        ws.merge_cells(f"B{i}:E{i}")
    return len(rows) + 2  # next row to write

def section_header(ws, row, title):
    ws[f"A{row}"] = title
    ws[f"A{row}"].font = SECTION_FONT
    ws[f"A{row}"].fill = SECTION_FILL
    ws.merge_cells(f"A{row}:E{row}")
    return row + 1

def col_header(ws, row):
    ws[f"A{row}"] = "Item #"
    ws[f"B{row}"] = "Check Item"
    ws[f"C{row}"] = "Status"
    ws[f"D{row}"] = "Notes / Evidence"
    ws[f"E{row}"] = "Verifier Reference"
    for c in "ABCDE":
        ws[f"{c}{row}"].font = HEADER_FONT
        ws[f"{c}{row}"].fill = HEADER_FILL
        ws[f"{c}{row}"].border = BORDER
        ws[f"{c}{row}"].alignment = Alignment(horizontal="center", vertical="center")
    return row + 1

def write_item(ws, row, item_num, desc, status, notes, ref):
    ws[f"A{row}"] = item_num
    ws[f"B{row}"] = desc
    ws[f"C{row}"] = status
    ws[f"D{row}"] = notes
    ws[f"E{row}"] = ref
    for c in "ABCDE":
        ws[f"{c}{row}"].border = BORDER
        ws[f"{c}{row}"].alignment = WRAP
    if status == "Fail":
        for c in "ABCDE":
            ws[f"{c}{row}"].fill = FAIL_FILL
    elif status == "Flag":
        for c in "ABCDE":
            ws[f"{c}{row}"].fill = FLAG_FILL
            ws[f"{c}{row}"].font = FLAG_FONT
    elif status == "N/A":
        for c in "ABCDE":
            ws[f"{c}{row}"].font = NA_FONT
    # Pass = default no fill
    ws.row_dimensions[row].height = 60
    return row + 1


# ---- Build the checklist ----
r = write_title_block(ws)
r += 1  # blank spacer

# SECTION 1
r = section_header(ws, r, "Section 1: Batch Folder Completeness")
r = col_header(ws, r)
r = write_item(ws, r, "1.1",
    "Batch folder exists with correct naming",
    "Pass",
    "Drive folder titled 'K-002' (id: 1PNK-dv_W-PABRX_asipplRBgWoLlMqNr) under Charm batches structure. Naming follows Kerry batch convention K-NNN. [log:drive_folder_url]",
    "VR Appendix 1")
r = write_item(ws, r, "1.2",
    "Mass quantification subfolder present and populated (tote mode: tote weighing log + per-tote weighing records)",
    "Fail",
    "No 'Tote Weighing' subfolder in Drive. K-002 sheet records gross/net/leftover per tote BUT does NOT contain measured tare per tote, weighing date per tote, or weighing location per tote — required by VR Appendix 1 for mass quantification evidence in tote mode. No standalone weighing log present in Drive or in Certify Sources. [log:drive_folder_contents.subfolders_present=false] [log:K-002_sheet_content.math_checks.missing_fields]",
    "VR Appendix 1")
r = write_item(ws, r, "1.3",
    "Testing subfolder present with all required lab results (per-tote SGS CoA in tote mode)",
    "Pass",
    "SGS CoA SR26-02553 (filename '2-192-2-198+Kc0005-Kc-0008.pdf') uploaded to Certify Sources tab as Public source mapped to 'Carbon Content of Bio-oil Injected (per injection)'. Pages 5-8 cover all 4 K-002 totes (KC-0005, KC-0006, KC-0007, KC-0008) with carbon, hydrogen, nitrogen, viscosity, acid number, density. Per Document Location Rule, Drive-or-Certify presence is sufficient. [log:sgs_coa.kerry_per_tote]",
    "VR Appendix 1")
r = write_item(ws, r, "1.4",
    "Ops Notes spreadsheet present and complete (Kerry mass IN, AECN mass IN, completion date, pH)",
    "Fail",
    "K-002 Google Sheet present in Drive (id 1HbUqSoD_iSkA8fSdARwDXYS7LYc-yGX77U1CD3Mc2Fg). Sheet records per-tote mass, carbon content, weighted-avg, uncertainty, totals (4,953 transported / 4,618 oil injected). MISSING: completion date (only in COBB at L83-86), pH (only in COBB at M83-86), AECN mass IN context (no homogenization-tank breakdown documented), no carrier/tractor/trailer/driver fields, no weighing date+location per tote. Sheet is a math-only artifact, not a complete operational record. [log:K-002_sheet_content.math_checks.missing_fields]",
    "VR Appendix 1")

r += 1  # spacer
# SECTION 2
r = section_header(ws, r, "Section 2: Mass Quantification (Tote Mode)")
r = col_header(ws, r)
r = write_item(ws, r, "2.1",
    "Per-tote weighing records present (gross, tare, net) for every tote with weighing date+location",
    "Fail",
    "K-002 sheet records per-tote: KC-0005 (gross 1185 / net 1092 / leftover 93), KC_0006 (1284 / 1209 / 75), KC_0007 (1261 / 1176 / 86), KC_0008 (1223 / 1142 / 82). Math checks: gross − net = leftover within ±1 kg ✓. NO measured tare per tote, NO weighing date, NO weighing location, NO scale identifier. The sheet is a derived summary, not a primary weighing record. Per VR Sec 5.1, mass quantification requires primary weighing evidence with full provenance — absent here. [log:K-002_sheet_content.totes] [log:K-002_sheet_content.math_checks.missing_fields]",
    "VR Sec 5.1")
r = write_item(ws, r, "2.2",
    "Tote tare measured (not assumed nominal) — measured value paired with weighing date for THIS batch's totes (FC-K1)",
    "Fail",
    "K-002 sheet has NO tare column at all. The 'Gross Tote Weight' values (1185–1284 kg) reconcile with 'Net Oil In Tote' + 'Leftover' (oil staying in tote post-offload), so 'Gross' = oil-only mass delivered, NOT (oil + measured tote tare). Tote tare is therefore not documented anywhere in available evidence — neither measured nor nominal. Per FC-K1, measured tare with per-batch weighing date is required; absence is FAIL territory. [log:K-002_sheet_content.totes] [log:interpretation_of_gross_tote_weight]",
    "VR Sec 5.1, FC-K1")
r = write_item(ws, r, "2.3",
    "Net mass per tote: net = gross − tare; sum of nets = MOBI",
    "Flag",
    "Independent recompute of summed nets: 1092 + 1209 + 1176 + 1142 = 4,619 kg. Sheet displays 'Total Oil Injected: 4,618' (1 kg rounding). Certify Mass of Injectant = 4.619 t = 4,619 kg ✓. Math reconciles with sheet, but the 'tare' inputs are unverifiable (see 2.2) — flagging because while computed nets match sheet/Certify, the underlying tare is undocumented. [log:K-002_sheet_content.totes] [log:certify_datapoints.Mass of Injectant]",
    "VR Sec 5.1")
r = write_item(ws, r, "2.4",
    "Conservative MOBI: MIN(Σ tote nets, BOL oil-only mass) used as Kerry credit basis",
    "Fail",
    "Σ tote nets = 4,619 kg; Certify Mass of Injectant = 4,619 kg ✓ as the credit basis. HOWEVER: there is no BOL to compare against (see 3.1) → cannot confirm MIN() rule was applied. The conservative-mass test (selecting smaller of two independent measures) is structurally impossible without the second measure. Status FAIL due to missing BOL reference, not due to wrong arithmetic. [log:certify_datapoints] [log:certify_sources_missing]",
    "GHG Stmt Sec C")
r = write_item(ws, r, "2.5",
    "Per-tote reconciliation: count(totes on BOL) = count(serials weighed) = count(nets in Ops Notes) (FC-K3); transport-EF Mass of Load includes tote tares (FC-K2)",
    "Fail",
    "Per-tote count K-002 sheet: 4 (KC-0005, KC_0006, KC_0007, KC_0008). SGS CoA: 4 totes (matching serials). COBB rows 83-86: 4 totes (matching serials). Three-way count match ✓. HOWEVER: count(totes on BOL) is unknowable — no BOL exists to compare. FC-K2 check on transport-EF basis: Certify 'Allocated Kerry Oil Gross Mass' = 4,953 kg = Σ K-002-sheet 'Gross Tote Weight' = oil-only delivered mass (oil + leftover residue, NO physical tote tare or deck items). Per FC-K2 specification, transport process Mass of Load should be BOL gross including tote tares. Current basis under-counts hauled mass by ≥4×~60 kg (typical IBC tare) ≈ 240 kg → under-reports transport process emissions. [log:cobb_tracker.rows_for_K-002] [log:interpretation_of_gross_tote_weight] [log:certify_per_component_uncertainty.transport_process_inputs.FC_K2_check]",
    "VR Sec 5.1, FC-K2/FC-K3")

r += 1
# SECTION 3
r = section_header(ws, r, "Section 3: Bills of Lading & Transportation")
r = col_header(ws, r)
r = write_item(ws, r, "3.1",
    "BOL(s) present for each shipment",
    "Fail",
    "No BOL filename in Certify Sources tab (8 sources enumerated; none is a BOL). No BOL in Drive folder root (only K-001 sheet, K-002 sheet, prior-batch SGS PDF). No BOL_number populated in COBB tracker for any of rows 83-86 (column E blank for all 4 K-002 entries). Per Document Location Rule, BOL must be in Drive OR Certify Sources — absent from BOTH. VR Appendix 1 requires BOL for transport leg verification. [log:certify_sources_enumeration] [log:certify_sources_missing] [log:cobb_tracker.rows_for_K-002]",
    "VR Appendix 1")
r = write_item(ws, r, "3.2",
    "BOL mass consistent with mass artifacts (Tote: BOL gross ≥ MOBI + Σ tote tares)",
    "N/A",
    "No BOL exists; cannot perform BOL-vs-mass-artifact reconciliation. See 3.1.",
    "VR Sec 3.5")
r = write_item(ws, r, "3.3",
    "Origin and destination correct on BOL (Kerry origin → Basco)",
    "N/A",
    "No BOL exists; cannot verify origin/destination on BOL. Certify reports Route Distance Kerry to Basco = 1,090 mi (per Datapoints tab); Garrett confirmed single Kerry origin + 1,090 mi distance on 2026-05-04 (per skill SKILL.md route reference table). [log:certify_datapoints.Route Distance Kerry to Basco]",
    "PDD Sec A")
r = write_item(ws, r, "3.4",
    "Transportation distance consistent with confirmed Kerry route (1,090 mi single origin)",
    "Pass",
    "Certify Datapoint 'Route Distance Kerry to Basco' = 1,090 mi (Reported, source 'Kerry to Basco Distance.png'). Matches SKILL.md/KB reference distance for current single Kerry origin (confirmed by Garrett 2026-05-04). Process modal Inputs: Distance traveled = 1,090 mi; Embodied modal Inputs: Distance traveled = 1,090 mi. Both legs of transport reference the same distance value ✓. [log:certify_datapoints.Route Distance Kerry to Basco] [log:certify_per_component_uncertainty]",
    "Transport Emissions Acctg v1.1")

r += 1
# SECTION 4
r = section_header(ws, r, "Section 4: Bio-Oil Composition & Testing")
r = col_header(ws, r)
r = write_item(ws, r, "4.1",
    "CHN analysis PDF(s) present (per-tote CoA in current tote-mode practice)",
    "Pass",
    "SGS CoA SR26-02553 (issued 2026-04-21, Lab: SGS North America Inc., 151 James Drive West, St. Rose, LA 70087) consolidated PDF in Certify Sources. 9 pages total: pages 5-8 cover all 4 K-002 totes (KC-0005, KC-0006, KC-0007, KC-0008) with full ASTM D5291 (Method C) carbon, hydrogen, nitrogen, ASTM D445 viscosity, ASTM D664 acid number, ASTM D4052 density. All samples received 2026-04-15 07:30 at SGS Louisiana facility. Charm Industrial as client. [log:sgs_coa]",
    "VR Appendix 1")
r = write_item(ws, r, "4.2",
    "Carbon content reported and used in Certify (Kerry baseline N/A — n=8 cumulative across K-001+K-002, FLAG-with-note per skill spec)",
    "Flag",
    "Per-tote carbon (from SGS): KC-0005 56.6%, KC-0006 55.0%, KC-0007 57.5%, KC-0008 55.0%. Independent weighted-avg recompute: (56.6×1092 + 55.0×1209 + 57.5×1176 + 55.0×1142) / 4619 = 258,732.2 / 4619 = 56.013% → reported 56.01% ✓. Certify 'Carbon Content of Bio-oil Injected (per injection)' = 56.01% with full-precision 56.01477 (per Sequestration modal Inputs). FLAG-with-note: Kerry historical baseline NOT yet established (only n=8 samples cumulative across K-001 and K-002); cannot be evaluated against AECN's range (mean 39.42%, SD 6.09%) which is feedstock-chain-different. Per skill spec, item starts as FLAG-with-note until n ≥ ~30 Kerry samples accumulate. [log:sgs_coa.kerry_per_tote] [log:certify_per_component_uncertainty.injection_sequestration_inputs.carbon_content] [log:K-002_sheet_content.wt_avg_carbon_pct_reported]",
    "VR Appendix 1, Biomass Feedstock Acctg v1.2")
r = write_item(ws, r, "4.3",
    "Lab report references correct batch/sample IDs (per-tote sample-to-serial match in tote mode)",
    "Pass",
    "SGS CoA per-page CLIENT ID 'BATCH ' field: page 5='KC-0005', page 6='KC-0006', page 7='KC-0007', page 8='KC-0008'. Each matches K-002 sheet's tote serials exactly (modulo dash/underscore: sheet uses 'KC-0005','KC_0006','KC_0007','KC_0008' — minor stylistic FLAG-worthy inconsistency but not material). All 4 Kerry totes received same time 2026-04-15 07:30, sampled by Charm. [log:sgs_coa.kerry_per_tote]",
    "VR Appendix 1")

r += 1
# SECTION 5
r = section_header(ws, r, "Section 5: Ops Notes Accuracy & Internal Consistency")
r = col_header(ws, r)
r = write_item(ws, r, "5.1",
    "Correct batch ID and injection date",
    "Flag",
    "K-002 sheet header reads 'Kerry Contents, 4/1, C Content Uncertainty' — sheet has NO explicit batch-ID header field, NO injection date field. The COBB tracker is the canonical batch-ID reference (cell A83-A86 = 'K-002' verified ✓). Injection completion date = 2026-04-21 from COBB column L (matches Certify removal date). Recommendation: add explicit Batch ID + Injection Date fields to Ops Notes sheet. [log:cobb_tracker.cobb_cell_A_value] [log:cobb_tracker.rows_for_K-002.0.injection_completion_date]",
    "Ops Notes")
r = write_item(ws, r, "5.2",
    "Mass values correctly transcribed (Kerry mass IN, AECN mass IN, total injectate)",
    "Flag",
    "K-002 sheet 'Total Oil Injected' = 4,618 kg → Certify Mass of Injectant = 4,619 kg (1 kg rounding gap, immaterial). 'Total Transported' = 4,953 kg → Certify Allocated Kerry Oil Gross Mass = 4,953 kg ✓ exact. AECN mass IN: NOT recorded in K-002 sheet — the sheet does not document the homogenization-tank context (Kerry vol mixed with sparged AECN before injection). Total injectate mass (Kerry + sparged-AECN) is NOT computed here either. For Kerry-only credit accounting this is acceptable (Kerry credit math doesn't require AECN context), but the operational record is incomplete. FLAG (not FAIL) since the values used IN Certify reconcile correctly. [log:K-002_sheet_content.total_oil_injected_kg] [log:K-002_sheet_content.total_transported_kg] [log:certify_datapoints]",
    "Ops Notes vs mass artifacts")
r = write_item(ws, r, "5.3",
    "COBB values present and match (consumed_lot follows kc_NNNN; origin_qty, BOL_number align with Ops Notes)",
    "Fail",
    "ROW VERIFICATION: COBB cell A83='K-002', A84='K-002', A85='K-002', A86='K-002' — all 4 rows verified. consumed_lot pattern: 'KC_0005', 'KC_0008', 'KC_0007', 'KC_0006' — uppercase variant of skill spec's lowercase 'kc_NNNN' (acceptable). BOL_number column E is BLANK for all 4 K-002 rows (no BOL recorded in COBB). origin_quantity_kg in COBB: KC_0005 1185 ✓, KC_0007 1261 ✓, KC_0006 1284 ✓, KC_0008 1232 ≠ K-002 sheet's 1,223 (9 kg discrepancy on KC_0008 only). offloaded_quantity_kg blank for all 4 (vs AECN comparator 2-199 which has it populated). order_line_completed_on blank for all 4. Status FAIL due to: (a) missing BOL_number across all 4 rows (links back to Section 3.1 BOL absence), and (b) KC_0008 origin_quantity 9-kg discrepancy with K-002 sheet — Certify uses sheet value (4,953 kg total) not COBB sum (4,962 kg), so the discrepancy is a COBB-side audit-trail issue but should be reconciled. [log:cobb_tracker.rows_for_K-002] [log:cobb_tracker.cobb_vs_K002_sheet_discrepancy]",
    "COBB tracker vs K-002 sheet")
r = write_item(ws, r, "5.4",
    "Completion date present and matches COBB",
    "Pass",
    "COBB rows 83-86 column L 'injection_completion_date' = 2026-04-21 (uniform across all 4 K-002 rows). Matches Certify removal header date '21 Apr 2026' exactly. K-002 sheet itself has no completion-date field (see 5.1 FLAG). [log:cobb_tracker.rows_for_K-002.0.injection_completion_date]",
    "Ops Notes vs COBB")

r += 1
# SECTION 6
r = section_header(ws, r, "Section 6: Gross-to-Net Calculation")
r = col_header(ws, r)
r = write_item(ws, r, "6.1",
    "Gross CO₂e = MOBI (tonnes) × carbon wt% (decimal) × 3.667; no sparging deduction applied to Kerry mass (FC-K4)",
    "Pass",
    "Independent recompute: 4.619 t × 0.5601 × 3.667 = 9.488 tCO₂e (Certify Sequestration shows 9.488 tCO₂e exactly ✓). Confirmed: NO 1.29% sparging deduction applied to Kerry mass anywhere — Certify Calculation View shows Sequestration directly from Mass of Injectant × Carbon Content × 3.667 with no intermediate sparging step. FC-K4 PASS — Kerry enters un-sparged per VVB-validated waste boundary. [log:certify_components.sequestrations_total_tCO2e] [log:certify_per_component_uncertainty.injection_sequestration_inputs]",
    "GHG Stmt Sec E, Bio-oil Geol Storage v1.1")
r = write_item(ws, r, "6.2",
    "Process emissions correctly deducted (transport process + embodied; in-scope only — no upstream Kerry production component, FC-K5)",
    "Fail",
    "Components present: (a) Kerry to Basco Flatbed Process 0.747 tCO₂e, (b) Kerry to Basco Flatbed Embodied 0.070 tCO₂e. Per-batch site emissions (electricity/pumping/sitework) NOT in this removal — confirmed amortized at GHG-statement level via Project emissions section ('managed according to amortization rules; will show here once added to a GHG statement'). NO upstream Kerry production component (FC-K5 PASS). HOWEVER: status FAIL because Embodied EF used (0.0401 kgCO₂e/km) is from 'Truck and Trailer EE 2025.pdf' source — the LIVE 2026 Standard EF Sheet, Truck and Trailer EE 2026 tab, row 24 shows Non-Tanker Truck Embodied = 0.1197949629 kgCO₂e/km (~3× higher). KB confirms 2026 update for Tanker Truck (0.0490 → 0.1476 kgCO₂e/km on 3/30/26); by analogy Non-Tanker should be 0.1198 for 2026. K-002 RP is 2026-04-21, so 2026 EFs apply. Material under-reporting: should-be Embodied = 0.1198 × 1,090 × 1.609344 = 210.18 kg vs current 70.34 kg — gap of ~140 kg = 0.140 tCO₂e (~1.5% of Net CDR). [log:certify_per_component_uncertainty.transport_embodied_inputs.live_2026_sheet_check]",
    "Certify Calc View vs Standard EFs 2026")
r = write_item(ws, r, "6.3",
    "Uncertainty discount correctly applied (per Isometric Standard S3.7 Option B variance propagation)",
    "Pass",
    "Certify Calculation View shows Uncertainty Discount = 39 kg, deducted from net (Net = Sequestration 9,488 + Adjustments 747 − Activities 818 − Uncertainty 39 = 9,378 kg = 9.378 t ✓). Discount is present and the math reconciles. Per-component inputs verified — see Section 6a. [log:certify_components.net_co2e_tonnes] [log:certify_components.uncertainty_discount_kg]",
    "Certify Calc View, Isometric Standard v2.0 S3.7")

r += 1
# SECTION 6a
r = section_header(ws, r, "Section 6a: Per-Component Uncertainty Verification")
r = col_header(ws, r)
r = write_item(ws, r, "6a.1",
    "Layer 1 — Completeness: every Kerry component has uncertainty on correct datapoint",
    "Pass",
    "Per-component uncertainty inputs read directly from Certify component modals (mandatory verification sub-task). Sequestration: Carbon ±0.2241%, Mass of product ±4.54 kg. Transport Process: Mass of load ±4.54 kg. Transport Embodied: no Mass of load datapoint (correct — embodied is per-mile only, FC8). BCU: Constant 747 kg, no uncertainty input (correct for fixed-allocation BCU). All components requiring uncertainty have it on the correct datapoint. [log:certify_per_component_uncertainty]",
    "Certify component modals")
r = write_item(ws, r, "6a.2",
    "Layer 2 — Correctness: reported ± matches expected value",
    "Pass",
    "Carbon Content uncertainty: Certify ±0.2241%; Expected (per Charm convention 0.4% × Wt Avg) = 0.004 × 56.01 = 0.2240% → match within rounding ✓. Mass of Product uncertainty: Certify ±4.54 kg; Expected for tote mode (4 × 1.13 kg) = 4.52 kg → match (uses sheet's 1.13398 per-tote precision: 4 × 1.13398 = 4.53592 ≈ 4.54) ✓. Mass of Load (Process) uncertainty: same ±4.54 kg, same expected basis ✓. Per-weighing precision = 1.13 kg (Kerry totes weighed on same scale class as Charm WODO; not the AECN 72.57 kg rail-truck spec). [log:certify_per_component_uncertainty.injection_sequestration_inputs] [log:certify_per_component_uncertainty.transport_process_inputs.mass_of_load_uncertainty_status]",
    "Charm Uncertainty Reference, Isometric S3.7.3")
r = write_item(ws, r, "6a.3",
    "Layer 3 — Discount presence and consistency",
    "Pass",
    "Certify Calc View shows aggregate Uncertainty Discount = 39 kg deducted from net. Math: Net = 9,488 + 747 − 818 − 39 = 9,378 kg = 9.378 t ✓ (matches displayed header 9.38 t). Discount is present in net calculation, applied AFTER Sequestrations + Adjustments − Activities, and consistent with displayed total. [log:certify_components.uncertainty_discount_kg]",
    "Certify Calc View vs GHG Statement")

r += 1
# SECTION 7
r = section_header(ws, r, "Section 7: Data Integrity & Anomaly Check")
r = col_header(ws, r)
r = write_item(ws, r, "7.1",
    "No unexplained outliers in mass, carbon content, or calculated values",
    "Flag",
    "Mass: per-tote nets 1092–1209 kg (typical IBC tote range) ✓. Carbon: 55.0–57.5% (no Kerry baseline yet — see 4.2 FLAG). DENSITY ANOMALY: KC-0008 density per SGS = 1.035 g/cm³, vs other 3 totes 1.232–1.242 g/cm³. KC-0008 density is markedly low (~17% below others). Above the 1.01 g/cm³ floor for safe injection vs formation fluid, but unexplained within-batch divergence. Possible causes: lab measurement error, water contamination of the sample, or actual lighter oil in this tote. Surface to Garrett to confirm whether the value is real or a sampling/analysis artifact. Viscosity: 1632–3280 cSt, much higher than typical AECN (~22 cSt) — Kerry oil is markedly more viscous; not anomalous within Kerry. Acid number: 46–79 mg KOH/g (wide range but no historical Kerry baseline yet). [log:sgs_coa.anomalies] [log:sgs_coa.kerry_per_tote]",
    "GHG Stmt Sec C, VR Sec 3.5")
r = write_item(ws, r, "7.2",
    "Any spills/losses documented, unrecovered mass deducted, incident report filed",
    "N/A",
    "No spill/loss event documented in any source for K-002 (Drive folder, Certify Sources, COBB, K-002 sheet, SGS CoA — none reference an incident or spill). Standard tote-mode 'leftover' residue (~336 kg total = 6.8% of delivered) is documented in K-002 sheet and is NOT credited (Mass of Injectant 4,619 < Total Transported 4,953 — leftover is excluded from credit). No formal spill incident applies. [log:K-002_sheet_content.totes] [log:K-002_sheet_content.math_checks.sum_leftover_kg]",
    "Ops Notes, batch folder")

r += 1
# SECTION 8
r = section_header(ws, r, "Section 8: Isometric Certify — Removal Component Mapping")
r = col_header(ws, r)
r = write_item(ws, r, "8.1",
    "SEQUESTRATIONS: Gross injection sequestration matches internal gross",
    "Pass",
    "Certify Components → Sequestrations → Injection sequestration: 9.488 tCO₂e. Independent recompute: 4.619 t × 0.5601 × 3.667 = 9.488 tCO₂e ✓ exact match. Single Kerry sequestration entry; no commingling with AECN sequestration in this removal. [log:certify_components.sequestrations_total_tCO2e]",
    "Certify Components")
r = write_item(ws, r, "8.2",
    "ACTIVITIES — Transport: Process + embodied for flatbed Kerry → Basco; Mass of Load includes tote tares (FC-K2); 2026 EFs applied (FC8/embodied per-mile)",
    "Fail",
    "TWO simultaneous issues: (a) FC-K2: Process 'Mass of Load' = 4,953 kg = K-002 sheet's 'Allocated Kerry Oil Gross Mass' = sum of per-tote Gross weights interpreted as oil-only (oil + leftover residue). Excludes physical IBC tote tare (~60 kg/tote × 4 ≈ 240 kg) and any deck items. Per FC-K2, transport process EF basis should be BOL gross (oil + tote tares + deck) — current basis under-counts hauled load. (b) Embodied EF: Certify uses 0.0401 kgCO₂e/km from 'Truck and Trailer EE 2025.pdf'. Live 2026 Standard EF Sheet → Truck and Trailer EE 2026 tab → row 24 (Non-Tanker Truck Embodied, kgCO₂e/km) = 0.1197949629. KB confirms Tanker Embodied was updated for 2026 (0.0490 → 0.1476); by analogy Non-Tanker should be 0.1198. Material under-reporting: 0.0401 × 1090 × 1.609344 = 70.34 kg current vs should-be 0.1198 × 1090 × 1.609344 = 210.18 kg → ~140 kg gap = 0.140 tCO₂e ≈ 1.5% of Net CDR. Process EF native value 0.1384 kgCO₂e/(t·mi) verified against 2026 sheet F6 = 0.000138 MTCO₂e/(t·mi) ✓ (Process is current; only Embodied is stale). [log:certify_per_component_uncertainty.transport_process_inputs.FC_K2_check] [log:certify_per_component_uncertainty.transport_embodied_inputs.live_2026_sheet_check]",
    "Certify Components, Transport Emissions Acctg v1.1")
r = write_item(ws, r, "8.3",
    "ACTIVITIES — Injection: electricity, pumping, embodied present",
    "N/A",
    "Per-batch injection-site emissions are NOT present as components in this removal — by design. Certify 'Project establishment and end of life' section explicitly states: 'Project emissions which are set to an automated amortization strategy will show here once this removal is added to a GHG statement.' Site emissions (electricity, pumping, embodied) are amortized at GHG-statement level, not per-batch. Confirmed convention across recent Basco removals. [log:certify_components.no_injection_site_components_per_batch]",
    "Certify Components")
r = write_item(ws, r, "8.4",
    "NO upstream Kerry production components (waste boundary — FC-K5)",
    "Pass",
    "Certify Components tab reviewed end-to-end. NO 'Kerry Production Process Emissions', NO 'Kerry Production Embodied Emissions', NO 'Kerry Biomass Sourcing', NO pyrolysis or feedstock-tagged Kerry component. Only in-scope components present: Sequestration, Flatbed Process, Flatbed Embodied, BCU adjustment. System-boundary respected per VVB-validated waste-product methodology (FC-K5 ✓). [log:certify_components.no_upstream_kerry_production_components] [log:certify_components.no_kerry_biomass_sourcing_component]",
    "Certify Components, Biomass Feedstock Acctg v1.2")
r = write_item(ws, r, "8.5",
    "COUNTERFACTUALS: zero (waste boundary — FC-K6)",
    "Pass",
    "Certify Components → Removal counterfactuals → Total = 0 tCO₂e. No counterfactual sub-component present. Consistent with waste-boundary methodology (excludes upstream entirely; does NOT add an offsetting counterfactual). FC-K6 ✓. [log:certify_components.counterfactuals_tCO2e]",
    "Certify Components, Biomass Feedstock Acctg v1.2")
r = write_item(ws, r, "8.6",
    "REDUCTIONS: BCU Transport quantity ≥ transport process activity",
    "Pass",
    "Certify Components → Adjustments → BCUs - Applied to Bio-Oil Transport Process Emissions: 0.747 tCO₂e (= 747 kgCO₂e). Transport Process emission: 0.747 tCO₂e. Equality: 747 ≥ 747 ✓ (exactly offsetting; Process fully offset by BCU). Embodied emissions (0.070 tCO₂e) NOT BCU-offset (per protocol — embodied is not BCU-eligible) ✓. BCU certificate Charm-A3-2024-10-24-64-1976 has 70.68 MT total CO₂e reduction available; K-002 draws 0.747 MT (1.06% of cert capacity). [log:certify_components.adjustments] [log:bcu_certificate]",
    "Certify Components")
r = write_item(ws, r, "8.7",
    "Net removal matches internal calculation",
    "Pass",
    "Certify Calc View formula: Net = Sequestrations (9,488) + Adjustments (747) − Activities (818) − Uncertainty Discount (39) = 9,378 kg = 9.378 t. Header displays 9.38 tCO₂e ✓. Independent reconstruction: Sequestration 9.488 = 4.619 × 0.5601 × 3.667 ✓; Process 0.747 = 0.1384 × 1090 × 4.953 ✓; Embodied 0.070 = 0.0401 × 1090 × 1.609344 ✓ (against the EF Certify uses, though that EF is itself disputed in 8.2); BCU = Process exact ✓. Math reconciles end-to-end against the inputs Certify is using. [log:certify_components.net_co2e_tonnes] [log:certify_components.uncertainty_discount_kg]",
    "Certify Calc View")
r = write_item(ws, r, "8.8",
    "All Datapoints have descriptive source names (no 'Measurement XXXX' or 'IMG_001')",
    "Flag",
    "Source filenames in Certify Sources tab: 'K-002 (1).xlsx' and 'K-002.xlsx' (one is the duplicate (1) variant of the same K-002 sheet — both serve different datapoints but identical underlying file → mildly confusing for verifier), 'Standard Emission Factors + Calculations - Truck and Trailer EE 2025.pdf' (descriptive but year-suffix is wrong for a 2026 batch — see 8.2), '2-192-2-198+Kc0005-Kc-0008.pdf' (descriptive bundling), 'Kerry to Basco Distance.png' (descriptive), 'Charm-A3-2024-10-24-64-1976.pdf' (BCU certificate ID — descriptive), 'GLEC NA Road Transport.png' (descriptive). All names are descriptive (no generic 'Measurement XXXX' / 'IMG_001'). Two FLAG-worthy hygiene items: (a) duplicate K-002 / K-002 (1) filename pair, (b) 2025-suffix EF source on a 2026 RP batch. [log:certify_sources_enumeration]",
    "Certify Datapoints")
r = write_item(ws, r, "8.9",
    "All Sources match files in batch folder on Drive",
    "Flag",
    "Document Location Rule: documents may be in Drive OR Certify Sources. Of 8 Certify Sources: K-002.xlsx & K-002 (1).xlsx — Drive K-002 sheet present ✓. 2-192-2-198+Kc0005-Kc-0008.pdf — Drive copy exists at id 1v6kbqtjUi-r0O6jYA4gRXDKq9DjFrcnb (separate location, not this batch's folder). Standard EFs PDF, Kerry to Basco Distance.png, GLEC NA Road Transport.png, Charm-A3-...pdf — only in Certify Sources, not in Drive batch folder. Per skill spec, only-in-Certify is acceptable (not a FAIL). FLAG because batch folder is sparse: Drive K-002 folder contains only 3 items (K-001 sheet, K-002 sheet, prior-batch SGS PDF) — would be cleaner to mirror Certify Sources to Drive batch folder for verifier convenience. [log:drive_folder_contents] [log:certify_sources_enumeration]",
    "Certify Sources vs Drive")

# Final formatting + save
ws.freeze_panes = "A2"
wb.save(OUT)
print("Saved:", OUT)
print("Total rows:", r)
