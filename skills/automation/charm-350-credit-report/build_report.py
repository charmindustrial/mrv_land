#!/usr/bin/env python3
"""
Deterministic report builder for the Charm 350 bio-oil credit report.

Given the statement period + EVERY removal on the statement, already classified by the skill
(the classification itself needs judgment and is done by the agent from the Certify component
names), this writes the .xlsx to the Drive folder and prints the saved path + per-source
totals. It does NOT decide classification and does NOT send anything.

Since 2026-10-09 the report covers all three oil sources (Charm-produced / AECN / Kerry),
each in its own section with a subtotal, plus a grand total. JZ asked for the AECN and Kerry
numbers after the earlier Charm-only version.

Usage:
    build_report.py '<json payload>'

Payload shape:
{
  "statement_id": "ggs_...",
  "period_label": "01 Jun 2026 to 30 Jun 2026",
  "period_slug": "2026-06",
  "certify_status": "under verification",
  "removals": [
     {"removal_id": "rmv_...", "date": "20-27 Jun 2026", "oil_type": "QOWV", "tcdr": 5.19},
     {"removal_id": "rmv_...", "date": "20-27 Jun 2026", "oil_type": "AECN", "tcdr": 362.11},
     {"removal_id": "rmv_...", "date": "20-27 Jun 2026", "oil_type": "Kerry", "tcdr": 40.19},
     ...
  ]
}
oil_type must be one of the Charm keywords in config (QOWV / WODO / Aqueous), "AECN", or
"Kerry". Anything else fails loud. Every value must trace to what the agent read in Certify
this run. No defaults, no guesses.

The file is written in place (open(path, "wb")) — never removed, renamed or copied — so an
existing Drive file ID and any link already shared survive a rebuild.
"""
import io, json, os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(HERE, "config.json")


def load_config():
    if not os.path.exists(CONFIG_PATH):
        sys.exit(f"FAIL: config.json not found at {CONFIG_PATH}")
    with open(CONFIG_PATH) as f:
        cfg = json.load(f)
    for k in ("certify_removal_url_prefix", "drive_output_folder", "charm_oil_keywords",
              "aecn_keywords", "kerry_keywords"):
        if not cfg.get(k):
            sys.exit(f"FAIL: config.json missing required key '{k}'")
    return cfg


def main():
    if len(sys.argv) != 2:
        sys.exit("FAIL: expected one JSON payload argument")
    cfg = load_config()
    p = json.loads(sys.argv[1])

    for k in ("statement_id", "period_label", "period_slug", "certify_status", "removals"):
        if k not in p:
            sys.exit(f"FAIL: payload missing '{k}'")
    if not p["removals"]:
        sys.exit("FAIL: no removals passed — nothing to build (this is a guard, not an error you should force past)")

    charm_types = set(cfg["charm_oil_keywords"])
    sections = [
        ("Charm-produced oil (" + " / ".join(cfg["charm_oil_keywords"]) + ")", "Charm-produced", charm_types),
        ("AECN", "AECN", set(cfg["aecn_keywords"])),
        ("Kerry", "Kerry", set(cfg["kerry_keywords"])),
    ]
    allowed = set().union(*(s[2] for s in sections))
    bad = [r for r in p["removals"] if r.get("oil_type") not in allowed]
    if bad:
        sys.exit(f"FAIL: oil_type not in {sorted(allowed)} for: {[r.get('removal_id') for r in bad]}")

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    base = cfg["certify_removal_url_prefix"]
    outdir = cfg["drive_output_folder"]
    os.makedirs(outdir, exist_ok=True)

    HDR = PatternFill("solid", fgColor="D6E4F0"); TOT = PatternFill("solid", fgColor="EAF1FB")
    GRAND = PatternFill("solid", fgColor="C9DAF8")
    bold = Font(bold=True); link = Font(color="0563C1", underline="single")
    thin = Side(style="thin", color="BBBBBB"); border = Border(thin, thin, thin, thin)
    center = Alignment(horizontal="center", vertical="center"); left = Alignment(horizontal="left", vertical="center")

    wb = Workbook(); ws = wb.active; ws.title = "Bio-oil Credits"
    ws["A1"] = f"Bio-oil Credits Submitted — {p['period_label']}"; ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:D1")
    n = len(p["removals"])
    ws["A2"] = (f"GHG Statement {p['statement_id']} • Certify status: {p['certify_status']} • "
                f"all {n} removals, grouped by oil source")
    ws["A2"].font = Font(italic=True, size=9, color="666666"); ws.merge_cells("A2:D2")

    r = 4
    subtotal_cells = []
    totals = {}
    for title, short, types in sections:
        rows = [x for x in p["removals"] if x["oil_type"] in types]
        totals[short] = (len(rows), round(sum(float(x["tcdr"]) for x in rows), 2))
        ws.cell(r, 1, title).font = Font(bold=True, size=11)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); r += 1
        for c, h in enumerate(["Certify Removal ID", "Removal Date", "Oil Type", "Removal Quantity (tCDR)"], 1):
            cell = ws.cell(r, c, h); cell.font = bold; cell.fill = HDR; cell.border = border; cell.alignment = center
        r += 1; first = r
        for x in rows:
            a = ws.cell(r, 1, x["removal_id"]); a.hyperlink = base + x["removal_id"]; a.font = link; a.alignment = left; a.border = border
            b = ws.cell(r, 2, x["date"]); b.alignment = center; b.border = border
            c = ws.cell(r, 3, x["oil_type"]); c.alignment = center; c.border = border
            d = ws.cell(r, 4, float(x["tcdr"])); d.number_format = "0.00"; d.alignment = center; d.border = border
            r += 1
        if not rows:
            ws.cell(r, 1, "none this period").alignment = left; r += 1
        ws.cell(r, 1, f"{short} subtotal").font = bold
        for c in range(1, 5): ws.cell(r, c).fill = TOT; ws.cell(r, c).border = border
        sv = ws.cell(r, 4, f"=SUM(D{first}:D{r-1})" if rows else 0)
        sv.font = bold; sv.number_format = "0.00"; sv.alignment = center
        subtotal_cells.append(f"D{r}"); r += 2

    ws.cell(r, 1, "TOTAL — all removals on statement").font = bold
    for c in range(1, 5): ws.cell(r, c).fill = GRAND; ws.cell(r, c).border = border
    gv = ws.cell(r, 4, "=" + "+".join(subtotal_cells)); gv.font = bold; gv.number_format = "0.00"; gv.alignment = center

    for col, w in zip("ABCD", (30, 26, 10, 24)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    path = os.path.join(outdir, f"Charm_Bio-oil_Credits_{p['period_slug']}.xlsx")
    buf = io.BytesIO(); wb.save(buf)
    with open(path, "wb") as f:   # in place: preserves the Drive file ID on a rebuild
        f.write(buf.getvalue())

    grand = round(sum(t for _, t in totals.values()), 2)
    print(json.dumps({"saved_path": path, "n_rows": n, "by_source": totals, "total_tcdr": grand}))


if __name__ == "__main__":
    main()
