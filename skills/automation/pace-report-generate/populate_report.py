#!/usr/bin/env python3
"""
Populates a freshly-copied Pace Report xlsx for the new week:
  1. Writes the "Week of" date (B2) on the Max and Garrett tabs.
  2. Copies each person's "Top 3 I'll Focus On Next Week" items (B18:B20 on
     their tab) from the previous week's file into "Last Week's Focus - Did
     It Land" (B7:B9) on the same tab of the new file. Status/why columns are
     left blank -- those get filled in after the week happens.

Usage:
  python3 populate_report.py <previous_local_path> <new_local_path> <thursday_date_iso>

Does not touch the "Pace Report" tab -- its cells are live formulas that pull
from Max/Garrett B12:B14 (accomplished) and B18:B20 (next-week focus), and
none of those are written here.
"""
import sys
import openpyxl

PERSON_TABS = ["Max", "Garrett"]


def main():
    if len(sys.argv) != 4:
        print("FAILED: usage: populate_report.py <previous_local_path> <new_local_path> <thursday_date_iso>")
        sys.exit(1)

    previous_path, new_path, thursday_iso = sys.argv[1], sys.argv[2], sys.argv[3]

    prev_wb = openpyxl.load_workbook(previous_path, data_only=False)
    new_wb = openpyxl.load_workbook(new_path, data_only=False)

    copied_counts = {}
    for tab in PERSON_TABS:
        if tab not in prev_wb.sheetnames or tab not in new_wb.sheetnames:
            print(f"FAILED: tab '{tab}' missing from previous or new workbook")
            sys.exit(1)

        prev_ws = prev_wb[tab]
        new_ws = new_wb[tab]

        new_ws["B2"] = thursday_iso

        count = 0
        for i in range(3):
            src_cell = f"B{18 + i}"
            dst_cell = f"B{7 + i}"
            val = prev_ws[src_cell].value
            if val not in (None, ""):
                new_ws[dst_cell] = val
                count += 1
        copied_counts[tab] = count

    new_wb.save(new_path)

    # Read the file back from disk and confirm the write actually took. A
    # prior run silently produced a file where B2 still held the template's
    # literal "YYYY-MM-DD" placeholder despite this function believing it had
    # written the real date -- this check turns that into a loud FAILED
    # instead of a quietly wrong report shipped to the team.
    check_wb = openpyxl.load_workbook(new_path, data_only=False)
    for tab in PERSON_TABS:
        check_ws = check_wb[tab]
        actual_b2 = check_ws["B2"].value
        if actual_b2 != thursday_iso:
            print(
                f"FAILED: wrote B2={thursday_iso!r} on tab '{tab}' but reading "
                f"{new_path} back shows B2={actual_b2!r} -- write did not persist"
            )
            sys.exit(1)

    summary = ", ".join(f"{tab}: {n} item(s)" for tab, n in copied_counts.items())
    print(f"OK: populated {new_path} -- week of {thursday_iso} -- carried forward {summary}")


if __name__ == "__main__":
    main()
