#!/usr/bin/env python3
"""
Basco Thermal Oxidizer (ToX) — CatOx reporting package builder.

Produces the verified reporting package for a reporting period, matching the precedent
established in the May 2026 reporting:

  - "{Month} CatOx -- Raw.xlsx"        raw readings exactly as received (one sheet)
  - "{Month} CatOx -- Processed.xlsx"  verifier-facing: the tabs listed in the config's
                                       processed_tabs — one clean chart per train (inlet +
                                       outlet lines, setpoint line, grey offline shading)
                                       plus a Raw Data tab. Nothing else.

The interpretation/controls narrative lives in the static CatOx Engineering Brief
(Google Drive), NOT in these workbooks. Do NOT add Operational Log / Transition Events /
Combustion Signal / Notes tabs to the Processed workbook — those are TMI for a
verifier (confirmed by Max, July 2026).

Usage:
    # combined single-file export (May/June 2026 style)
    python3 generate_tox_charts.py --input <csv> --output <dir> [--month "June 2026"]
    # per-train split export (July 2026 onward): pass both files, either order
    python3 generate_tox_charts.py --input <train1.csv> <train2.csv> --output <dir> --month "July 2026"

Dependencies: pandas, numpy, matplotlib, openpyxl
"""

import argparse, json, sys
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import matplotlib.patches as mpatches
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Drift-prone facts are NOT hardcoded here — they live in the single-source-of-truth config.
# If the config is missing we STOP rather than fall back to a buried default (that fallback is
# exactly how a stale 650F literal survived once). The skill holds logic; the config holds facts.
_CONFIG_PATH = Path(__file__).resolve().parent.parent / 'references' / 'catox_config.json'
try:
    _CFG = json.loads(_CONFIG_PATH.read_text())
    THRESH_F = _CFG['outlet_setpoint_f']
    OFFLINE_SENTINELS = _CFG['offline_sentinels']
    LIM = _CFG['manual_limits']
    ANOM = _CFG['anomaly_checks']
except (FileNotFoundError, KeyError, json.JSONDecodeError) as e:
    sys.exit(f"ERROR: cannot load canonical config {_CONFIG_PATH} ({e}). "
             f"Refusing to run with hardcoded defaults — fix the config first.")


# ── Data loading ────────────────────────────────────────────────────────────
def _parse_time(s):
    """
    Parse the timestamp column. Exports have arrived in three shapes:
      'M/D/YY H:MM'            (May/June 2026 combined export)
      'YYYY-MM-DD HH:MM:SS.fff' (July/August 2026 per-train export)
      'M/D/YYYY H:MM'          (September 2026 combined export, 'timestamp_central')
    Try each declared shape, then fall back to pandas' own detection. Unparseable
    values become blank so the caller can drop those rows.
    """
    for fmt in ('%m/%d/%y %H:%M', '%Y-%m-%d %H:%M:%S.%f', '%m/%d/%Y %H:%M'):
        out = pd.to_datetime(s, format=fmt, errors='coerce')
        if out.notna().mean() > 0.9:
            return out
    return pd.to_datetime(s, errors='coerce')


def _merge_per_train(paths):
    """
    Combine a per-train split export (July 2026 onward) into the same 5-column table the
    combined export produces. Each file is timestamp + that train's inlet + outlet.

    Files are assigned to trains by what their own column headers say ('Train1'/'2100' vs
    'Train2'/'2200'), NOT by the order they were passed, and the merge is an inner join on
    the timestamp so a reading is only kept when both trains report the same instant.
    Values are never altered — only placed side by side.
    """
    frames = {}
    for p in paths:
        d = pd.read_csv(p)
        if len(d.columns) != 3:
            sys.exit(f"ERROR: {Path(p).name} has {len(d.columns)} columns; a per-train export "
                     f"must be timestamp + inlet + outlet.")
        blob = ' '.join(d.columns).lower()
        t1 = ('train1' in blob.replace(' ', '') or '2100' in blob)
        t2 = ('train2' in blob.replace(' ', '') or '2200' in blob)
        if t1 == t2:
            sys.exit(f"ERROR: cannot tell which train {Path(p).name} is from its headers: "
                     f"{list(d.columns)}")
        n = 1 if t1 else 2
        if n in frames:
            sys.exit(f"ERROR: two files both look like Train {n}. Pass one file per train.")
        ic, oc = d.columns[1], d.columns[2]
        if 'inlet' not in ic.lower() or 'outlet' not in oc.lower():
            sys.exit(f"ERROR: {Path(p).name} columns are not inlet-then-outlet: {list(d.columns)}")
        frames[n] = d.rename(columns={d.columns[0]: 't_stamp'})
        print(f"  Train {n}: {Path(p).name} — {len(d)} rows")

    if set(frames) != {1, 2}:
        sys.exit(f"ERROR: need one Train 1 file and one Train 2 file; got trains {sorted(frames)}.")

    merged = frames[1].merge(frames[2], on='t_stamp', how='inner')
    dropped = max(len(frames[1]), len(frames[2])) - len(merged)
    if dropped:
        print(f"  NOTE: {dropped} reading(s) present for only one train were dropped "
              f"(timestamps must match to pair the trains).")
    return merged


def load_data(paths):
    """
    Load the temperature export(s) and return (df, raw_df, cols).

    Accepts either:
      - ONE combined file (May/June 2026 style):
        t_stamp, Train1 InletTemp, Train1 OutletTemp, Train2 InletTemp, Train2 OutletTemp
      - TWO per-train files (July 2026 onward): each timestamp + inlet + outlet, merged here
        on the timestamp into the same 5-column table. Train identity comes from the headers.

    5-minute readings, 0 marks offline/no-reading, columns correctly labeled (NO swap).

    Also tolerates the legacy combined format (7 columns with sparge binaries, -238 sentinel,
    hourly, and the April-era Train 2 header swap) as a fallback.

    Orientation is ALWAYS verified empirically after loading, and the code warns loudly if
    Train 2 looks reversed rather than blindly trusting any historical swap rule.
    """
    if isinstance(paths, (str, Path)):
        paths = [paths]
    if len(paths) == 2:
        raw = _merge_per_train(paths)
    elif len(paths) == 1:
        raw = pd.read_csv(paths[0])
    else:
        sys.exit(f"ERROR: --input takes one combined file or two per-train files; got {len(paths)}.")

    cols = list(raw.columns)
    df = raw.copy()

    if len(cols) >= 5 and any('inlet' in c.lower() for c in cols) and any('outlet' in c.lower() for c in cols):
        # Current format — positional map, no swap, 0 = offline
        df = df.rename(columns={cols[0]: 'time', cols[1]: 'T1_in', cols[2]: 'T1_out',
                                cols[3]: 'T2_in', cols[4]: 'T2_out'})
    elif len(cols) == 7:
        # Legacy combined format — apply the April-era T2 header swap
        df = df.rename(columns={cols[0]: 'time', cols[1]: 'sparge1', cols[2]: 'T1_in',
                                cols[3]: 'T1_out', cols[4]: 'sparge2',
                                cols[5]: 'T2_out', cols[6]: 'T2_in'})  # note swap: idx5=out, idx6=in
    else:
        sys.exit(f"ERROR: unrecognized column layout ({len(cols)} cols): {cols}")

    df['time'] = _parse_time(df['time'])
    df = df.dropna(subset=['time']).sort_values('time').reset_index(drop=True)
    # Two kinds of "no reading" arrive: an offline sentinel (0 / -238) and, since September
    # 2026, empty cells. Both become blank, but they are counted apart for the chat report —
    # sentinel stretches are the unit offline, scattered empties are dropped single readings.
    for c in ['T1_in', 'T1_out', 'T2_in', 'T2_out']:
        v = pd.to_numeric(df[c], errors='coerce')
        n_empty, n_sent = int(v.isna().sum()), int(v.isin(OFFLINE_SENTINELS).sum())
        print(f"  {c}: {n_sent} offline-sentinel reading(s), {n_empty} empty cell(s)")
        # Kept so the anomaly checks can tell an offline marker from a dropped (empty) reading.
        df[f'{c}_offline'] = v.isin(OFFLINE_SENTINELS)
        df[c] = v.replace(OFFLINE_SENTINELS, np.nan)

    # ── Orientation self-check (never trust a historical swap blindly) ──
    for n, ic, oc in [(1, 'T1_in', 'T1_out'), (2, 'T2_in', 'T2_out')]:
        m = (df[ic] > 100) & (df[oc] > 100)
        if m.sum() > 50:
            frac = (df.loc[m, oc] > df.loc[m, ic]).mean()
            if frac < 0.5:
                print(f"  *** WARNING: Train {n} outlet reads COLDER than inlet on "
                      f"{(1-frac)*100:.0f}% of operating rows — columns may be reversed. "
                      f"Verify orientation before delivering. ***")
            else:
                print(f"  Train {n} orientation OK (outlet hotter than inlet on {frac*100:.0f}% of operating rows).")

    return df, raw, cols


def report_stuck_readings(df, min_readings=12):
    """
    Flag stretches where a channel repeats one value exactly, reading after reading.

    Short repeats are normal — the historian only logs a change once it exceeds a deadband.
    What matters is a repeat spanning ALL FOUR channels on BOTH trains starting and ending at
    the same instant: two independent oxidizer trains cannot physically hold four temperatures
    constant to the degree, so that pattern is the logger holding last-good values, and those
    rows are not live readings. It matters for the numbers because a value frozen ABOVE the
    setpoint counts as at-temperature and inflates the at-temperature share.

    Printed as a diagnostic for the chat report only. Nothing here goes in the workbook.
    Returns the list of all-channel frozen blocks as (start, end, n_readings).
    """
    chans = ['T1_in', 'T1_out', 'T2_in', 'T2_out']
    print("\n── Stuck-reading check (diagnostic — not in the workbook) ──")

    for c in chans:
        s = df[c]
        run = (s != s.shift()).cumsum()
        g = df.groupby(run).agg(val=(c, 'first'), n=(c, 'size'))
        g = g[(g.n >= min_readings) & g.val.notna()]
        if len(g):
            print(f"  {c}: {len(g)} run(s) of >={min_readings*5/60:.0f}h holding one value "
                  f"({g.n.sum()} readings, {g.n.sum()/len(df)*100:.1f}% of period); "
                  f"longest {g.n.max()*5/60:.1f}h at {g.loc[g.n.idxmax(),'val']:.0f}°F")
        else:
            print(f"  {c}: no held-value run >= {min_readings*5/60:.0f}h")

    # Blocks where every channel is simultaneously repeating its previous value.
    allfrozen = pd.concat([df[c] == df[c].shift() for c in chans], axis=1).all(axis=1)
    blocks = []
    for _, d in df.groupby((allfrozen != allfrozen.shift()).cumsum()):
        if allfrozen.loc[d.index[0]] and len(d) >= min_readings:
            blocks.append((d['time'].iloc[0], d['time'].iloc[-1], len(d)))

    report_train_holds(df, min_readings)

    if not blocks:
        print("  No all-channel frozen block found. Readings track independently — good.")
        return blocks

    print(f"\n  *** {len(blocks)} ALL-CHANNEL FROZEN BLOCK(S) — treat as logger hold, not readings ***")
    for st, en, n in blocks:
        held = {c: df.loc[df['time'] == st, c].iloc[0] for c in chans}
        above = [c for c in ('T1_out', 'T2_out') if held[c] >= THRESH_F]
        print(f"    {n} readings ({n*5/60:.1f}h)  {st:%b %d %H:%M} -> {en:%b %d %H:%M}")
        print(f"      held at: " + ", ".join(f"{c} {held[c]:.0f}°F" for c in chans))
        if above:
            print(f"      NOTE: {' and '.join(above)} frozen AT/ABOVE {THRESH_F}°F — these rows "
                  f"inflate the at-temperature share. Report the stat both ways.")

    keep = ~pd.concat([(df['time'] >= s) & (df['time'] <= e) for s, e, _ in blocks],
                      axis=1).any(axis=1)
    _print_share(df, keep)
    return blocks


def _print_share(df, keep):
    print(f"\n  Excluding held readings — {keep.sum()} of {len(df)} readings remain:")
    for n, oc in [(1, 'T1_out'), (2, 'T2_out')]:
        s = df.loc[keep, oc]
        at = (s >= THRESH_F).sum(); bl = ((s < THRESH_F) & s.notna()).sum(); off = s.isna().sum()
        print(f"    Train {n}: at/above {THRESH_F}°F {at/keep.sum()*100:.1f}% | "
              f"below {bl/keep.sum()*100:.1f}% | offline {off/keep.sum()*100:.1f}%")


def report_train_holds(df, min_readings=12):
    """
    One-train version of the frozen check (added September 2026, when a Train 2 hold at
    600/694°F ran 9.8 h and the all-channel check missed it because Train 1 kept moving).

    Flags stretches where BOTH sensors of one train repeat their previous value together for
    >= min_readings with the outlet held at/above the setpoint — the case that inflates the
    at-temperature share. Dropouts of up to 3 readings inside a stretch are bridged; longer gaps
    (offline stretches) end it. Prints the jump on
    release: a multi-hundred-degree move in one 5-minute step is not physical cooling, so the
    held values were stale. No release threshold is applied here; read the steps and judge.
    Diagnostic for the chat report only.
    """
    print("\n  One-train holds with outlet at/above setpoint (both sensors frozen together):")
    held = pd.Series(False, index=df.index)
    for n, ic, oc in [(1, 'T1_in', 'T1_out'), (2, 'T2_in', 'T2_out')]:
        x = df[[ic, oc]].ffill(limit=3)  # bridge short dropouts only, never an offline stretch
        fr = (x[ic] == x[ic].shift()) & (x[oc] == x[oc].shift())
        for _, g in df[fr].groupby((fr != fr.shift()).cumsum()[fr]):
            i0, i1 = g.index[0] - 1, g.index[-1]
            if len(g) < min_readings or not x.loc[i0, oc] >= THRESH_F:
                continue
            held.loc[i0:i1] = True
            after = (f"{df.loc[i1+1, ic]:.0f}/{df.loc[i1+1, oc]:.0f}" if i1 + 1 in df.index else "end of file")
            print(f"    Train {n}: {df.time[i0]:%b %d %H:%M} -> {df.time[i1]:%b %d %H:%M} "
                  f"({(len(g)+1)*5/60:.1f}h) held {x.loc[i0, ic]:.0f}/{x.loc[i0, oc]:.0f}°F "
                  f"(inlet/outlet), next reading {after}")
    if held.any():
        _print_share(df, ~held)
    else:
        print("    none")


# ── Anomaly checks (added October 2026) — diagnostics for the chat report only ──
# Limits and check settings come from catox_config.json (manual_limits, anomaly_checks).
_MAX_LINES = 20  # per list, so a flickering sensor cannot bury the report


def _episodes(mask):
    """Consecutive True rows of a boolean Series, as (first_index, last_index) pairs."""
    return [(g.index[0], g.index[-1])
            for _, g in mask[mask].groupby((mask != mask.shift()).cumsum()[mask])]


def _step(df):
    """The export's normal reading interval (the most common gap between readings)."""
    d = df['time'].diff()
    return d[d > pd.Timedelta(0)].mode().iloc[0]


def _print_capped(lines):
    for ln in lines[:_MAX_LINES]:
        print(ln)
    if len(lines) > _MAX_LINES:
        print(f"      … and {len(lines) - _MAX_LINES} more")


def report_time_coverage(df, month):
    """Missing readings, repeated timestamps, and whether the file covers the whole month."""
    print("\n  Time coverage:")
    t, step = df['time'], _step(df)
    print(f"    Reading interval {step.total_seconds()/60:.0f} min")

    dup = t[t.duplicated(keep=False)]
    if len(dup):
        print(f"    *** {dup.nunique()} timestamp(s) appear more than once ({len(dup)} rows) ***")
        if (dup.dt.month == 11).all() and (dup.dt.hour == 1).all() and dup.dt.date.nunique() == 1:
            print("      All fall in one November 1 am hour — matches the clock change back to "
                  "standard time (Central time repeats that hour).")
        _print_capped([f"      {v:%b %d %H:%M}" for v in dup.drop_duplicates()])
    else:
        print("    No repeated timestamps")

    gaps = [(t.iloc[i - 1], t.iloc[i]) for i in np.where(t.diff() > step * 1.5)[0]]
    if gaps:
        n_missing = sum(round((b - a) / step) - 1 for a, b in gaps)
        print(f"    *** {len(gaps)} gap(s) in the timestamps, ~{n_missing} reading(s) missing ***")
        _print_capped([f"      {a:%b %d %H:%M} -> {b:%b %d %H:%M}  ({(b - a).total_seconds()/3600:.1f} h)"
                       for a, b in gaps])
    else:
        print("    No gaps in the timestamps")

    try:
        start = pd.Timestamp(month)
    except ValueError:
        return
    end = start + pd.offsets.MonthBegin(1)
    if t.iloc[0] > start + step:
        print(f"    *** File starts {t.iloc[0]:%b %d %H:%M}, after the period start {start:%b %d %H:%M} ***")
    if t.iloc[-1] < end - 2 * step:
        print(f"    *** File ends {t.iloc[-1]:%b %d %H:%M}, before the period end {end:%b %d %H:%M} ***")
    outside = ((t < start) | (t >= end)).sum()
    if outside:
        print(f"    *** {outside} reading(s) fall outside {month} ***")


def report_isolated_offline(df):
    """
    Offline markers (0 / -238) inside a run: a short offline stretch with the same sensor
    reading hot on both sides. A real shutdown leaves the unit cold; a 0 between two hot
    readings is a dropped reading the export labeled offline, and the chart shades it grey.
    """
    k, floor = ANOM['isolated_offline_max_readings'], ANOM['running_floor_f']
    print(f"\n  Offline markers inside a run (<= {k} reading(s), sensor above {floor}°F on both sides):")
    lines = []
    for c in ['T1_in', 'T1_out', 'T2_in', 'T2_out']:
        for i0, i1 in _episodes(df[f'{c}_offline']):
            if i1 - i0 + 1 > k or i0 == 0 or i1 + 1 not in df.index:
                continue
            before, after = df.loc[i0 - 1, c], df.loc[i1 + 1, c]
            if before > floor and after > floor:
                lines.append(f"    {c}: {df.time[i0]:%b %d %H:%M} ({i1 - i0 + 1} reading(s)), "
                             f"{before:.0f}°F before, {after:.0f}°F after")
    if lines:
        print(f"    *** {len(lines)} found ***")
        _print_capped(lines)
    else:
        print("    none")


def report_limit_exceedances(df):
    """Readings above the manufacturer's alarm and warranty limits (O&M Guide N-24-2702A)."""
    print("\n  Readings above the manufacturer's limits (O&M Guide N-24-2702A §6.9, §4.3):")
    checks = [(c, 'inlet alarm', LIM['inlet_alarm_f']) for c in ('T1_in', 'T2_in')] + \
             [(c, lab, LIM[k]) for c in ('T1_out', 'T2_out')
              for lab, k in (('outlet alarm', 'outlet_alarm_f'), ('outlet never-exceed', 'outlet_never_exceed_f'))]
    lines = []
    for c, lab, lim in checks:
        for i0, i1 in _episodes(df[c] > lim):
            lines.append(f"    {c} above {lab} {lim}°F: {df.time[i0]:%b %d %H:%M} -> {df.time[i1]:%b %d %H:%M} "
                         f"({i1 - i0 + 1} reading(s)), peak {df.loc[i0:i1, c].max():.0f}°F")

    # §4.3: no more than N minutes above the continuous limit in any 24 hours (rolling window).
    lim, cap, step_min = (LIM['outlet_continuous_max_f'], LIM['outlet_continuous_max_minutes_per_24h'],
                          _step(df).total_seconds() / 60)
    for c in ('T1_out', 'T2_out'):
        hot = pd.Series((df[c] > lim).astype(float).values, index=df['time'])
        mins = pd.Series(hot.rolling('24h').sum().values * step_min, index=df.index)
        for i0, i1 in _episodes(mins > cap):
            lines.append(f"    {c} above {lim}°F for more than {cap} min in 24 h: from {df.time[i0]:%b %d %H:%M}, "
                         f"up to {mins.loc[i0:i1].max():.0f} min in the 24 h to {df.time[mins.loc[i0:i1].idxmax()]:%b %d %H:%M}")
    if lines:
        print(f"    *** {len(lines)} found ***")
        _print_capped(lines)
    else:
        print("    none")


def report_sudden_drops(df):
    """
    Possible trips: an outlet at/above trip_from_f that falls by trip_drop_f or more within
    trip_window_min. The temperature record cannot say WHY (thresholds.md: do not call any of
    these an E-stop without the control-system log). A drop that follows repeated identical
    readings may be the logger releasing a hold, not a real fall — flagged so it gets checked.
    """
    frm, drop, win = ANOM['trip_from_f'], ANOM['trip_drop_f'], pd.Timedelta(minutes=ANOM['trip_window_min'])
    n_win = max(1, int(win / _step(df)))
    print(f"\n  Sudden outlet drops (from >= {frm}°F, down >= {drop}°F within {win.total_seconds()/60:.0f} min):")
    for n, oc in [(1, 'T1_out'), (2, 'T2_out')]:
        s, t, lines, last_j = df[oc], df['time'], [], -1
        for i in np.where(s >= frm)[0]:
            if i <= last_j:
                continue  # still inside the previous fall
            j = next((j for j in range(i + 1, min(i + n_win, len(df) - 1) + 1)
                      if s.iloc[j] <= s.iloc[i] - drop and t.iloc[j] - t.iloc[i] <= win), None)
            if j is None:
                continue
            last_j = j
            held = i >= 2 and s.iloc[i - 2] == s.iloc[i - 1] == s.iloc[i]
            lines.append(f"    {t.iloc[i]:%b %d %H:%M}  {s.iloc[i]:.0f} -> {s.iloc[j]:.0f}°F "
                         f"in {(t.iloc[j] - t.iloc[i]).total_seconds()/60:.0f} min"
                         + ("  (follows repeated identical readings — check for a logger hold)" if held else ""))
        print(f"    Train {n}: {len(lines)} found")
        _print_capped(lines)


def report_anomalies(df, month):
    print("\n── Anomaly checks (diagnostic — not in the workbook) ──")
    report_time_coverage(df, month)
    report_isolated_offline(df)
    report_limit_exceedances(df)
    report_sudden_drops(df)


# ── Processed chart (matches the May precedent style) ─────────────────────────
def make_chart(df, n, inc, outc, month, path):
    fig, ax = plt.subplots(figsize=(16, 6.5))
    off = df[outc].isna(); ins = False; sx = None
    for t, v in zip(df['time'], off):
        if v and not ins:
            ins, sx = True, t
        elif not v and ins:
            ins = False
            ax.axvspan(sx, t, color='#e6e6e6', zorder=0)
    if ins and sx is not None:
        ax.axvspan(sx, df['time'].iloc[-1], color='#e6e6e6', zorder=0)

    ax.plot(df['time'], df[inc], color='#9ecae1', lw=1.1, label='Catalyst Inlet (°F)', zorder=3)
    ax.plot(df['time'], df[outc], color='#1f4e79', lw=1.3, label='Catalyst Outlet (°F)', zorder=3)
    ax.axhline(THRESH_F, color='darkorange', lw=2.0, ls='--', zorder=2,
               label=f'{THRESH_F}°F minimum setpoint')
    ax.set_title(f'Train {n} — Catalyst Temperatures — {month}', fontsize=15, fontweight='bold')
    ax.set_ylabel('Temperature (°F)', fontsize=12, fontweight='bold')
    ax.set_xlabel(f'Date ({month})', fontsize=11)
    ax.set_ylim(-50, 1250)
    ax.grid(axis='y', alpha=0.3)
    off_patch = mpatches.Patch(color='#e6e6e6', label='Offline')
    h, _ = ax.get_legend_handles_labels()
    ax.legend(handles=h + [off_patch], loc='upper right', fontsize=10, framealpha=0.95)
    span_days = (df['time'].iloc[-1] - df['time'].iloc[0]).days
    ax.xaxis.set_major_formatter(mdates.DateFormatter('%b %d'))
    ax.xaxis.set_major_locator(mdates.DayLocator(interval=max(1, span_days // 10)))
    plt.setp(ax.get_xticklabels(), rotation=0, ha='center')
    plt.tight_layout()
    plt.savefig(path, dpi=150, bbox_inches='tight')
    plt.close()
    print(f"  Saved chart: {path}")
    return path


def _write_raw_sheet(ws, raw, cols):
    """Populate a worksheet with the raw readings exactly as received."""
    for j, h in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=str(h))
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D6E4F0')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(j)].width = 20 if j == 1 else 28
    for i, row in raw.iterrows():
        for j, h in enumerate(cols, 1):
            # Empty cells in the export arrive as NaN; write them as truly empty, not as a
            # blank number (which Excel can flag as damaged content).
            v = row[h]
            ws.cell(row=i + 2, column=j, value=None if pd.isna(v) else v)
    ws.freeze_panes = 'A2'


def build_processed(df, raw, cols, month, out_dir):
    # Two chart tabs + a Raw Data tab. Tab set is declared in the config (processed_tabs);
    # we assert the built workbook matches it so code and config cannot drift.
    c1 = make_chart(df, 1, 'T1_in', 'T1_out', month, str(out_dir / 'catox_T1.png'))
    c2 = make_chart(df, 2, 'T2_in', 'T2_out', month, str(out_dir / 'catox_T2.png'))
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for name, img in [('Train 1 — Temperature', c1), ('Train 2 — Temperature', c2)]:
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        im = XLImage(img); im.anchor = 'A1'; ws.add_image(im)
    _write_raw_sheet(wb.create_sheet('Raw Data'), raw, cols)
    expected = _CFG['processed_tabs']
    if wb.sheetnames != expected:
        sys.exit(f"ERROR: Processed tabs {wb.sheetnames} != config processed_tabs {expected}. "
                 f"Reconcile the script and catox_config.json before delivering.")
    path = out_dir / f'{month.split()[0]} CatOx -- Processed.xlsx'
    wb.save(path)
    return path


def build_raw(raw, cols, month, out_dir):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Raw'
    _write_raw_sheet(ws, raw, cols)
    path = out_dir / f'{month.split()[0]} CatOx -- Raw.xlsx'
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True, nargs='+',
                    help='one combined export CSV, or two per-train export CSVs')
    ap.add_argument('--output', required=True)
    ap.add_argument('--month', help='e.g. "June 2026"; derived from data if omitted')
    a = ap.parse_args()
    out_dir = Path(a.output); out_dir.mkdir(parents=True, exist_ok=True)

    print(f"\nSetpoint {THRESH_F}°F  (source: {_CFG['setpoint_source']})")
    print(f"Loading: {', '.join(a.input)}")
    df, raw, cols = load_data(a.input)
    month = a.month or df['time'].dt.strftime('%B %Y').iloc[0]
    print(f"  {len(df)} rows, {df['time'].iloc[0]:%b %d %Y %H:%M} – {df['time'].iloc[-1]:%b %d %Y %H:%M}  ({month})")

    proc = build_processed(df, raw, cols, month, out_dir)
    rawx = build_raw(raw, cols, month, out_dir)
    print(f"\nProcessed (verifier-facing): {proc}")
    print(f"Raw:                          {rawx}")

    # Summary for the chat report to Max (NOT written into any workbook)
    total = len(df)
    print("\n── Summary (for chat report — not in the workbook) ──")
    for n, oc in [(1, 'T1_out'), (2, 'T2_out')]:
        s = df[oc]
        at = (s >= THRESH_F).sum(); below = ((s < THRESH_F) & s.notna()).sum(); off = s.isna().sum()
        peak = s.max()
        print(f"  Train {n}: at/above {THRESH_F}°F {at/total*100:.1f}% | below {below/total*100:.1f}% | "
              f"offline {off/total*100:.1f}% | peak {peak:.0f}°F")

    report_stuck_readings(df)
    report_anomalies(df, month)
    print(f"\n  open \"{proc}\"")


if __name__ == '__main__':
    main()
