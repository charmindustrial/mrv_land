import re, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ---- build corpus from the SAME sources the Stop hook reloads ----
FILES = ["knowledge-base/biochar_protocol_v1.2.txt", "knowledge-base/soil_module_v1.3.txt", "knowledge-base/pdd_full.txt"]
def normalize_punct(t):
    return (t.replace('‘',"'").replace('’',"'").replace('“','"').replace('”','"')
             .replace('–','-').replace('—','-').replace('…','...'))
raw = "\n".join(open(f, encoding='utf-8', errors='replace').read() for f in FILES)
raw = normalize_punct(raw)
norm_corpus = re.sub(r'\s+', ' ', raw).strip()

SIGNAL_RE = re.compile(r'(?:§\s*\d[\d.]*[a-z]?|[Ss]ection\s+\d[\d.]*|[Cc]lause\s+\d[\d.]*|[Aa]rticle\s+\d[\d.]*|[Tt]able\s+[A-Z]?\d+)')
def in_corpus(q):
    return re.sub(r'\s+',' ',normalize_punct(q)).strip() in norm_corpus

# ---- items: ('S', title) section | ('I', item, status_notes_ref_quote...) ----
# I rows: (text, ref, quote)   quote='' allowed only if ref has no citation signal
S=[
 ('S','1. Delivery & Removal Identification'),
 ('I','Delivery folder contains the calculation workbook plus all supporting evidence (scale tickets, pickup/BOL, scale certification, end-use evidence)','Protocol §8.4.1',
   'The Project Proponent must maintain the following records as evidence of gross'),
 ('I','Recipient/offtaker name populated in the workbook (no "[Offtaker]" placeholder) and matches the folder name','House rule',''),
 ('I','Delivery date consistent across workbook, folder name, scale ticket and pickup/BOL documents','House rule',''),
 ('I','Each production period (PP) in the delivery is identified, and CO2e-stored is computed per PP','Protocol §8.3.1.1',
   'Where biochar Production Batches are not blended prior to storage'),
 ('I','Removal granularity (per production batch vs blended storage batch) follows the batch-definition approach justified in the PDD','Protocol §8.3.1; PDD v1.12',
   'Project Proponents must justify their intended approach for defining a production batch in the PDD'),
 ('I','Production batch defined as <= 1 month and applied consistently','Protocol §8.3.1',
   'The definition of a batch is project specific but must be a period of less than one month'),
 ('I','Durability option (200-yr Option 1) matches the validated PDD','PDD v1.12',''),
 ('I','Production-side data per PP traceable to the PP summary verified under bio-oil crediting','PDD v1.12',''),

 ('S','2. Mass Determination'),
 ('I','Scale tickets (gross / tare / net) present and legible for each weighed load','Protocol §8.4.1',
   'Weigh scale tickets showing arrival and departure weights for each biochar delivery'),
 ('I','Weighing used a certified legal-for-trade scale (NIST HB 44); scale certification document present','Protocol §8.4',
   'All truck scales must maintain current certification under applicable local, state, or federal legal-for-trade regulations'),
 ('I','Mass determined by difference between loaded and unloaded weights; bags/pallets/added-water backed out','Protocol §8.4',
   'Accurate mass determination via difference between loaded and unloaded truck weights'),
 ('I','lbs to kg conversion (x 0.453592) applied and shown explicitly','House rule',''),
 ('I','Mass used for CO2e-stored is the dry biochar mass (bone-dry, less ambient moisture + unprocessed material)','Module §5.1.1',
   'is the dry mass of biochar applied'),
 ('I','Added water reconciles with the wet-to-30%-moisture delivery spec','PDD v1.12',''),
 ('I','Multi-PP load: mass split across PPs so each PP stored mass is correct','Protocol §8.3.1',
   'multiple production batches'),
 ('I','Missing/estimated scale data handled conservatively; data-loss ladder followed (re-take, then conservatively estimate, then forego)','PDD v1.12',''),
 ('I','Any biochar spills/losses quantified and deducted from delivered mass','Protocol §8.4.1.2',
   'Deduct the quantified loss from the delivered biochar amount recorded on delivery weigh tickets'),

 ('S','3. Carbon Content & Characterization (CO2e-stored inputs)'),
 ('I','Total Carbon Content measured per PP from 3 replicates (Method A); Total C via ASTM D5373','Protocol §8.3.3; PDD v1.12',
   'Total C content measured via ASTM D5373'),
 ('I','Inorganic carbon measured (ASTM D4373); only organic carbon (Total C - Cinorg) credited','Module §5.1.2',
   'This Module only Credits for the durably stored'),
 ('I','Moisture content measured per PP (ASTM D1762-84); wet-sent lab samples handled so dry vs moisture-adjusted masses reconcile','Module Table 2; PDD v1.12',
   'Moisture content is reported for every production batch'),
 ('I','Molar H/Corg < 0.5 for each PP (3 replicates; H via ASTM D5373, H:Corg by calculation)','Module Table 2',
   'For the 200 year crediting option, this is used to model biochar durability'),
 ('I','Molar O/Corg < 0.2 (O by difference; reported once at validation)','Module Table 2',
   'C-O bonds are more labile than C-C bonds'),
 ('I','Sampling follows Method A with >= 3 independent replicates per PP','Protocol §8.3.3; PDD v1.12',
   'Charm will submit 3 replicates per Production Period for analysis'),
 ('I','Analytical labs accredited: Soil Control Labs, Eurofins (ISO 17025), BSK (NELAP)','Module §4.1.2; PDD v1.12',
   'EuroFins holds ISO 17025 accreditations for all analytes'),
 ('I','All credited carbon & moisture results reported on a dry basis','Module §3.3',
   'All results pertaining to the calculation of carbon removal must be reported on a dry basis'),

 ('S','4. CO2e Stored Calculation (per PP removal)'),
 ('I','CO2e-stored = Cbiochar x m(dry biochar) x Fdurable x 44.01/12.01','Module §5.1.1',
   'is the fraction of durable biochar that remains in the soil for the full duration of the crediting timeline'),
 ('I','Cbiochar = measured Corg for that PP; mass = Dry Biochar mass (bone-dry)','Module §5.1.2',
   'Only organic carbon is credited for under this Protocol and Module'),
 ('I','Fdurable(200) per Woolf model: min(0.95, 1 - [c + (a + b ln Tsoil) H/Corg]); a=-0.383, b=0.350, c=-0.048','Module §5.1.2.2.1',
   'is based on the modelling approach set out by Woolf et al'),
 ('I','Tsoil from Lembrechts/equivalent or direct measurement; >= 7C floor; air temperature NOT used as proxy','Module §5.1.2.2.1; PDD v1.12',
   'Air temperatures must not be used as a proxy for average annual soil temperatures'),
 ('I','Soil-temp variation > 1C within the application boundary handled (subdivide or use highest)','Module §5.1.2.2.1',
   'If the soil temperature variation within a project boundary exceeds'),
 ('I','Stoichiometric ratio 44.01/12.01 applied; result ties to workbook Gross CO2e Sequestered','Module §5.1.1',
   'is the mass fraction of carbon dioxide and elemental carbon'),

 ('S','5. Durability & Storage Eligibility'),
 ('I','200-yr Option 1 applied to every batch; 1000-yr R0/inertinite method correctly marked N/A','Module §5.1.2.2.1; PDD v1.12',
   'Option 1 results in Credits issued at 200 year durability'),
 ('I','Heavy metals, PAH, PCDD/F and PCB within Module limits at process/validation level','Module Table 2; PDD v1.12',
   'Quantification of heavy metals in biochar is essential to ensure environmental and human health safety'),
 ('I','Fuel-diversion risk mitigated: buyers restricted to ag soil amendment; particle-size/sieve check where material risk','Module §7.2; PDD v1.12',
   'with at least 95% by weight passing through a 10 mm sieve'),
 ('I','Stockpiling <= 12 months; stored wet/cooled/covered/away from watercourses; inventory aging tracked','Module §7.3; PDD v1.12',
   'Biochar must not be stored for more than 12 months after production'),
 ('I','Biochar SDS provided to buyer/freight; sealed vinyl sacks used for transport','PDD v1.12',''),
 ('I','Buffer pool = 0% applied (No Observable Risk for this storage module)','Module §7.1',
   'are typically deemed to have No Observable Risk of reversal'),

 ('S','6. Proof of End Use & Chain of Custody'),
 ('I','End-use pathway identified: direct-to-soil / compost mixing / third-party sale','Module §8.2',
   'There are three pathways to durable biochar storage in soils'),
 ('I','Direct-to-soil: geotagged photos (stockpile, spreading, incorporation) + boundary map / application records','Module §8.5.1',
   'Geotagged photos or videos are critical for visually confirming the application of biochar'),
 ('I','Compost/third-party: signed affidavit of intended use & location; biochar < 50% v/v of final product','Module §8.2.2; §8.7.1',
   'A legally binding affidavit or declaration signed by the purchaser of the biochar'),
 ('I','Unique sequential batch IDs on all documents; documentation present at each custody handoff','Module §8.9',
   'Every batch of biochar produced must be assigned a unique, sequential identification number'),
 ('I','Application rate <= jurisdictional maximum loading rate','Module §8.5.2',
   'The application rate should not exceed the maximum loading rates established by the relevant jurisdiction'),
 ('I','Downstream transport to end user accounted for where biochar is mixed or sold to a third party','Module §8.8',
   'The average transport distance for biochar-amended compost is greater than the average transport distance'),
 ('I','Evidence retained / retainable for >= 5 years','Module §8.9',
   'must be retained for a minimum period'),

 ('S','7. Emissions Accounting'),
 ('I','Production/process emissions scaled to the correct biochar mass; PP CI traceable to bio-oil-verified PP summary','PDD v1.12',''),
 ('I','Transport (process + embodied) emissions scaled appropriately; multi-PP loads allocated by load proportion','Protocol §8.6.5.6',
   'The Transportation Emissions Accounting Module v1.1 provides requirements on how transportation-related emissions must be calculated'),
 ('I','Spreading/application emissions scaled to applied mass','Protocol §8.6.2',
   'the Reporting Period begins when the activity associated with a batch of Removals begins'),
 ('I','Co-product allocation between bio-oil and biochar applied (energy-content basis), no double-counting','Protocol §11',
   'will be allocated proportionately to each storage pathway for the purposes of calculating net-CO2e removal'),
 ('I','Inventory reconciliation: unreconciled inventory at PP exhaustion accounted as an emission to the correct RP','Protocol §8.4.1.2',
   'Account for all documented losses in the GHG Statement for the relevant Project batch'),
 ('I','Emission factors sourced, dated and version-referenced','Protocol §6.5.1',
   'Emission factors utilized, as published in public, or other databases, used'),

 ('S','8. Net CDR, Reporting Period & Materiality'),
 ('I','Net removal = CO2e-stored - CO2e-counterfactual - CO2e-emissions','Protocol §8.2',
   'Net CO2e removal for the production of biochar and its durable storage for each Reporting Period'),
 ('I','Counterfactual storage quantified per Biomass Feedstock Accounting Module v1.3','Protocol §8.5',
   'The counterfactual storage of all feedstocks must be quantified according to the Biomass Feedstock Accounting Module'),
 ('I','RP bounded sourcing -> application; removal date not earlier than dates already credited in Certify','Protocol §8.6.2',
   'the Reporting Period begins when the activity associated with a batch of Removals begins'),
 ('I','Excluded SSRs individually and collectively below the materiality threshold of net removals','Protocol §7.1.1',
   'Negligible SSRs are those which fall below a Materiality threshold based on environmental significance of'),
 ('I','Uncertainty & sensitivity analysis present; immaterial parameters may be omitted, others quantified','Protocol §6.5',
   'Input variables may be omitted from an uncertainty analysis if they contribute to a'),

 ('S','9. Certify Entry Consistency (Phase 2 - N/A until entered)'),
 ('I','Each PP removal entered in Certify at the batch granularity defined in the PDD','vs PDD / workbook',''),
 ('I','Mass, Corg, H/Corg, Tsoil, Fdurable, gross stored, emissions and net removal match the workbook','vs workbook',''),
 ('I','Recipient, delivery/removal date and end-use pathway match the evidence','vs evidence',''),
 ('I','Supporting evidence uploaded to Certify Sources','vs evidence',''),

 ('S','10. Cross-Document Consistency & Hygiene'),
 ('I','No template placeholders left unfilled anywhere (recipient name, dates, locations)','House rule',''),
 ('I','Dates consistent across folder name, workbook, scale ticket, BOL/pickup and end-use evidence','House rule',''),
 ('I','File names match their contents and sit in the correct delivery folder','House rule',''),
 ('I','Figures tie across workbook tabs and to the GHG Statement Report totals','House rule',''),
]

# ---- SELF-VERIFY: every row whose Reference has a citation signal must carry a quote that is verbatim in corpus ----
failures=[]
for row in S:
    if row[0]!='I': continue
    _,text,ref,quote=row
    if SIGNAL_RE.search(ref):
        if not quote:
            failures.append(("NO QUOTE for signal ref", ref, text[:50])); continue
        if len(quote)<15:
            failures.append(("QUOTE <15 chars", quote, ref)); continue
        if not in_corpus(quote):
            failures.append(("QUOTE NOT IN SOURCE", quote, ref))
if failures:
    print("BUILD ABORTED -", len(failures), "unverified citation(s):")
    for f in failures: print("  -", f[0], "::", f[1][:80], "::", f[2][:40])
    raise SystemExit(1)
print("All", sum(1 for r in S if r[0]=='I' and SIGNAL_RE.search(r[2])), "cited rows verified verbatim against source. Building...")

# ---- build workbook ----
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Biochar Removal QA Checklist"
TITLE=Font(bold=True,size=14,color="2F5496"); HF=Font(bold=True,size=11,color="FFFFFF")
HFILL=PatternFill("solid",fgColor="2F5496"); SF=Font(bold=True,size=11); SFILL=PatternFill("solid",fgColor="D6E4F0")
ML=Font(bold=True,size=10); IT=Font(size=10); WRAP=Alignment(wrap_text=True,vertical="top"); TOP=Alignment(vertical="top")
REF=Font(size=9,color="404040"); QF=Font(size=9,italic=True,color="375623")
thin=Side(style="thin",color="BFBFBF"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
for c,w in zip("ABCDE",[66,11,34,24,52]): ws.column_dimensions[c].width=w
r=1
ws.cell(r,1,"Biochar Removal QA Checklist").font=TITLE
ws.cell(r,2,"v1.0").font=Font(italic=True,size=10,color="808080"); r+=1
for a,b_,c_,d_ in [("Project:","Range and Plains Biochar","Reviewer:","[Name / Date]"),
 ("Delivery / Recipient:","[Recipient + date]","Reporting Period:","[RP dates]"),
 ("Production Period(s):","[PP #s; PP = production batch, <=1 month]","Durability:","200-yr Option 1 (PDD v1.12)"),
 ("Sampling:","Method A, 3 replicates/PP (PDD v1.12)","Soil temp:","Lembrechts/direct (PDD v1.12)"),
 ("Protocol:","Biochar Production & Storage v1.2","Storage Module:","Soil Environments v1.3")]:
    ws.cell(r,1,a).font=ML; ws.cell(r,2,b_).font=IT; ws.cell(r,3,c_).font=ML; ws.cell(r,4,d_).font=IT; r+=1
ws.cell(r,1,'Each citation in the Reference column carries the verbatim source text it rests on in the next column. Numeric eligibility bounds that live in module data tables (e.g. H/Corg below 0.5) are stated in the item; the adjacent quote is the supporting Module prose.').font=Font(italic=True,size=9,color="595959"); ws.cell(r,1).alignment=WRAP; r+=1
ws.cell(r,1,"Status: Pass = met & evidenced | Fail = not met | Flag = needs Charm action | N/A = not applicable").font=Font(italic=True,size=9,color="595959"); r+=1
hdr=r
for c,t in enumerate(["Checklist Item","Status","Notes / Evidence","Reference","Verbatim source text"],1):
    cell=ws.cell(r,c,t); cell.font=HF; cell.fill=HFILL; cell.alignment=TOP; cell.border=B
r+=1
for row in S:
    if row[0]=='S':
        ws.cell(r,1,row[1]).font=SF
        for c in range(1,6): ws.cell(r,c).fill=SFILL; ws.cell(r,c).border=B
        r+=1
    else:
        _,text,ref,quote=row
        ws.cell(r,1,text).font=IT; ws.cell(r,1).alignment=WRAP
        ws.cell(r,2,"").alignment=TOP; ws.cell(r,3,"").alignment=WRAP
        ws.cell(r,4,ref).font=REF; ws.cell(r,4).alignment=WRAP
        ws.cell(r,5,('"'+quote+'"') if quote else "").font=QF; ws.cell(r,5).alignment=WRAP
        for c in range(1,6): ws.cell(r,c).border=B
        r+=1
ws.freeze_panes=f"A{hdr+1}"
dv=DataValidation(type="list",formula1='"Pass,Fail,Flag,N/A"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"B{hdr+1}:B{r}")
out="skills/qa/batch-qa-biochar/Biochar_Removal_QA_Checklist_Template.xlsx"
wb.save(out); print("SAVED", out)
