"""
batch_step_ledger.py — the required-step check for EVERY batch QA type, called by run_gates.py.

Origin: INJ-09-2026-W3-QOWV (2026-09-28). Max opened with "DO. NOT. SKIP. STEPS." Every
existing gate passed and eight required steps were still skipped, because the gates test
citations, the recompute, attachments and the tracker row, not the rest of SKILL.md.
The first version (charm_step_ledger.py) covered Charm batches only. 2026-09-29: AECN and
Kerry added, and a batch whose oil type is not recognised is blocked rather than let through.

Three layers:
  1. REQUIRED_STEPS[kind] is fixed here, one list per QA skill. The batch folder must hold
     step_ledger.json with an entry for every id, done=true (or "n/a" where allowed), and
     evidence that resolves. Evidence counts only if it is (a) a [log:KEY] whose value in
     the source-read log is not empty, or (b) a file in the batch folder (or a registry /
     the performance tracker) made or changed at or after T0. A timestamp or URL alone
     does not count (2026-09-29: a bare timestamp let a step pass with nothing behind it).
     A [log:KEY] or path that does not exist fails the entry. The lock's own files
     (step_ledger.json, ledger_audit.json, gate_receipt.json, the source-read log) never count.
     Pre-flight reads are therefore recorded in the source-read log under `preflight`.
  2. ledger_audit.json, written by an independent agent that did not do the QA: one
     verdict per step id after opening what the evidence names, and the sha256 of the
     step_ledger.json it audited. Any later change to the ledger breaks the match.
  3. Checks that need no ledger, read straight from the files: one checklist under its
     canonical name (no _vN), no process text quoting Max, status-row and header
     formatting, explicit heights on merged sidebar rows, T0-T4 stamped, empty workspace root.
"""
import datetime, hashlib, json, math, os, re
from pathlib import Path

# (id, what SKILL.md / the monitor skill requires, n/a allowed)
_PREFLIGHT = [
    ("P1_knowledge_base", "Pre-flight: read the Knowledge Base (knowledge-base/Charm_Isometric_Knowledge_Base.md; in the GCP container, its generated slice for this QA mode, knowledge-base/slices/<mode>.md)", False),
    ("P2_lessons_learned", "Pre-flight: read references/lessons_learned.md, every Failure Category", False),
    ("P3_checklist_structure", "Pre-flight: read references/checklist_structure.md", False),
    ("P4_failure_categories_printed", "Pre-flight: print the applicable Failure Categories to chat before Gate 1 begins", False),
]
_CLOSE = [
    ("G3_reviewer_independent", "Gate 3: adversarial reviewer run as an independent agent on the checklist", False),
    ("G3_challenges_resolved", "Gate 3 exit: every challenged item accepted or rebutted with evidence; checklist updated", False),
    ("G4_render_in_excel", "Gate 4: checklist opened in Excel and Excel's own PDF export read (not LibreOffice)", False),
    ("G4_status_counts", "Gate 4: final verification pass on status counts", False),
    ("G4_monitor_independent", "Performance tracker row written by an independent monitor agent, not the QA agent", False),
]

REQUIRED_STEPS = {
    "CHARM": _PREFLIGHT + [
        ("P5_registry_checked", "Pre-flight: check qa_active_charm_batches.json for a resuming batch", False),
        ("T0_stamped_monitor_triggered", "Run date at kickoff (T0) and trigger the performance monitor", False),
        ("G1_0_vvb_guidance_full", "Gate 1 step 0: VVB guidance doc for the period read in full, before any other source", False),
        ("G1_0_cell_comments", "Gate 1: Ops Notes cell comments read (openpyxl without data_only)", False),
        ("G1_1_batch_folder", "Gate 1 step 1: batch Drive folder enumerated (testing, Ops Notes, tracker)", False),
        ("G1_2_totes", "Gate 1 step 2: every tote recorded (tag, gross, tare, fraction, production period)", False),
        ("G1_3_ops_notes", "Gate 1 step 3: Ops Notes totals, uncertainty and LCS use extracted", False),
        ("G1_4_shipment_tracker", "Gate 1 step 4: shipment allocation read; computed cells confirmed as formulas, formula text quoted", False),
        ("G1_5_testing", "Gate 1 step 5: lab report PDF opened; certificate, C, H, N, density read", False),
        ("G1_6_pp_summaries", "Gate 1 step 6: each PP Summary: PP id, WODO CI, Aqueous CI, mass per fraction, date range", False),
        ("G1_7_charm_contents", "Gate 1 step 7: Charm Contents sheet read, or its absence shown from folder + Certify sources", False),
        ("G1_8_certify_four_tabs", "Gate 1 step 8: Certify removal fetched via API (certify_read.py): components, datapoints, sources and blueprint recompute", False),
        ("G1_8_sources_tab_total", "Gate 1: certify_sources_tab_total with full enumeration of source filenames", False),
        ("G1_registry_updated", "Gate 1 exit: batch added to qa_active_charm_batches.json", False),
        ("G1_injection_record", "Bio-oil Injection Tracker (COBB) opened: batch rows, tags, completion date (monitor T1 = Drive, Certify AND COBB)", False),
        ("G2_1_tare_measured", "Gate 2 sub-task 1: verify gross/tare were measured (not nominal)", False),
        ("G2_2_tracker_cells_clicked", "Gate 2 sub-task 2: each computed allocation cell's formula quoted (drive_read.py rows formula field / openpyxl)", False),
        ("G2_3_component_modals", "Gate 2 sub-task 3: every component's input datapoints and ± (standard_deviation) read from the Certify API", False),
        ("G2_4_each_pp_summary", "Gate 2 sub-task 4: open EACH PP Summary referenced by the production-emissions component", False),
        ("G2_5_recompute_pasted", "Gate 2 sub-task 5: recompute run, exit 0, full output pasted into the chat before delivery", False),
        ("G2_lcs_search", "LCS use searched in Ops Notes (Aqueous); matching Certify component if reported", True),
    ] + _CLOSE,

    "AECN": _PREFLIGHT + [
        ("P5_registry_checked", "Pre-flight 5: check qa_active_batches.json for a resuming batch", False),
        ("P6_inputs_confirmed", "Batch ID, reporting period, Certify removal URL and Drive folder URL in hand; batch confirmed AECN-only", False),
        ("T0_stamped_monitor_triggered", "Run date at kickoff (T0) and trigger the performance monitor", False),
        ("G1_0_vvb_guidance_full", "Gate 1 step 0: VVB guidance doc for the period read in full, before any other source; verbatim quotes logged", False),
        ("G1_1_scale_tickets", "Gate 1 step 1: BOTH full and empty ticket images opened; gross, date, tractor, trailer, scale, location from each", False),
        ("G1_2_bols", "Gate 1 step 2: every BOL opened (Drive or Certify Sources); BOL #, origin, destination, ship date, mass in all units, carrier", False),
        ("G1_3_lab_results", "Gate 1 step 3: SGS CoA opened on the right page; certificate #, batch ID, sample date/location, C, H, N, viscosity, acid number, density", False),
        ("G1_4_ops_notes", "Gate 1 step 4: Ops Notes batch ID, completion date, origin lot, every mass value, sparging rate + loss, pH, total injectate", False),
        ("G1_4_cell_comments", "Gate 1 step 4: every Ops Notes cell comment on the offload table dumped (openpyxl without data_only)", False),
        ("G1_5_cobb_row_verified", "Gate 1 step 5: 'Basco Injection -- COBB' tab, cell A of the row equals the batch ID before any value is read", False),
        ("G1_6_certify_four_tabs", "Gate 1 step 6: Certify removal fetched via API (certify_read.py): components, datapoints, sources and blueprint recompute", False),
        ("G1_6_sources_tab_total", "Gate 1: certify_sources_tab_total with full enumeration of source filenames and linked datapoints", False),
        ("G1_mass_chain_script", "Gate 1 exit 4 (rail): compute_mass_chain.py run; every computation FAIL either traced to a log transcription error and fixed, or confirmed against the primary source and carried as a Fail or Flag row in the checklist", True),
        ("G1_registry_updated", "Gate 1 exit: batch added to qa_active_batches.json", False),
        ("G2_0_findings_routed", "Gate 2 start (rail): mass_chain_findings.json read; each finding routed by check_type", True),
        ("G2_1_ticket_pairs_compared", "Gate 2 sub-task 1: full and empty ticket PDFs from Certify Sources opened; tractor and trailer numbers compared", False),
        ("G2_2_cobb_values", "Gate 2 sub-task 2: COBB tab name and row number recorded; offloaded qty, origin qty, completion date extracted", False),
        ("G2_3_component_modals", "Gate 2 sub-task 3: every component's input datapoints and ± (standard_deviation) read from the Certify API; every value recorded against the uncertainty reference", False),
        ("G2_ef_live_sheet", "Primary source: Standard Emission Factors sheet, 2026 tab, opened; every EF read from the live cell", False),
        ("G2_carbon_range_source", "Primary source: MRV_AECN_Analysis.xlsx opened for carbon mean / SD / n", False),
        ("G2_bcu_sources", "Primary source: BCU certificate PDF opened (coverage confirmed) and protocol opened for BCU eligibility", True),
        ("G2_route_map", "Primary source: route map opened in Certify Sources, not the Distance datapoint", False),
        ("G2_rail_allocation", "Rail: allocation fraction recomputed for AECN-SOPOR truck and SOPOR-Opelousas rail legs; railcar-balance check run", True),
        ("G2_recompute_pasted", "recompute_*.py in the batch folder, exit 0, full output pasted into the chat before delivery", False),
    ] + _CLOSE,

    "KERRY": _PREFLIGHT + [
        ("P5_registry_checked", "Pre-flight 5: check qa_active_kerry_batches.json for a resuming batch", False),
        ("P6_inputs_confirmed", "Batch ID, Certify removal URL and Drive folder URL in hand; batch confirmed Kerry-only", False),
        ("T0_stamped_monitor_triggered", "Run date at kickoff (T0) and trigger the performance monitor", False),
        ("G1_0_vvb_guidance_full", "Gate 1 step 0: VVB guidance doc for the period read in full, before any other source; verbatim quotes logged", False),
        ("G1_0_cell_comments", "Gate 1 step 0: Ops Notes cell comments read (openpyxl without data_only)", False),
        ("G1_delivery_mode", "Gate 1: delivery mode (tote / tanker) detected from Drive contents and recorded", False),
        ("G1_1_mass_artifacts", "Gate 1 step 1: every tote (serial, gross, tare, net, date, location) or both tanker tickets recorded", False),
        ("G1_2_bol", "Gate 1 step 2: BOL opened (Drive or Certify Sources); mass, origin, destination, carrier recorded", False),
        ("G1_3_lab_results", "Gate 1 step 3: SGS CoA(s) opened (per tote or per tanker), or absence recorded as lab_results.status", False),
        ("G1_4_ops_notes", "Gate 1 step 4: Ops Notes Kerry mass IN, sparged-AECN mass IN, pH, total injectate", False),
        ("G1_5_cobb_row_verified", "Gate 1 step 5: COBB cell A equals the batch ID before any value is read; consumed_lot follows kc_NNNN", False),
        ("G1_6_certify_four_tabs", "Gate 1 step 6: Certify removal fetched via API (certify_read.py): components, datapoints, sources and blueprint recompute", False),
        ("G1_6_sources_tab_total", "Gate 1: certify_sources_tab_total with full enumeration of source filenames", False),
        ("G1_registry_updated", "Gate 1 exit: batch added to qa_active_kerry_batches.json", False),
        ("G2_1_mass_artifacts_opened", "Gate 2 sub-task 1: every weighing artifact opened; tote tare measured not nominal, or tanker tractor/trailer match", False),
        ("G2_2_cobb_values", "Gate 2 sub-task 2: COBB row read; kc_NNNN pattern, offloaded qty, origin qty, completion date", False),
        ("G2_3_component_modals", "Gate 2 sub-task 3: every component's input datapoints and ± (standard_deviation) read from the Certify API; scale on the ticket checked before choosing 1.13 or 72.57 kg", False),
        ("G2_system_boundary", "No upstream Kerry production or biomass component in Certify; counterfactual = 0; no sparging on Kerry mass", False),
        ("G2_transport_mass_basis", "Mass of Load checked against BOL gross (tote) or oil-only mass (tanker)", False),
        ("G2_ef_live_sheet", "Primary source: Standard Emission Factors sheet, 2026 tab, opened; every EF read from the live cell", False),
        ("G2_required_evidence_defined", "certify.required_evidence defined from this batch's actual components for the attachment gate", False),
        ("G2_recompute_pasted", "recompute_*.py in the batch folder, exit 0, full output pasted into the chat before delivery", False),
    ] + _CLOSE,
}

# Steps that cannot run in the GCP QA container (skills/qa/container), where they may be
# marked done="n/a" with a reason. G4_render_in_excel checks how the checklist looks in
# Excel; there is no Excel in a Linux container. Every QA check and the code-enforced
# formatting checks in check_files() still run.
CONTAINER_NA = {"G4_render_in_excel"}

ANCHOR_LOG = re.compile(r"\[log:([^\]]+)\]")
ANCHOR_PATH = re.compile(r"[\w~./-]+\.(?:xlsx|json|txt|csv|pdf|png|md|py|yml|jsonl)\b")
PROCESS_TEXT = re.compile(r"\bMax\b|Per Max|per Max")
VERSIONED = re.compile(r"_v\d+\.xlsx$")
HEADER = re.compile(r"^(SECTION\b|Section\s|\d+[a-z]?\.\s|SIDEBAR|Sidebar|CALCULATION CHECK)")
LOCK_FILES = {"step_ledger.json", "ledger_audit.json", "gate_receipt.json"}
WRITTEN_WORKSPACE_FILES = {"qa_active_batches.json", "qa_active_charm_batches.json",
                           "qa_active_kerry_batches.json", "Claude_QA_Performance_Tracker.xlsx"}
_MISSING = object()


def batch_kind(log: dict, folder: Path):
    """Which QA skill's step list applies. None = not recognised, and run_gates blocks."""
    oil = str(log.get("oil_type", "")).strip().upper()
    name = folder.name.upper()
    if oil.startswith(("AQUEOUS", "WODO", "QOWV")) or name.endswith(("-QOWV", "-WODO")):
        return "CHARM"
    if oil.startswith("KERRY") or "-KERRY" in name or re.fullmatch(r"BATCH-K-\d{3}", name):
        return "KERRY"
    if oil.startswith("AECN") or "-AECN" in name or re.fullmatch(r"BATCH-2-\d{3}", name):
        return "AECN"
    return None


def is_charm(log: dict, folder: Path) -> bool:
    return batch_kind(log, folder) == "CHARM"


def t0_epoch(folder: Path):
    try:
        ts = json.loads((folder / "batch_timestamps.json").read_text()).get("timestamps", {})
        raw = ts.get("T0_initiation")
        return datetime.datetime.fromisoformat(raw.replace("Z", "+00:00")).timestamp() if raw else None
    except Exception:
        return None


def _log_get(log: dict, key: str):
    node = log
    for part in key.strip().split("."):
        if isinstance(node, dict) and part in node:
            node = node[part]
        else:
            return _MISSING
    return node


def _evidence_problems(ev: str, log: dict, folder: Path, ws: Path, t0, log_name: str) -> list:
    probs, anchors = [], 0
    for grp in ANCHOR_LOG.findall(ev):
        for key in grp.split(","):
            v = _log_get(log, key)
            if v is _MISSING:
                probs.append(f"[log:{key.strip()}] not in the source-read log")
            elif v not in (None, "", [], {}):
                anchors += 1    # an empty value is not a false citation, but it proves nothing
    folder_r = folder.resolve()
    # a log key such as [log:certify.sources_tab.yml] is a key, not a file path
    for p in ANCHOR_PATH.findall(ANCHOR_LOG.sub(" ", ev)):
        hit = next((c for c in (folder / p, ws / p, Path(p).expanduser()) if c.exists()), None)
        if hit is None:
            probs.append(f"path '{p}' does not exist")
            continue
        hit = hit.resolve()
        if hit.name in LOCK_FILES or hit.name == log_name:
            continue            # the lock's own files never count as evidence
        # A file counts only if this QA made or changed it: inside the batch folder, or one of
        # the workspace files a step writes. The KB or a skill reference exists (and may be
        # edited for other reasons) whether or not it was read, so it proves nothing.
        written_here = folder_r in hit.parents or hit.name in WRITTEN_WORKSPACE_FILES
        if written_here and t0 is not None and hit.stat().st_mtime >= t0:
            anchors += 1
    if anchors == 0:
        probs.append("no evidence that resolves: needs a [log:KEY] with a value, or a file made "
                     "or changed during this QA (a timestamp or URL alone does not count)")
    return probs


def check_ledger(folder: Path, log: dict, ws: Path, kind: str, log_name: str = "") -> list:
    path = folder / "step_ledger.json"
    if not path.exists():
        return [f"no step_ledger.json in {folder.name}: every {kind} required step needs an entry with evidence"]
    steps = json.loads(path.read_text()).get("steps", {})
    t0 = t0_epoch(folder)
    fails = []
    for sid, text, na_ok in REQUIRED_STEPS[kind]:
        na_ok = na_ok or (sid in CONTAINER_NA and os.environ.get("QA_CONTAINER") == "1")
        e = steps.get(sid)
        if not e:
            fails.append(f"{sid} missing ({text})")
            continue
        done, ev = e.get("done"), str(e.get("evidence", ""))
        if done is True:
            pass
        elif done == "n/a" and na_ok and len(ev) >= 20:
            continue
        else:
            fails.append(f"{sid} not done (done={done!r}) ({text})")
            continue
        for p in _evidence_problems(ev, log, folder, ws, t0, log_name):
            fails.append(f"{sid}: {p}")
    return fails


def check_audit(folder: Path, kind: str) -> list:
    """ledger_audit.json: an independent agent's verdict on every step, tied to this exact ledger."""
    ledger, audit = folder / "step_ledger.json", folder / "ledger_audit.json"
    if not ledger.exists():
        return []               # check_ledger already reports this
    if not audit.exists():
        return ["no ledger_audit.json: an independent agent must open the evidence for every step "
                "and record a verdict before delivery"]
    try:
        a = json.loads(audit.read_text())
    except Exception as e:
        return [f"ledger_audit.json unreadable ({e})"]
    fails = []
    live = hashlib.sha256(ledger.read_bytes()).hexdigest()
    if a.get("ledger_sha256") != live:
        fails.append("ledger_audit.json was written against a different step_ledger.json "
                     "(the ledger changed after the audit); re-run the independent audit")
    verdicts = a.get("steps", {})
    for sid, text, _ in REQUIRED_STEPS[kind]:
        v = verdicts.get(sid) or {}
        if v.get("verdict") != "verified":
            fails.append(f"audit {sid}: verdict {v.get('verdict')!r}, not 'verified' ({text})")
        elif len(str(v.get("checked", ""))) < 30:
            fails.append(f"audit {sid}: 'checked' does not say what the auditor opened and saw")
    return fails


def check_files(folder: Path, checklist: Path, ws: Path) -> list:
    import openpyxl
    fails = []
    if VERSIONED.search(checklist.name):
        fails.append(f"checklist '{checklist.name}' carries a version suffix; one checklist per batch, original name")
    ts_path = folder / "batch_timestamps.json"
    if ts_path.exists():
        ts = json.loads(ts_path.read_text()).get("timestamps", {})
        for k in ("T0_initiation", "T1_data_gathering_complete", "T2_checklist_draft_complete",
                  "T3_adversarial_review_complete", "T4_checklist_finalized"):
            if not ts.get(k):
                fails.append(f"batch_timestamps.json: {k} not stamped")
    else:
        fails.append("no batch_timestamps.json")
    stray = [p.name for p in ws.iterdir() if p.is_file() and p.suffix.lower() in (".png", ".jpg", ".jpeg", ".yml", ".yaml")]
    if stray:
        fails.append(f"workspace root not clean: {stray}")

    ws_ = openpyxl.load_workbook(checklist).active
    def fill(c):
        return (c.fill.fgColor.rgb or "")[-6:] if c.fill is not None and c.fill.fill_type == "solid" else ""
    for r in range(1, ws_.max_row + 1):
        cells = [ws_.cell(r, c) for c in range(1, 6)]
        for c in cells:
            if isinstance(c.value, str) and PROCESS_TEXT.search(c.value):
                fails.append(f"{c.coordinate}: process text naming Max in the checklist")
        a, st = cells[0].value, cells[2].value
        if isinstance(a, str) and HEADER.match(a) and st in (None, "") and any(fill(c) != "D6E4F0" for c in cells):
            fails.append(f"row {r}: section header not #D6E4F0 across A-E")
        if st == "Fail" and any(fill(c) in ("", "FFF2CC", "D6E4F0") for c in cells):
            fails.append(f"row {r}: Fail row without red fill across A-E")
        if st == "Flag" and (any(fill(c) != "FFF2CC" for c in cells) or
                             any((c.font.color is None or (c.font.color.rgb or "")[-6:] != "996600") for c in cells)):
            fails.append(f"row {r}: Flag row not #FFF2CC fill with #996600 font across A-E")
        if st == "Pass" and any(fill(c) for c in cells):
            fails.append(f"row {r}: Pass row carries a fill")
        if st == "N/A" and (any(fill(c) for c in cells) or not all(c.font.italic for c in cells)):
            fails.append(f"row {r}: N/A row not gray italic without fill")
    for rng in ws_.merged_cells.ranges:
        v = ws_.cell(rng.min_row, rng.min_col).value
        if isinstance(v, str) and len(v) > 150:
            h = ws_.row_dimensions[rng.min_row].height
            if not h or h < 15 * math.ceil(len(v) / 200):
                fails.append(f"row {rng.min_row}: merged cell text without a row height; Excel shows one line")
    return fails
