"""
Charm-Produced Bio-Oil Injection Batch QA Checklist Template Builder.

Structured around Max's actual QA question framework for Charm WODO and Aqueous
batches. The structure is shared across oil types except:
  - Section 2.3 (Total Oil in Batch math) branches by oil_type:
      WODO    -> Σ (Net-in-tote − Standard Loss)
      Aqueous -> Σ (Pre-Unload − Post-Unload)  [NO Standard Loss]
  - Section 9 (LCS buffering) renders only for Aqueous; N/A for WODO.

Template structure (canonical):
  Section 1: Batch Folder                          (3 items)
  Section 2: Ops Notes                             (4 items)
  Section 3: Bio-Oil Shipment Tracker              (5 items)
  Section 4: Certify - Injection Sequestration     (4 items)
  Section 5: Certify - Transport Process           (5 items)
  Section 6: Certify - Transport Embodied          (3 items)
  Section 7: Certify - Production Emissions        (6 items, applied per PP)
  Section 8: Other (BCUs)                          (1 item)
  Section 9: Aqueous Buffering (LCS)               (1 item, Aqueous-only)

Out-of-frame side observations (system-boundary checks, density anomalies, etc.)
go in a "Sidebar Observations" section at the end with no Pass/Fail status,
just a flagged-for-attention listing.

Header color discipline (per workspace memory feedback_checklist_header_one_color.md):
  All section and sub-section headers use ONE neutral color (#D6E4F0).
  Never yellow (conflicts with FLAG), never green, never any second tone.

Usage:
    Caller passes a `findings` dict mapping item-id -> {status, evidence, ref}.
    Builder writes the xlsx to OUT path, applies fail/flag/pass formatting,
    and saves.

This is a TEMPLATE - caller fills in findings per-batch, builder applies
canonical formatting and structure. Do NOT hard-code per-batch values into
this template; that defeats the point of having a template.
"""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# Branch text by oil type for item 2.3
_TOTAL_OIL_IN_BATCH_Q = {
    "WODO": (
        "Is the 'Total Oil in Batch' value correctly computed as the sum of "
        "(Net Weight of Oil in Tote − Standard Loss) for all totes?"
    ),
    "Aqueous": (
        "Is the 'Total Oil in Batch' value correctly computed as the sum of "
        "(Measured Pre-Unload − Measured Post-Unload) per tote? (No Standard "
        "Loss; Aqueous tote loss is measured directly.)"
    ),
}


def _structure(oil_type: str) -> list:
    """Build the canonical checklist structure for the given oil type."""
    q23 = _TOTAL_OIL_IN_BATCH_Q[oil_type]
    sections = [
        {
            "section": "Section 1: Batch Folder",
            "items": [
                ("1.1", "Is there testing in the testing folder for the batch?"),
                ("1.2", "Are Ops Notes in the folder and fully completed?"),
                ("1.3", "Is the Bio-Oil Shipment Tracker in the batch folder and fully completed?"),
            ],
        },
        {
            "section": "Section 2: Ops Notes",
            "items": [
                ("2.1", "Is the 'Total Mass Transported' value correctly computed as the sum of the gross weight of all totes?"),
                ("2.2", "Is the 'Total Oil Produced' value correctly computed as the sum of the net weight of oil in tote for all totes (gross − tare per tote)?"),
                ("2.3", q23),
                ("2.4", "Is uncertainty correctly computed as 2.5 lb per tote, summed across totes, converted to kg?"),
            ],
        },
        {
            "section": "Section 3: Bio-Oil Shipment Tracker",
            "items": [
                ("3.1", "Are totes for the RP correctly identified?"),
                ("3.2", "Do all totes identified for the RP match the Ops Notes for the applicable batch?"),
                ("3.3", "Are computed outputs (RP Share of Truck, RP Allocated Miles, Batch Mileage Allocation) auditable formulas rather than hard-coded values?"),
                ("3.4", "Are computed outputs computed correctly? "
                        "(RP Share of Truck = N_totes_in_RP / N_totes_on_Truck; "
                        "RP Allocated Miles = Mileage(one-way) × RP Share of Truck; "
                        "Batch Mileage Allocation = (N_totes_in_Batch / N_totes_on_Truck(s)) × Mileage(one-way))"),
                ("3.5", "Do the totes assigned to each batch match the Ops Notes on tag number, fraction, gross mass, and production period?"),
            ],
        },
        {
            "section": "Section 4: Certify - Injection Sequestration",
            "items": [
                ("4.1", "Does the C content in the testing match the C content in the Sequestration component?"),
                ("4.2", "Is the C content in the Sequestration component properly evidenced (test copy + Ops Notes copy)?"),
                ("4.3", "Does the 'Mass of Product' measurement uncertainty in the Sequestration component match the uncertainty value in Ops Notes?"),
                ("4.4", "Does the C content have the appropriately computed 0.4% relative uncertainty (0.004 × reported value)?"),
            ],
        },
        {
            "section": "Section 5: Certify - Transport Process Emissions",
            "items": [
                ("5.1", "Does the Transport Process emission use the correct and up-to-date Flatbed Truck emission factor?"),
                ("5.2", "Is the Transport Process distance value correct and evidenced by a Google Maps route where the longest given value is used?"),
                ("5.3", "Does the Mass of Load match the 'Total Mass Transported' value in Ops Notes?"),
                ("5.4", "Is the Mass of Load evidenced by a copy of the batch Ops Notes?"),
                ("5.5", "Does the uncertainty reported for Mass of Load match the uncertainty value in Ops Notes?"),
            ],
        },
        {
            "section": "Section 6: Certify - Transport Embodied Emissions",
            "items": [
                ("6.1", "Does the Transport Embodied emission use the correct and up-to-date Non-Tanker Truck emission factor?"),
                ("6.2", "Does the aggregate distance transported match the Batch Mileage Allocation in the RP Bio-Oil Shipment Tracker?"),
                ("6.3", "Is the aggregate distance evidenced with the RP Bio-Oil Shipment Tracker, batch Charm Contents, and applicable route Google Map?"),
            ],
        },
        {
            "section": "Section 7: Certify - Production Emissions (per Production Period)",
            "items": [
                ("7.1", "Are there Production Emissions components for ALL Production Periods in the batch?"),
                ("7.2", "Does the CI for each Production Emissions component match the CI for that fraction in the applicable PP Summary?"),
                ("7.3", "Is the CI evidenced by the completed PP Summary for the applicable Production Period?"),
                ("7.4", "Does the mass match the sum of Net Weight of Bio-Oil in Tote values for the PP that applies to the component (from Ops Notes)?"),
                ("7.5", "Does the uncertainty match the sum of uncertainties for the PP that applies to the component (from Ops Notes)?"),
                ("7.6", "Is the mass evidenced with a copy of the Charm Contents sheet for the batch?"),
            ],
        },
        {
            "section": "Section 8: Other",
            "items": [
                ("8.1", "If BCUs are applied, does the total match the sum of Transport Process emissions in the batch?"),
            ],
        },
    ]

    if oil_type == "Aqueous":
        sections.append({
            "section": "Section 9: Aqueous Buffering (LCS)",
            "items": [
                ("9.1", "If LCS use is reported in Ops Notes, does a matching Certify component exist (correct mass, correct EF, evidenced)?"),
            ],
        })
    return sections


# Style helpers
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="D6E4F0")  # one neutral color for ALL header rows
SIDEBAR_FILL = PatternFill("solid", fgColor="D6E4F0")
FLAG_FONT = Font(color="996600", bold=False)
NA_FONT = Font(italic=True, color="808080")
SECTION_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(border_style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def _write_title_block(ws, batch_meta: dict) -> int:
    ws["A1"] = "Charm Injection Batch QA Checklist"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:E1")
    rows = [
        ("Batch ID", batch_meta.get("batch_id", "")),
        ("Reporting Period", batch_meta.get("reporting_period", "")),
        ("Reviewer", batch_meta.get("reviewer", "Claude QA Agent")),
        ("Date", batch_meta.get("date", "")),
        ("Injection Site", batch_meta.get("injection_site", "Basco")),
        ("Oil Type", batch_meta.get("oil_type", "")),
        ("Production Periods", ", ".join(batch_meta.get("production_periods", []))),
        ("Certify Removal", batch_meta.get("certify_removal", "")),
    ]
    for i, (k, v) in enumerate(rows, start=2):
        ws[f"A{i}"] = k
        ws[f"B{i}"] = v
        ws[f"A{i}"].font = HEADER_FONT
        ws.merge_cells(f"B{i}:E{i}")
    return len(rows) + 2


def _section_header(ws, row: int, title: str) -> int:
    ws[f"A{row}"] = title
    ws[f"A{row}"].font = SECTION_FONT
    ws[f"A{row}"].fill = HEADER_FILL
    ws.merge_cells(f"A{row}:E{row}")
    return row + 1


def _col_header(ws, row: int) -> int:
    ws[f"A{row}"] = "Item #"
    ws[f"B{row}"] = "Question"
    ws[f"C{row}"] = "Status"
    ws[f"D{row}"] = "Evidence / Finding"
    ws[f"E{row}"] = "Reference"
    for c in "ABCDE":
        ws[f"{c}{row}"].font = HEADER_FONT
        ws[f"{c}{row}"].fill = HEADER_FILL
        ws[f"{c}{row}"].border = BORDER
        ws[f"{c}{row}"].alignment = CENTER
    return row + 1


def _write_item(ws, row: int, item_id: str, question: str, finding: dict) -> int:
    status = finding.get("status", "")
    evidence = finding.get("evidence", "")
    ref = finding.get("ref", "")
    ws[f"A{row}"] = item_id
    ws[f"B{row}"] = question
    ws[f"C{row}"] = status
    ws[f"D{row}"] = evidence
    ws[f"E{row}"] = ref
    for c in "ABCDE":
        ws[f"{c}{row}"].border = BORDER
        ws[f"{c}{row}"].alignment = WRAP
    if status == "Fail":
        for c in "ABCDE":
            ws[f"{c}{row}"].fill = FAIL_FILL
    elif status == "Flag":
        for c in "ABCDE":
            ws[f"{c}{row}"].fill = FLAG_FILL
            ws[f"{c}{row}"].font = FLAG_FONT
    elif status == "N/A":
        for c in "ABCDE":
            ws[f"{c}{row}"].font = NA_FONT
    ws.row_dimensions[row].height = 60
    return row + 1


def _write_sidebar(ws, row: int, observations: list) -> int:
    if not observations:
        return row
    row += 1
    ws[f"A{row}"] = "Sidebar Observations (out-of-frame, no status)"
    ws[f"A{row}"].font = SECTION_FONT
    ws[f"A{row}"].fill = SIDEBAR_FILL
    ws.merge_cells(f"A{row}:E{row}")
    row += 1
    ws[f"A{row}"] = "#"
    ws[f"B{row}"] = "Observation"
    ws[f"C{row}"] = ""
    ws[f"D{row}"] = "Detail"
    ws[f"E{row}"] = "Reference"
    for c in "ABCDE":
        ws[f"{c}{row}"].font = HEADER_FONT
        ws[f"{c}{row}"].fill = HEADER_FILL
        ws[f"{c}{row}"].border = BORDER
        ws[f"{c}{row}"].alignment = CENTER
    row += 1
    for i, obs in enumerate(observations, start=1):
        ws[f"A{row}"] = f"S.{i}"
        ws[f"B{row}"] = obs.get("title", "")
        ws[f"C{row}"] = ""
        ws[f"D{row}"] = obs.get("detail", "")
        ws[f"E{row}"] = obs.get("ref", "")
        for c in "ABCDE":
            ws[f"{c}{row}"].border = BORDER
            ws[f"{c}{row}"].alignment = WRAP
        ws.row_dimensions[row].height = 50
        row += 1
    return row


def build_charm_checklist(
    out_path: str,
    batch_meta: dict,
    findings: dict,
    sidebar: list | None = None,
) -> dict:
    """
    Build a Charm WODO/Aqueous QA checklist xlsx.

    Args:
        out_path: absolute path for output xlsx
        batch_meta: dict with batch_id, reporting_period, date, certify_removal,
                    oil_type ('WODO' or 'Aqueous'), production_periods (list of PP IDs),
                    reviewer, injection_site
        findings: dict mapping item_id (e.g. '1.1', '2.3', '7.4') ->
                  {'status': 'Pass'|'Fail'|'Flag'|'N/A', 'evidence': str, 'ref': str}
        sidebar: optional list of {'title', 'detail', 'ref'} for out-of-frame observations

    Returns:
        dict with status counts and item id -> status mapping for caller's records.
    """
    oil_type = batch_meta.get("oil_type")
    if oil_type not in ("WODO", "Aqueous"):
        raise ValueError(
            f"oil_type must be 'WODO' or 'Aqueous'; got {oil_type!r}. "
            "AECN and Kerry batches use different skills."
        )
    sidebar = sidebar or []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QA Checklist"

    for col, w in {"A": 8, "B": 60, "C": 9, "D": 100, "E": 28}.items():
        ws.column_dimensions[col].width = w

    r = _write_title_block(ws, batch_meta)
    r += 1

    counts = {"Pass": 0, "Fail": 0, "Flag": 0, "N/A": 0, "Unfilled": 0}
    item_status = {}

    for section in _structure(oil_type):
        r = _section_header(ws, r, section["section"])
        r = _col_header(ws, r)
        for item_id, question in section["items"]:
            finding = findings.get(item_id, {})
            if not finding:
                finding = {"status": "Unfilled", "evidence": "[No finding provided by caller]", "ref": ""}
            r = _write_item(ws, r, item_id, question, finding)
            status = finding.get("status", "Unfilled")
            counts[status] = counts.get(status, 0) + 1
            item_status[item_id] = status
        r += 1

    r = _write_sidebar(ws, r, sidebar)

    ws.freeze_panes = "A2"
    wb.save(out_path)
    return {"counts": counts, "item_status": item_status, "out_path": out_path}


if __name__ == "__main__":
    print("Charm WODO/Aqueous checklist template — call build_charm_checklist(out_path, batch_meta, findings, sidebar)")
    print()
    print("WODO structure:")
    for s in _structure("WODO"):
        print(f"  {s['section']}")
        for item_id, q in s["items"]:
            print(f"    {item_id}: {q[:100]}{'...' if len(q) > 100 else ''}")
    print()
    print("Aqueous structure (adds Section 9):")
    for s in _structure("Aqueous")[-1:]:
        print(f"  {s['section']}")
        for item_id, q in s["items"]:
            print(f"    {item_id}: {q[:100]}{'...' if len(q) > 100 else ''}")
