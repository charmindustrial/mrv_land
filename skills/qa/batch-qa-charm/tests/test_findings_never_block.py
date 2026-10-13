"""Regression tests for the findings rule (2026-09-29, INJ-09-2026-W3-AECN).

A problem in Certify or the evidence is a finding for Garrett. It may block delivery only while the
checklist has not recorded it; once a Fail/Flag row carries it, it must go through. Each test builds a
tiny checklist in memory and calls the same helpers run_gates.py uses.
Run: python3 skills/qa/batch-qa-charm/tests/test_findings_never_block.py
"""
import sys
from pathlib import Path
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import run_gates as rg

GAPS = ["Allocated Distance AECN to Sopor Quebec (truck)", "Allocated Distance Opelousas to B6 (truck)"]


def sheet(rows):
    ws = openpyxl.Workbook().active
    for r, (item, status, text) in enumerate(rows, start=1):
        ws.cell(r, 1, item); ws.cell(r, 3, status); ws.cell(r, 4, text)
    return ws


def test_attachment_gap_recorded_as_fail_goes_through():
    ws = sheet([("3.5", "Fail", "missing maps on " + "; ".join(GAPS) + " [log:certify.attachment_gate]")])
    row, why = rg.attachment_gaps_routed(GAPS, ws)
    assert row == "3.5", why


def test_attachment_gap_not_recorded_blocks():
    ws = sheet([("3.5", "Pass", "distances fine [log:route_evidence]")])
    row, _ = rg.attachment_gaps_routed(GAPS, ws)
    assert row is None


def test_attachment_gap_partly_recorded_blocks():
    ws = sheet([("3.5", "Fail", GAPS[0] + " [log:certify.attachment_gate]")])
    row, why = rg.attachment_gaps_routed(GAPS, ws)
    assert row is None and GAPS[1] in why


def test_recompute_finding_under_fail_goes_through():
    ws = sheet([("2.7", "Fail", "x [log:a]"), ("9.1", "Flag", "y [log:b]")])
    assert rg.findings_unrouted("OK ...\nFINDING 2.7/9.1 Mass of Injectant: ...\n", ws) == []


def test_recompute_finding_under_pass_blocks():
    ws = sheet([("2.7", "Pass", "x [log:a]")])
    assert rg.findings_unrouted("FINDING 2.7 Mass of Injectant: ...", ws)


def test_recompute_finding_with_missing_row_blocks():
    ws = sheet([("2.7", "Fail", "x [log:a]")])
    assert rg.findings_unrouted("FINDING 9.6 BCU: ...", ws)


def test_informational_finding_lines_ignored():
    ws = sheet([])
    assert rg.findings_unrouted("FINDING corrections A (2.7: ...)\nFINDING 120509: tote mass ...", ws) == []


def test_skipped_step_needs_authorisation():
    ledger = {"P4_failure_categories_printed": {"done": False}}
    ok, _ = rg.step_exception_ok({"id": "step:P4_failure_categories_printed"}, "P4_failure_categories_printed", ledger, None)
    assert not ok
    ok, _ = rg.step_exception_ok({"id": "step:P4_failure_categories_printed", "authorized_by": "YES (MRV lead)"},
                                 "P4_failure_categories_printed", ledger, None)
    assert ok


def test_skipped_step_marked_done_cannot_be_excepted():
    ledger = {"P4_failure_categories_printed": {"done": True}}
    ok, _ = rg.step_exception_ok({"authorized_by": "YES"}, "P4_failure_categories_printed", ledger, None)
    assert not ok


if __name__ == "__main__":
    tests = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for t in tests:
        t()
        print("PASS", t.__name__)
    print(f"{len(tests)} tests passed")
