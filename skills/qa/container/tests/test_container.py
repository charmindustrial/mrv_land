#!/usr/bin/env python3
"""
Tests for the QA container's cost changes. Each test pins down one way a change
could weaken the QA, and fails if it does.

  python3 skills/qa/container/tests/test_container.py

Needs openpyxl only (no Google libraries, no network).
"""
from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import tempfile
import types
from pathlib import Path

HERE = Path(__file__).resolve().parent
C = HERE.parent
WS = C.parents[2]
sys.path[:0] = [str(C), str(C / "tools"), str(WS / "skills/qa/batch-qa-charm")]

import kb_slices  # noqa: E402
import parity_check  # noqa: E402
import qa_inputs  # noqa: E402
import certify_read  # noqa: E402
import drive_read  # noqa: E402

RESULTS: list[tuple[str, bool, str]] = []


def test(fn):
    try:
        fn()
        RESULTS.append((fn.__name__, True, ""))
    except AssertionError as e:
        RESULTS.append((fn.__name__, False, str(e)))
    except Exception as e:  # noqa: BLE001
        RESULTS.append((fn.__name__, False, f"{type(e).__name__}: {e}"))
    return fn


# ------------------------------------------------------------- Knowledge Base slices

@test
def kb_every_exclusion_matches_exactly_one_heading():
    lines = kb_slices.KB.read_text(encoding="utf-8").splitlines(keepends=True)
    for mode, excl in kb_slices.SLICES.items():
        _, errors = kb_slices.plan(lines, excl)
        assert not errors, f"{mode}: {errors}"


@test
def kb_slice_is_full_kb_minus_only_the_named_sections():
    """Every line outside an excluded span survives, in order, unmodified."""
    lines = kb_slices.KB.read_text(encoding="utf-8").splitlines(keepends=True)
    for mode, excl in kb_slices.SLICES.items():
        spans, _ = kb_slices.plan(lines, excl)
        drop = {i for s, e, _, _ in spans for i in range(s, e)}
        expected = "".join(l for i, l in enumerate(lines) if i not in drop)
        with tempfile.TemporaryDirectory() as d:
            old = kb_slices.OUT
            kb_slices.OUT = Path(d)
            try:
                kb_slices.build(mode, write=True)
                text = (Path(d) / f"{mode}.md").read_text(encoding="utf-8")
            finally:
                kb_slices.OUT = old
        assert text.endswith(expected), f"{mode}: slice body differs from KB minus exclusions"
        for _, _, title, _ in spans:
            assert f"**{title}**" in text.split(expected)[0], f"{mode}: banner does not list {title!r}"


@test
def kb_rule_sections_are_in_every_slice():
    """The sections that carry QA rules must never be sliced out."""
    must_keep = ["0. Core QA Rule", "0A. Progressive Gate System", "0B. Reward & Penalty Framework",
                 "6. Emission Factor Reference", "10A. Injection Batch QA Checklist",
                 "13A. Data Verification", "13A-II. Primary Source Verification",
                 "Batch QA Workflow", "QA Performance Self-Tracking"]
    for mode, excl in kb_slices.SLICES.items():
        for keep in must_keep:
            assert not any(keep.startswith(e) or e.startswith(keep) for e in excl), \
                f"{mode} excludes rule section {keep!r}"


@test
def kb_new_section_is_kept_by_default():
    """A section added to the KB later appears in every slice without config changes."""
    lines = ["# KB\n", "## 1. A\n", "a\n", "## 99. Brand new rule\n", "keep me\n", "## 7. Service Level Agreements\n", "x\n"]
    spans, errors = kb_slices.plan(lines, {"7. Service Level Agreements": "r"})
    assert not errors
    drop = {i for s, e, _, _ in spans for i in range(s, e)}
    kept = "".join(l for i, l in enumerate(lines) if i not in drop)
    assert "Brand new rule" in kept and "keep me" in kept and "Service Level" not in kept


# ------------------------------------------------------------- narrowed spreadsheet reads

def _workbook(path: Path):
    import openpyxl
    from openpyxl.comments import Comment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Basco Injection -- COBB"
    ws.append(["Batch ID", "Completion date", "Offloaded kg", "Origin kg", "Ratio"])
    for i in range(1, 301):
        ws.append([f"2-{i:03d}", "2026-09-01", 1000 + i, 2000 + i, None])
    ws["E202"] = "=C202/D202"
    ws["C202"].comment = Comment("Last offload back-calculated to BOL aggregate", "ops")
    ws["A202"].hyperlink = "https://drive.google.com/file/d/abc"
    wb.save(path)
    Path(str(path) + ".meta.json").write_text(json.dumps({"live_source_url": "https://docs.google.com/x"}))


@test
def rows_returns_match_with_formula_comment_link_and_location():
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "cobb.xlsx"
        _workbook(p)
        r = drive_read.cmd_rows(p, "Basco Injection -- COBB", r"^2-201$", None, 0, False)
        assert r["matched_rows"] == [202], r["matched_rows"]
        cells = r["rows"][0]["cells"]
        assert cells["A202"]["value"] == "2-201"
        assert cells["C202"]["comment"].startswith("Last offload"), cells["C202"]
        assert cells["E202"]["formula"] == "=C202/D202", cells["E202"]
        assert cells["A202"]["hyperlink"].startswith("https://drive.google.com")
        assert r["header"]["A1"]["value"] == "Batch ID"
        assert r["live_source_url"] == "https://docs.google.com/x"


@test
def rows_absence_is_a_full_tab_scan():
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "cobb.xlsx"
        _workbook(p)
        r = drive_read.cmd_rows(p, "Basco Injection -- COBB", r"^2-999$", None, 0, False)
        assert r["match_count"] == 0 and r["rows_scanned"] == 301, (r["match_count"], r["rows_scanned"])


@test
def rows_without_match_returns_whole_tab_and_comments_flag_dumps_all():
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "cobb.xlsx"
        _workbook(p)
        r = drive_read.cmd_rows(p, "Basco Injection -- COBB", None, None, 0, True)
        assert len(r["rows"]) == 301
        assert r["all_comments"] == {"C202": "Last offload back-calculated to BOL aggregate"}


@test
def full_workbook_stays_on_disk():
    """Narrowed reads never replace the downloaded file."""
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "cobb.xlsx"
        _workbook(p)
        before = hashlib.sha256(p.read_bytes()).hexdigest()
        drive_read.cmd_rows(p, "Basco Injection -- COBB", r"2-1", None, 2, True)
        assert hashlib.sha256(p.read_bytes()).hexdigest() == before


# ------------------------------------------------------------- Certify summary

def _dump():
    dp = lambda i, mag, sd, src: {"id": f"dp_{i}", "display_name": f"Input {i}", "project_id": "prj_1",
                                  "quantity": {"magnitude": mag, "unit": "kg", "standard_deviation": sd},
                                  "source_ids": src}
    return {
        "fetched_at_utc": "t", "env": "prod", "removal_id": "rmv_1", "project_id": "prj_1",
        "live_source_url": "u", "removal": {"co2e_net_removed_kg": 12345.6, "co2e_net_removed_standard_deviation_kg": 12.3},
        "components": [
            {"attribution": {"attribution_factor": 1.0}, "component": {"display_name": "Injection Sequestration"},
             "blueprint_key": "bp_inj", "inputs": [{"input_key": "mass", "datapoints": [dp(1, 21844.99, 1.13, ["src_a"])]},
                                                   {"input_key": "carbon", "datapoints": [dp(2, 0.4321, 0.0017, [])]}]},
            {"attribution": {"attribution_factor": 0.1542}, "component": {"display_name": "AECN->SOPOR Transport"},
             "blueprint_key": "bp_tr", "inputs": [{"input_key": "distances", "datapoints": [dp(3, 512.0, 0, ["src_b"]), dp(4, 77.5, 0, ["src_b"])]}]},
        ],
        "blueprints": {}, "sources": [{"id": "src_a", "original_filename": "TKT-1.pdf", "linked_datapoint_ids": ["dp_1"]},
                                       {"id": "src_b", "original_filename": "route.png", "linked_datapoint_ids": ["dp_3", "dp_4"]}],
        "certify_sources_via_datapoints_total": 2, "gaps": ["g1", "g2"],
    }


@test
def summary_carries_every_datapoint_value_sd_and_path():
    s = certify_read.summarize(_dump(), "out/certify_removal.json")
    for needle in ["21844.99", "1.13", "0.4321", "0.0017", "512", "77.5", "0.1542",
                   "components.0.inputs.0.datapoints.0", "components.1.inputs.0.datapoints.1",
                   "TKT-1.pdf", "route.png", "12345.6", "out/certify_removal.json", "g1", "g2"]:
        assert needle in s, f"summary missing {needle!r}"
    assert s.count("| mass |") + s.count("| carbon |") + s.count("| distances |") == 4, "a datapoint row is missing"
    exact = certify_read.summarize({**_dump(), "removal": {"co2e_net_removed_kg": 362112.2896088354}}, "x")
    assert "362112.2896088354" in exact, "summary must not round values"


@test
def summary_marks_datapoint_with_no_source():
    s = certify_read.summarize(_dump(), "x.json")
    line = next(l for l in s.splitlines() if "| carbon |" in l)
    assert "**none**" in line, line


# ------------------------------------------------------------- delivery gate

@test
def checklist_upload_needs_matching_green_receipt():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d) / "outputs/qa/batch-2-201"
        root.mkdir(parents=True)
        ck = root / "Batch_2-201_QA_Checklist.xlsx"
        ck.write_bytes(b"v1")
        for setup, allowed in [(None, False), ("match", True), ("edited", False), ("not_green", False)]:
            if setup in ("match", "edited", "not_green"):
                (root / "gate_receipt.json").write_text(json.dumps(
                    {"all_gates_green": setup != "not_green", "checklist_sha256": hashlib.sha256(b"v1").hexdigest()}))
            ck.write_bytes(b"v2" if setup == "edited" else b"v1")
            try:
                drive_read._check_delivery_gate(ck)
                ok = True
            except SystemExit:
                ok = False
            assert ok == allowed, f"{setup}: allowed={ok}, expected {allowed}"


# ------------------------------------------------------------- step ledger

@test
def excel_render_na_only_in_container():
    import batch_step_ledger as bsl
    with tempfile.TemporaryDirectory() as d:
        folder = Path(d)
        (folder / "step_ledger.json").write_text(json.dumps({"steps": {
            "G4_render_in_excel": {"done": "n/a", "evidence": "No Excel in the GCP QA container (CONTAINER_NA)"}}}))
        def g4_fails():
            return [f for f in bsl.check_ledger(folder, {}, folder, "AECN") if f.startswith("G4_render_in_excel")]
        old = os.environ.pop("QA_CONTAINER", None)
        try:
            assert g4_fails(), "on the Mac, n/a for the Excel render must still fail"
            os.environ["QA_CONTAINER"] = "1"
            assert not g4_fails(), f"in the container, n/a with a reason must pass: {g4_fails()}"
        finally:
            os.environ.pop("QA_CONTAINER", None)
            if old is not None:
                os.environ["QA_CONTAINER"] = old


@test
def container_na_covers_only_the_excel_render():
    import batch_step_ledger as bsl
    assert bsl.CONTAINER_NA == {"G4_render_in_excel"}, bsl.CONTAINER_NA


# ------------------------------------------------------------- parity check

@test
def parity_flags_any_softer_grade_as_regression():
    base = {i: {"status": s, "evidence": "e"} for i, s in
            {"2.3": "Fail", "3.1": "Flag", "4.1": "Pass", "5.1": "Pass", "6.1": "Fail", "7.1": "N/A"}.items()}
    cont = {i: {"status": s, "evidence": "e"} for i, s in
            {"2.3": "Flag", "3.1": "Pass", "4.1": "N/A", "5.1": "Fail", "7.1": "N/A"}.items()}
    r = parity_check.compare(base, cont)
    reg = {x["item"] for x in r["regressions"]}
    assert reg == {"2.3", "3.1", "4.1", "6.1"}, reg          # softer, softer, skipped, missing
    assert {x["item"] for x in r["stricter"]} == {"5.1"}
    assert r["verdict"] == "FAIL"
    same = parity_check.compare(base, base)
    assert same["verdict"] == "PASS" and same["same"] == 6


# ------------------------------------------------------------- inputs and gate sessions

@test
def every_gated_mode_matches_the_skill_gate_headings():
    import re
    for mode, gates in qa_inputs.GATES.items():
        if not gates:
            continue
        skill = (WS / "skills/qa" / qa_inputs.MODES[mode][0] / "SKILL.md").read_text(encoding="utf-8")
        found = sorted({int(n) for n in re.findall(r"^### GATE (\d)", skill, re.M)})
        assert found == gates, f"{mode}: skill has gates {found}, container runs {gates}"


@test
def every_mode_has_a_slice_entry_and_valid_skill():
    for mode in qa_inputs.MODES:
        assert mode in qa_inputs.KB_SLICE and mode in qa_inputs.GATES, mode
        sl = qa_inputs.KB_SLICE[mode]
        assert sl is None or sl in kb_slices.SLICES, f"{mode}: unknown slice {sl}"
        if mode in ("smoke", "post", "discover"):
            continue                      # container check / Slack delivery, no skill
        assert (WS / "skills/qa" / qa_inputs.MODES[mode][0] / "SKILL.md").exists(), mode


@test
def input_validation():
    assert qa_inputs.validate({"mode": "batch-qa-aecn", "batch_id": "2-201", "certify_url": "x", "drive_folder": "y"}) == []
    assert qa_inputs.validate({"mode": "batch-qa-aecn", "batch_id": "2-201"})
    assert qa_inputs.validate({"mode": "batch-qa-aecn", "batch_id": "2", "certify_url": "x",
                               "drive_folder": "y", "start_gate": "9"})
    assert qa_inputs.validate({"mode": "nope"})
    assert qa_inputs.validate({"mode": "post", "post_execution": "qa-runner-qv6rk"}) == []
    assert qa_inputs.validate({"mode": "post"}), "post needs post_execution"
    ok = {"mode": "re-review", "batch_id": "INJ-09-2026-W4"}
    for m in qa_inputs.MODELS:
        assert qa_inputs.validate({**ok, "model": m}) == [], m
    assert qa_inputs.validate({**ok, "model": "claude-opus-5"}), "unlisted model must be rejected"


@test
def gate_prompts_enforce_previous_exit_condition_and_scope():
    sys.modules.setdefault("google", types.ModuleType("google"))
    for name in ["google.api_core", "google.api_core.exceptions", "google.cloud", "google.cloud.storage"]:
        sys.modules.setdefault(name, types.ModuleType(name))
    sys.modules["google.api_core.exceptions"].NotFound = Exception
    sys.modules["google.api_core.exceptions"].PreconditionFailed = Exception
    sys.modules["google.cloud"].storage = sys.modules["google.cloud.storage"]
    import entrypoint
    entrypoint.WS = Path("/workspace")
    rf = Path("/workspace/outputs/qa/_container/batch-qa-aecn__2-201")
    inputs = {"mode": "batch-qa-aecn", "batch_id": "2-201", "certify_url": "u", "drive_folder": "d"}
    g1 = entrypoint.session_prompt(inputs, 1, [1, 2, 3, 4], rf)
    g3 = entrypoint.session_prompt(inputs, 3, [1, 2, 3, 4], rf)
    g4 = entrypoint.session_prompt(inputs, 4, [1, 2, 3, 4], rf)
    assert "Pre-Flight Gate first" in g1 and "Do not start Gate 2" in g1
    assert "confirm Gate 2's exit condition" in g3 and "Do not start Gate 4" in g3
    assert "run_gates.py" in g4 and "CONTAINER_MODE.md" in g4


# ------------------------------------------------------------- Slack posting limits

def _delivery(tmp: Path, data: dict, folder: str = "outputs/qa/batch-INJ-10-2026-W3") -> Path:
    d = tmp / folder
    d.mkdir(parents=True, exist_ok=True)
    p = d / "delivery_messages.json"
    p.write_text(json.dumps(data), encoding="utf-8")
    return p


@test
def slack_only_team_mrv_alerts_no_mentions_no_dms():
    import slack_post
    ok = {"posts": [{"text": "QA complete: INJ-10-2026-W3", "thread": ["detail"]}]}
    with tempfile.TemporaryDirectory() as d:
        assert slack_post.load(_delivery(Path(d), ok))
        for bad in [{**ok, "channel": "C0XXXXXXXXX"},                        # another channel
                    {**ok, "channel": "D08KEFVRJDP"},                        # a DM
                    {"posts": [{"text": "done <@U08JW9HQG5B>"}]},           # mention
                    {"posts": [{"text": "ok", "thread": ["<!here> look"]}]},  # broadcast in a reply
                    {"posts": [{"text": "x", "thread_ts": "1.2", "thread": ["y"]}]},
                    {"posts": []}]:
            try:
                slack_post.load(_delivery(Path(d), bad))
                raise AssertionError(f"not refused: {bad}")
            except slack_post.Refused:
                pass


@test
def slack_batch_post_needs_matching_green_receipt():
    import slack_post
    with tempfile.TemporaryDirectory() as d:
        p = _delivery(Path(d), {"posts": [{"text": "QA complete"}]})
        ck = p.parent / "Batch_INJ-10-2026-W3_QA_Checklist.xlsx"
        ck.write_bytes(b"v1")
        def gate_ok():
            try:
                slack_post.check_delivery_gate(p); return True
            except slack_post.Refused:
                return False
        assert not gate_ok(), "no receipt must refuse"
        (p.parent / "gate_receipt.json").write_text(json.dumps(
            {"all_gates_green": True, "checklist_file": ck.name, "checklist_sha256": hashlib.sha256(b"v1").hexdigest()}))
        assert gate_ok(), "matching green receipt must pass"
        ck.write_bytes(b"v2")
        assert not gate_ok(), "checklist edited after gates must refuse"


@test
def slack_posts_replies_into_the_new_thread():
    import slack_post
    calls = []
    orig = slack_post._post
    slack_post._post = lambda token, text, thread_ts=None: calls.append((text, thread_ts)) or f"ts{len(calls)}"
    try:
        slack_post.post_all({"posts": [{"text": "top", "thread": ["r1", "r2"]}, {"thread_ts": "9.9", "text": "re"}]}, "xoxb-test")
    finally:
        slack_post._post = orig
    assert calls == [("top", None), ("r1", "ts1"), ("r2", "ts1"), ("re", "9.9")], calls


# ------------------------------------------------------------- Slack post built from the checklist

@test
def checklist_becomes_the_slack_post():
    import openpyxl
    import checklist_to_slack as c2s
    import slack_post
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "Batch_X_QA_Checklist.xlsx"
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Batch ID", "INJ-10-2026-W3"])
        ws.append(["Item", "Check", "Status", "Evidence", "Verifier Reference"])
        ws.append(["1. Batch Folder Completeness"])
        ws.append(["1.1", "Week folder exists", "Flag", "Testing/ missing at Gate 1", "VR App. 1"])
        ws.append(["1.2", "Scale tickets present", "Pass", "26 tickets <@U08JW9HQG5B> & co", "VR App. 1"])
        ws.append(["2. Scale Tickets & Mass Determination"])
        ws.append(["2.3", "Tractor/trailer match", "Fail", "CH1 full vs CH2 empty", "VR Sec 5.1"])
        ws.append(["2.4", "Not applicable here", "N/A", "", ""])
        wb.save(p)
        out = c2s.build(p, "INJ-10-2026-W3", ["Evidence changed after review: Testing/ added"], None)
        top = out["posts"][0]["text"]
        assert "1 Pass · 1 Fail · 1 Flag · 1 N/A" in top, top
        assert "2.3" in top and "Tractor/trailer match" in top, "every Fail must be in the top post"
        assert "Evidence changed after review" in top
        thread = out["posts"][0]["thread"]
        assert len(thread) == 2 and "Batch Folder Completeness" in thread[0] and "2.3 Fail" in thread[1], thread
        assert "&lt;@U08JW9HQG5B&gt;" in thread[0], "evidence text must not become a real mention"
        dj = Path(d) / "delivery_messages.json"
        dj.write_text(json.dumps(out), encoding="utf-8")
        slack_post.load(dj)                                   # passes the poster's own checks


# ------------------------------------------------------------- discovery

@test
def discover_classifies_batch_types_from_component_names():
    import discover as d
    cases = {
        "AECN": ["Injection sequestration", "AECN pyrolysis (from rail loads) - process emissions", "SOPOR to Baton Rouge Rail Transport Process Emissions"],
        "AECN-Truck": ["AECN pyrolysis (from truck loads) - process emissions", "AECN to LA Tanker Truck Transport Process Emissions - AECN"],
        "QOWV": ["CO to LA Tanker Truck Transport Process Emissions - QOWV", "QOWV Pyrolysis Emissions"],
        "WODO": ["CO to LA Flatbed Truck Transport Process Emissions - WODO", "Pyrolysis - Charm WODO PP Emissions"],
        "KERRY": ["Kerry to Basco Flatbed Truck Transport Process Emissions"],
    }
    for want, names in cases.items():
        got, _ = d.classify(names + ["B6 Injection site diesel emissions"])      # site emissions ignored
        assert got == want, (want, got)
    mixed, note = d.classify(["AECN pyrolysis (from rail loads) - x", "AECN pyrolysis (from truck loads) - y"])
    assert mixed == "AECN" and "rail and truck" in note
    none, why = d.classify(["QOWV Pyrolysis Emissions", "Kerry to Basco Flatbed Truck"])
    assert none is None and "several" in why


@test
def discover_week_labels_match_known_batches():
    import discover as d
    known = {"2026-10-16": "INJ-10-2026-W3", "2026-10-02": "INJ-10-2026-W1", "2026-09-25": "INJ-09-2026-W4",
             "2026-08-29": "INJ-08-2026-W5", "2026-06-06": "INJ-06-2026-W1", "2026-07-11": "INJ-07-2026-W2"}
    for completed, label in known.items():
        assert d.week_label(completed) == label, (completed, d.week_label(completed))
    assert d.week_label("2026-10-16", -1) == "INJ-10-2026-W2"
    assert d.rp_period("October 1 - October 31 RP", "INJ-10-2026-W3") == "2026-10-01 to 2026-10-31"


@test
def discover_rail_plus_truck_pairs_only_exact_sums():
    import discover as d
    rail, truck = [153160.0, 12.5, 999.0], [156495.07, 3.2, 88.0]
    assert d.mass_as_sum(rail, truck, [309655.07])[1:] == (153160.0, 156495.07)
    assert d.mass_as_sum(rail, truck, [309655.10]) is None, "0.03 kg off must not match"


# ------------------------------------------------------------- image contents

@test
def image_shell_scripts_have_unix_line_endings():
    """In the image, a CRLF shell script breaks every hook: bash fails on the stray \\r and the
    fail-closed hooks then block every tool call (2026-10-22 smoke test). Only checked during
    the image build (QA_IMAGE_BUILD=1); a Windows checkout legitimately has CRLF on disk."""
    if os.environ.get("QA_IMAGE_BUILD") != "1":
        return
    bad = [str(p.relative_to(WS)) for p in list((WS / ".claude/hooks").glob("*.sh")) + list((WS / "skills/qa").rglob("*.sh"))
           if b"\r" in p.read_bytes()]
    assert not bad, f"CRLF line endings in the image: {bad}"


# ------------------------------------------------------------- Max's own gate test, unchanged

@test
def max_findings_never_block_test_still_passes():
    t = WS / "skills/qa/batch-qa-charm/tests/test_findings_never_block.py"
    r = subprocess.run([sys.executable, str(t)], capture_output=True, text=True, cwd=WS)
    assert r.returncode == 0, (r.stdout + r.stderr)[-800:]


if __name__ == "__main__":
    for name, ok, msg in RESULTS:
        print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"\n      {msg}" if msg else ""))
    failed = sum(not ok for _, ok, _ in RESULTS)
    print(f"\n{len(RESULTS) - failed}/{len(RESULTS)} passed")
    sys.exit(1 if failed else 0)
