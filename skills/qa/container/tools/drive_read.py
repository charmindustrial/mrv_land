#!/usr/bin/env python3
"""
drive_read.py — Google Drive access for the QA container.

The QA skills read and publish Drive files through this file (the Google
Drive MCP connector is the fallback on a workstation without Google
credentials; the Mac Drive mount is no longer used). Authenticates as the Cloud Run job's service
account (Application Default Credentials); there is no key file. The batch
folders must be shared with that service account (or it must be a member of
the shared drive).

Commands:
  ls <folder id | URL> [--recursive] [--out FILE]
      Every file in the folder (and subfolders with --recursive), with path,
      name, mimeType, size, modifiedTime, owners and webViewLink. Saved with
      --out, this listing is the artifact an absence claim cites
      (`artifact_path`), the way a Drive snapshot YAML was before.

  search "<Drive query>" [--out FILE]
      e.g. search "name contains 'VVB Guidance'". Trashed files are excluded.

  get <file id | URL> --out DIR [--format xlsx|csv|pdf|docx|txt]
      Download a file. Google Sheets export to xlsx (notes and links kept),
      Docs to docx, Slides to pdf, unless --format says otherwise.

  tabs <local xlsx>
      Tab names and sizes of a workbook downloaded with `get`.

  rows <local xlsx> --tab NAME [--match REGEX] [--context N] [--comments] [--out FILE]
      Only the rows of one tab that match REGEX (e.g. the batch ID in COBB),
      plus the header row, with each cell's coordinate, cached value, formula,
      comment and link. Scans every row and reports rows_scanned and
      match_count, so "not found" is a full-tab result. --comments adds every
      cell comment on the tab. The full workbook stays on disk from `get`.

  put <local file> --folder <folder id | URL>
      Upload (or overwrite a same-named file in) a folder. Only folders listed
      in DRIVE_WRITE_FOLDERS may be written. A QA checklist xlsx under
      outputs/qa/batch-*/ is refused unless that folder's gate_receipt.json is
      green and its checklist_sha256 matches the file being uploaded — the
      container's equivalent of qa_delivery_gate_guard.sh. Never deletes.
"""
from __future__ import annotations

import argparse
import datetime
import hashlib
import io
import json
import os
import re
import sys
from pathlib import Path

SCOPES = ["https://www.googleapis.com/auth/drive"]
FOLDER_MIME = "application/vnd.google-apps.folder"
FIELDS = "id,name,mimeType,size,modifiedTime,owners(displayName,emailAddress),webViewLink,parents"
EXPORT_DEFAULTS = {
    "application/vnd.google-apps.spreadsheet": "xlsx",
    "application/vnd.google-apps.document": "docx",
    "application/vnd.google-apps.presentation": "pdf",
}
EXPORT_MIME = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "csv": "text/csv",
    "txt": "text/plain",
}


def _now() -> str:
    return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def file_id(ref: str) -> str:
    for pat in (r"/folders/([\w-]+)", r"/d/([\w-]+)", r"[?&]id=([\w-]+)"):
        m = re.search(pat, ref)
        if m:
            return m.group(1)
    return ref.strip()


def service():
    # Imported here so the local-file commands (tabs, rows) need no Google libraries.
    global MediaFileUpload, MediaIoBaseDownload
    import google.auth
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload, MediaIoBaseDownload
    creds, _ = google.auth.default(scopes=SCOPES)
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def _list_children(svc, folder: str) -> list[dict]:
    out, token = [], None
    while True:
        resp = svc.files().list(
            q=f"'{folder}' in parents and trashed = false",
            fields=f"nextPageToken, files({FIELDS})",
            pageSize=1000, pageToken=token,
            supportsAllDrives=True, includeItemsFromAllDrives=True,
        ).execute()
        out.extend(resp.get("files", []))
        token = resp.get("nextPageToken")
        if not token:
            return out


def cmd_ls(svc, ref: str, recursive: bool) -> dict:
    root = file_id(ref)
    meta = svc.files().get(fileId=root, fields=FIELDS, supportsAllDrives=True).execute()
    rows, stack = [], [(root, meta["name"])]
    while stack:
        fid, path = stack.pop()
        for f in _list_children(svc, fid):
            f["path"] = f"{path}/{f['name']}"
            rows.append(f)
            if recursive and f["mimeType"] == FOLDER_MIME:
                stack.append((f["id"], f["path"]))
    rows.sort(key=lambda r: r["path"])
    return {"fetched_at_utc": _now(), "folder": meta, "live_source_url": meta.get("webViewLink"),
            "recursive": recursive, "file_count": sum(r["mimeType"] != FOLDER_MIME for r in rows),
            "files": rows}


def cmd_search(svc, query: str) -> dict:
    out, token = [], None
    while True:
        resp = svc.files().list(
            q=f"({query}) and trashed = false", fields=f"nextPageToken, files({FIELDS})",
            pageSize=200, pageToken=token, supportsAllDrives=True,
            includeItemsFromAllDrives=True, corpora="allDrives",
        ).execute()
        out.extend(resp.get("files", []))
        token = resp.get("nextPageToken")
        if not token:
            return {"fetched_at_utc": _now(), "query": query, "files": out}


def cmd_get(svc, ref: str, out_dir: Path, fmt: str | None) -> dict:
    fid = file_id(ref)
    meta = svc.files().get(fileId=fid, fields=FIELDS, supportsAllDrives=True).execute()
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[\\/:*?\"<>|]", "_", meta["name"])
    native = meta["mimeType"].startswith("application/vnd.google-apps.")
    if native:
        fmt = fmt or EXPORT_DEFAULTS.get(meta["mimeType"], "pdf")
        request = svc.files().export_media(fileId=fid, mimeType=EXPORT_MIME[fmt])
        dest = out_dir / f"{safe}.{fmt}"
    else:
        request = svc.files().get_media(fileId=fid, supportsAllDrives=True)
        dest = out_dir / safe
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = dl.next_chunk()
    dest.write_bytes(buf.getvalue())
    result = {"fetched_at_utc": _now(), "file": meta, "live_source_url": meta.get("webViewLink"),
              "saved_to": str(dest), "bytes": dest.stat().st_size}
    # Sidecar so `tabs` / `rows` can cite the live Drive URL of the file they read.
    Path(str(dest) + ".meta.json").write_text(json.dumps(result, indent=1))
    return result


def _xlsx(local: Path):
    import openpyxl  # only the sheet commands need it
    formulas = openpyxl.load_workbook(local, data_only=False)   # formulas, comments, links
    values = openpyxl.load_workbook(local, data_only=True)      # cached values Sheets/Excel computed
    meta_path = Path(str(local) + ".meta.json")
    meta = json.loads(meta_path.read_text()) if meta_path.exists() else {}
    return formulas, values, meta.get("live_source_url")


def cmd_tabs(local: Path) -> dict:
    wb, _, url = _xlsx(local)
    return {"file": str(local), "live_source_url": url,
            "tabs": [{"name": ws.title, "dimensions": ws.dimensions, "max_row": ws.max_row,
                      "max_column": ws.max_column} for ws in wb.worksheets]}


def _cell(fc, vc) -> dict | None:
    d = {}
    if vc.value is not None:
        d["value"] = vc.value if isinstance(vc.value, (int, float, str, bool)) else str(vc.value)
    if isinstance(fc.value, str) and fc.value.startswith("="):
        d["formula"] = fc.value
    if fc.comment is not None:
        d["comment"] = fc.comment.text
    if fc.hyperlink is not None and fc.hyperlink.target:
        d["hyperlink"] = fc.hyperlink.target
    return d or None


def _row(wsf, wsv, r: int) -> dict:
    cells = {}
    for fc in wsf[r]:
        d = _cell(fc, wsv[fc.coordinate])
        if d:
            cells[fc.coordinate] = d
    return cells


def cmd_rows(local: Path, tab: str, match: str | None, header_row: int | None,
             context: int, comments: bool) -> dict:
    """Rows of one tab whose cells match a regex (or the whole tab when no
    --match), with the header row, so values can be cited by tab + cell."""
    wbf, wbv, url = _xlsx(local)
    if tab not in wbf.sheetnames:
        sys.exit(f"No tab {tab!r}. Tabs: {wbf.sheetnames}")
    wsf, wsv = wbf[tab], wbv[tab]
    if header_row is None:      # first row with 3+ non-empty cells
        header_row = next((r for r in range(1, min(wsv.max_row, 30) + 1)
                           if sum(c.value is not None for c in wsv[r]) >= 3), 1)
    pat = re.compile(match) if match else None
    hits = []
    for r in range(1, wsv.max_row + 1):
        if pat is None or any(c.value is not None and pat.search(str(c.value)) for c in wsv[r]):
            hits.append(r)
    keep = sorted({x for r in hits for x in range(max(1, r - context), min(wsv.max_row, r + context) + 1)})
    out = {
        "fetched_at_utc": _now(), "file": str(local), "live_source_url": url, "tab": tab,
        "match": match, "rows_scanned": wsv.max_row, "match_count": len(hits), "matched_rows": hits,
        "header_row": header_row, "header": _row(wsf, wsv, header_row),
        "rows": [{"row": r, "cells": _row(wsf, wsv, r)} for r in keep],
    }
    if comments:
        out["all_comments"] = {c.coordinate: c.comment.text for row in wsf.iter_rows()
                               for c in row if c.comment is not None}
    return out


def _check_delivery_gate(local: Path) -> None:
    """Refuse to publish a batch checklist that has no matching green receipt."""
    if not re.search(r"QA_Checklist.*\.xlsx$", local.name):
        return
    if not re.search(r"outputs/qa/batch-[^/]+/", local.resolve().as_posix() + "/"):
        return
    receipt_path = local.resolve().parent / "gate_receipt.json"
    if not receipt_path.exists():
        sys.exit(f"DELIVERY BLOCKED: no gate_receipt.json next to {local.name}. Run run_gates.py first.")
    receipt = json.loads(receipt_path.read_text())
    sha = hashlib.sha256(local.read_bytes()).hexdigest()
    if not receipt.get("all_gates_green") or receipt.get("checklist_sha256") != sha:
        sys.exit(f"DELIVERY BLOCKED: gate_receipt.json does not match {local.name} "
                 f"(edited after run_gates.py, or gates not green). Re-run run_gates.py.")


def cmd_put(svc, local: Path, folder_ref: str) -> dict:
    folder = file_id(folder_ref)
    allowed = {file_id(f) for f in os.environ.get("DRIVE_WRITE_FOLDERS", "").split(",") if f.strip()}
    if folder not in allowed:
        sys.exit(f"WRITE BLOCKED: folder {folder} is not in DRIVE_WRITE_FOLDERS for this run.")
    if not local.is_file():
        sys.exit(f"No such file: {local}")
    _check_delivery_gate(local)
    name = local.name.replace("'", "\\'")
    existing = svc.files().list(
        q=f"'{folder}' in parents and name = '{name}' and trashed = false",
        fields="files(id)", supportsAllDrives=True, includeItemsFromAllDrives=True,
    ).execute().get("files", [])
    media = MediaFileUpload(str(local), resumable=True)
    if existing:
        meta = svc.files().update(fileId=existing[0]["id"], media_body=media,
                                  fields=FIELDS, supportsAllDrives=True).execute()
        action = "overwritten"
    else:
        meta = svc.files().create(body={"name": local.name, "parents": [folder]}, media_body=media,
                                  fields=FIELDS, supportsAllDrives=True).execute()
        action = "created"
    return {"uploaded_at_utc": _now(), "action": action, "file": meta}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ls"); p.add_argument("ref"); p.add_argument("--recursive", action="store_true"); p.add_argument("--out")
    p = sub.add_parser("search"); p.add_argument("query"); p.add_argument("--out")
    p = sub.add_parser("get"); p.add_argument("ref"); p.add_argument("--out", required=True)
    p.add_argument("--format", choices=sorted(EXPORT_MIME))
    p = sub.add_parser("put"); p.add_argument("local"); p.add_argument("--folder", required=True)
    p = sub.add_parser("tabs"); p.add_argument("local"); p.add_argument("--out")
    p = sub.add_parser("rows"); p.add_argument("local"); p.add_argument("--tab", required=True)
    p.add_argument("--match"); p.add_argument("--header-row", type=int)
    p.add_argument("--context", type=int, default=0); p.add_argument("--comments", action="store_true")
    p.add_argument("--out")
    args = ap.parse_args()

    if args.cmd in ("tabs", "rows"):    # local xlsx from `get`; no Drive call
        result = (cmd_tabs(Path(args.local)) if args.cmd == "tabs" else
                  cmd_rows(Path(args.local), args.tab, args.match, args.header_row, args.context, args.comments))
        text = json.dumps(result, indent=1, default=str)
        if args.out:
            Path(args.out).parent.mkdir(parents=True, exist_ok=True)
            Path(args.out).write_text(text)
            print(f"wrote {args.out}" + (f" ({result['match_count']} matching rows of "
                                         f"{result['rows_scanned']})" if args.cmd == "rows" else ""))
        else:
            print(text)
        return 0

    svc = service()
    if args.cmd == "ls":
        result = cmd_ls(svc, args.ref, args.recursive)
    elif args.cmd == "search":
        result = cmd_search(svc, args.query)
    elif args.cmd == "get":
        print(json.dumps(cmd_get(svc, args.ref, Path(args.out), args.format), indent=1))
        return 0
    else:
        print(json.dumps(cmd_put(svc, Path(args.local), args.folder), indent=1))
        return 0

    text = json.dumps(result, indent=1)
    if args.out:
        Path(args.out).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out).write_text(text)
        print(f"wrote {args.out} ({len(result.get('files', []))} entries)")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
